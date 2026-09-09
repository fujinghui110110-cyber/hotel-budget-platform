from decimal import Decimal, InvalidOperation, ROUND_HALF_UP

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from budgeting.models import ValidationIssue


CHANNEL_REVENUE_ROWS = (101, 102, 103, 105, 107, 108, 110, 111, 112, 114, 115,
                        117, 118, 119, 120, 121, 122, 124, 126, 127, 128, 129, 130)


def validate_channel_values(workbook_path, run):
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheet_name = "A1客房收入(新)"
    if sheet_name not in workbook.sheetnames:
        workbook.close()
        return
    grid = list(workbook[sheet_name].iter_rows(min_row=1, max_row=130, max_col=22, values_only=True))
    def number(row, column):
        raw = grid[row - 1][column - 1]
        if raw is None or isinstance(raw, bool):
            return None
        try:
            value = Decimal(str(raw))
            return value if value.is_finite() else None
        except InvalidOperation:
            return None

    for row in CHANNEL_REVENUE_ROWS:
        for column in range(11, 23):
            revenue, rate, nights = number(row, column), number(row - 64, column), number(row - 32, column)
            if any(value is None for value in (revenue, rate, nights)):
                continue
            expected = (rate * nights).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
            actual = revenue.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
            if actual != expected:
                letter = get_column_letter(column)
                ValidationIssue.objects.get_or_create(run=run, code="CHANNEL_REVENUE_RECONCILIATION",
                    location=f"{sheet_name}!{letter}{row}", defaults={"severity": "P2",
                    "message": f"渠道收入不等于渠道房价 × 房晚（{letter}{row - 64} × {letter}{row - 32}）。",
                    "actual_value": str(actual), "expected_value": str(expected)})
    workbook.close()
