from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook

from budgeting.models import ValidationIssue


def validate_legacy_balancing_inputs(workbook_path, run):
    workbook = load_workbook(workbook_path, read_only=True, data_only=False, keep_links=False)
    try:
        if "经营补充指标" not in workbook.sheetnames:
            return
        sheet = workbook["经营补充指标"]
        if "尾差" not in str(sheet["A8"].value or ""):
            return
        for row in sheet.iter_rows(min_row=8, max_row=8, min_col=4, max_col=15):
            for cell in row:
                if cell.value in (None, ""):
                    continue
                try:
                    value = Decimal(str(cell.value))
                    is_zero = value.is_finite() and value == 0
                except InvalidOperation:
                    is_zero = False
                if is_zero:
                    continue
                ValidationIssue.objects.get_or_create(
                    run=run,
                    code="UNSUPPORTED_BALANCING_INPUT",
                    location=f"{sheet.title}!{cell.coordinate}",
                    defaults={
                        "severity": "P0",
                        "message": "送餐服务费尾差调整缺少已确认业务依据，不能用于平账。原件和金额保留；请管理员核实来源并通过有痕修订处理。",
                    },
                )
    finally:
        workbook.close()
