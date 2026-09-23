from __future__ import annotations

import hashlib
import json
import re
import re
import shutil
import subprocess
import tempfile
import zipfile
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter

from budgeting.excel.ooxml import BLOCKED_PREFIXES, cached_errors, formula_manifest, repair_drawing_namespaces, validate_xlsx_zip
from budgeting.management.commands.build_template_v2 import (
    HISTORY_SHEET,
    _existing_input_cells,
    _load_base_manifest,
    _merge_input_cells,
    _sheet_manifest,
)
from budgeting.models import REPORTS
from budgeting.services.metrics import METRICS


TEMPLATE_VERSION = "V3"
RULE_VERSION = "R3"
SOURCE_YEAR = 2026
# Coordinates verified against the controlled V1 source. Never search arbitrary
# business input cells for numbers that happen to resemble a year.
YEAR_LABEL_CELLS = {
    "预算编制说明": ("A1", "A2", "A4", "A9", "A11", "A13"),
    "利润表 ": ("F26", "G26", "H26"),
    "五年铺排": ("J21", "L21", "N21", "P21"),
    "工资福利费": ("N159",),
    **{name: ("Y21", "AA21", "AD21", "AF21", "AH21", "AK21") for name in (
        "酒店损益总表（含名酒）", "酒店损益总表（不含名酒）",
    )},
    **{name: ("F4",) for name in (
        "损益表（含名酒）（拆中智）", "损益表（不含名酒）（拆中智）",
    )},
    **{name: ("X21", "Z21", "AC21", "AE21", "AG21", "AJ21", "AK21") for name in (
        "A房务部", "A1客房收入(新)", "B餐饮部汇总", "B营业点汇总", "B2宴会收入",
    )},
    **{name: ("X21", "Z21", "AA21", "AC21", "AE21", "AG21", "AH21", "AJ21") for name in (
        "E1行政管理", "E2信息及通讯", "E3市场销售", "E4工程部", "E5能耗",
    )},
}
SUPPLEMENTARY_SHEET = "经营补充指标"
MONTH_COLUMNS = tuple("KLMNOPQRSTUV")
SUPPLEMENTARY_MONTH_COLUMNS = tuple("DEFGHIJKLMNO")
REPORT_CONTRACTS = {
    "PL_TOTAL_WINE": ("酒店损益总表（含名酒）", "J", tuple("LMNOPQRSTUVW"), {29: (28, 24), 30: (41, 28)}),
    "PL_TOTAL_NOWINE": ("酒店损益总表（不含名酒）", "J", tuple("LMNOPQRSTUVW"), {29: (28, 24), 30: (41, 28)}),
    "PL_ZZ_WINE": ("损益表（含名酒）（拆中智）", "S", tuple("FGHIJKLMNOPQ"), {12: (11, 10), 13: (21, 11)}),
    "PL_ZZ_NOWINE": ("损益表（不含名酒）（拆中智）", "S", tuple("FGHIJKLMNOPQ"), {12: (11, 10), 13: (21, 11)}),
}
B12_REQUIRED_INPUTS = tuple(f"{column}{row}" for row in (98, 99) for column in ("K", "L", "M", "N", "R"))
ZZ_REQUIRED_INPUTS = tuple(f"{column}{row}" for row in (121, 122, 123) for column in tuple("FGHIJKLMNOPQ"))
ZERO_BASELINE_DIRECT_INPUTS = {
    "A21前台": ("V26",),
    "B12中餐厅": ("K39", "L39", "Q39", "S39", "V39"),
    "B3宴会厅": ("K49", "V41"),
    "C3康体中心": ("O112", "P93", "R65", "V76"),
    "C5车队": ("Q47", "R47"),
    "E11行政办": ("L63", "O52"),
    "E12人力资源": ("N59", "O59"),
    "E2信息及通讯": ("N52",),
    "E31市场销售部": tuple(f"{column}45" for column in "LMNOPQRSTUV"),
    "E32公关部": ("M73", "P73", "U63", "V63"),
    "E4工程部": ("Q71", "V26"),
    "E5能耗": ("K25", "L25", "U24"),
    "G1员工餐厅分摊": tuple(f"{column}32" for column in MONTH_COLUMNS),
}
SUPPLEMENTARY_ADJUSTMENTS = (
    (5, "餐饮部印刷与文具调整（元）", "B餐饮部汇总", 68),
    (6, "固定支出杂项调整（元）", "F固定支出", 34),
    (7, "前台旅行社佣金调整（元）", "A21前台", 66),
)


def sha256_file(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def input_manifest_hash(input_cells):
    payload = json.dumps(input_cells, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def rolling_history_columns(budget_year):
    columns = []
    index = 2
    for year, kind in ((budget_year - 3, "ACTUAL"), (budget_year - 2, "ACTUAL"), (budget_year - 1, "FORECAST")):
        prefix = "A" if kind == "ACTUAL" else "F"
        for month in range(1, 13):
            columns.append({
                "column": get_column_letter(index),
                "period": f"{prefix}{year}M{month:02d}",
                "year": year,
                "kind": kind,
                "month": month,
            })
            index += 1
    return columns


def history_rows():
    rows = {}
    for metric_code, metric in METRICS.items():
        aggregation = metric.get("aggregation", "SUM")
        unit = metric.get("unit", "MONEY")
        if unit != "RATIO" and aggregation != "DERIVED":
            for report_code, row_code in (metric.get("rows") or {}).items():
                if not row_code or "+" in row_code:
                    continue
                key = (report_code, row_code)
                rows.setdefault(key, {
                    "report_code": report_code,
                    "metric_code": metric_code,
                    "row_code": row_code,
                    "row_label": metric.get("label", metric_code),
                    "unit": unit,
                    "aggregation": aggregation,
                    "input": True,
                })
        if aggregation in {"RATIO", "DERIVED"}:
            if metric_code == "occ":
                component_units = {"numerator": "COUNT", "denominator": "COUNT"}
            elif metric_code in {"adr", "revpar", "cost_room_operating_per_room"}:
                component_units = {"numerator": "MONEY", "denominator": "COUNT"}
            else:
                component_units = {"numerator": "MONEY", "denominator": "MONEY"}
            for suffix, field in (("numerator", "numerator_rows"), ("denominator", "denominator_rows")):
                for report_code, row_code in (metric.get(field) or {}).items():
                    if not row_code or "+" in row_code:
                        continue
                    key = (report_code, row_code)
                    rows.setdefault(key, {
                        "report_code": report_code,
                        "metric_code": f"{metric_code}_{suffix}",
                        "row_code": row_code,
                        "row_label": f"{metric.get('label', metric_code)}（{'分子' if suffix == 'numerator' else '分母'}）",
                        "unit": component_units[suffix],
                        "aggregation": "SUM",
                        "input": True,
                    })
    result = list(rows.values())
    for row_number, row in enumerate(result, start=3):
        row["row_number"] = row_number
    return result


def add_history_sheet(workbook, budget_year):
    if HISTORY_SHEET in workbook.sheetnames:
        del workbook[HISTORY_SHEET]
    sheet = workbook.create_sheet(HISTORY_SHEET)
    columns, rows = rolling_history_columns(budget_year), history_rows()
    sheet["A1"] = f"历史月度输入（{budget_year - 3}—{budget_year - 1}，仅填黄色单元格）"
    sheet["A2"] = "报表 / 指标"
    for cell in (sheet["A1"], sheet["A2"]):
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    input_cells = []
    for item in columns:
        cell = sheet[f"{item['column']}2"]
        cell.value = item["period"]
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="5B9BD5")
        cell.alignment = Alignment(horizontal="center")
    for row in rows:
        unit_label = {"MONEY": "元", "COUNT": "房晚/数量", "RATIO": "%"}.get(row["unit"], row["unit"])
        sheet[f"A{row['row_number']}"] = f"{REPORTS.get(row['report_code'], row['report_code'])} · {row['row_label']}（{unit_label}）"
        for item in columns:
            cell = sheet[f"{item['column']}{row['row_number']}"]
            cell.fill = PatternFill("solid", fgColor="FFF2CC")
            input_cells.append(cell.coordinate)
    sheet.freeze_panes = "B3"
    sheet.column_dimensions["A"].width = 48
    for item in columns:
        sheet.column_dimensions[item["column"]].width = 13
    return columns, rows, input_cells


def add_supplementary_sheet(workbook, wine_source=None):
    if SUPPLEMENTARY_SHEET in workbook.sheetnames:
        del workbook[SUPPLEMENTARY_SHEET]
    sheet = workbook.create_sheet(SUPPLEMENTARY_SHEET)
    sheet["A1"] = "经营补充指标"
    sheet["A2"] = "名酒收入已含在其他收入中，名酒成本已含在其他销售成本中；仅单列名酒明细，不得重复加总，也不得用整个科目代替"
    sheet["A3"] = "名酒收入（元）"
    sheet["C2"] = "年度"
    sheet["C3"] = '=IF(COUNT(D3:O3)=12,SUM(D3:O3),"")' if wine_source else "=SUM(D3:O3)"
    for month, column in enumerate(range(4, 16), start=1):
        sheet.cell(row=2, column=column, value=f"{month:02d}月")
        cell = sheet.cell(row=3, column=column)
        if wine_source:
            source_sheet, source_row, first_column = wine_source
            ref = f"'{source_sheet}'!{get_column_letter(first_column + month - 1)}{source_row}"
            cell.value = f'=IF(COUNT({ref})=1,{ref},"")'
        else:
            cell.fill = PatternFill("solid", fgColor="FFF2CC")
            cell.font = Font(color="0000FF")
    if wine_source:
        sheet["A2"] = "名酒收入自动读取原明细，请在 OOD-其他 的名酒明细行填报；本页只展示，不重复填报或计入收入"
    sheet["A4"] = "模板清理调整项（默认留空，按项目实际需要填报）"
    input_cells = [] if wine_source else [f"{get_column_letter(column)}3" for column in range(4, 16)]
    for row, label, _, _ in SUPPLEMENTARY_ADJUSTMENTS:
        sheet.cell(row=row, column=1, value=label)
        sheet.cell(row=row, column=3, value=f"=SUM(D{row}:O{row})")
        for column in range(4, 16):
            cell = sheet.cell(row=row, column=column)
            cell.fill = PatternFill("solid", fgColor="FFF2CC")
            cell.font = Font(color="0000FF")
            input_cells.append(cell.coordinate)
    sheet["A1"].font = Font(bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    sheet.column_dimensions["A"].width = 64
    sheet.column_dimensions["C"].width = 15
    for column in range(4, 16):
        sheet.column_dimensions[get_column_letter(column)].width = 12
    return input_cells


def _replace_year(value, budget_year):
    shifts = {SOURCE_YEAR - offset: budget_year - offset for offset in range(4)}
    if isinstance(value, datetime):
        try:
            return value.replace(year=shifts.get(value.year, value.year))
        except ValueError:
            return value
    if isinstance(value, date):
        try:
            return value.replace(year=shifts.get(value.year, value.year))
        except ValueError:
            return value
    if isinstance(value, int) and not isinstance(value, bool) and value in shifts:
        return shifts[value]
    if isinstance(value, str):
        if value.startswith("="):
            return value
        # A single substitution prevents a shifted year being shifted again
        # when generating a previous budget year. Only year-label syntax moves.
        pattern = r"(?<!\d)(?:2023|2024|2025|2026)(?=年|预算|实际|预测|vs|s20)|(?<=vs)(?:2023|2024|2025|2026)(?!\d)|(?<=s)(?:2023|2024|2025|2026)(?=说明)|(?<=求和项:)(?:2023|2024|2025|2026)(?!\d)"
        value = re.sub(pattern, lambda match: str(shifts[int(match.group())]), value)
    return value


def parameterize_year(workbook, budget_year):
    changed = []
    for sheet_name, coordinates in YEAR_LABEL_CELLS.items():
        if sheet_name not in workbook.sheetnames:
            continue
        sheet = workbook[sheet_name]
        for coordinate in coordinates:
            cell = sheet[coordinate]
            new = _replace_year(cell.value, budget_year)
            if new != cell.value:
                cell.value = new
                changed.append(f"{sheet.title}!{cell.coordinate}")
    meta = workbook["SYS_META"]
    meta["A6"] = "project_signature"
    values = {meta[f"A{row}"].value: row for row in range(1, meta.max_row + 1)}
    for key, value in (("template_version", f"{TEMPLATE_VERSION}-{budget_year}"), ("budget_year", budget_year), ("rule_version", RULE_VERSION)):
        row = values.get(key)
        if row:
            meta[f"B{row}"] = value
    return changed


def repair_confirmed_findings(workbook):
    required_inputs = {"B12中餐厅": []}
    repaired = []
    # Earlier purification emitted TM1 header labels as bare formula operands.
    # Preserve their text, using valid Excel string literals at known headers.
    header_literals = {"H1": "!", "H2": "不区分区域", "H3": "金额",
                       "H5": "cny", "H9": "本年度预算", "H11": "2020一上版"}
    for sheet in workbook.worksheets:
        for ref, label in header_literals.items():
            if sheet[ref].value == f"=IFERROR({label},0)":
                sheet[ref] = f'=IFERROR("{label}",0)'
                repaired.append(f"{sheet.title}!{ref}:恢复TM1标题文本的Excel字符串引号")
    for ref in B12_REQUIRED_INPUTS:
        value = workbook["B12中餐厅"][ref].value
        if not isinstance(value, str) or "#REF!" not in value.upper():
            continue
        required_inputs["B12中餐厅"].append(ref)
        workbook["B12中餐厅"][ref] = None
        workbook["B12中餐厅"][ref].fill = PatternFill("solid", fgColor="FFF2CC")
        repaired.append(f"B12中餐厅!{ref}:缺失业务源改为待填输入")
    for sheet_name in ("损益表（含名酒）（拆中智）", "损益表（不含名酒）（拆中智）"):
        required_inputs[sheet_name] = list(ZZ_REQUIRED_INPUTS)
        sheet = workbook[sheet_name]
        for ref in ZZ_REQUIRED_INPUTS:
            sheet[ref] = None
            sheet[ref].fill = PatternFill("solid", fgColor="FFF2CC")
        sheet["S121"] = "=ROUND(SUM(F121:Q121),2)"
        sheet["S122"] = "=ROUND(SUM(F122:Q122),2)"
        sheet["S123"] = "=ROUND(SUM(F123:Q123),2)"
        sheet["S124"] = "=ROUND(S113+S121-SUM(S115:S120,S122)+S123,2)"
        repaired.append(f"{sheet_name}!S124:年度税前利润纳入S123投资收益")

    no_wine_total = workbook["酒店损益总表（不含名酒）"]
    for column in tuple("LMNOPQRSTUVW"):
        ref = f"{column}66"
        no_wine_total[ref] = f"={column}44+{column}56+{column}61+{column}65"
        repaired.append(f"酒店损益总表（不含名酒）!{ref}:恢复酒店经营利润派生公式")

    room_revenue = workbook["A1客房收入(新)"]
    room_revenue["AE38"] = "=IF(AE69=0,0,AE100/AE69)"
    room_revenue["X43"] = "=IF(X74=0,0,X105/X74)"
    repaired.extend(
        [
            "A1客房收入(新)!AE38:同比率按实际分母判零",
            "A1客房收入(新)!X43:同比率按实际分母判零",
        ]
    )

    for sheet_name, refs in ZERO_BASELINE_DIRECT_INPUTS.items():
        required_inputs.setdefault(sheet_name, []).extend(refs)
        sheet = workbook[sheet_name]
        for ref in refs:
            sheet[ref] = None
            sheet[ref].fill = PatternFill("solid", fgColor="FFF2CC")
            repaired.append(f"{sheet_name}!{ref}:项目金额残留改为待填输入")

    for month_index, source_column in enumerate(MONTH_COLUMNS):
        supplement_column = SUPPLEMENTARY_MONTH_COLUMNS[month_index]
        workbook["B餐饮部汇总"][f"{source_column}68"] = (
            f"=ROUND((B1餐厅汇总!{source_column}68+B3宴会厅!{source_column}68),2)"
            f"+'{SUPPLEMENTARY_SHEET}'!{supplement_column}5"
        )
        wage_column = get_column_letter(month_index + 2)
        workbook["F固定支出"][f"{source_column}34"] = (
            f"=(+工资福利费!{wage_column}129)+'{SUPPLEMENTARY_SHEET}'!{supplement_column}6"
        )
        workbook["A21前台"][f"{source_column}66"] = (
            f"=ROUND('A1客房收入(新)'!{source_column}99*A21前台!$W$66,2)"
            f"+'{SUPPLEMENTARY_SHEET}'!{supplement_column}7"
        )
        workbook["B14送餐"][f"{source_column}31"] = (
            f"=ROUND({source_column}25*$W$31,2)"
        )
    repaired.extend(
        f"{sheet_name}!{MONTH_COLUMNS[0]}{source_row}:{MONTH_COLUMNS[-1]}{source_row}:"
        "硬编码月度调整改为经营补充指标显式输入"
        for _, _, sheet_name, source_row in SUPPLEMENTARY_ADJUSTMENTS
    )
    return required_inputs, repaired


def clear_stale_inputs(workbook, input_cells):
    cleared = []
    for sheet_name, refs in input_cells.items():
        sheet = workbook[sheet_name]
        for ref in refs:
            cell = sheet[ref]
            if cell.data_type != "f" and cell.value not in (None, ""):
                cell.value = None
                cleared.append(f"{sheet_name}!{ref}")
    return cleared


def recover_dense_row_numeric_inputs(workbook, input_cells):
    recovered = {}
    for sheet_name, refs in input_cells.items():
        if sheet_name not in workbook.sheetnames:
            continue
        sheet = workbook[sheet_name]
        rows = {}
        for ref in refs:
            cell = sheet[ref]
            rows.setdefault(cell.row, []).append(cell.column)
        for row, columns in rows.items():
            span = max(columns) - min(columns) + 1
            if len(columns) < 8 or span > 14:
                continue
            known = set(columns)
            for column in range(min(columns), max(columns) + 1):
                cell = sheet.cell(row=row, column=column)
                if column in known or cell.data_type == "f":
                    continue
                if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                    recovered.setdefault(sheet_name, []).append(cell.coordinate)
    return {sheet: sorted(set(refs)) for sheet, refs in recovered.items()}


def close_report_mapping_inputs(workbook, reports):
    mapped_inputs = {}
    for report in (reports or {}).values():
        sheet_name = report.get("sheet")
        if sheet_name not in workbook.sheetnames:
            continue
        sheet = workbook[sheet_name]
        for item in report.get("mapping") or []:
            ref = item.get("cell")
            if not ref:
                continue
            cell = sheet[ref]
            if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")):
                continue
            mapped_inputs.setdefault(sheet_name, []).append(ref)
            item["classification"] = "input"
            item["formula"] = None
            item["shared_attributes"] = {}
            item["allowed_functions"] = []
    return {sheet: sorted(set(refs)) for sheet, refs in mapped_inputs.items()}



def refresh_report_mapping_formulas(reports, formulas):
    """Refresh report mapping formulas from the final saved workbook."""
    lookup = {(item["sheet"], item["cell"]): item for item in formulas}
    for report in (reports or {}).values():
        sheet_name = report.get("sheet")
        for item in report.get("mapping") or []:
            source = item.get("source") or {}
            cell = source.get("cell") or item.get("cell")
            sheet = source.get("sheet") or sheet_name
            observed = lookup.get((sheet, cell))
            if observed:
                formula = observed.get("formula") or ""
                item["classification"] = "formula"
                item["formula"] = formula
                item["shared_attributes"] = observed.get("shared_attributes") or {}
                item["allowed_functions"] = sorted({
                    match.group(1).upper()
                    for match in re.finditer(r"([A-Z][A-Z0-9.]*)\(", formula)
                })
            else:
                item["classification"] = "input"
                item["formula"] = None
                item["shared_attributes"] = {}
                item["allowed_functions"] = []

def normalize_blank_template_formulas(workbook):
    denominator_pattern = re.compile(r"/(\s*(?:(?:'[^']+'|[A-Za-z0-9_\u4e00-\u9fff ]+)!)?\$?[A-Z]{1,3}\$?\d+)")
    denominator_sum_pattern = re.compile(r"/(\s*SUM\([^)]*\))", re.I)
    average_pattern = re.compile(r"AVERAGE\(([^()]*)\)", re.I)
    guarded = []
    blanked = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                value = cell.value
                if value == "=":
                    cell.value = None
                    blanked.append(f"{sheet.title}!{cell.coordinate}")
                    continue
                if not isinstance(value, str) or not value.startswith("="):
                    continue
                upper = value.upper()
                if "IFERROR(" in upper or re.search(r"IF\s*\([^,]*(?:=|<>|<=)\s*0", upper):
                    continue
                averages = [match.group(1).strip() for match in average_pattern.finditer(value)]
                if averages:
                    condition = averages[0] if len(averages) == 1 else None
                    count_check = f"COUNT({condition})=0" if condition else "OR(" + ",".join(f"COUNT({item})=0" for item in averages) + ")"
                    cell.value = f"=IF({count_check},0,{value[1:]})"
                    guarded.append(f"{sheet.title}!{cell.coordinate}")
                    continue
                if "/" not in value:
                    continue
                denominators = []
                for match in denominator_pattern.finditer(value):
                    denominator = match.group(1).strip()
                    if denominator not in denominators:
                        denominators.append(denominator)
                for match in denominator_sum_pattern.finditer(value):
                    denominator = match.group(1).strip()
                    if denominator not in denominators:
                        denominators.append(denominator)
                if not denominators:
                    continue
                condition = denominators[0] + "=0" if len(denominators) == 1 else "OR(" + ",".join(f"{item}=0" for item in denominators) + ")"
                cell.value = f"=IF({condition},0,{value[1:]})"
                guarded.append(f"{sheet.title}!{cell.coordinate}")
    return guarded, blanked


def protect_workbook(workbook, input_cells):
    for sheet in workbook.worksheets:
        for ref in input_cells.get(sheet.title, []):
            sheet[ref].protection = Protection(locked=False)
        sheet.protection.sheet = True
        sheet.protection.objects = True
        sheet.protection.scenarios = True
        sheet.protection.selectLockedCells = True
        sheet.protection.selectUnlockedCells = False


def scan_package(path):
    with zipfile.ZipFile(path) as zf:
        names = zf.namelist()
        blocked = [name for name in names if name.startswith(BLOCKED_PREFIXES)]
        external_rels = []
        for name in names:
            if not name.endswith(".rels"):
                continue
            text = zf.read(name).decode("utf-8", errors="ignore")
            if 'TargetMode="External"' in text:
                external_rels.append(name)
        workbook_xml = zf.read("xl/workbook.xml").decode("utf-8", errors="ignore")
    formulas, _ = formula_manifest(path)
    proprietary = [item for item in formulas if any(token in item["formula"].upper() for token in ("DBS(", "VIEW(", "SUBNM("))]
    return {
        "blocked_parts": blocked,
        "external_relationships": external_rels,
        "invalid_defined_names": "#REF!" in workbook_xml,
        "proprietary_formulas": proprietary,
        "cached_errors": cached_errors(path),
        "zip_issues": validate_xlsx_zip(path),
    }


def recalc_libreoffice(source, output_dir, soffice="/opt/homebrew/bin/soffice"):
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    profile = Path(tempfile.mkdtemp(prefix="budget-v3-lo-profile-"))
    command = [soffice, f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "xlsx", "--outdir", str(output_dir), str(source)]
    run = subprocess.run(command, capture_output=True, text=True, timeout=300)
    target = output_dir / Path(source).name
    return target, {"command": command, "returncode": run.returncode, "stdout": run.stdout.strip(), "stderr": run.stderr.strip(), "output_exists": target.exists()}


def validate_management_contract(path):
    workbook = load_workbook(path, read_only=False, data_only=False, keep_links=False)
    checks = []
    for report_code, (sheet_name, annual_col, months, ratios) in REPORT_CONTRACTS.items():
        sheet = workbook[sheet_name]
        annual_rows = range(23, 108) if annual_col == "J" else range(9, 145)
        for row in annual_rows:
            formula = sheet[f"{annual_col}{row}"].value
            populated_months = any(sheet[f"{column}{row}"].value not in (None, "") for column in months)
            if populated_months and row not in ({47, 58, 66, 91, 104, 132, 137} if annual_col == "S" else set()):
                checks.append({"check": "annual_formula", "report": report_code, "cell": f"{annual_col}{row}", "pass": isinstance(formula, str) and formula.startswith("=")})
        for row, (numerator, denominator) in ratios.items():
            formula = str(sheet[f"{annual_col}{row}"].value or "").replace("$", "").upper()
            checks.append({"check": "occ_or_adr", "report": report_code, "cell": f"{annual_col}{row}", "pass": f"{annual_col}{numerator}" in formula and f"{annual_col}{denominator}" in formula})
        if annual_col == "S":
            tax_formula = str(sheet["S124"].value or "").replace("$", "").upper()
            checks.append({"check": "profit_before_tax", "report": report_code, "cell": "S124", "pass": "S113" in tax_formula and "S123" in tax_formula and "S115:S120" in tax_formula and "S122" in tax_formula})
        profit_rows = (33, 90, 105) if annual_col == "J" else (16, 113, 129)
        for row in profit_rows:
            formula = str(sheet[f"{annual_col}{row}"].value or "").upper()
            checks.append({"check": "profit_annual", "report": report_code, "cell": f"{annual_col}{row}", "pass": "SUM(" in formula})
    workbook.close()
    return {"status": "PASS" if all(item["pass"] for item in checks) else "FAIL", "checks": checks, "failed": [item for item in checks if not item["pass"]]}


def repair_wine_total_bridge(workbook):
    """Repair only the verified V3 OOD subset bridge, never infer fee allocation."""
    included = "酒店损益总表（含名酒）"
    excluded = "酒店损益总表（不含名酒）"
    detail = "OOD-其他"
    aggregate = "C其他运营汇总"
    result = {"status": "UNKNOWN", "scope": "V3 OOD明细至总表收入及直接成本",
              "repairs": [], "issues": [], "expense_allocation": "UNKNOWN"}
    if any(name not in workbook.sheetnames for name in (included, excluded, detail, aggregate)):
        result["issues"].append("缺少名酒总表桥接所需工作表，未改写公式")
        return result
    for row, word in ((55, "收入"), (69, "成本")):
        label = str(workbook[detail].cell(row, 8).value or "")
        if "名酒" not in label or word not in label:
            result["issues"].append(f"{detail}!H{row}未明确标为名酒{word}")
    changes = []
    for month in range(12):
        source_column = get_column_letter(11 + month)
        report_column = get_column_letter(12 + month)
        # These exact additive structures are observed in the controlled base.
        income_total = f"=ROUND(({source_column}25+{source_column}26+{source_column}35+{source_column}41+{source_column}45+{source_column}49+{source_column}55+{source_column}61+{source_column}56),2)"
        cost_total = f"=ROUND(SUM({source_column}63:{source_column}71),2)"
        for sheet in (detail, aggregate):
            for row, expected in ((24, income_total), (62, cost_total)):
                if workbook[sheet].cell(row, 11 + month).value != expected:
                    result["issues"].append(f"{sheet}!{source_column}{row}不符合已验证的加总路径")
        for parent, leaf, total in ((57, 55, 24), (58, 69, 62)):
            combined = str(workbook[aggregate].cell(leaf, 11 + month).value or "")
            parts = combined[len("=ROUND(("):-len("),2)")].split("+")
            source_ref = f"'{detail}'!{source_column}{leaf}"
            if not (combined.startswith("=ROUND((") and combined.endswith("),2)")
                    and parts.count(source_ref) == 1 and len(parts) == len(set(parts))
                    and all(re.fullmatch(r"'[^']+'!" + source_column + str(leaf), part) for part in parts)):
                result["issues"].append(f"{aggregate}!{source_column}{leaf}未确认仅一次加总独立名酒明细")
            cell = f"{report_column}{parent}"
            base = workbook[included][cell].value
            current = workbook[excluded][cell].value
            expected_excluded = f"=ROUND('{included}'!{cell}-'{detail}'!{source_column}{leaf},2)"
            expected_included = f"=ROUND({aggregate}!{source_column}{total},2)"
            # The controlled base already deducts wine on the adjustment sheet.
            # Inspect this edge too: checking only the difference of two reports
            # would miss both reports being understated by the same wine amount.
            adjustment_name = "酒店损益总表-明细表（调整用第一轮不用）"
            base_with_adjustment = f"=({aggregate}!{source_column}{total})+('{adjustment_name}'!{cell})"
            adjustment = workbook[adjustment_name][cell].value if adjustment_name in workbook.sheetnames else None
            if adjustment != f"=-'{detail}'!{source_column}{leaf}":
                result["issues"].append(f"{adjustment_name}!{cell}不是已核实的独立名酒扣减，禁止推定调整口径")
            if base not in (base_with_adjustment, expected_included):
                result["issues"].append(f"{included}!{cell}不是已验证的其他经营汇总来源")
            if current not in (base_with_adjustment, expected_excluded):
                result["issues"].append(f"{excluded}!{cell}已有不同口径公式，禁止覆盖")
            changes.append((included, cell, expected_included, base))
            changes.append((excluded, cell, expected_excluded, current))
    if result["issues"]:
        return result
    for sheet, cell, expected, current in changes:
        if current != expected:
            workbook[sheet][cell] = expected
            result["repairs"].append(f"{sheet}!{cell}:按已核实调整页还原含名酒并仅剔除一次明确明细")
    result["input_cells"] = {detail: [
        f"{get_column_letter(column)}{row}"
        for column in range(11, 23) for row in (55, 69)
        if workbook[detail].cell(row, column).data_type != "f"]}
    result["status"] = "PASS"
    result["rows"] = [{"report_row": "R0057", "detail_row": 55, "metric": "revenue_wine"},
                      {"report_row": "R0058", "detail_row": 69, "metric": "cost_wine"}]
    result["note"] = "只修复直接收入与成本；不改变明细成本公式，不推定共同费用和拆中智口径"
    return result


def validate_wine_zz_bridge(workbook, total_bridge, reports):
    """Verify controlled split-view cells, including any differing helper dependencies."""
    from openpyxl.formula import Tokenizer
    from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries

    wine = "损益表（含名酒）（拆中智）"
    nowine = "损益表（不含名酒）（拆中智）"
    total = "酒店损益总表（不含名酒）"
    adjustment = "酒店损益总表-明细表（调整用第一轮不用）"
    result = {"status": "UNKNOWN", "issues": [], "checked_formulas": 0,
              "cross_scope_references": 0, "shared_expenses": "UNCHANGED_SOURCE_FORMULAS",
              "unmapped_formula_differences": []}
    if total_bridge["status"] != "PASS" or any(name not in workbook.sheetnames for name in (wine, nowine)):
        result["issues"].append("未通过总表名酒桥接或缺少拆中智表")
        return result
    refs = {item["cell"] for code in ("PL_ZZ_WINE", "PL_ZZ_NOWINE")
            for item in (reports or {}).get(code, {}).get("mapping", [])}
    mandatory = {f"{get_column_letter(column)}{row}" for column in range(6, 18) for row in (48, 49)}
    if not mandatory.issubset(refs):
        result["issues"].append("受控报表映射缺少完整名酒桥接单元格")
        return result
    for ref in sorted(refs):
        row, column = coordinate_to_tuple(ref)
        a, b = workbook[wine][ref].value, workbook[nowine][ref].value
        formula_a = isinstance(a, str) and a.startswith("=")
        formula_b = isinstance(b, str) and b.startswith("=")
        if not (formula_a or formula_b):
            if a != b or ref in mandatory:
                result["issues"].append(f"{ref}存在不同固定值或缺少必要桥接公式")
            continue
        result["checked_formulas"] += 1
        if formula_a:
            result["cross_scope_references"] += a.count(f"'{total}'!")
        if ref in mandatory:
            parent = 57 if row == 48 else 58
            source = f"{get_column_letter(column + 6)}{parent}"
            expected_a = f"=+ROUND('{total}'!{source}-'{adjustment}'!{source},2)"
            expected_b = f"=+ROUND('{total}'!{source},2)"
            if a != expected_a or b != expected_b:
                result["issues"].append(f"{ref}不是已验证的名酒明细加回关系")
        elif (not formula_a or not formula_b
              or a.replace(f"'{wine}'!", "'SELF'!") != b.replace(f"'{nowine}'!", "'SELF'!")):
            result["issues"].append(f"{ref}存在未确认的跨口径公式差异")
    # Unlabelled scratch cells are not business metrics. A difference may be
    # ignored only if no controlled formula refers to it, including via ranges.
    for row in range(9, 145):
        for column in (*range(6, 18), 19, 20):
            ref = f"{get_column_letter(column)}{row}"
            a, b = workbook[wine][ref].value, workbook[nowine][ref].value
            if ref in refs or a == b or not any(isinstance(v, str) and v.startswith("=") for v in (a, b)):
                continue
            result["unmapped_formula_differences"].append(ref)
    targets = {(sheet, ref) for sheet in (wine, nowine) for ref in result["unmapped_formula_differences"]}
    pending = [(sheet, ref) for sheet in (wine, nowine) for ref in refs]
    visited = set()
    while pending:
        sheet, ref = pending.pop()
        if (sheet, ref) in visited:
            continue
        visited.add((sheet, ref))
        if sheet in (wine, nowine) and ref not in refs:
            a, b = workbook[wine][ref].value, workbook[nowine][ref].value
            a = a.replace(f"'{wine}'!", "'SELF'!") if isinstance(a, str) else a
            b = b.replace(f"'{nowine}'!", "'SELF'!") if isinstance(b, str) else b
            if a != b:
                targets.add((sheet, ref))
                if ref not in result["unmapped_formula_differences"]:
                    result["unmapped_formula_differences"].append(ref)
        if (sheet, ref) in targets:
            result["issues"].append(f"受控报表依赖未确认的辅助差异{sheet}!{ref}")
            continue
        if len(visited) > 200000:
            result["issues"].append("辅助依赖范围超出受控扫描上限，未确认桥接")
            break
        formula = workbook[sheet][ref].value
        if not isinstance(formula, str) or not formula.startswith("="):
            continue
        # The source workbook builds expense references from a fixed sheet-name
        # label and a literal cell address. Resolve only this exact, static form.
        def resolve_static_indirect(match):
            label_ref, target_ref = (value.replace("$", "") for value in match.groups())
            label = workbook[sheet][label_ref]
            target_sheet = label.value
            if (isinstance(target_sheet, str) and label.data_type != "f"
                    and label.protection.locked and target_sheet in workbook.sheetnames):
                pending.append((sheet, label_ref))
                result["resolved_static_indirects"] = result.get("resolved_static_indirects", 0) + 1
                return "'" + target_sheet.replace("'", "''") + "'!" + target_ref
            return match.group(0)
        formula = re.sub(
            r'(?<![A-Z0-9_.])INDIRECT\((\$?[A-Z]{1,3}\$?[1-9][0-9]*)&"!(\$?[A-Z]{1,3}\$?[1-9][0-9]*)"\)',
            resolve_static_indirect, formula, flags=re.IGNORECASE,
        )
        for token in Tokenizer(formula).items:
            if token.type == "FUNC" and token.subtype == "OPEN" and token.value.upper().rstrip("(").split(".")[-1] in {"INDIRECT", "OFFSET"}:
                result["issues"].append(f"{sheet}!{ref}含动态引用，不能证明费用口径一致")
            if token.type != "OPERAND" or token.subtype != "RANGE":
                continue
            reference, target_sheet = token.value, sheet
            if "!" in reference:
                target_sheet, reference = reference.rsplit("!", 1)
                target_sheet = target_sheet.strip("'").replace("''", "'")
            if reference in workbook.defined_names:
                result["issues"].append(f"{sheet}!{ref}含未解析定义名称{reference}")
                continue
            try:
                left, top, right, bottom = range_boundaries(reference.replace("$", ""))
            except ValueError:
                result["issues"].append(f"{sheet}!{ref}含无法确认的依赖{token.value}")
                continue
            if target_sheet not in workbook.sheetnames:
                result["issues"].append(f"{sheet}!{ref}依赖缺失工作表{target_sheet}")
                continue
            target = workbook[target_sheet]
            left, right = left or 1, right or target.max_column
            top, bottom = top or 1, bottom or target.max_row
            if (right - left + 1) * (bottom - top + 1) > 200000:
                result["issues"].append(f"{sheet}!{ref}依赖范围过大，未确认桥接")
                continue
            for r in range(top, bottom + 1):
                for c in range(left, right + 1):
                    pending.append((target_sheet, f"{get_column_letter(c)}{r}"))
    if not result["issues"]:
        result["status"] = "PASS"
        result["note"] = "含名酒行加回调整页的负名酒明细；受控费用公式保持原模板一致，不新增分摊"
    return result


def build_v3(clean_base, original_source, base_manifest_path, output, manifest_path, report_path, budget_year):
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(clean_base, output)
    base = _load_base_manifest(clean_base, base_manifest_path)
    workbook = load_workbook(output, keep_links=False)
    source_inputs = _existing_input_cells(workbook)
    recovered_inputs = recover_dense_row_numeric_inputs(workbook, source_inputs)
    year_changes = parameterize_year(workbook, budget_year)
    required_inputs, repaired = repair_confirmed_findings(workbook)
    wine_bridge = repair_wine_total_bridge(workbook)
    repaired.extend(wine_bridge["repairs"])
    columns, rows, history_inputs = add_history_sheet(workbook, budget_year)
    supplementary_inputs = add_supplementary_sheet(
        workbook, wine_source=("OOD-其他", 55, 11) if wine_bridge["status"] == "PASS" else None)
    denominator_guards, blank_formulas = normalize_blank_template_formulas(workbook)
    zz_bridge = validate_wine_zz_bridge(workbook, wine_bridge, base.get("reports"))
    mapped_inputs = close_report_mapping_inputs(workbook, base.get("reports"))
    input_cells = _merge_input_cells(
        source_inputs,
        recovered_inputs,
        mapped_inputs,
        required_inputs,
        wine_bridge.get("input_cells", {}),
        {HISTORY_SHEET: history_inputs},
        {SUPPLEMENTARY_SHEET: supplementary_inputs},
    )
    if wine_bridge["status"] == "PASS":
        # Older input manifests may still list a restored wine cost formula.
        input_cells["OOD-其他"] = [
            ref for ref in input_cells.get("OOD-其他", [])
            if workbook["OOD-其他"][ref].data_type != "f"
        ]
        for row in (55, 69):
            for column in range(11, 23):
                cell = workbook["OOD-其他"].cell(row, column)
                if cell.data_type == "f":
                    cell.protection = Protection(locked=True)
        wine_display_cells = {f"{get_column_letter(column)}3" for column in range(4, 16)}
        input_cells[SUPPLEMENTARY_SHEET] = [
            ref for ref in input_cells.get(SUPPLEMENTARY_SHEET, []) if ref not in wine_display_cells]
    cleared = clear_stale_inputs(workbook, input_cells)
    protect_workbook(workbook, input_cells)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(output)
    workbook.close()
    drawing_namespace_repairs = repair_drawing_namespaces(output)
    formulas, fingerprint = formula_manifest(output)
    history = {"sheet": HISTORY_SHEET, "header_row": 2, "label_column": "A", "columns": columns, "rows": rows}
    manifest = dict(base)
    refresh_report_mapping_formulas(manifest.get("reports"), formulas)
    manifest.update({
        "template_version": f"{TEMPLATE_VERSION}-{budget_year}",
        "budget_year": budget_year,
        "year": budget_year,
        "rule_version": RULE_VERSION,
        "management_v2": True,
        "management_v3": True,
        "input_cells": input_cells,
        "input_manifest_hash": input_manifest_hash(input_cells),
        "history": history,
        "history_layout": history,
        "history_month_count": len(columns),
        "supplementary": {
            "sheet": SUPPLEMENTARY_SHEET,
            "row_code": "R9003",
            "annual_cell": "C3",
            "month_cells": [f"{get_column_letter(column)}3" for column in range(4, 16)],
            "editable": wine_bridge["status"] != "PASS",
            "included_in_pnl_total": False,
        },
        "formula_count": len(formulas),
        "formula_manifest_hash": fingerprint,
        "strict_formula_fingerprint": fingerprint,
        "source_identity": {"path": str(original_source), "sha256": sha256_file(original_source), "business_sheet_count": 67},
        "clean_base": {"path": str(clean_base), "sha256": sha256_file(clean_base)},
        "required_business_inputs": required_inputs,
        "blank_input_policy": "ZERO_FOR_MAPPED_NON_FORMULA_BUDGET_INPUTS",
        "wine_total_bridge": wine_bridge,
        "wine_zz_bridge": zz_bridge,
        "workbook_sha256": sha256_file(output),
    })
    check_wb = load_workbook(output, read_only=True, keep_links=False)
    manifest["sheets"] = _sheet_manifest(check_wb)
    check_wb.close()
    Path(manifest_path).write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    package = scan_package(output)
    contract = validate_management_contract(output)
    blockers = []
    if any((package["blocked_parts"], package["external_relationships"], package["invalid_defined_names"], package["proprietary_formulas"], package["zip_issues"])):
        blockers.append("OOXML 清理检查未通过")
    if wine_bridge["status"] != "PASS":
        blockers.append("名酒收入/直接成本桥接来源未确认，禁止自动剔除")
    if zz_bridge["status"] != "PASS":
        blockers.append("拆中智名酒桥接存在未确认差异")
    if contract["status"] != "PASS":
        blockers.append("四张管理报表公式契约未通过")
    business_confirmations = []
    if zz_bridge["status"] != "PASS":
        business_confirmations.append("拆中智名酒桥接存在未确认差异")
    report = {
        "status": "LOCAL_TRIAL" if not blockers else "BLOCKED",
        "publication_status": "BLOCKED",
        "publication_blockers": list(dict.fromkeys(blockers + business_confirmations + [
            "尚未完成 Microsoft Excel 黄金样本逐单元格对账",
        ])),
        "budget_year": budget_year,
        "source_sha256_unchanged": sha256_file(original_source) == manifest["source_identity"]["sha256"],
        "year_changes": len(year_changes),
        "stale_input_cells_cleared": len(cleared),
        "confirmed_formula_repairs": repaired,
        "drawing_namespace_repairs": drawing_namespace_repairs,
        "explicit_zero_denominator_guards": len(denominator_guards),
        "empty_formula_nodes_blanked": len(blank_formulas),
        "business_confirmation_required": business_confirmations,
        "history_month_count": len(columns),
        "input_cell_count": sum(len(value) for value in input_cells.values()),
        "package_scan": package,
        "management_contract": contract,
        "wine_total_bridge": wine_bridge,
        "wine_zz_bridge": zz_bridge,
        "template": str(output),
        "manifest": str(manifest_path),
    }
    Path(report_path).write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return manifest, report
