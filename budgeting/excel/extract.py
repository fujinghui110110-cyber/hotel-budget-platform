import functools
import hashlib
import json
import re
from collections import defaultdict
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

from django.conf import settings
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter

from budgeting.excel.money import yuan_to_cents
from budgeting.models import NormalizedValue, REPORTS, ValidationIssue


MONTH_RE = re.compile(r"^(0[1-9]|1[0-2])$")
RATIO_RE = re.compile(r"(OCC|出租率|入住率|率)", re.I)
COUNT_RE = re.compile(r"(人数|房间数|员工|数量|间数|天数|人次)", re.I)
ANNUAL_RE = re.compile(r"(全年合计|年合计|全年|年度|合计|总计|FY)", re.I)
YEAR_COL_RE = re.compile(r"(20\d{2})\s*年\s*(实际|预测|预估|预算)?")
CJK_RE = re.compile(r"[一-鿿]")
SLUG_UNSAFE_RE = re.compile(r"[^\w一-鿿]+")
SKIP_LABELS = {"小计", "合计", "总计", "平衡检查", "备注", "说明", "不用填"}
VALID_AGGREGATIONS = {"SUM", "AVERAGE", "DERIVED", "RATIO", "EXCLUDE"}


def _cell_key(cell_ref):
    return str(cell_ref).replace("$", "").upper()


def _mapped_cell_refs(mapping):
    refs = set()
    for item in mapping:
        cell_ref = item.get("cell") or item.get("source_cell")
        if not cell_ref:
            continue
        refs.add(_cell_key(cell_ref))
        numerator_cell = item.get("numerator_cell") or item.get("ratio_num_cell")
        denominator_cell = item.get("denominator_cell") or item.get("ratio_den_cell")
        if numerator_cell:
            refs.add(_cell_key(numerator_cell))
        if denominator_cell:
            refs.add(_cell_key(denominator_cell))
        unit = (item.get("unit") or NormalizedValue.Unit.MONEY).upper()
        row_code = str(item.get("row_code") or "")
        row_label = str(item.get("row_label") or item.get("label") or "")
        if unit == NormalizedValue.Unit.RATIO and (
            row_code.upper() == "OCC" or RATIO_RE.search(row_label)
        ):
            column = re.sub(r"\d+", "", _cell_key(cell_ref))
            refs.update({f"{column}24", f"{column}28"})
    return refs


def _read_mapped_cells(worksheet, refs, *, formulas=False):
    if not refs:
        return {}
    coordinates = [coordinate_to_tuple(ref) for ref in refs]
    values = {}
    min_row = min(row for row, _ in coordinates)
    min_col = min(column for _, column in coordinates)
    for row_index, row in enumerate(
        worksheet.iter_rows(
            min_row=min_row,
            max_row=max(row for row, _ in coordinates),
            min_col=min_col,
            max_col=max(column for _, column in coordinates),
        ),
        start=min_row,
    ):
        for column_index, cell in enumerate(row, start=min_col):
            cell_ref = f"{get_column_letter(column_index)}{row_index}"
            if cell_ref not in refs:
                continue
            values[cell_ref] = (
                cell.value if not formulas or cell.data_type == "f" else ""
            )
    return values


def extract_report_values(upload, workbook_path, validation_run=None):
    manifest = _load_manifest(upload)
    wb_values = load_workbook(workbook_path, data_only=True, read_only=True)
    wb_formulas = load_workbook(workbook_path, data_only=False, read_only=True)
    strict = manifest.get("management_v2", False)

    NormalizedValue.objects.filter(upload=upload).delete()
    rows = []
    by_report_row_unit = defaultdict(list)
    mapped_annual = {}

    for report_code, default_sheet in REPORTS.items():
        report = (manifest.get("reports") or {}).get(report_code) or {}
        sheet = report.get("sheet") or default_sheet
        mapping = report.get("mapping") or report.get("cells") or []
        if sheet not in wb_values.sheetnames:
            continue
        row_meta = {}
        for item in mapping:
            rc = item.get("row_code")
            cell = item.get("cell") or item.get("source_cell")
            if not rc or not cell:
                continue
            unit = (item.get("unit") or NormalizedValue.Unit.MONEY).upper()
            label = str(item.get("row_label") or item.get("label") or rc)[:240]
            row_meta.setdefault(rc, (coordinate_to_tuple(cell)[0], unit, label))
        header_row = (report.get("region") or {}).get("header_row")
        history_cols = (
            []
            if manifest.get("management_v3") and manifest.get("history_layout")
            else _history_columns(wb_values[sheet], header_row)
        )
        refs = _mapped_cell_refs(mapping)
        for col_letter, nature, _year in history_cols:
            if nature == "B":
                continue
            for row_number, _unit, _label in row_meta.values():
                refs.add(f"{col_letter}{row_number}")
        values = _read_mapped_cells(wb_values[sheet], refs)
        formulas = _read_mapped_cells(wb_formulas[sheet], refs, formulas=True)
        for item in mapping:
            cell_ref = item.get("cell") or item.get("source_cell")
            row_code = item.get("row_code")
            period = _period(item.get("period"))
            unit = (item.get("unit") or NormalizedValue.Unit.MONEY).upper()
            aggregation = str(item.get("aggregation") or "SUM").upper()
            if aggregation not in VALID_AGGREGATIONS:
                aggregation = "SUM"
            if aggregation == "EXCLUDE":
                continue
            if not cell_ref or not row_code or not period:
                continue
            value = values.get(_cell_key(cell_ref))
            formula = formulas.get(_cell_key(cell_ref), "")
            if strict and (value is None or value == "" or isinstance(value, str) and value.startswith("#")):
                if validation_run is not None:
                    ValidationIssue.objects.create(
                        run=validation_run, severity="P0", code="REPORT_VALUE_MISSING",
                        message="重算后的报表值缺失或错误，不能按零处理。",
                        location=f"{sheet}!{cell_ref}", actual_value=str(value),
                    )
                continue
            if strict:
                try:
                    numeric = Decimal(str(value))
                    valid = numeric.is_finite() and not isinstance(value, bool)
                    if unit == NormalizedValue.Unit.COUNT:
                        valid = valid and numeric == numeric.to_integral_value()
                except InvalidOperation:
                    valid = False
                if not valid:
                    if validation_run is not None:
                        ValidationIssue.objects.create(
                            run=validation_run, severity="P0", code="REPORT_VALUE_INVALID",
                            message="报表值必须为有限数值，数量必须为整数。",
                            location=f"{sheet}!{cell_ref}", actual_value=str(value),
                        )
                    continue
            row_label = item.get("row_label") or item.get("label") or row_code
            base = {
                "upload": upload,
                "report_code": report_code,
                "row_code": row_code,
                "row_label": str(row_label)[:240],
                "period": period,
                "unit": unit,
                "source_sheet": sheet,
                "source_cell": cell_ref,
                "source_formula": formula or "",
            }
            normalized = _build_value(base, value, item, values, aggregation)
            if normalized is None:
                continue
            if period != "YEAR" and unit == NormalizedValue.Unit.RATIO:
                _monthly_ratio_issue(validation_run, normalized, value)
            key = (
                report_code,
                row_code,
                unit,
                str(row_label)[:240],
                sheet,
                aggregation,
            )
            if period == "YEAR":
                mapped_annual[key] = normalized
            else:
                rows.append(normalized)
                by_report_row_unit[key].append(normalized)
        rows.extend(
            _history_rows(upload, report_code, sheet, row_meta, history_cols, values)
        )

    rows.extend(
        _annual_rows_from_months(
            upload, by_report_row_unit, mapped_annual, validation_run
        )
    )
    for row in rows:
        assign_dimensions(row, upload.cycle.budget_year)
    NormalizedValue.objects.bulk_create(rows)
    wb_values.close()
    wb_formulas.close()
    return len(rows)


def assign_dimensions(value, budget_year):
    historical = re.fullmatch(r"([AFB])(20\d{2})(?:M(0[1-9]|1[0-2]))?", value.period)
    if historical:
        value.data_year = int(historical[2])
        value.data_kind = {"A": "ACTUAL", "F": "FORECAST", "B": "BUDGET"}[historical[1]]
        value.month = int(historical[3]) if historical[3] else None
    elif value.period == "YEAR" or MONTH_RE.fullmatch(value.period):
        value.data_year = budget_year
        value.data_kind = "BUDGET"
        value.month = int(value.period) if value.period != "YEAR" else None


def _load_manifest(upload):
    if not upload.template or not upload.template.manifest_path:
        return {"reports": {}}
    manifest_path = Path(upload.template.manifest_path)
    if not manifest_path.is_absolute():
        manifest_path = Path(settings.BASE_DIR) / manifest_path
    if not manifest_path.exists():
        storage_path = (
            Path(settings.BUDGET_STORAGE_ROOT) / upload.template.manifest_path
        )
        manifest_path = storage_path if storage_path.exists() else manifest_path
    with manifest_path.open(encoding="utf-8") as fh:
        return json.load(fh)


def _period(value):
    if value is None:
        return ""
    text = str(value).strip().upper()
    if text in {"YEAR", "FY", "全年", "年度", "合计", "年合计"}:
        return "YEAR"
    month = re.fullmatch(r"(?:20\d{2}年)?\s*(1[0-2]|0?[1-9])\s*月?", text)
    if month:
        return f"{int(month.group(1)):02d}"
    return text


def _build_value(base, value, item, ws_values, aggregation="SUM"):
    unit = base["unit"]
    if unit == NormalizedValue.Unit.MONEY:
        kwargs = {"value_int": yuan_to_cents(value) or 0}
        if aggregation == "DERIVED":
            numerator_cell = item.get("numerator_cell") or item.get("ratio_num_cell")
            denominator_cell = item.get("denominator_cell") or item.get(
                "ratio_den_cell"
            )
            if numerator_cell and denominator_cell:
                kwargs["ratio_num"] = yuan_to_cents(
                    ws_values.get(_cell_key(numerator_cell))
                ) or 0
                kwargs["ratio_den"] = _to_int(
                    ws_values.get(_cell_key(denominator_cell))
                )
        return NormalizedValue(**base, **kwargs)
    if unit == NormalizedValue.Unit.COUNT:
        return NormalizedValue(**base, value_int=_to_int(value))
    if unit == NormalizedValue.Unit.RATIO:
        ratio_num, ratio_den = _ratio_parts(
            item,
            ws_values,
            base["source_cell"],
            value,
            base["row_code"],
            base["row_label"],
        )
        return NormalizedValue(
            **base, value_int=0, ratio_num=ratio_num, ratio_den=ratio_den
        )
    return None


def _to_int(value):
    try:
        return int(Decimal(str(value)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    except (InvalidOperation, TypeError, ValueError):
        return 0


def _ratio_parts(item, values, cell_ref, value, row_code, row_label):
    numerator_cell = item.get("numerator_cell") or item.get("ratio_num_cell")
    denominator_cell = item.get("denominator_cell") or item.get("ratio_den_cell")
    if numerator_cell and denominator_cell:
        return _to_int(values.get(_cell_key(numerator_cell))), _to_int(
            values.get(_cell_key(denominator_cell))
        )
    if row_code.upper() == "OCC" or RATIO_RE.search(str(row_label)):
        column = re.sub(r"\d+", "", _cell_key(cell_ref))
        den = _to_int(values.get(f"{column}24"))
        num = _to_int(values.get(f"{column}28"))
        return num, den
    try:
        scaled = int(
            (Decimal(str(value)) * Decimal("1000000")).quantize(
                Decimal("1"), rounding=ROUND_HALF_UP
            )
        )
    except (InvalidOperation, TypeError, ValueError):
        scaled = 0
    return scaled, 1000000


def _annual_rows_from_months(upload, by_report_row_unit, mapped_annual, validation_run):
    annual = []
    for (
        report_code,
        row_code,
        unit,
        row_label,
        sheet,
        aggregation,
    ), rows in by_report_row_unit.items():
        months = {row.period: row for row in rows if MONTH_RE.match(row.period)}
        if len(months) != 12:
            continue
        mapped = mapped_annual.get(
            (report_code, row_code, unit, row_label, sheet, aggregation)
        )
        if unit == NormalizedValue.Unit.RATIO:
            ratio_num = sum(int(row.ratio_num or 0) for row in months.values())
            ratio_den = sum(int(row.ratio_den or 0) for row in months.values())
            if mapped and _ratios_differ(
                mapped.ratio_num, mapped.ratio_den, ratio_num, ratio_den
            ):
                _annual_issue(
                    validation_run,
                    mapped,
                    f"{mapped.ratio_num}/{mapped.ratio_den}",
                    f"{ratio_num}/{ratio_den}",
                )
            annual.append(
                NormalizedValue(
                    upload=upload,
                    report_code=report_code,
                    row_code=row_code,
                    row_label=row_label,
                    period="YEAR",
                    unit=unit,
                    value_int=0,
                    ratio_num=ratio_num,
                    ratio_den=ratio_den,
                    source_sheet=sheet,
                    source_cell="PLATFORM",
                    source_formula="sum(monthly numerator) / sum(monthly denominator)",
                )
            )
        elif aggregation == "DERIVED":
            ratio_num = sum(int(row.ratio_num or 0) for row in months.values())
            ratio_den = sum(int(row.ratio_den or 0) for row in months.values())
            value_int = _divide_round_half_up(ratio_num, ratio_den)
            if mapped and int(mapped.value_int or 0) != value_int:
                _annual_issue(
                    validation_run, mapped, str(mapped.value_int or 0), str(value_int)
                )
            annual.append(
                NormalizedValue(
                    upload=upload,
                    report_code=report_code,
                    row_code=row_code,
                    row_label=row_label,
                    period="YEAR",
                    unit=unit,
                    value_int=value_int,
                    ratio_num=ratio_num,
                    ratio_den=ratio_den,
                    source_sheet=sheet,
                    source_cell="PLATFORM",
                    source_formula=(
                        "sum(monthly numerator) / sum(monthly denominator)"
                    ),
                )
            )
        else:
            value_int = sum(int(row.value_int or 0) for row in months.values())
            if aggregation == "AVERAGE":
                value_int = _divide_round_half_up(value_int, len(months))
            if mapped and int(mapped.value_int or 0) != value_int:
                _annual_issue(
                    validation_run, mapped, str(mapped.value_int or 0), str(value_int)
                )
            annual.append(
                NormalizedValue(
                    upload=upload,
                    report_code=report_code,
                    row_code=row_code,
                    row_label=row_label,
                    period="YEAR",
                    unit=unit,
                    value_int=value_int,
                    source_sheet=sheet,
                    source_cell="PLATFORM",
                    source_formula=(
                        "AVERAGE(monthly 01:12)"
                        if aggregation == "AVERAGE"
                        else "SUM(monthly 01:12)"
                    ),
                )
            )
    return annual


def _divide_round_half_up(numerator, denominator):
    if not denominator:
        return 0
    return int(
        (Decimal(numerator) / Decimal(denominator)).quantize(
            Decimal("1"), rounding=ROUND_HALF_UP
        )
    )


def _ratios_differ(actual_num, actual_den, expected_num, expected_den):
    if not actual_den or not expected_den:
        return (int(actual_num or 0), int(actual_den or 0)) != (
            int(expected_num or 0),
            int(expected_den or 0),
        )
    actual = (Decimal(actual_num) / Decimal(actual_den)).quantize(
        Decimal("0.000001"), rounding=ROUND_HALF_UP
    )
    expected = (Decimal(expected_num) / Decimal(expected_den)).quantize(
        Decimal("0.000001"), rounding=ROUND_HALF_UP
    )
    return actual != expected


def _monthly_ratio_issue(validation_run, normalized, displayed_value):
    if validation_run is None or displayed_value in (None, ""):
        return
    try:
        displayed_num = int(
            (Decimal(str(displayed_value)) * Decimal("1000000")).quantize(
                Decimal("1"), rounding=ROUND_HALF_UP
            )
        )
    except (InvalidOperation, TypeError, ValueError):
        return
    if normalized.ratio_den:
        differs = _ratios_differ(
            displayed_num,
            1000000,
            normalized.ratio_num,
            normalized.ratio_den,
        )
    else:
        differs = displayed_num != 0 or int(normalized.ratio_num or 0) != 0
    if not differs:
        return
    ValidationIssue.objects.create(
        run=validation_run,
        severity=ValidationIssue.Severity.P0,
        code="MONTHLY_RATIO_RECONCILIATION",
        message="工作簿月度比率与平台按分子、分母重算结果不一致。",
        location=f"{normalized.source_sheet}!{normalized.source_cell}",
        actual_value=str(displayed_value),
        expected_value=f"{normalized.ratio_num}/{normalized.ratio_den}",
    )


def _annual_issue(validation_run, mapped, actual, expected):
    if validation_run is None:
        return
    ValidationIssue.objects.create(
        run=validation_run,
        severity=ValidationIssue.Severity.P0,
        code="ANNUAL_RECONCILIATION",
        message="工作簿年度值与平台按 12 个月重算结果不一致。",
        location=f"{mapped.source_sheet}!{mapped.source_cell}",
        actual_value=actual,
        expected_value=expected,
    )


def sheet_slug(name):
    slug = SLUG_UNSAFE_RE.sub("_", str(name)).strip("_") or "sheet"
    if len(slug) > 40:
        slug = slug[:33] + "_" + hashlib.md5(slug.encode("utf-8")).hexdigest()[:6]
    return slug


def _month_number(value):
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        if float(value).is_integer() and 1 <= int(value) <= 12:
            return int(value)
        return None
    m = re.fullmatch(r"(0?[1-9]|1[0-2])\s*月?", str(value).strip())
    return int(m.group(1)) if m else None


def _month_columns(cells):
    by_month = {}
    for col, value in cells:
        m = _month_number(value)
        if m is not None:
            by_month.setdefault(m, col)
    if set(by_month) != set(range(1, 13)):
        return []
    return [by_month[m] for m in range(1, 13)]


def _label_column(all_rows, header_idx, min_month_col):
    scores = {}
    for row in all_rows[header_idx + 1:]:
        for cell in row:
            if getattr(cell, "column", None) is None:
                continue
            if cell.column >= min_month_col:
                break
            if isinstance(cell.value, str) and CJK_RE.search(cell.value):
                scores[cell.column] = scores.get(cell.column, 0) + 1
    if not scores:
        return None
    return max(scores, key=lambda c: (scores[c], c))


def _detect_grid(ws):
    all_rows = list(ws.iter_rows(values_only=False))
    header_idx = None
    month_cols = None
    for idx, row in enumerate(all_rows):
        cells = [(c.column, c.value) for c in row if c.value is not None]
        mc = _month_columns(cells)
        if mc:
            header_idx = idx
            month_cols = mc
            break
    if header_idx is None:
        return None
    label_col = _label_column(all_rows, header_idx, min(month_cols))
    if label_col is None:
        return None
    return label_col, month_cols, header_idx + 1


def _clean_label(value):
    if value is None:
        return ""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return ""
    text = str(value).strip()
    if not text or re.fullmatch(r"[\d.,%\s]+", text):
        return ""
    if text in SKIP_LABELS or any(m in text for m in ("小计", "合计", "总计", "不用填", "平衡检查")):
        return ""
    return text[:240]


def _unit_for_label(label):
    if RATIO_RE.search(label):
        return NormalizedValue.Unit.RATIO
    if COUNT_RE.search(label):
        return NormalizedValue.Unit.COUNT
    return NormalizedValue.Unit.MONEY


def _ratio_scaled(value):
    try:
        return int(
            (Decimal(str(value)) * Decimal("10000")).quantize(
                Decimal("1"), rounding=ROUND_HALF_UP
            )
        )
    except (InvalidOperation, TypeError, ValueError):
        return 0


def _sub_number(value, unit):
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, str):
        text = value.strip()
        if not text or text.startswith("#"):
            return None
        is_percent = text.endswith(("%", "％"))
        if is_percent:
            text = text[:-1].strip()
        text = text.replace(",", "").replace("，", "").replace(" ", "")
    else:
        text = str(value)
        is_percent = False
    try:
        number = Decimal(text)
    except (InvalidOperation, TypeError, ValueError):
        return None
    if not number.is_finite():
        return None
    try:
        if unit == NormalizedValue.Unit.RATIO:
            scale = Decimal("100") if is_percent else Decimal("10000")
            return int((number * scale).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
        if is_percent:
            return None
        if unit == NormalizedValue.Unit.COUNT:
            return int(number.quantize(Decimal("1"), rounding=ROUND_HALF_UP))
        if unit == NormalizedValue.Unit.MONEY:
            return int((number * Decimal("100")).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    except (InvalidOperation, TypeError, ValueError, OverflowError):
        return None
    return None


def _cell_at(row, col):
    idx = col - 1
    cell = row[idx] if 0 <= idx < len(row) else None
    if cell is not None and getattr(cell, "column", None) is None:
        return None
    return cell


def _sub_value(upload, report_code, sheet, row_code, label, unit, period, cell_ref, value):
    base = {
        "upload": upload,
        "report_code": report_code,
        "row_code": row_code,
        "row_label": label,
        "period": period,
        "unit": unit,
        "source_sheet": sheet,
        "source_cell": cell_ref,
        "source_formula": "",
    }
    parsed = _sub_number(value, unit)
    if parsed is None:
        return None
    if unit == NormalizedValue.Unit.RATIO:
        return NormalizedValue(**base, value_int=0, ratio_num=parsed, ratio_den=10000)
    return NormalizedValue(**base, value_int=parsed)


def _history_columns(ws, header_row):
    """Return [(column_letter, nature, year)] for year-named header columns.

    nature ∈ {"A"=实际, "F"=预测/预估, "B"=预算}. Only columns literally titled
    ``20xx年…`` are picked up, so variance columns like ``%``/``2026vs2025``/``说明``
    are skipped.
    """
    if not header_row:
        return []
    rows = list(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=False))
    if not rows:
        return []
    out = []
    for cell in rows[0]:
        if cell.value is None:
            continue
        m = YEAR_COL_RE.search(str(cell.value).strip())
        if not m:
            continue
        year = int(m.group(1))
        word = m.group(2)
        nature = "A" if word == "实际" else "F" if word in ("预测", "预估") else "B"
        out.append((get_column_letter(cell.column), nature, year))
    return out


def _history_rows(upload, report_code, sheet, row_meta, history_cols, values):
    """Build NormalizedValue rows for historical year columns (period like ``A2024``)."""
    out = []
    for col_letter, nature, year in history_cols:
        if nature == "B":
            continue
        period = f"{nature}{year}"
        for row_code, (row_number, unit, label) in row_meta.items():
            cell_ref = f"{col_letter}{row_number}"
            value = values.get(_cell_key(cell_ref))
            if value is None or value == "":
                continue
            out.append(
                _sub_value(upload, report_code, sheet, row_code, label, unit, period, cell_ref, value)
            )
    return out


def _sub_annual(upload, report_code, sheet, row_code, label, unit, monthly):
    """Roll monthly sub-table rows into a single YEAR value (SUM; ratio weighted-sum)."""
    base = {
        "upload": upload,
        "report_code": report_code,
        "row_code": row_code,
        "row_label": label,
        "period": "YEAR",
        "unit": unit,
        "source_sheet": sheet,
        "source_cell": "PLATFORM",
    }
    if unit == NormalizedValue.Unit.RATIO:
        num = sum(int(m.ratio_num or 0) for m in monthly)
        den = sum(int(m.ratio_den or 0) for m in monthly)
        return NormalizedValue(**base, value_int=0, ratio_num=num, ratio_den=den,
                               source_formula="sum(monthly numerator) / sum(monthly denominator)")
    total = sum(int(m.value_int or 0) for m in monthly)
    return NormalizedValue(**base, value_int=total, source_formula="SUM(monthly 01:12)")


def extract_sub_table_values(upload, workbook_path):
    wb = load_workbook(workbook_path, data_only=True, read_only=True)
    skip = set(REPORTS.values()) | {"SYS_META"}
    rows = []
    for name in wb.sheetnames:
        if name in skip:
            continue
        grid = _detect_grid(wb[name])
        if not grid:
            continue
        label_col, month_cols, header_idx = grid
        report_code = sheet_slug(name)
        for row_idx, row in enumerate(
            wb[name].iter_rows(min_row=header_idx + 1, values_only=False),
            start=header_idx + 1,
        ):
            label_cell = _cell_at(row, label_col)
            label = _clean_label(label_cell.value if label_cell is not None else None)
            if not label:
                continue
            unit = _unit_for_label(label)
            monthly = []
            for month, col in enumerate(month_cols, start=1):
                cell = _cell_at(row, col)
                value = cell.value if cell is not None else None
                parsed = _sub_value(
                    upload, report_code, name, f"R{row_idx:04d}", label, unit,
                    f"{month:02d}", f"{get_column_letter(col)}{row_idx}", value,
                )
                if parsed is not None:
                    monthly.append(parsed)
            rows.extend(monthly)
            if len(monthly) == len(month_cols):
                rows.append(_sub_annual(upload, report_code, name, f"R{row_idx:04d}", label, unit, monthly))
    NormalizedValue.objects.bulk_create(rows)
    return len(rows)


@functools.lru_cache(maxsize=8)
def _sub_table_specs_cached(path_str, mtime):
    wb = load_workbook(path_str, data_only=True, read_only=True)
    skip = set(REPORTS.values()) | {"SYS_META"}
    out = []
    for name in wb.sheetnames:
        if name in skip:
            continue
        if _detect_grid(wb[name]) is not None:
            out.append((sheet_slug(name), name))
    wb.close()
    out.sort(key=lambda x: x[1])
    return tuple(out)


def sub_table_specs(template_path):
    """Return sorted (slug, original_sheet_name) for every data sub-table in the template."""
    path = Path(template_path)
    if not path.exists():
        return []
    return list(_sub_table_specs_cached(str(path), path.stat().st_mtime))
