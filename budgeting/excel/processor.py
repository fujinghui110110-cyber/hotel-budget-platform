import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from django.conf import settings
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from budgeting.models import NormalizedValue, REPORTS, UploadVersion, ValidationIssue, ValidationRun
from budgeting.services.files import check_xlsx_package, rel, sha256_file, write_json

PROPRIETARY = re.compile(r"\b(DBS|VIEW|SUBNM)\s*\(", re.I)
ERROR_VALUES = {"#N/A", "#VALUE!", "#REF!", "#DIV/0!", "#NAME?", "#NUM!", "#NULL!"}


def sanitize_template(source: Path, out_file: Path, manifest_file: Path):
    wb = load_workbook(source, keep_links=False)
    manifest = {"template_version": "V1", "reports": {}, "sheets": []}
    removed = 0
    formula_cells = 0
    for ws in wb.worksheets:
        manifest["sheets"].append({"name": ws.title, "state": ws.sheet_state, "max_row": ws.max_row, "max_column": ws.max_column})
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "f" and isinstance(cell.value, str):
                    formula_cells += 1
                    if PROPRIETARY.search(cell.value):
                        cell.value = None
                        removed += 1
                elif cell.value in ERROR_VALUES:
                    cell.value = None
                    removed += 1
    meta = wb["SYS_META"] if "SYS_META" in wb.sheetnames else wb.create_sheet("SYS_META")
    meta.sheet_state = "hidden"
    rows = [
        ("template_version", "V1"),
        ("budget_year", 2026),
        ("project_code", ""),
        ("rule_version", "R1"),
        ("formula_manifest_hash", ""),
        ("signature_token", ""),
    ]
    for idx, (key, value) in enumerate(rows, 1):
        meta.cell(idx, 1, key)
        meta.cell(idx, 2, value)
    out_file.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_file)
    manifest["removed_proprietary_or_error_cells"] = removed
    manifest["formula_cells"] = formula_cells
    manifest["formula_manifest_hash"] = sha256_file(out_file)
    for code, sheet in REPORTS.items():
        manifest["reports"][code] = {"sheet": sheet, "mapping": infer_report_mapping(out_file, sheet)}
    write_json(manifest_file, manifest)
    return manifest


def infer_report_mapping(path: Path, sheet_name: str):
    wb = load_workbook(path, data_only=False, read_only=True)
    if sheet_name not in wb.sheetnames:
        return []
    ws = wb[sheet_name]
    period_cols = []
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 20)):
        for cell in row:
            text = str(cell.value or "").strip()
            m = re.fullmatch(r"(?:2026年)?\s*(1[0-2]|[1-9])\s*月?", text)
            if m:
                period_cols.append((cell.column, f"{int(m.group(1)):02d}"))
    if not period_cols:
        for col in range(2, min(ws.max_column, 13) + 1):
            period_cols.append((col, f"{col - 1:02d}"))
    mapping = []
    for row_idx in range(1, ws.max_row + 1):
        label = next((str(ws.cell(row_idx, c).value).strip() for c in range(1, min(ws.max_column, 6) + 1) if ws.cell(row_idx, c).value not in (None, "")), "")
        if not label:
            continue
        for col_idx, period in period_cols[:12]:
            mapping.append({"row_code": f"R{row_idx:04d}", "row_label": label[:120], "period": period, "cell": f"{get_column_letter(col_idx)}{row_idx}", "unit": "MONEY"})
    return mapping


def scan_workbook(path: Path):
    check_xlsx_package(path)
    wb = load_workbook(path, data_only=False, read_only=True, keep_links=False)
    issues = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "f" and isinstance(cell.value, str) and PROPRIETARY.search(cell.value):
                    issues.append(("P0", "PROPRIETARY_FORMULA", f"{ws.title}!{cell.coordinate}", "发布模板不得包含 DBS/VIEW/SUBNM。"))
                if cell.value in ERROR_VALUES:
                    issues.append(("P0", "EXCEL_ERROR", f"{ws.title}!{cell.coordinate}", f"发现错误值 {cell.value}。"))
    return issues


def recalculate_with_soffice(path: Path) -> Path:
    soffice = Path(settings.SOFFICE_BIN)
    if not soffice.exists():
        raise RuntimeError("LibreOffice/soffice 不存在。")
    tmp = Path(tempfile.mkdtemp(prefix="budget-lo-"))
    work = tmp / "work.xlsx"
    shutil.copy2(path, work)
    subprocess.run(
        [str(soffice), "--headless", "--convert-to", "xlsx", "--outdir", str(tmp), str(work)],
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        timeout=180,
        check=True,
    )
    return work


def process_upload(upload: UploadVersion):
    path = settings.BASE_DIR / upload.original_path
    run = ValidationRun.objects.create(upload=upload, rule_version=upload.template.rule_version if upload.template else "R1")
    issues = scan_workbook(path)
    for severity, code, location, message in issues:
        ValidationIssue.objects.create(run=run, severity=severity, code=code, location=location, message=message)
    if any(i[0] == "P0" for i in issues):
        upload.status = UploadVersion.Status.REJECTED
        upload.note = "存在 P0 阻断问题。"
        upload.save(update_fields=["status", "note"])
        return
    recalc = recalculate_with_soffice(path)
    dest = path.parent / "recalculated.xlsx"
    shutil.copy2(recalc, dest)
    upload.recalculated_path = rel(dest)
    NormalizedValue.objects.filter(upload=upload).delete()
    extract_normalized_values(upload, dest)
    upload.status = UploadVersion.Status.VALIDATED
    upload.save(update_fields=["status", "recalculated_path"])


def extract_normalized_values(upload: UploadVersion, workbook_path: Path):
    from budgeting.excel.extract import extract_report_values

    return extract_report_values(upload, workbook_path)
