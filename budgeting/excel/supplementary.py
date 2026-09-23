from decimal import Decimal, InvalidOperation, ROUND_HALF_UP

from django.db import transaction
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from budgeting.models import NormalizedValue, REPORTS, ValidationIssue
from budgeting.excel.wine_sources import discover_wine_sources


SUPPLEMENTARY_SOURCES = {
    "R9001": ("餐厅收入", "B1餐厅汇总", 24, 11),
    "R9002": ("宴会收入", "B2宴会收入", 24, 11),
    "R9005": ("季节性产品收入（月饼亭）", "B16月饼亭", 24, 11),
}


def _is_formula_cell(cell):
    """Return whether a formula-view cell identifies an Excel formula."""

    if cell is None:
        return False
    return getattr(cell, "data_type", None) == "f" or (
        isinstance(getattr(cell, "value", None), str)
        and getattr(cell, "value", "").startswith("=")
    )


@transaction.atomic
def extract_supplementary_values(upload, workbook_path, validation_run=None):
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    formula_workbook = None
    try:
        # A data-only workbook cannot distinguish an ordinary blank input from
        # a formula whose cached result is absent. Keep a second read-only
        # view solely for that distinction; the source workbook is unchanged.
        formula_workbook = load_workbook(workbook_path, read_only=True, data_only=False)
        pending = []
        wine_sources, diagnostics = discover_wine_sources(workbook)
        sources = dict(SUPPLEMENTARY_SOURCES)
        for row_code, spec in wine_sources.items():
            sources[row_code] = (spec['label'], spec['sheet'], spec['row'], spec['first_column'])
        for row_code, (label, sheet, row_number, first_column) in sources.items():
            if sheet not in workbook.sheetnames:
                continue
            row = next(workbook[sheet].iter_rows(min_row=row_number, max_row=row_number,
                                               min_col=first_column, max_col=first_column + 11))
            formula_row = next(formula_workbook[sheet].iter_rows(
                min_row=row_number,
                max_row=row_number,
                min_col=first_column,
                max_col=first_column + 11,
            )) if sheet in formula_workbook.sheetnames else ()
            monthly = []
            for month, cell in enumerate(row, 1):
                formula_cell = formula_row[month - 1] if formula_row else None
                if cell.value in (None, ""):
                    # Only an identified, non-formula blank is a business
                    # zero. Formula cache misses and Excel error cells stay
                    # missing so they cannot be silently converted to zero.
                    cents = None if cell.data_type == "e" or _is_formula_cell(formula_cell) else 0
                elif cell.data_type == "e" or isinstance(cell.value, bool):
                    cents = None
                else:
                    try:
                        number = Decimal(str(cell.value))
                        cents = int((number * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP)) if number.is_finite() else None
                    except (ValueError, InvalidOperation):
                        cents = None
                monthly.append(cents)
                if cents is None:
                    continue
                for report in REPORTS:
                    pending.append(NormalizedValue(upload=upload, report_code=report, row_code=row_code,
                        row_label=label, period=f"{month:02d}", data_year=upload.cycle.budget_year,
                        data_kind="BUDGET", month=month, unit="MONEY", value_int=cents,
                        source_sheet=sheet, source_cell=f"{get_column_letter(first_column + month - 1)}{row_number}"))
            if row_code in ('R9003', 'R9006') and any(value is None for value in monthly):
                missing = '、'.join(str(month) for month, value in enumerate(monthly, 1) if value is None)
                diagnostics.append({'code': 'WINE_DETAIL_INCOMPLETE',
                                    'message': f'{label}第{missing}月未读到有效金额（空白、错误值或非数字）；年度保持缺失，不按零计算。',
                                    'location': f'{sheet}!{get_column_letter(first_column)}{row_number}:{get_column_letter(first_column + 11)}{row_number}'})
            if all(value is not None for value in monthly):
                for report in REPORTS:
                    pending.append(NormalizedValue(upload=upload, report_code=report, row_code=row_code,
                        row_label=label, period="YEAR", data_year=upload.cycle.budget_year,
                        data_kind="BUDGET", month=None, unit="MONEY", value_int=sum(monthly),
                        source_sheet=sheet, source_cell=f"{get_column_letter(first_column)}{row_number}",
                        source_formula="SUM(12 months); analysis subset, do not add to report total"))
    finally:
        workbook.close()
        if formula_workbook is not None:
            formula_workbook.close()
    NormalizedValue.objects.filter(upload=upload, row_code__in=["R9003", "R9006"]).delete()
    if validation_run is not None:
        validation_run.issues.filter(code__in=["WINE_SOURCE_MISSING", "WINE_SOURCE_AMBIGUOUS", "WINE_DETAIL_INCOMPLETE"]).delete()
        for diagnostic in diagnostics:
            ValidationIssue.objects.create(run=validation_run, severity='P2',
                                           code=diagnostic['code'], message=diagnostic['message'],
                                           location=diagnostic.get('location', '')[:200])
    for report in REPORTS:
        numbers = (23, 24, 25, 26, 27) if report.startswith("PL_ZZ") else (42, 43)
        codes = {f"R{number:04d}" for number in numbers}
        source = NormalizedValue.objects.filter(upload=upload, report_code=report, row_code__in=codes)
        groups = {}
        for value in source:
            groups.setdefault(value.period, []).append(value)
        for period, values in groups.items():
            if {value.row_code for value in values} != codes:
                continue
            first = values[0]
            pending.append(NormalizedValue(upload=upload, report_code=report, row_code="R9004",
                row_label="客房运营成本合计（含人工）", period=period, data_year=first.data_year,
                data_kind=first.data_kind, month=first.month, unit="MONEY",
                value_int=sum(value.value_int for value in values),
                source_sheet=first.source_sheet, source_cell=first.source_cell,
                source_formula=" + ".join(sorted(codes))))
    NormalizedValue.objects.bulk_create(pending, update_conflicts=True,
        unique_fields=["upload", "report_code", "row_code", "period"],
        update_fields=["row_label", "data_year", "data_kind", "month", "value_int",
                       "source_sheet", "source_cell", "source_formula"])
    return len(pending)
