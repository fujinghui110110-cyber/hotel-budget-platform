from __future__ import annotations

import json
import re
from collections import defaultdict
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

from django.conf import settings
from django.db import transaction
from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter

from budgeting.models import NormalizedValue, ValidationIssue
from budgeting.services.metrics import METRICS


MONTHS = tuple(range(1, 13))
_PERIOD_RE = re.compile(r"^([AFB])\s*(20\d{2})(?:[-_/]?M?\s*(0?[1-9]|1[0-2]))?$", re.I)
_YEAR_MONTH_RE = re.compile(r"^(20\d{2})[-_/]?M?\s*(0?[1-9]|1[0-2])$", re.I)
_YEAR_RE = re.compile(r"^20\d{2}$", re.I)
_CELL_RE = re.compile(r"^\$?[A-Za-z]{1,3}\$?\d+$")
_COLUMN_RE = re.compile(r"^\$?([A-Za-z]{1,3})\$?$")
_KIND_ALIASES = {
    "A": "ACTUAL",
    "ACTUAL": "ACTUAL",
    "实际": "ACTUAL",
    "F": "FORECAST",
    "FORECAST": "FORECAST",
    "预测": "FORECAST",
    "预估": "FORECAST",
    "B": "BUDGET",
    "BUDGET": "BUDGET",
    "预算": "BUDGET",
}
_KIND_PREFIX = {value: key for key, value in _KIND_ALIASES.items() if key in {"A", "F", "B"}}
_UNSET = object()


def _cell_key(value):
    return str(value or "").replace("$", "").strip().upper()


def _blank(value):
    if value is None:
        return True
    if isinstance(value, str):
        return not value.strip()
    return False


def _workbook_sheet(workbook, sheet):
    if workbook is None or not sheet:
        return None
    try:
        return workbook[sheet]
    except (KeyError, TypeError):
        return None


def _cell_value(workbook, sheet, cell):
    worksheet = _workbook_sheet(workbook, sheet)
    if worksheet is None or not cell:
        return None
    try:
        return worksheet[_cell_key(cell)].value
    except (KeyError, TypeError, AttributeError):
        return None


def _kind(value):
    text = str(value or "").strip().upper()
    return _KIND_ALIASES.get(text)


def _prefix(kind):
    return _KIND_PREFIX.get(kind, "")


def _period(period=None, year=None, kind=None, month=None):
    raw = str(period or "").strip().upper()
    resolved_kind = _kind(kind)
    resolved_year = None
    resolved_month = None

    if raw:
        match = _PERIOD_RE.fullmatch(raw)
        if match:
            resolved_kind = _KIND_ALIASES.get(match.group(1).upper())
            resolved_year = int(match.group(2))
            resolved_month = int(match.group(3)) if match.group(3) else None
        else:
            match = _YEAR_MONTH_RE.fullmatch(raw)
            if match:
                resolved_year = int(match.group(1))
                resolved_month = int(match.group(2))
            elif _YEAR_RE.fullmatch(raw):
                resolved_year = int(raw)
            elif raw in {"YEAR", "FY", "全年", "年度", "年合计", "合计"}:
                resolved_month = None
            else:
                month_match = re.fullmatch(r"(?:20\d{2}年)?\s*(1[0-2]|0?[1-9])\s*月?", raw)
                if month_match:
                    resolved_month = int(month_match.group(1))

    if year not in (None, ""):
        try:
            resolved_year = int(year)
        except (TypeError, ValueError):
            resolved_year = None
    if month not in (None, ""):
        try:
            resolved_month = int(month)
        except (TypeError, ValueError):
            resolved_month = None

    if not resolved_year or not resolved_kind:
        return "", resolved_year, resolved_kind, resolved_month
    prefix = _prefix(resolved_kind)
    if resolved_month is None:
        return f"{prefix}{resolved_year}", resolved_year, resolved_kind, None
    if resolved_month not in MONTHS:
        return "", resolved_year, resolved_kind, resolved_month
    return f"{prefix}{resolved_year}M{resolved_month:02d}", resolved_year, resolved_kind, resolved_month


def _as_decimal(value):
    if _blank(value) or isinstance(value, bool):
        return None
    if isinstance(value, str):
        text = value.strip().replace(",", "")
        if text in {"", "-", "—", "–", "N/A", "NA", "#N/A", "#VALUE!"}:
            return None
        if text.endswith("%"):
            text = text[:-1].strip()
            try:
                return Decimal(text) / Decimal(100)
            except (InvalidOperation, ValueError):
                return None
    try:
        return Decimal(str(value))
    except (InvalidOperation, TypeError, ValueError):
        return None


def _to_int(value):
    number = _as_decimal(value)
    if number is None:
        return None
    return int(number.quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def _money(value):
    number = _as_decimal(value)
    if number is None:
        return None
    return int((number * Decimal(100)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def _ratio_value(value):
    number = _as_decimal(value)
    if number is None:
        return None
    return int((number * Decimal(10000)).quantize(Decimal("1"), rounding=ROUND_HALF_UP)), 10000


def _unit(value, metric=None):
    text = str(value or (metric or {}).get("unit") or NormalizedValue.Unit.MONEY).upper()
    if text not in {NormalizedValue.Unit.MONEY, NormalizedValue.Unit.COUNT, NormalizedValue.Unit.RATIO}:
        return NormalizedValue.Unit.MONEY
    return text


def _aggregation(value, metric=None):
    text = str(value or (metric or {}).get("aggregation") or "SUM").upper()
    return text if text in {"SUM", "AVERAGE", "DERIVED", "RATIO", "EXCLUDE"} else "SUM"


def _metric_code(report_code, row_code, declared=None):
    if declared and declared in METRICS:
        return declared
    for code, spec in METRICS.items():
        if (spec.get("rows") or {}).get(report_code) == row_code:
            return code
    return ""


def _metric_spec(report_code, row_code, metric_code=""):
    code = _metric_code(report_code, row_code, metric_code)
    return code, METRICS.get(code) or {}


def _mapping_value(mapping, key, default=_UNSET):
    if isinstance(mapping, dict):
        if key in mapping:
            return mapping[key]
        key_text = str(key)
        if key_text in mapping:
            return mapping[key_text]
    return default


def _normal_column(value):
    if isinstance(value, int):
        return get_column_letter(value)
    text = str(value or "").strip()
    match = _COLUMN_RE.fullmatch(text)
    if match:
        return match.group(1).upper()
    if _CELL_RE.fullmatch(text):
        return _cell_key(text).rstrip("0123456789")
    return ""


def _column_specs(layout):
    columns = None
    for key in ("period_columns", "columns", "month_columns", "history_columns"):
        if isinstance(layout, dict) and layout.get(key) is not None:
            columns = layout.get(key)
            break
    if columns is None:
        return []

    out = []
    if isinstance(columns, dict):
        iterator = columns.items()
        for key, declaration in iterator:
            declaration = declaration if isinstance(declaration, dict) else {"period": declaration}
            item = dict(declaration)
            item.setdefault("column", key)
            out.append(item)
    elif isinstance(columns, (list, tuple)):
        for declaration in columns:
            if isinstance(declaration, str):
                out.append({"period": declaration})
            elif isinstance(declaration, dict):
                out.append(dict(declaration))

    periods = layout.get("periods") if isinstance(layout, dict) else None
    if isinstance(periods, (list, tuple)):
        for index, item in enumerate(out):
            if not item.get("period") and index < len(periods):
                item["period"] = periods[index]

    normalized = []
    for item in out:
        column = item.get("column") or item.get("col") or item.get("column_letter")
        if not column and item.get("cell"):
            column = _normal_column(item["cell"])
        column = _normal_column(column)
        if not column:
            continue
        period, year, kind, month = _period(
            item.get("period"),
            item.get("year") or item.get("data_year"),
            item.get("kind") or item.get("nature") or item.get("data_kind"),
            item.get("month"),
        )
        if not period:
            continue
        item.update({"column": column, "period": period, "year": year, "kind": kind, "month": month})
        normalized.append(item)
    return normalized


def _history_layout(manifest):
    if not isinstance(manifest, dict):
        return {}
    for key in ("history_layout", "history"):
        value = manifest.get(key)
        if isinstance(value, dict):
            return value
    for key in ("history_inputs", "history_input"):
        value = manifest.get(key)
        if isinstance(value, dict):
            return value
    return {}


def _history_sheet(manifest, layout):
    return (
        layout.get("sheet")
        or layout.get("sheet_name")
        or manifest.get("history_sheet")
        or "历史月度输入"
    )


def _row_meta(item, layout):
    item = dict(item or {})
    report_code = item.get("report_code") or item.get("report") or layout.get("report_code") or "PL_TOTAL_WINE"
    row_code = item.get("row_code") or item.get("code")
    if not row_code and item.get("row") is not None and not str(item.get("row")).isdigit():
        row_code = item.get("row")
    metric_code = item.get("metric_code") or item.get("metric") or ""
    if not row_code and metric_code in METRICS:
        row_code = (METRICS[metric_code].get("rows") or {}).get(report_code)
    if not row_code:
        return None
    metric_code, metric = _metric_spec(report_code, str(row_code), metric_code)
    label = item.get("row_label") or item.get("label") or metric.get("label") or str(row_code)
    unit = _unit(item.get("unit"), metric)
    aggregation = _aggregation(item.get("aggregation"), metric)
    numerator_row = (
        item.get("numerator_row")
        or item.get("numerator_row_code")
        or item.get("ratio_num_row")
        or (metric.get("numerator_rows") or {}).get(report_code)
    )
    denominator_row = (
        item.get("denominator_row")
        or item.get("denominator_row_code")
        or item.get("ratio_den_row")
        or (metric.get("denominator_rows") or {}).get(report_code)
    )
    return {
        "report_code": str(report_code),
        "row_code": str(row_code),
        "row_label": str(label)[:240],
        "metric_code": metric_code,
        "unit": unit,
        "aggregation": aggregation,
        "numerator_row": numerator_row,
        "denominator_row": denominator_row,
        "declaration": item,
    }


def _cell_from_declaration(row, period_info, column_info=None):
    declaration = row.get("declaration") or {}
    period, year, kind, month = period_info
    candidates = []
    for key in ("cells", "month_cells", "period_cells", "inputs", "input_cells", "values"):
        mapping = declaration.get(key)
        if isinstance(mapping, dict):
            candidates.append(mapping)
    for mapping in candidates:
        for key in (period, str(month or ""), f"{month:02d}" if month else "", str(year), column_info or ""):
            if key and key in mapping:
                value = mapping[key]
                if isinstance(value, dict):
                    return value
                return {"cell": value}
    if isinstance(declaration.get("cells"), (list, tuple)):
        for item in declaration["cells"]:
            if isinstance(item, dict):
                candidate = _period(item.get("period"), item.get("year"), item.get("kind"), item.get("month"))[0]
                if candidate == period:
                    return dict(item)
    direct_period = _period(declaration.get("period"), declaration.get("year"), declaration.get("kind"), declaration.get("month"))[0]
    if direct_period == period or (not direct_period and declaration.get("cell")):
        return {
            key: declaration[key]
            for key in (
                "cell",
                "numerator_cell",
                "denominator_cell",
                "ratio_num_cell",
                "ratio_den_cell",
            )
            if key in declaration
        }
    if column_info:
        row_number = declaration.get("row_number") or declaration.get("row")
        if row_number is not None and str(row_number).isdigit():
            return {"cell": f"{column_info}{row_number}"}
    return None


def _iter_rows(layout):
    rows = layout.get("rows") if isinstance(layout, dict) else None
    if rows is None:
        rows = layout.get("mapping") if isinstance(layout, dict) else None
    if rows is None:
        rows = layout.get("input_rows") if isinstance(layout, dict) else None
    if isinstance(rows, dict):
        rows = [dict(value, row_code=key) if isinstance(value, dict) else {"row_code": key, "cells": value} for key, value in rows.items()]
    return rows if isinstance(rows, (list, tuple)) else []


def _iter_input_records(manifest, layout, sheet):
    records = []
    sources = []
    for container in (layout, manifest):
        for key in ("input_cells", "history_input_cells", "inputs"):
            value = container.get(key) if isinstance(container, dict) else None
            if isinstance(value, (list, tuple)):
                sources.extend(value)
            elif isinstance(value, dict):
                value = value.get(sheet) if sheet in value else []
                if isinstance(value, (list, tuple)):
                    sources.extend(value)
    for item in sources:
        if isinstance(item, str):
            continue
        if not isinstance(item, dict) or not item.get("cell"):
            continue
        if item.get("sheet") and item.get("sheet") != sheet:
            continue
        row = _row_meta(item, layout)
        if not row:
            continue
        period = _period(item.get("period"), item.get("year"), item.get("kind") or item.get("nature"), item.get("month"))
        if not period[0]:
            continue
        records.append((row, period, dict(item)))
    return records


def _declarations(manifest):
    layout = _history_layout(manifest)
    sheet = _history_sheet(manifest, layout)
    columns = _column_specs(layout)
    records = _iter_input_records(manifest, layout, sheet)
    for item in _iter_rows(layout):
        row = _row_meta(item, layout)
        if not row:
            continue
        declaration = row["declaration"]
        direct = _period(declaration.get("period"), declaration.get("year"), declaration.get("kind") or declaration.get("nature"), declaration.get("month"))
        if direct[0] and declaration.get("cell"):
            records.append((row, direct, dict(declaration)))
        for column in columns:
            period_info = (column["period"], column["year"], column["kind"], column["month"])
            cell = _cell_from_declaration(row, period_info, column.get("column"))
            if not cell or not cell.get("cell"):
                continue
            cell.setdefault("period", period_info[0])
            cell.setdefault("year", period_info[1])
            cell.setdefault("kind", period_info[2])
            cell.setdefault("month", period_info[3])
            records.append((row, period_info, cell))
    deduped = {}
    for row, period_info, cell in records:
        period = period_info[0]
        if not period or not cell.get("cell"):
            continue
        key = (sheet, row["report_code"], row["row_code"], period)
        deduped.setdefault(key, (row, period_info, cell))
    return sheet, layout, list(deduped.values())


def _raw_parts(workbook, sheet, cell):
    value = _cell_value(workbook, sheet, cell)
    return value, _cell_key(cell)


def _record_value(row, value, num_value=None, den_value=None):
    unit = row["unit"]
    if num_value is not None or den_value is not None:
        numerator = _to_int(num_value)
        denominator = _to_int(den_value)
        if numerator is None or denominator is None:
            return None
        if unit == NormalizedValue.Unit.RATIO:
            return {"value_int": 0, "ratio_num": numerator, "ratio_den": denominator}
        converted = _money(value) if unit == NormalizedValue.Unit.MONEY else _to_int(value)
        return {"value_int": converted, "ratio_num": numerator, "ratio_den": denominator}
    if unit == NormalizedValue.Unit.MONEY:
        converted = _money(value)
        return {"value_int": converted, "ratio_num": None, "ratio_den": None}
    if unit == NormalizedValue.Unit.COUNT:
        converted = _to_int(value)
        return {"value_int": converted, "ratio_num": None, "ratio_den": None}
    ratio = _ratio_value(value)
    if ratio is None:
        return None
    return {"value_int": 0, "ratio_num": ratio[0], "ratio_den": ratio[1]}


def _numeric_record(row, value):
    if row["unit"] == NormalizedValue.Unit.RATIO:
        if value.get("ratio_num") is not None and value.get("ratio_den") is not None:
            return int(value["ratio_num"]), int(value["ratio_den"])
    return int(value.get("value_int") or 0), 1


def _formula_value(workbook, sheet, cell):
    value = _cell_value(workbook, sheet, cell)
    return "" if value is None else str(value)


def _annual_value(row, monthly):
    if len(monthly) != 12 or {item["month"] for item in monthly} != set(MONTHS):
        return None
    unit = row["unit"]
    if unit == NormalizedValue.Unit.RATIO or row["aggregation"] == "DERIVED":
        if all(item["value"].get("ratio_num") is not None and item["value"].get("ratio_den") is not None for item in monthly):
            numerator = sum(int(item["value"]["ratio_num"] or 0) for item in monthly)
            denominator = sum(int(item["value"]["ratio_den"] or 0) for item in monthly)
            value_int = 0 if unit == NormalizedValue.Unit.RATIO or not denominator else int((Decimal(numerator) / Decimal(denominator)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
            return {"value_int": value_int, "ratio_num": numerator, "ratio_den": denominator}
    if unit == NormalizedValue.Unit.RATIO:
        numerator = sum(int(item["value"]["ratio_num"] or 0) for item in monthly)
        denominator = sum(int(item["value"]["ratio_den"] or 0) for item in monthly)
        return {"value_int": 0, "ratio_num": numerator, "ratio_den": denominator}
    total = sum(int(item["value"]["value_int"] or 0) for item in monthly)
    if row["aggregation"] == "AVERAGE":
        total = int((Decimal(total) / Decimal(12)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    return {"value_int": total, "ratio_num": None, "ratio_den": None}


def _source_formula(row, period, cell, formula, annual=False):
    if annual:
        return f"SUM(history {row['row_code']} {period}M01:{period}M12)"
    return formula or ""


def _make_value(upload, row, period, value, source_sheet, source_cell, source_formula):
    _, year, kind, month = _period(period)
    kwargs = {
        "upload": upload,
        "report_code": row["report_code"],
        "row_code": row["row_code"],
        "row_label": row["row_label"],
        "period": period,
        "data_year": year,
        "data_kind": kind or "",
        "month": month,
        "unit": row["unit"],
        "value_int": int(value.get("value_int") or 0),
        "ratio_num": value.get("ratio_num"),
        "ratio_den": value.get("ratio_den"),
        "source_sheet": source_sheet,
        "source_cell": source_cell or "HISTORY",
        "source_formula": source_formula or "",
    }
    return NormalizedValue(**kwargs)


def _upload_from_run(validation_run):
    return getattr(validation_run, "upload", None) if validation_run is not None else None


def _sync_dimensions(upload):
    if upload is None or not hasattr(NormalizedValue, "objects"):
        return 0
    budget_year = getattr(getattr(upload, "cycle", None), "budget_year", None)
    pending = []
    for value in NormalizedValue.objects.filter(upload=upload):
        period = value.period
        parsed = _period(period)
        if not parsed[0] and budget_year:
            if re.fullmatch(r"0[1-9]|1[0-2]", str(period or "")):
                parsed = _period(period, budget_year, "B")
            elif str(period or "").upper() in {"YEAR", "FY"}:
                parsed = _period("YEAR", budget_year, "B")
        _, year, kind, month = parsed
        if not year or not kind:
            continue
        updates = {}
        if value.data_year != year:
            updates["data_year"] = year
        if value.data_kind != kind:
            updates["data_kind"] = kind
        if value.month != month:
            updates["month"] = month
        if updates:
            for key, item in updates.items():
                setattr(value, key, item)
            pending.append(value)
    if pending:
        NormalizedValue.objects.bulk_update(pending, ["data_year", "data_kind", "month"], batch_size=500)
    return len(pending)


def _sync_metric_metadata(upload, manifest):
    if upload is None:
        return 0
    reports = (manifest or {}).get("reports") or {}
    metadata = {}
    for report_code, report in reports.items():
        for item in (report.get("mapping") or report.get("cells") or []):
            if not isinstance(item, dict):
                continue
            row_code = item.get("row_code")
            if not row_code:
                continue
            metadata[(report_code, str(row_code))] = (
                str(item.get("row_label") or item.get("label") or "")[:240],
                _unit(item.get("unit")),
            )
    pending = []
    for value in NormalizedValue.objects.filter(upload=upload):
        row_label, unit = metadata.get((value.report_code, value.row_code), ("", ""))
        updates = {}
        if row_label and value.row_label != row_label:
            updates["row_label"] = row_label
        if unit and value.unit != unit:
            updates["unit"] = unit
        if updates:
            for key, item in updates.items():
                setattr(value, key, item)
            pending.append(value)
    if pending:
        NormalizedValue.objects.bulk_update(pending, ["row_label", "unit"], batch_size=500)
    return len(pending)


def _load_manifest(upload):
    template = getattr(upload, "template", None)
    path = getattr(template, "manifest_path", None)
    if not path:
        return {}
    path = Path(path)
    if not path.is_absolute():
        path = Path(getattr(settings, "BASE_DIR", Path.cwd())) / path
    if not path.exists():
        storage_root = Path(getattr(settings, "BUDGET_STORAGE_ROOT", Path.cwd()))
        candidate = storage_root / str(getattr(template, "manifest_path", ""))
        path = candidate if candidate.exists() else path
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError, TypeError):
        return {}


def _persist(upload, values):
    if upload is None:
        return len(values)
    count = 0
    for value in values:
        defaults = {
            "row_label": value.row_label,
            "data_year": value.data_year,
            "data_kind": value.data_kind,
            "month": value.month,
            "unit": value.unit,
            "value_int": value.value_int,
            "ratio_num": value.ratio_num,
            "ratio_den": value.ratio_den,
            "source_sheet": value.source_sheet,
            "source_cell": value.source_cell,
            "source_formula": value.source_formula,
        }
        NormalizedValue.objects.update_or_create(
            upload=upload,
            report_code=value.report_code,
            row_code=value.row_code,
            period=value.period,
            defaults=defaults,
        )
        count += 1
    return count


def load_history(upload, wb_values, wb_formulas=None, manifest=None):
    if manifest is None:
        manifest = _load_manifest(upload)
    sheet, layout, declarations = _declarations(manifest or {})
    if not declarations:
        _sync_dimensions(upload)
        _sync_metric_metadata(upload, manifest)
        return 0
    values = []
    by_row_period = {}
    pending_derived = []
    for row, period_info, cell in declarations:
        period, year, kind, month = period_info
        if month is None or kind == "BUDGET":
            continue
        raw = _cell_value(wb_values, sheet, cell.get("cell"))
        has_derived_sources = bool(row.get("numerator_row") and row.get("denominator_row"))
        if _blank(raw) and not has_derived_sources:
            continue
        numerator_cell = cell.get("numerator_cell") or cell.get("ratio_num_cell")
        denominator_cell = cell.get("denominator_cell") or cell.get("ratio_den_cell")
        numerator = _cell_value(wb_values, sheet, numerator_cell) if numerator_cell else None
        denominator = _cell_value(wb_values, sheet, denominator_cell) if denominator_cell else None
        converted = _record_value(row, raw, numerator, denominator)
        if converted is None:
            if _blank(raw) and has_derived_sources:
                pending_derived.append((row, period_info, cell))
            continue
        formula = _formula_value(wb_formulas, sheet, cell.get("cell"))
        record = {
            "row": row,
            "period": period,
            "year": year,
            "kind": kind,
            "month": month,
            "value": converted,
            "cell": _cell_key(cell.get("cell")),
            "formula": formula,
            "numerator_cell": numerator_cell,
            "denominator_cell": denominator_cell,
        }
        by_row_period[(row["report_code"], row["row_code"], period)] = record

    for row, period_info, cell in pending_derived:
        period, year, kind, month = period_info
        numerator = by_row_period.get((row["report_code"], str(row["numerator_row"]), period))
        denominator = by_row_period.get((row["report_code"], str(row["denominator_row"]), period))
        if not numerator or not denominator:
            continue
        ratio_num = int(numerator["value"].get("value_int") or 0)
        ratio_den = int(denominator["value"].get("value_int") or 0)
        value = {"value_int": 0, "ratio_num": ratio_num, "ratio_den": ratio_den}
        if row["aggregation"] == "DERIVED" and ratio_den:
            value["value_int"] = int(
                (Decimal(ratio_num) / Decimal(ratio_den)).quantize(
                    Decimal("1"), rounding=ROUND_HALF_UP
                )
            )
        by_row_period[(row["report_code"], row["row_code"], period)] = {
            "row": row,
            "period": period,
            "year": year,
            "kind": kind,
            "month": month,
            "value": value,
            "cell": _cell_key(cell.get("cell")),
            "formula": _formula_value(wb_formulas, sheet, cell.get("cell")),
        }

    for record in by_row_period.values():
        row = record["row"]
        if not row.get("numerator_row") or not row.get("denominator_row"):
            continue
        numerator_row = row.get("numerator_row")
        denominator_row = row.get("denominator_row")
        if not numerator_row or not denominator_row:
            continue
        numerator = by_row_period.get((row["report_code"], str(numerator_row), record["period"]))
        denominator = by_row_period.get((row["report_code"], str(denominator_row), record["period"]))
        if numerator and denominator:
            record["value"]["ratio_num"] = int(numerator["value"].get("value_int") or 0)
            record["value"]["ratio_den"] = int(denominator["value"].get("value_int") or 0)
            if row["aggregation"] == "DERIVED":
                denominator_value = record["value"]["ratio_den"]
                record["value"]["value_int"] = (
                    0
                    if not denominator_value
                    else int(
                        (Decimal(record["value"]["ratio_num"]) / Decimal(denominator_value)).quantize(
                            Decimal("1"), rounding=ROUND_HALF_UP
                        )
                    )
                )

    for record in by_row_period.values():
        values.append(
            _make_value(
                upload,
                record["row"],
                record["period"],
                record["value"],
                sheet,
                record["cell"],
                record["formula"],
            )
        )

    grouped = defaultdict(list)
    for record in by_row_period.values():
        key = (record["row"]["report_code"], record["row"]["row_code"], record["year"], record["kind"])
        grouped[key].append(record)
    for records in grouped.values():
        annual = _annual_value(records[0]["row"], records)
        if annual is None:
            continue
        row = records[0]["row"]
        period = f"{_prefix(records[0]['kind'])}{records[0]['year']}"
        values.append(
            _make_value(
                upload,
                row,
                period,
                annual,
                sheet,
                "HISTORY_DERIVED",
                _source_formula(row, period, "", "", annual=True),
            )
        )

    with transaction.atomic():
        count = _persist(upload, values)
        _sync_dimensions(upload)
        _sync_metric_metadata(upload, manifest)
    return count


def extract_management_values(upload, recalculated_path, validation_run=None):
    manifest = _load_manifest(upload)
    if not manifest.get("management_v2"):
        return 0
    path = Path(recalculated_path)
    if not path.exists() and not path.is_absolute():
        for root in (
            Path(getattr(settings, "BASE_DIR", Path.cwd())),
            Path(getattr(settings, "BUDGET_STORAGE_ROOT", Path.cwd())),
        ):
            candidate = root / path
            if candidate.exists():
                path = candidate
                break
    values = load_workbook(path, data_only=True, read_only=True)
    formulas = None
    try:
        formulas = load_workbook(path, data_only=False, read_only=True)
        count = load_history(upload, values, formulas, manifest)
        validate_history(validation_run, upload, values, formulas, manifest)
        return count
    finally:
        if formulas is not None:
            formulas.close()
        values.close()


def _resolve_validate_args(validation_run, args, kwargs):
    upload = kwargs.pop("upload", None) or _upload_from_run(validation_run)
    wb_values = kwargs.pop("wb_values", None) or kwargs.pop("workbook_values", None)
    wb_formulas = kwargs.pop("wb_formulas", None) or kwargs.pop("workbook_formulas", None)
    manifest = kwargs.pop("manifest", None)
    paths = []
    for item in args:
        if isinstance(item, dict):
            manifest = item
        elif isinstance(item, (str, Path)):
            paths.append(item)
        elif hasattr(item, "template") and hasattr(item, "cycle"):
            upload = item
        elif hasattr(item, "sheetnames") or hasattr(item, "worksheets"):
            if wb_values is None:
                wb_values = item
            elif wb_formulas is None:
                wb_formulas = item
        elif item is not None and hasattr(item, "objects") and hasattr(item, "_meta"):
            upload = item
    if manifest is None and upload is not None:
        manifest = _load_manifest(upload)
    if paths and wb_values is None:
        path = paths[0]
        wb_values = load_workbook(path, data_only=True, read_only=True)
        if wb_formulas is None:
            wb_formulas = load_workbook(path, data_only=False, read_only=True)
    return upload, wb_values, wb_formulas, manifest or {}


def _issue(validation_run, out, severity, code, message, location="", actual="", expected=""):
    if validation_run is None:
        out.append((severity, code, message, location))
        return
    out.append(
        ValidationIssue.objects.create(
            run=validation_run,
            severity=severity,
            code=code,
            message=message,
            location=location,
            actual_value=str(actual)[:120],
            expected_value=str(expected)[:120],
        )
    )


def validate_history(validation_run, *args, **kwargs):
    upload, wb_values, wb_formulas, manifest = _resolve_validate_args(validation_run, args, kwargs)
    sheet, layout, declarations = _declarations(manifest)
    issues = []
    if not declarations:
        return issues
    if wb_values is None or _workbook_sheet(wb_values, sheet) is None:
        _issue(validation_run, issues, "P0", "HISTORY_SHEET_MISSING", "历史输入工作表不存在。", sheet)
        return issues

    by_group = defaultdict(set)
    for row, period_info, cell in declarations:
        period, year, kind, month = period_info
        location = f"{sheet}!{_cell_key(cell.get('cell'))}"
        if month is None or kind == "BUDGET":
            continue
        formula = _cell_value(wb_formulas, sheet, cell.get("cell")) if wb_formulas is not None else None
        if isinstance(formula, str) and formula.startswith("="):
            _issue(validation_run, issues, "P0", "HISTORY_INPUT_FORMULA", "历史输入单元格不得包含公式。", location)
        raw = _cell_value(wb_values, sheet, cell.get("cell"))
        if _blank(raw):
            continue
        converted = _record_value(
            row,
            raw,
            _cell_value(wb_values, sheet, cell.get("numerator_cell") or cell.get("ratio_num_cell")),
            _cell_value(wb_values, sheet, cell.get("denominator_cell") or cell.get("ratio_den_cell")),
        )
        if converted is None:
            _issue(validation_run, issues, "P0", "HISTORY_VALUE_INVALID", "历史输入值不是可识别的数值。", location, raw)
            continue
        by_group[(row["report_code"], row["row_code"], year, kind)].add(month)

    for (report_code, row_code, year, kind), months in by_group.items():
        if len(months) < 12:
            continue
        if upload is None:
            continue
        period_prefix = f"{_prefix(kind)}{year}"
        monthly = list(
            NormalizedValue.objects.filter(
                upload=upload,
                report_code=report_code,
                row_code=row_code,
                data_year=year,
                data_kind=kind,
                month__isnull=False,
            )
        )
        if len({item.month for item in monthly}) < 12:
            continue
        annual = NormalizedValue.objects.filter(upload=upload, report_code=report_code, row_code=row_code, period=period_prefix).first()
        if annual is None:
            continue
        spec = _row_meta({"report_code": report_code, "row_code": row_code}, layout)
        expected = _annual_value(
            spec,
            [
                {
                    "month": item.month,
                    "value": {"value_int": item.value_int, "ratio_num": item.ratio_num, "ratio_den": item.ratio_den},
                }
                for item in monthly
            ],
        )
        if expected is None:
            continue
        differs = (
            int(annual.value_int or 0) != int(expected.get("value_int") or 0)
            or int(annual.ratio_num or 0) != int(expected.get("ratio_num") or 0)
            or int(annual.ratio_den or 0) != int(expected.get("ratio_den") or 0)
        )
        if differs:
            _issue(
                validation_run,
                issues,
                "P0",
                "HISTORY_ANNUAL_RECONCILIATION",
                "历史年度值与完整月度值重算结果不一致。",
                f"{sheet}!{row_code}",
                annual.value_int or f"{annual.ratio_num}/{annual.ratio_den}",
                expected.get("value_int") or f"{expected.get('ratio_num')}/{expected.get('ratio_den')}",
            )
    return issues


__all__ = ["extract_management_values", "load_history", "validate_history"]
