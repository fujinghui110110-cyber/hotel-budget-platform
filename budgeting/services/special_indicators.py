"""Strict, cached-value-only import of the four special indicator worksheets."""
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl
from django.core.files import File
from django.db import transaction

from budgeting.excel.ooxml import sha256_file, validate_xlsx_zip
from budgeting.models import IndicatorProject, Project, SpecialIndicatorBatch, SpecialIndicatorValue

INDICATORS = {'BANQUET': '宴会收入', 'RESTAURANT': '餐厅收入', 'RENT': '租赁收入', 'ENERGY': '第三方能耗收入'}
UNITS = {'YUAN': Decimal(1), 'WAN': Decimal(10000)}


def preview_import(path, *, year, data_type, unit='YUAN', cycle=None, allowed_project_ids=None):
    path = Path(path)
    result = dict(valid=False, errors=[], warnings=[], rows=[], projects=[], year=year,
                  data_type=data_type, unit=unit, cycle_id=getattr(cycle, 'pk', None),
                  original_name=path.name, sha256='')
    errors = result['errors']
    if not isinstance(year, int) or not 1900 <= year <= 2200:
        errors.append('年度必须为1900至2200之间的整数')
    if data_type not in ('ACTUAL', 'FORECAST', 'BUDGET'):
        errors.append('请选择实际、预测或预算口径')
    if unit not in UNITS:
        errors.append('单位必须为元或万元')
    if data_type == 'BUDGET' and (cycle is None or cycle.budget_year != year):
        errors.append('预算数据必须选择同年度预算版本')
    if data_type != 'BUDGET' and cycle is not None:
        errors.append('实际与预测数据不应绑定预算版本')
    if errors:
        return result
    issues = validate_xlsx_zip(path)
    if issues:
        errors.extend(f'{code}: {message}' for _, code, message in issues)
        return result
    result['sha256'] = sha256_file(path)
    try:
        cached = openpyxl.load_workbook(path, data_only=True, read_only=True)
        formulas = openpyxl.load_workbook(path, data_only=False, read_only=True)
    except Exception as exc:
        errors.append(f'无法读取工作簿：{type(exc).__name__}')
        return result
    all_projects = set()
    sheet_projects = {}
    known = {p.name: p.pk for p in Project.objects.all()}
    if '深圳凯骊酒店' in known:
        known['深圳凯骊'] = known['深圳凯骊酒店']
    allowed = None if allowed_project_ids is None else set(allowed_project_ids)
    try:
        for code, title in INDICATORS.items():
            if title not in cached.sheetnames:
                errors.append(f'{title}：缺少工作表')
                continue
            sheet, raw = cached[title], formulas[title]
            if sheet.max_row > 2000 or sheet.max_column > 100:
                errors.append(f'{title}：工作表尺寸超出模板范围')
                continue
            if sheet.cell(2, 3).value != '项目' or [sheet.cell(2, c).value for c in range(4, 16)] != [f'{m}月' for m in range(1, 13)] or sheet.cell(2, 16).value != '合计':
                errors.append(f'{title}：表头应为项目、1月至12月、合计；月份不可缺失或重复')
                continue
            names = set()
            # Materialize once: read-only worksheet random access is quadratic.
            values = list(sheet.iter_rows(min_row=3, max_col=16, values_only=True))
            raw_values = list(raw.iter_rows(min_row=3, max_col=16, values_only=True))
            for rownum, (cells, source) in enumerate(zip(values, raw_values), 3):
                name = str(cells[2] or '').strip()
                if not name:
                    if any(v is not None for v in cells[3:]):
                        errors.append(f'{title}!C{rownum}：存在数据但项目名称为空')
                    continue
                if name in names:
                    errors.append(f'{title}!C{rownum}：项目“{name}”重复')
                    continue
                names.add(name)
                if allowed is not None and known.get(name) not in allowed:
                    errors.append(f'{title}!C{rownum}：项目“{name}”不在当前账号授权范围')
                    continue
                all_projects.add(name)
                parsed = []
                for col in range(3, 16):
                    value, original = cells[col], source[col]
                    address = f'{title}!{openpyxl.utils.get_column_letter(col + 1)}{rownum}'
                    if value is None:
                        if isinstance(original, str) and original.startswith('='):
                            errors.append(f'{address}：公式没有缓存值，请在Excel计算并保存后上传')
                        parsed.append(None)
                        continue
                    try:
                        if isinstance(value, bool):
                            raise InvalidOperation
                        number = Decimal(str(value)) * UNITS[unit]
                        if not number.is_finite() or abs(number) >= Decimal('1e20'):
                            raise InvalidOperation
                        if number != number.quantize(Decimal('.0001')):
                            raise InvalidOperation
                        parsed.append(number)
                    except (InvalidOperation, ValueError):
                        errors.append(f'{address}：非有效数字或超过允许精度（元保留4位小数）')
                        parsed.append(None)
                months, annual = parsed[:12], parsed[12]
                if annual is not None and all(v is not None for v in months):
                    if abs(sum(months) - annual) > Decimal('.01'):
                        errors.append(f'{title}!P{rownum}：年度合计与12个月之和不一致')
                elif annual is not None:
                    result['warnings'].append(f'{title}／{name}：月度不完整，保留填写的年度合计；无法核对全年勾稽')
                for month, number in enumerate(parsed, 1):
                    result['rows'].append(dict(project_name=name, indicator=code, month=month if month <= 12 else 0,
                                               value=None if number is None else str(number)))
            sheet_projects[title] = names
        for title, names in sheet_projects.items():
            missing = all_projects - names
            if missing:
                errors.append(f'{title}：缺少项目行：{", ".join(sorted(missing))}')
        result['projects'] = sorted(all_projects)
        if not all_projects:
            errors.append('模板未包含可导入的项目')
        if not any(row['value'] is not None for row in result['rows']):
            result['warnings'].append('模板指标全部为空：仅登记项目与缺数状态，不会生成零值或历史数据')
        result['valid'] = not errors
        return result
    finally:
        cached.close()
        formulas.close()


def confirm_import(preview, *, source_path, user=None, expected_active_batch_ids=None):
    """Revalidate server-owned source and replace the entire selected period atomically."""
    from budgeting.models import BudgetCycle
    if user is not None and not user.is_admin_role:
        raise PermissionError('只有管理端可以确认全项目指标导入')
    cycle = BudgetCycle.objects.get(pk=preview['cycle_id']) if preview.get('cycle_id') else None
    checked = preview_import(source_path, year=preview['year'], data_type=preview['data_type'], unit=preview['unit'], cycle=cycle)
    if not checked['valid'] or checked['sha256'] != preview['sha256']:
        raise ValueError('原件发生变化或校验未通过，请重新预览')
    batch = SpecialIndicatorBatch(year=checked['year'], data_type=checked['data_type'], cycle=cycle,
                                  unit=checked['unit'], original_name=preview.get('original_name', checked['original_name']),
                                  sha256=checked['sha256'], created_by=user)
    saved = False
    try:
        with transaction.atomic():
            # All imports lock the same stable row on database engines supporting row locks.
            list(Project.objects.select_for_update().order_by('pk').values_list('pk', flat=True))
            active = SpecialIndicatorBatch.objects.filter(year=batch.year, data_type=batch.data_type, cycle=cycle, active=True)
            if expected_active_batch_ids is not None and set(active.values_list('pk', flat=True)) != set(expected_active_batch_ids):
                raise ValueError('当前批次已发生变化，请重新预览后确认替换范围')
            SpecialIndicatorBatch.objects.filter(year=batch.year, data_type=batch.data_type, cycle=cycle, active=True).update(active=False)
            with open(source_path, 'rb') as handle:
                batch.original.save(Path(batch.original_name).name, File(handle), save=False)
            saved = True
            batch.save()
            projects = {}
            for name in checked['projects']:
                match = Project.objects.filter(name='深圳凯骊酒店' if name == '深圳凯骊' else name).first()
                project, _ = IndicatorProject.objects.get_or_create(name=name, defaults={'project': match})
                projects[name] = project
            SpecialIndicatorValue.objects.bulk_create([
                SpecialIndicatorValue(batch=batch, project=projects[row['project_name']], indicator=row['indicator'],
                                      month=row['month'], value=row['value']) for row in checked['rows']
            ])
        return batch
    except Exception:
        if saved:
            batch.original.delete(save=False)
        raise


def comparison_data(*, year, data_type, cycle=None, project_ids=None):
    batches = SpecialIndicatorBatch.objects.filter(year=year, data_type=data_type, cycle=cycle, active=True)
    projects = IndicatorProject.objects.all()
    if project_ids is not None:
        projects = projects.filter(pk__in=project_ids)
    projects = list(projects)
    values = SpecialIndicatorValue.objects.filter(batch__in=batches, project__in=projects)
    lookup = {(v.project_id, v.indicator, v.month): v.value for v in values}
    rows = []
    for project in projects:
        for indicator in INDICATORS:
            months = [lookup.get((project.pk, indicator, m)) for m in range(1, 13)]
            complete = all(v is not None for v in months)
            annual = lookup.get((project.pk, indicator, 0))
            total = annual if annual is not None else sum(months) if complete else None
            rows.append(dict(project_id=project.pk, project_name=project.name, indicator=indicator,
                             values=[None if v is None else float(v) for v in months],
                             total=None if total is None else float(total), complete=complete,
                             source='special_indicator_template' if (project.pk, indicator, 1) in lookup else None))
    return dict(projects=[{'id': p.pk, 'name': p.name, 'project_id': p.project_id} for p in projects],
                indicators=INDICATORS, rows=rows, coverage={'complete': sum(r['complete'] for r in rows), 'expected': len(rows)})
