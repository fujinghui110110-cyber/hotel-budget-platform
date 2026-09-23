from budgeting.services.project_scope import cycle_projects
import json
import time
import shutil
import uuid
import zipfile
from collections import defaultdict
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from django.conf import settings
from budgeting.services.template_paths import resolve_template_path
from django.db import OperationalError, connection, transaction
from django.db.models import Count, F
from django.utils import timezone
from openpyxl import Workbook

from budgeting.excel.extract import _load_manifest, extract_report_values, extract_sub_table_values, sub_table_specs
from budgeting.excel.supplementary import extract_supplementary_values
from budgeting.excel.channel_checks import validate_channel_values
from budgeting.excel.business_checks import validate_management_values
from budgeting.excel.history import extract_management_values
from budgeting.excel.ooxml import (
    formula_comparison_rows,
    cached_errors,
    formula_manifest,
    sha256_file,
    validate_upload_contract,
    validate_xlsx_zip,
)
from budgeting.excel.legacy_adjustments import validate_legacy_balancing_inputs
from budgeting.excel.recalc import RecalcInfrastructureError, recalc_with_libreoffice
from budgeting.models import (
    AdjustmentBatch,
    AdjustmentLine,
    AuditEvent,
    BudgetCycle,
    FreezeSnapshot,
    NormalizedValue,
    PlanHistoryBinding,
    PlanProject,
    TargetEvaluation,
    ProcessingJob,
    ProcessingRun,
    Project,
    ProjectCycle,
    REPORTS,
    SnapshotArtifact,
    TemplateVersion,
    UploadVersion,
    ValidationIssue,
    ValidationRun,
)
from budgeting.services.processing_runs import (
    current_blocking_issues,
    enqueue_processing_job,
    ensure_job_run,
    has_current_p0,
    has_unacknowledged_current_p1,
    mark_run_failed,
    mark_run_running,
    mark_run_succeeded,
)
from budgeting.services.allocations import largest_remainder
from budgeting.services.drivers import ROOM_REV, simulate_driver
from budgeting.services.pnl_graph import (
    MONTHS,
    build_report_graph,
    load_external_values,
    load_manifest,
    recompute,
    required_external_refs,
)
from budgeting.services.report_context import ReportContext
from budgeting.services.report_query import freeze_selection_payload, query_report
from budgeting.services.targets import latest_targets


class InfrastructureProcessingError(RuntimeError):
    pass


RATIO_SCALE = 10_000


def _lock_mutable_cycle(cycle):
    if not BudgetCycle.objects.filter(pk=cycle.pk).exclude(status=BudgetCycle.Status.FROZEN).update(revision_no=F("revision_no")):
        raise ValueError("冻结周期不允许修改")
    cycle = BudgetCycle.objects.select_for_update().get(pk=cycle.pk)
    if cycle.status == BudgetCycle.Status.FROZEN:
        raise ValueError("冻结周期不允许修改")
    return cycle


def _verify_original_hash(upload, source_path=None):
    source_path = source_path or (settings.BUDGET_STORAGE_ROOT / upload.original_path)
    if not source_path.exists():
        raise ValueError("上传原始文件不存在，不能生成冻结快照")
    expected = (upload.sha256 or "").strip().lower()
    actual = sha256_file(source_path).lower()
    if not expected or actual != expected:
        raise ValueError(f"上传原始文件哈希不匹配：{upload.id}")
    return source_path


def _issue_parts(issue):
    severity, code, message = issue[:3]
    location = issue[3] if len(issue) > 3 else ""
    actual_value = issue[4] if len(issue) > 4 else ""
    expected_value = issue[5] if len(issue) > 5 else ""
    return severity, code, message, location, actual_value, expected_value


def active_cycle():
    return BudgetCycle.objects.exclude(status=BudgetCycle.Status.FROZEN).order_by("-budget_year", "-revision_no").first()


def active_template(cycle=None):
    if cycle and cycle.template_id:
        return cycle.template
    return TemplateVersion.objects.filter(is_active=True).order_by("-created_at").first()


def audit(actor, action, object_type, object_id, payload=None, project=None, cycle=None, upload=None):
    AuditEvent.objects.create(
        actor=actor,
        action=action,
        project=project,
        cycle=cycle,
        upload=upload,
        payload={"object_type": object_type, "object_id": str(object_id), **(payload or {})},
    )


def save_upload(project, cycle, uploaded_file):
    if cycle is None:
        raise ValueError("当前没有开放预算周期")
    cycle = BudgetCycle.objects.get(pk=cycle.pk)
    project = Project.objects.get(pk=project.pk)
    if cycle.status not in (BudgetCycle.Status.OPEN, BudgetCycle.Status.ADJUSTING):
        raise ValueError("当前周期不允许上传")
    project_cycle = ProjectCycle.objects.filter(project=project, cycle=cycle).first()
    if project_cycle and not project_cycle.is_open:
        raise ValueError("项目周期已关闭，不能上传")
    suffixes = (".xlsx", ".xlsm")
    if not uploaded_file.name.lower().endswith(suffixes):
        raise ValueError("请选择 .xlsx 或 .xlsm 格式的 Excel 文件")
    if uploaded_file.size > 50 * 1024 * 1024:
        raise ValueError("压缩文件超过 50 MiB")
    template = active_template(cycle)
    if template is None:
        raise ValueError("当前周期没有可用模板")

    upload_id = uuid.uuid4()
    rel_dir = Path("uploads") / project.code / str(upload_id)
    abs_dir = settings.BUDGET_STORAGE_ROOT / rel_dir
    abs_dir.mkdir(parents=True, exist_ok=True)
    dest = abs_dir / ("original" + Path(uploaded_file.name).suffix.lower())
    with dest.open("wb") as fh:
        for chunk in uploaded_file.chunks():
            fh.write(chunk)
    digest = sha256_file(dest)
    existing = UploadVersion.objects.filter(project=project, cycle=cycle, sha256=digest).order_by("-created_at").first()
    if existing:
        shutil.rmtree(abs_dir, ignore_errors=True)
        if existing.status == UploadVersion.Status.REJECTED:
            enqueue_processing_job(existing, reason="retry_duplicate_upload")
        return existing
    upload = UploadVersion.objects.create(
        id=upload_id,
        project=project,
        cycle=cycle,
        template=template,
        original_name=uploaded_file.name,
        original_path=str(rel_dir / dest.name),
        sha256=digest,
    )
    enqueue_processing_job(upload, reason="new_upload")
    audit(None, "UPLOAD_RECEIVED", "UploadVersion", upload.id, {"file": uploaded_file.name}, project=project, cycle=cycle, upload=upload)
    return upload


def _fail_processing_run(run, upload, note=""):
    mark_run_failed(run, note)
    upload.processing_current_run = run
    upload.status = UploadVersion.Status.REJECTED
    if note:
        upload.note = str(note)
        upload.save(update_fields=["status", "note", "processing_current_run"])
    else:
        upload.save(update_fields=["status", "processing_current_run"])


def _publish_processing_run(upload, run, recalculated_path, passed):
    run.passed = passed
    run.save(update_fields=["passed"])
    if passed:
        mark_run_succeeded(run.processing_run)
        upload.processing_current_run = run.processing_run
        upload.status = UploadVersion.Status.VALIDATED
        upload.recalculated_path = str(recalculated_path.relative_to(settings.BUDGET_STORAGE_ROOT))
        upload.save(update_fields=["status", "recalculated_path", "processing_current_run"])
        NormalizedValue.objects.filter(upload=upload).update(processing_run=run.processing_run)
        return True
    NormalizedValue.objects.filter(upload=upload).delete()
    mark_run_failed(run.processing_run, "validation_failed")
    upload.processing_current_run = run.processing_run
    upload.status = UploadVersion.Status.REJECTED
    upload.save(update_fields=["status", "processing_current_run"])
    return False


def _validation_has_p0(run):
    return run.issues.filter(severity=ValidationIssue.Severity.P0).exists()


def process_upload(upload, processing_run=None):
    source_path = settings.BUDGET_STORAGE_ROOT / upload.original_path
    if processing_run is None:
        job = (
            ProcessingJob.objects.filter(
                upload=upload,
                status=ProcessingJob.Status.RUNNING,
                processing_run__isnull=False,
            )
            .order_by("-updated_at", "-created_at")
            .first()
        )
        processing_run = job.processing_run if job else None
    processing_run = processing_run or ProcessingRun.objects.create(
        upload=upload,
        status=ProcessingRun.Status.RUNNING,
        rule_version=upload.template.rule_version if upload.template else "rules-v1",
        started_at=timezone.now(),
    )
    mark_run_running(processing_run)
    run = ValidationRun.objects.create(
        upload=upload,
        processing_run=processing_run,
        rule_version=upload.template.rule_version if upload.template else "rules-v1",
    )
    if not source_path.exists():
        ValidationIssue.objects.create(
            run=run,
            severity=ValidationIssue.Severity.P0,
            code="STORED_FILE_MISSING",
            message="上传原始文件不存在，不能处理。",
            location=upload.original_path,
        )
        run.passed = False
        run.save(update_fields=["passed"])
        _fail_processing_run(processing_run, upload)
        return False
    if sha256_file(source_path) != upload.sha256:
        ValidationIssue.objects.create(
            run=run, severity="P0", code="STORED_FILE_HASH_MISMATCH",
            message="上传原件哈希与接收记录不一致，不能处理。", location=upload.original_path,
        )
        run.passed = False
        run.save(update_fields=["passed"])
        _fail_processing_run(processing_run, upload)
        return False
    if upload.cycle.source_budget_year:
        validate_legacy_balancing_inputs(source_path, run)
        if _validation_has_p0(run):
            run.passed = False
            run.save(update_fields=["passed"])
            _fail_processing_run(processing_run, upload)
            return False
        from budgeting.services.legacy_rehearsal import process_legacy_rehearsal
        result = process_legacy_rehearsal(upload, source_path, run)
        if result:
            mark_run_succeeded(processing_run)
            upload.processing_current_run = processing_run
            upload.save(update_fields=["processing_current_run"])
            NormalizedValue.objects.filter(upload=upload).update(processing_run=processing_run)
        else:
            mark_run_failed(processing_run, "validation_failed")
            upload.processing_current_run = processing_run
            upload.save(update_fields=["processing_current_run"])
        from budgeting.services.historical_data import sync_history
        from budgeting.services.history_workbook import validate_workbook_history
        if result:
            history_path = settings.BUDGET_STORAGE_ROOT / upload.recalculated_path if upload.recalculated_path else source_path
            validate_workbook_history(upload, history_path, run)
            if _validation_has_p0(run):
                _fail_processing_run(processing_run, upload)
                result = False
        sync_history(upload)
        return result
    for issue in validate_xlsx_zip(source_path):
        severity, code, message, location, actual_value, expected_value = _issue_parts(issue)
        ValidationIssue.objects.create(
            run=run,
            severity=severity,
            code=code,
            message=message,
            location=location,
            actual_value=actual_value,
            expected_value=expected_value,
        )
    if _validation_has_p0(run):
        run.passed = False
        run.save(update_fields=["passed"])
        _fail_processing_run(processing_run, upload)
        return False
    for issue in validate_upload_contract(upload, source_path):
        severity, code, message, location, actual_value, expected_value = _issue_parts(issue)
        ValidationIssue.objects.create(
            run=run,
            severity=severity,
            code=code,
            message=message,
            location=location,
            actual_value=actual_value,
            expected_value=expected_value,
        )
    _record_formula_mapping_issues(run, upload, source_path)
    from budgeting.services.history_workbook import validate_history_identity
    validate_history_identity(upload, source_path, run)
    validate_legacy_balancing_inputs(source_path, run)
    if _validation_has_p0(run):
        run.passed = False
        run.save(update_fields=["passed"])
        _fail_processing_run(processing_run, upload)
        return False
    try:
        recalculated = recalc_with_libreoffice(source_path, settings.SOFFICE_BIN)
    except RecalcInfrastructureError as exc:
        raise InfrastructureProcessingError(str(exc)) from exc
    from budgeting.excel.ooxml import summary_sheet_names
    summary_sheets = summary_sheet_names(upload)
    for error in cached_errors(recalculated):
        ValidationIssue.objects.create(
            run=run, severity="P0" if error["sheet"] in summary_sheets else "P2", code="RECALCULATED_EXCEL_ERROR",
            message=f"服务端重算产生 Excel 错误：{error['value']}。",
            location=f"{error['sheet']}!{error['cell']}",
        )
    if _validation_has_p0(run):
        run.passed = False
        run.save(update_fields=["passed"])
        _fail_processing_run(processing_run, upload)
        return False
    with transaction.atomic():
        upload = UploadVersion.objects.select_for_update().get(pk=upload.pk)
        count = extract_report_values(upload, recalculated, validation_run=run)
        manifest = _load_manifest(upload)
        if not manifest.get("management_v2") or manifest.get("management_v3"):
            extract_sub_table_values(upload, recalculated)
        extract_management_values(upload, recalculated, validation_run=run, include_history=False)
        if manifest.get("management_v3"):
            extract_supplementary_values(upload, recalculated, validation_run=run)
            validate_channel_values(recalculated, run)
        from budgeting.services.historical_data import sync_history
        from budgeting.services.history_workbook import validate_workbook_history
        validate_workbook_history(upload, recalculated, run, manifest=manifest)
        sync_history(upload)
        if manifest.get("management_v2"):
            validate_management_values(upload, run)
        if count == 0:
            ValidationIssue.objects.create(
                run=run,
                severity=ValidationIssue.Severity.P0,
                code="NO_REPORT_VALUES",
                message="四张固定报表未抽取到标准化值",
            )
        passed = not _validation_has_p0(run)
        return _publish_processing_run(upload, run, recalculated, passed)


def _check_budget_governance(upload, *, approving=False):
    from django.core.exceptions import ValidationError
    from budgeting.services.plan_history import ensure_plan, assert_current_history
    from budgeting.services.targets import evaluate_upload, assert_approval_targets

    ensure_plan(upload.cycle)
    try:
        assert_current_history(upload)
        if approving:
            return assert_approval_targets(upload)
        return evaluate_upload(upload)
    except ValidationError as exc:
        raise ValueError("；".join(exc.messages)) from exc


def legacy_adjustment_blockers(upload):
    from budgeting.models import TargetConstraint
    reviewed = TargetConstraint.objects.filter(
        target_set__plan_id=upload.cycle.plan_id, target_set__project=upload.project,
    ) if upload.cycle.plan_id else TargetConstraint.objects.none()
    blockers = []
    for line in AdjustmentLine.objects.filter(
        project=upload.project, cycle=upload.cycle, status=AdjustmentLine.Status.OPEN,
        batch__status=AdjustmentBatch.Status.ISSUED,
    ).select_related("batch"):
        if (line.batch.cascade or {}).get("informational_only"):
            continue
        if reviewed.filter(report_code=line.report_code, row_code=line.row_code, period=line.period,
                           target_set__created_at__gte=line.batch.issued_at or line.batch.created_at).exists():
            continue
        if _adjustment_line_value(upload, line) != line.target_cents:
            blockers.append(f"旧下发任务 {line.batch_id} 尚未落实或复核：请管理员明确选为年度目标，不能自动忽略。")
    return blockers


@transaction.atomic
def submit_upload(upload, actor=None):
    upload = UploadVersion.objects.select_for_update().get(pk=upload.pk)
    if upload.cycle.status == BudgetCycle.Status.FROZEN:
        raise ValueError("冻结周期不能提交版本")
    if upload.status != UploadVersion.Status.VALIDATED or has_current_p0(upload):
        raise ValueError("只有无 P0 的已校验版本可以提交")
    if any(not issue.acknowledgement_note.strip() for issue in current_blocking_issues(upload, ValidationIssue.Severity.P1)):
        raise ValueError("请先填写所有 P1 波动事项的说明")
    _check_budget_governance(upload)
    upload.status = UploadVersion.Status.SUBMITTED
    upload.submitted_at = timezone.now()
    upload.save(update_fields=["status", "submitted_at"])
    update_adjustment_lines_for_upload(upload)
    audit(actor, "UPLOAD_SUBMITTED", "UploadVersion", upload.id, project=upload.project, cycle=upload.cycle, upload=upload)
    return upload


def _adjustment_line_value(upload, line):
    values = NormalizedValue.objects.filter(
        upload=upload,
        report_code=line.report_code,
        row_code=line.row_code,
        period=line.period,
    )
    if (line.batch.cascade or {}).get("kind") != "summary_annual":
        values = values.filter(unit=NormalizedValue.Unit.MONEY)
    value = values.first()
    if value is None:
        return None
    if value.unit == NormalizedValue.Unit.RATIO:
        if value.ratio_den:
            ratio = Decimal(value.ratio_num or 0) / Decimal(value.ratio_den)
            return int((ratio * 10_000).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
        return int(value.value_int or 0)
    return int(value.value_int or 0)


@transaction.atomic
def approve_upload(upload, actor=None):
    _lock_mutable_cycle(upload.cycle)
    upload = UploadVersion.objects.select_for_update().get(pk=upload.pk)
    if upload.cycle.status == BudgetCycle.Status.FROZEN:
        raise ValueError("冻结周期不能批准或替换版本")
    if upload.status != UploadVersion.Status.SUBMITTED:
        raise ValueError("只有已提交版本可以批准")
    if has_current_p0(upload):
        raise ValueError("存在 P0 问题，不能批准")
    if has_unacknowledged_current_p1(upload):
        raise ValueError("存在未确认 P1 问题，不能批准")
    _check_budget_governance(upload, approving=True)
    legacy_blockers = legacy_adjustment_blockers(upload)
    if legacy_blockers:
        raise ValueError("；".join(legacy_blockers))
    pc, _ = ProjectCycle.objects.select_for_update().get_or_create(project=upload.project, cycle=upload.cycle, defaults={"current_upload": upload})
    old = pc.current_upload if pc.current_upload_id != upload.id else None
    upload.status = UploadVersion.Status.APPROVED
    upload.approved_at = timezone.now()
    upload.save(update_fields=["status", "approved_at"])
    pc.current_upload = upload
    pc.save(update_fields=["current_upload"])
    if old:
        old.status = UploadVersion.Status.SUPERSEDED
        old.save(update_fields=["status"])
    update_adjustment_lines_for_upload(upload)
    audit(actor, "UPLOAD_APPROVED", "UploadVersion", upload.id, project=upload.project, cycle=upload.cycle, upload=upload)
    return upload


@transaction.atomic
def reject_upload(upload, actor=None, reason=""):
    upload = UploadVersion.objects.select_for_update().get(pk=upload.pk)
    _lock_mutable_cycle(upload.cycle)
    if upload.status != UploadVersion.Status.SUBMITTED:
        raise ValueError("只有已提交版本可以打回")
    upload.status = UploadVersion.Status.REJECTED
    upload.note = (reason or "").strip()
    upload.save(update_fields=["status", "note"])
    audit(actor, "UPLOAD_REJECTED", "UploadVersion", upload.id, {"reason": upload.note}, project=upload.project, cycle=upload.cycle, upload=upload)
    return upload


def _claim_inline_upload_job(upload):
    for attempt in range(25):
        try:
            return _claim_inline_upload_job_once(upload)
        except OperationalError as exc:
            if connection.vendor != "sqlite" or "locked" not in str(exc).lower() or attempt == 24:
                raise
            time.sleep(0.02 * (attempt + 1))


def _claim_inline_upload_job_once(upload):
    job = (
        ProcessingJob.objects
        .filter(upload_id=upload.pk)
        .order_by("-created_at", "-pk")
        .first()
    )
    if job is None:
        job, _ = enqueue_processing_job(upload, reason="inline_missing_job")
        job.refresh_from_db()
    if job.status != ProcessingJob.Status.QUEUED:
        return job, False

    now = timezone.now()
    claimed = ProcessingJob.objects.filter(
        pk=job.pk,
        status=ProcessingJob.Status.QUEUED,
    ).update(
        status=ProcessingJob.Status.RUNNING,
        attempts=F("attempts") + 1,
        lease_until=None,
        heartbeat_at=now,
    )
    if not claimed:
        job.refresh_from_db()
        upload.refresh_from_db(fields=["status", "note", "recalculated_path"])
        return job, False
    job.refresh_from_db()
    processing_run = ensure_job_run(job)
    mark_run_running(processing_run)
    job.processing_run = processing_run
    job.processing_run_id = processing_run.pk
    UploadVersion.objects.filter(pk=upload.pk).update(status=UploadVersion.Status.PROCESSING)
    upload.status = UploadVersion.Status.PROCESSING
    return job, True


def process_upload_now(upload):
    job, claimed = _claim_inline_upload_job(upload)
    if not claimed:
        upload.refresh_from_db(fields=["status", "note", "recalculated_path"])
        return job
    try:
        process_upload(upload)
    except InfrastructureProcessingError as exc:
        upload.status = UploadVersion.Status.REJECTED
        upload.note = str(exc)
        upload.save(update_fields=["status", "note"])
        mark_run_failed(job.processing_run, str(exc))
        job.status = ProcessingJob.Status.FAILED
        job.error = str(exc)
    except Exception as exc:
        upload.status = UploadVersion.Status.REJECTED
        upload.note = str(exc)
        upload.save(update_fields=["status", "note"])
        mark_run_failed(job.processing_run, str(exc))
        job.status = ProcessingJob.Status.FAILED
        job.error = str(exc)
    else:
        job.status = ProcessingJob.Status.DONE
        job.error = ""
    job.lease_until = None
    job.heartbeat_at = timezone.now()
    job.save(update_fields=["status", "error", "lease_until", "heartbeat_at"])
    upload.refresh_from_db(fields=["status", "note", "recalculated_path"])
    return job


def _round_ratio(num, den):
    if not den:
        return 0
    return int(
        (Decimal(num) * Decimal(RATIO_SCALE) / Decimal(den)).quantize(
            Decimal("1"), rounding=ROUND_HALF_UP
        )
    )


def _detail_for_rows(rows):
    """Aggregate one or more NormalizedValue rows into a single detail dict."""
    units = {row.unit for row in rows}
    unit = rows[0].unit
    if units == {NormalizedValue.Unit.RATIO}:
        ratio_num = sum(int(row.ratio_num or 0) for row in rows)
        ratio_den = sum(int(row.ratio_den or 0) for row in rows)
        return {
            "unit": unit,
            "value_int": _round_ratio(ratio_num, ratio_den),
            "ratio_num": ratio_num,
            "ratio_den": ratio_den,
        }
    if units == {NormalizedValue.Unit.MONEY} and all(
        row.ratio_num is not None and row.ratio_den is not None for row in rows
    ):
        ratio_num = sum(int(row.ratio_num or 0) for row in rows)
        ratio_den = sum(int(row.ratio_den or 0) for row in rows)
        return {
            "unit": unit,
            "value_int": (
                int(
                    (Decimal(ratio_num) / Decimal(ratio_den)).quantize(
                        Decimal("1"), rounding=ROUND_HALF_UP
                    )
                )
                if ratio_den
                else 0
            ),
            "ratio_num": ratio_num,
            "ratio_den": ratio_den,
        }
    return {
        "unit": unit,
        "value_int": sum(int(row.value_int or 0) for row in rows),
        "ratio_num": None,
        "ratio_den": None,
    }


def _company_value_details(cycle, report_code, data_scope="approved"):
    grouped = defaultdict(list)
    if cycle is None:
        return {}
    current_ids = ProjectCycle.objects.filter(
        cycle=cycle,
        current_upload__cycle=cycle,
        current_upload__project=F("project"),
        current_upload__status=UploadVersion.Status.APPROVED,
        project_id__in=cycle_projects(cycle).values("pk"),
    ).values_list("current_upload_id", flat=True)
    if data_scope == "latest":
        from budgeting.services.trends import latest_report_uploads
        current_ids = [item.current_upload_id for item in latest_report_uploads(cycle)]
    for value in NormalizedValue.objects.filter(upload_id__in=current_ids, report_code=report_code):
        grouped[(value.row_code, value.period)].append(value)
    return {key: _detail_for_rows(rows) for key, rows in grouped.items()}


def project_value_details(project, cycle, report_code, data_scope="approved"):
    """Return ``{(row_code, period): detail}`` for one project's approved upload."""
    if cycle is None:
        return {}
    pc = ProjectCycle.objects.filter(
        project=project,
        cycle=cycle,
        current_upload__cycle=cycle,
        current_upload__project=F("project"),
        current_upload__status=UploadVersion.Status.APPROVED,
        project_id__in=cycle_projects(cycle).values("pk"),
    ).select_related("current_upload").first()
    if data_scope == "latest":
        from budgeting.services.trends import latest_report_uploads
        pc = next(iter(latest_report_uploads(cycle, project_id=project.pk)), None)
    if not pc:
        return {}
    grouped = defaultdict(list)
    for value in NormalizedValue.objects.filter(upload=pc.current_upload, report_code=report_code):
        grouped[(value.row_code, value.period)].append(value)
    return {key: _detail_for_rows(rows) for key, rows in grouped.items()}


def company_values(cycle, report_code):
    return {
        key: detail["value_int"]
        for key, detail in _company_value_details(cycle, report_code).items()
    }


def company_trend(cycle, report_code, row_code):
    """Return ``{period: detail}`` for one row across all approved projects."""
    return {
        period: detail
        for (rc, period), detail in _company_value_details(cycle, report_code).items()
        if rc == row_code
    }


def project_trend_rows(cycle, report_code, row_code):
    """Return ``[{code, name, series}]`` where series is ``{period: detail}`` per project."""
    if cycle is None:
        return []
    pcs = list(
        ProjectCycle.objects.filter(
            cycle=cycle,
            current_upload__cycle=cycle,
            current_upload__project=F("project"),
            current_upload__status=UploadVersion.Status.APPROVED,
            project_id__in=cycle_projects(cycle).values("pk"),
        )
        .select_related("project")
        .order_by("project__code")
    )
    upload_to_project = {pc.current_upload_id: pc.project for pc in pcs}
    per = defaultdict(lambda: defaultdict(list))
    for value in NormalizedValue.objects.filter(
        upload_id__in=list(upload_to_project),
        report_code=report_code,
        row_code=row_code,
    ):
        per[value.upload_id][value.period].append(value)
    rows = []
    for pc in pcs:
        series = {
            period: _detail_for_rows(vals)
            for period, vals in per.get(pc.current_upload_id, {}).items()
        }
        rows.append({"code": pc.project.code, "name": pc.project.name, "series": series})
    return rows


def sub_table_reports(cycle, data_scope="approved"):
    template = active_template(cycle)
    specs = []
    if template and template.file_path:
        specs = sub_table_specs(resolve_template_path(template.file_path))
    counts = {}
    if cycle:
        current_ids = ProjectCycle.objects.filter(
            cycle=cycle,
            current_upload__cycle=cycle,
            current_upload__project=F("project"),
            current_upload__status=UploadVersion.Status.APPROVED,
            project_id__in=cycle_projects(cycle).values("pk"),
        ).values_list("current_upload_id", flat=True)
        if data_scope == "latest":
            from budgeting.services.trends import latest_report_uploads
            current_ids = [item.current_upload_id for item in latest_report_uploads(cycle)]
        counts = dict(
            NormalizedValue.objects.filter(upload_id__in=current_ids)
            .exclude(report_code__in=REPORTS)
            .values("report_code")
            .annotate(n=Count("row_code", distinct=True))
            .values_list("report_code", "n")
        )
    if data_scope == "latest":
        known = {code for code, _ in specs}
        specs.extend((code, code) for code in sorted(counts) if code not in known)
    return [
        {"code": code, "name": name, "rows": counts.get(code, 0)}
        for code, name in specs
    ]


def project_contributions(cycle, report_code, row_code=None, period=None):
    current_ids = ProjectCycle.objects.filter(
        cycle=cycle,
        current_upload__cycle=cycle,
        current_upload__project=F("project"),
        current_upload__status=UploadVersion.Status.APPROVED,
        project_id__in=cycle_projects(cycle).values("pk"),
    ).values_list("current_upload_id", flat=True)
    values = NormalizedValue.objects.filter(upload_id__in=current_ids, report_code=report_code)
    if row_code:
        values = values.filter(row_code=row_code)
    if period:
        values = values.filter(period=period)
    return values.select_related("upload", "upload__project").order_by("upload__project__code", "row_code", "period")


@transaction.atomic
def create_adjustment_batch(cycle, report_code, row_code, period, delta_cents, reason, actor=None, due_date=None):
    cycle = _lock_mutable_cycle(cycle)
    baselines = {}
    for pc in ProjectCycle.objects.filter(cycle=cycle, current_upload__status=UploadVersion.Status.APPROVED).select_related("project", "current_upload"):
        value = NormalizedValue.objects.filter(
            upload=pc.current_upload,
            report_code=report_code,
            row_code=row_code,
            period=period,
            unit=NormalizedValue.Unit.MONEY,
        ).first()
        baselines[pc.project.code] = int(value.value_int or 0) if value else 0
    if not baselines:
        raise ValueError("没有可用于调整分配的已批准项目版本")
    batch = AdjustmentBatch.objects.create(
        cycle=cycle,
        report_code=report_code,
        row_code=row_code,
        period=period,
        baseline_total_cents=sum(baselines.values()),
        delta_cents=delta_cents,
        reason=reason,
        due_date=due_date,
    )
    allocation = dict(largest_remainder(delta_cents, baselines.items()))
    projects = {project.code: project for project in Project.objects.filter(code__in=baselines)}
    for code, baseline in baselines.items():
        delta = allocation.get(code, 0)
        AdjustmentLine.objects.create(
            batch=batch,
            cycle=cycle,
            project=projects[code],
            report_code=report_code,
            row_code=row_code,
            period=period,
            baseline_cents=baseline,
            weight=abs(baseline),
            allocated_delta_cents=delta,
            target_cents=baseline + delta,
        )
    audit(actor, "ADJUSTMENT_DRAFTED", "AdjustmentBatch", batch.id, {"delta_cents": delta_cents}, cycle=cycle)
    return batch


@transaction.atomic
def create_driver_adjustment(cycle, project, report_code, driver, to_value, reason, actor=None, due_date=None):
    """Create a single-project driver adjustment and a matching R0041 line."""
    cycle = _lock_mutable_cycle(cycle)
    details = project_value_details(project, cycle, report_code)
    baseline = {rc: detail for (rc, period), detail in details.items() if period == "YEAR"}
    if not baseline:
        raise ValueError("该项目尚无已批准年度数据，无法测算")
    snapshot = simulate_driver(baseline, driver, to_value)
    rows = {r[0]: r for r in snapshot["rows"]}
    room_rev_baseline = rows["ROOM_REV"][3]
    room_rev_target = rows["ROOM_REV"][4]
    delta = snapshot["delta_room_rev"]
    batch = AdjustmentBatch.objects.create(
        cycle=cycle,
        project=project,
        driver=driver,
        driver_label=snapshot["driver_label"],
        from_value=Decimal(str(snapshot["from_value"])) if snapshot["from_value"] is not None else None,
        to_value=Decimal(str(snapshot["to_value"])),
        target_room_rev_cents=room_rev_target,
        cascade={
            "driver": snapshot["driver"],
            "driver_label": snapshot["driver_label"],
            "from_value": snapshot["from_value"],
            "to_value": snapshot["to_value"],
            "delta_room_rev": delta,
            "rows": [list(r) for r in snapshot["rows"]],
        },
        report_code=report_code,
        row_code=ROOM_REV,
        period="YEAR",
        baseline_total_cents=room_rev_baseline,
        delta_cents=delta,
        reason=reason,
        due_date=due_date,
    )
    AdjustmentLine.objects.create(
        batch=batch,
        cycle=cycle,
        project=project,
        report_code=report_code,
        row_code=ROOM_REV,
        period="YEAR",
        baseline_cents=room_rev_baseline,
        weight=1,
        allocated_delta_cents=delta,
        target_cents=room_rev_target,
    )
    audit(actor, "ADJUSTMENT_DRAFTED", "AdjustmentBatch", batch.id,
          {"driver": driver, "to_value": str(to_value), "delta_cents": delta},
          project=project, cycle=cycle)
    return batch


def _detail_numeric(detail):
    if detail.get("ratio_den"):
        return detail["ratio_num"] / detail["ratio_den"]
    return detail.get("value_int", 0)


def _stored_int(unit, value):
    if unit == "RATIO":
        return int(round(float(value) * RATIO_SCALE))
    return int(round(value))


def _cascade_value(unit, value):
    if unit == "RATIO":
        return float(value)
    return int(round(value))


def _project_upload(project, cycle, upload=None):
    if not cycle_projects(cycle).filter(pk=project.pk).exists():
        raise ValueError("项目不属于当前年度计划范围")
    if upload is not None:
        if upload.project_id != project.pk or upload.cycle_id != cycle.pk:
            raise ValueError("调整基准上传版本与项目或预算周期不一致")
        return UploadVersion.objects.select_related("template").get(pk=upload.pk)
    pc = ProjectCycle.objects.filter(
        project=project,
        cycle=cycle,
        current_upload__cycle=cycle,
        current_upload__project=F("project"),
        current_upload__status=UploadVersion.Status.APPROVED,
    ).select_related("current_upload", "current_upload__template").first()
    if not pc or not pc.current_upload_id:
        raise ValueError("该项目尚无同项目同周期的已批准基准版本")
    return pc.current_upload


def _upload_value_details(upload, report_code):
    grouped = defaultdict(list)
    for value in NormalizedValue.objects.filter(upload=upload, report_code=report_code):
        grouped[(value.row_code, value.period)].append(value)
    return {key: _detail_for_rows(rows) for key, rows in grouped.items()}


def _project_report(project, cycle, report_code, upload=None):
    """Baseline values + formula graph for one project report."""
    upload = _project_upload(project, cycle, upload)
    details = _upload_value_details(upload, report_code)
    manifest = load_manifest(upload.template)
    if not manifest.get("reports"):
        raise ValueError("基准上传版本没有可用模板清单，不能执行全表公式重算")
    rows, num_to_code = build_report_graph(manifest, report_code)
    values = {}
    for (rc, period), detail in details.items():
        if rc in rows:
            values.setdefault(rc, {})[period] = _detail_numeric(detail)
    order = sorted(rows, key=lambda rc: rows[rc]["row_num"])
    return rows, num_to_code, values, order, upload


def _validate_edits(rows, edits):
    for rc, per in edits.items():
        if rc not in rows:
            raise ValueError("未知行代码：" + rc)
        if rows[rc]["kind"] != "leaf":
            raise ValueError("该行为联动计算项，不可直接编辑：" + rows[rc]["label"])
        for period in per:
            if period != "YEAR" and period not in MONTHS:
                raise ValueError("未知期间：" + period)


def preview_adjustment(project, cycle, report_code, edits=None, upload=None):
    """Recompute a project report with leaf edits; returns before/after rows.

    ``edits``: ``{row_code: {period: engine_numeric}}`` (MONEY cents, COUNT int,
    RATIO fraction).  Leaf-only; derived rows are recomputed.
    """
    edits = edits or {}
    rows, num_to_code, values, order, upload = _project_report(project, cycle, report_code, upload=upload)
    _validate_edits(rows, edits)
    external_values = load_external_values(upload, required_external_refs(rows, edits))
    result = recompute(rows, num_to_code, values, edits, external_values=external_values)
    return {
        "report_code": report_code,
        "periods": MONTHS + ["YEAR"],
        "rows": [
            {
                "code": rc,
                "label": rows[rc]["label"],
                "unit": rows[rc]["unit"],
                "kind": rows[rc]["kind"],
                "before": values.get(rc, {}),
                "after": result.get(rc, {}),
            }
            for rc in order
        ],
        "edited": {rc: per for rc, per in edits.items()},
    }


@transaction.atomic
def create_full_adjustment(cycle, project, report_code, edits, reason, actor=None, due_date=None, batch=None, upload=None):
    """Issue a full-table adjustment: one batch + one line per edited leaf."""
    cycle = _lock_mutable_cycle(cycle)
    rows, num_to_code, values, order, upload = _project_report(project, cycle, report_code, upload=upload)
    _validate_edits(rows, edits)
    if not edits:
        raise ValueError("没有可下发的调整项")
    for rc, per in edits.items():
        for period in per:
            if AdjustmentLine.objects.filter(
                cycle=cycle,
                project=project,
                report_code=report_code,
                row_code=rc,
                period=period,
                status=AdjustmentLine.Status.OPEN,
            ).exists():
                raise ValueError(f"「{rows[rc]['label']}」{period} 已有未确认的调整任务，请先取消或完成")
    external_values = load_external_values(upload, required_external_refs(rows, edits))
    result = recompute(rows, num_to_code, values, edits, external_values=external_values)
    cascade_rows = [
        [
            rc,
            rows[rc]["label"],
            rows[rc]["unit"],
            _cascade_value(rows[rc]["unit"], values.get(rc, {}).get("YEAR", 0)),
            _cascade_value(rows[rc]["unit"], result.get(rc, {}).get("YEAR", 0)),
        ]
        for rc in order
    ]
    edited_rows = [
        [
            rc,
            rows[rc]["label"],
            rows[rc]["unit"],
            period,
            _cascade_value(rows[rc]["unit"], values.get(rc, {}).get(period, 0)),
            _cascade_value(rows[rc]["unit"], value),
        ]
        for rc, per in sorted(edits.items())
        for period, value in sorted(per.items())
    ]
    line_data = []
    total_delta = 0
    for rc, per in edits.items():
        unit = rows[rc]["unit"]
        for period, value in per.items():
            baseline_int = _stored_int(unit, values.get(rc, {}).get(period, 0))
            target_int = _stored_int(unit, value)
            line_data.append((rc, period, baseline_int, target_int))
            total_delta += target_int - baseline_int
    if batch is None:
        target_room_rev = int(round(result.get("R0041", {}).get("YEAR", 0))) if "R0041" in result else 0
        batch = AdjustmentBatch.objects.create(
            cycle=cycle,
            project=project,
            driver="",
            driver_label="全表调整",
            target_room_rev_cents=target_room_rev,
            cascade={
                "kind": "full",
                "report_code": report_code,
                "rows": cascade_rows,
                "edited": edited_rows,
                "reports": {
                    report_code: {"rows": cascade_rows, "edited": edited_rows},
                },
            },
            report_code=report_code,
            row_code="",
            period="",
            baseline_total_cents=0,
            delta_cents=total_delta,
            reason=reason,
            due_date=due_date,
        )
    else:
        batch = AdjustmentBatch.objects.select_for_update().get(pk=batch.pk)
        if batch.cycle_id != cycle.pk or batch.project_id != project.pk:
            raise ValueError("追加全表调整的项目或预算周期不一致")
        if batch.status != AdjustmentBatch.Status.DRAFT:
            raise ValueError("只有草稿调整可以追加全表变更")
        cascade = dict(batch.cascade or {})
        reports = dict(cascade.get("reports") or {})
        reports[report_code] = {"rows": cascade_rows, "edited": edited_rows}
        cascade["reports"] = reports
        batch.cascade = cascade
        batch.delta_cents += total_delta
        batch.save(update_fields=["cascade", "delta_cents"])
    for rc, period, baseline_int, target_int in line_data:
        AdjustmentLine.objects.create(
            batch=batch,
            cycle=cycle,
            project=project,
            report_code=report_code,
            row_code=rc,
            period=period,
            baseline_cents=baseline_int,
            weight=1,
            allocated_delta_cents=target_int - baseline_int,
            target_cents=target_int,
        )
    audit(
        actor,
        "ADJUSTMENT_DRAFTED",
        "AdjustmentBatch",
        batch.id,
        {"kind": "full", "report_code": report_code, "edited": sorted(edits), "delta_cents": total_delta},
        project=project,
        cycle=cycle,
    )
    return batch


@transaction.atomic
def override_adjustment_line(line, allocated_delta_cents, actor=None):
    line = AdjustmentLine.objects.select_for_update().get(pk=line.pk)
    _lock_mutable_cycle(line.cycle)
    if line.batch.status != AdjustmentBatch.Status.DRAFT:
        raise ValueError("只有草稿调整可以手工覆盖")
    line.allocated_delta_cents = allocated_delta_cents
    line.target_cents = line.baseline_cents + allocated_delta_cents
    line.save(update_fields=["allocated_delta_cents", "target_cents"])
    audit(actor, "ADJUSTMENT_LINE_OVERRIDDEN", "AdjustmentLine", line.id, {"allocated_delta_cents": allocated_delta_cents}, project=line.project, cycle=line.cycle)
    return line


@transaction.atomic
def issue_adjustment(batch, actor=None):
    batch = AdjustmentBatch.objects.select_for_update().get(pk=batch.pk)
    batch.cycle = _lock_mutable_cycle(batch.cycle)
    if batch.status != AdjustmentBatch.Status.DRAFT:
        raise ValueError("只有草稿调整可以下发")
    cascade_kind = (batch.cascade or {}).get("kind")
    if cascade_kind == "fixed_cost_scenario":
        expected = batch.cascade.get("report_deltas_cents", {})
        actual = {code: 0 for code in expected}
        for line in batch.lines.all():
            actual[line.report_code] = actual.get(line.report_code, 0) + line.allocated_delta_cents
            if line.target_cents - line.baseline_cents != line.allocated_delta_cents:
                raise ValueError("场景目标与明细差额不一致")
        if not expected or actual != expected or actual.get(batch.report_code) != batch.delta_cents:
            raise ValueError("场景各报表差额必须与测算结果逐分一致")
    elif cascade_kind == "summary_annual":
        lines = list(batch.lines.all())
        if any(line.target_cents - line.baseline_cents != line.allocated_delta_cents for line in lines):
            raise ValueError("汇总表目标与明细差额不一致")
        reference = next(
            (
                line
                for line in lines
                if line.report_code == batch.report_code
                and line.row_code == batch.row_code
                and line.period == batch.period
            ),
            None,
        )
        if reference is None or reference.allocated_delta_cents != batch.delta_cents:
            raise ValueError("汇总表批次差额与主调整项不一致")
    elif sum(line.allocated_delta_cents for line in batch.lines.all()) != batch.delta_cents:
        raise ValueError("项目差额合计必须精确等于批次差额")
    batch.status = AdjustmentBatch.Status.ISSUED
    batch.issued_at = timezone.now()
    batch.save(update_fields=["status", "issued_at"])
    batch.cycle.status = BudgetCycle.Status.ADJUSTING
    batch.cycle.save(update_fields=["status"])
    audit(actor, "ADJUSTMENT_ISSUED", "AdjustmentBatch", batch.id, cycle=batch.cycle)
    return batch


@transaction.atomic
def cancel_adjustment(batch, actor=None):
    batch = AdjustmentBatch.objects.select_for_update().get(pk=batch.pk)
    batch.cycle = _lock_mutable_cycle(batch.cycle)
    if batch.status == AdjustmentBatch.Status.COMPLETED or batch.lines.filter(status=AdjustmentLine.Status.CONFIRMED).exists():
        raise ValueError("已有项目版本确认后不能直接取消")
    batch.status = AdjustmentBatch.Status.CANCELLED
    batch.save(update_fields=["status"])
    batch.lines.filter(status=AdjustmentLine.Status.OPEN).update(status=AdjustmentLine.Status.CANCELLED)
    audit(actor, "ADJUSTMENT_CANCELLED", "AdjustmentBatch", batch.id, cycle=batch.cycle)
    return batch


@transaction.atomic
def update_adjustment_lines_for_upload(upload):
    cycle = _lock_mutable_cycle(upload.cycle)
    open_lines = AdjustmentLine.objects.filter(project=upload.project, cycle=cycle, status=AdjustmentLine.Status.OPEN, batch__status=AdjustmentBatch.Status.ISSUED).select_related("batch")
    for line in open_lines:
        line.latest_upload = upload
        line.latest_value_cents = _adjustment_line_value(upload, line)
        line.difference_cents = None if line.latest_value_cents is None else line.latest_value_cents - line.target_cents
        if (line.difference_cents == 0 or (line.batch.cascade or {}).get("informational_only")) and upload.status == UploadVersion.Status.APPROVED:
            line.status = AdjustmentLine.Status.CONFIRMED
        line.save(update_fields=["latest_upload", "latest_value_cents", "difference_cents", "status"])
    for batch in AdjustmentBatch.objects.filter(lines__project=upload.project, cycle=upload.cycle, status=AdjustmentBatch.Status.ISSUED).distinct():
        if not batch.lines.filter(status=AdjustmentLine.Status.OPEN).exists():
            batch.status = AdjustmentBatch.Status.COMPLETED
            batch.save(update_fields=["status"])

def _freeze_project_ids(cycle):
    if cycle.plan_id:
        return list(
            PlanProject.objects.filter(plan_id=cycle.plan_id)
            .order_by("project__code", "project_id")
            .values_list("project_id", flat=True)
        )
    return list(Project.objects.filter(is_active=True).order_by("code").values_list("id", flat=True))


def _freeze_project_uploads(cycle, project_ids):
    rows = (
        ProjectCycle.objects.filter(
            cycle=cycle,
            project_id__in=project_ids,
            current_upload__status=UploadVersion.Status.APPROVED,
        )
        .select_related("project", "current_upload")
        .order_by("project__code", "project_id")
    )
    return {row.project_id: row.current_upload for row in rows if row.current_upload_id}


def _freeze_history_bindings(cycle, project_ids):
    if not cycle.plan_id:
        return {}
    bindings = {}
    for project_id, binding_id, binding_hash, revision, baseline_id in (
        PlanHistoryBinding.objects.filter(plan_id=cycle.plan_id, project_id__in=project_ids)
        .order_by("project_id", "-revision", "-id")
        .values_list("project_id", "id", "binding_hash", "revision", "baseline_id")
    ):
        key = str(project_id)
        if key not in bindings:
            bindings[key] = {
                "binding_id": binding_id,
                "binding_hash": binding_hash,
                "revision": revision,
                "baseline_id": baseline_id,
            }
    return bindings


def _freeze_target_context(cycle, uploads_by_project):
    if not cycle.plan_id:
        return {}
    context = {}
    for project_id, upload in uploads_by_project.items():
        target_set = latest_targets(cycle.plan, upload.project)
        if not target_set:
            context[str(project_id)] = None
            continue
        context[str(project_id)] = {
            "target_set_id": target_set.pk,
            "revision": target_set.revision,
            "constraint_count": target_set.constraints.count(),
            "evaluation_id": (
                TargetEvaluation.objects.filter(upload=upload, target_set=target_set)
                .order_by("-created_at", "-id")
                .values_list("id", flat=True)
                .first()
            ),
        }
    return context


def _freeze_selection(cycle, report_code=None):
    project_ids = _freeze_project_ids(cycle)
    uploads_by_project = _freeze_project_uploads(cycle, project_ids)
    if cycle.plan_id and cycle.template_id:
        context = ReportContext(
            budget_year=cycle.budget_year,
            cycle_id=cycle.pk,
            report_code=report_code or next(iter(REPORTS)),
            source_mode="APPROVED",
            project_ids=project_ids,
        )
        selection = freeze_selection_payload(context)
    else:
        selection = {
            "schema": 1,
            "budget_year": cycle.budget_year,
            "cycle_id": cycle.pk,
            "project_ids": project_ids,
            "upload_ids": {str(project_id): upload.pk for project_id, upload in uploads_by_project.items()},
            "manifest_hash": cycle.template.formula_manifest_hash if cycle.template_id else "",
            "processing_runs": {
                str(project_id): str(upload.processing_current_run_id) if upload.processing_current_run_id else None
                for project_id, upload in uploads_by_project.items()
            },
        }
    selection = dict(selection)
    selection["project_ids"] = [int(project_id) for project_id in selection.get("project_ids", [])]
    selection["upload_ids"] = {
        str(project_id): str(upload_id)
        for project_id, upload_id in selection.get("upload_ids", {}).items()
    }
    selection["processing_runs"] = {
        str(project_id): str(run_id) if run_id is not None else None
        for project_id, run_id in selection.get("processing_runs", {}).items()
    }
    selection["query_report_enabled"] = bool(cycle.plan_id and cycle.template_id)
    selection["report_codes"] = list(REPORTS)
    selection["history_bindings"] = _freeze_history_bindings(cycle, project_ids)
    selection["target_context"] = _freeze_target_context(cycle, uploads_by_project)
    selection["plan_revision_token"] = cycle.plan.revision_token if cycle.plan_id else None
    return selection


def _selection_upload_ids(selection):
    return list(selection.get("upload_ids", {}).values())


def _selection_expected_run_id(selection, upload):
    return selection.get("processing_runs", {}).get(str(upload.project_id))


def _run_matches(actual_run_id, expected_run_id):
    if expected_run_id is None:
        return actual_run_id is None
    return str(actual_run_id) == str(expected_run_id)


def _selection_uploads(selection):
    upload_ids = _selection_upload_ids(selection)
    return {upload.pk: upload for upload in UploadVersion.objects.filter(pk__in=upload_ids).select_related("project")}


def _selection_periods(selection, report_code):
    uploads = _selection_uploads(selection)
    periods = set()
    for value in NormalizedValue.objects.filter(upload_id__in=list(uploads), report_code=report_code).select_related("upload"):
        upload = uploads.get(value.upload_id)
        if upload and _run_matches(value.processing_run_id, _selection_expected_run_id(selection, upload)):
            periods.add(value.period)
    return sorted(periods, key=lambda period: (period != "YEAR", period))


def _metric_value_int(metric):
    if metric.value is None:
        return None
    if metric.unit == NormalizedValue.Unit.MONEY:
        return int((metric.value * Decimal(100)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    if metric.unit == NormalizedValue.Unit.RATIO:
        return int((metric.value * Decimal(RATIO_SCALE)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    return int(metric.value.quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def _decimal_payload(value):
    return None if value is None else str(value)


def _metric_payload(metric):
    return {
        "unit": metric.unit,
        "value_int": _metric_value_int(metric),
        "ratio_num": metric.numerator,
        "ratio_den": metric.denominator,
        "value": _decimal_payload(metric.value),
        "by_project": {str(project_id): _decimal_payload(value) for project_id, value in sorted(metric.by_project.items())},
        "missing": {str(project_id): reason for project_id, reason in sorted(metric.missing.items())},
    }


def _selected_value_details(cycle, selection, report_code):
    if not selection.get("query_report_enabled"):
        return _company_value_details(cycle, report_code)
    details = {}
    periods = _selection_periods(selection, report_code)
    if not periods:
        periods = ["YEAR"]
    for period in periods:
        context = ReportContext(
            budget_year=selection["budget_year"],
            cycle_id=selection["cycle_id"],
            report_code=report_code,
            source_mode="APPROVED",
            project_ids=selection["project_ids"],
            period=period,
        )
        result = query_report(context)
        expected_uploads = {str(project_id): str(upload_id) for project_id, upload_id in selection.get("upload_ids", {}).items()}
        actual_uploads = {str(project_id): str(upload_id) for project_id, upload_id in result.upload_ids.items()}
        if actual_uploads != expected_uploads:
            raise ValueError("冻结导出集合与固定选择不一致。")
        runs = UploadVersion.objects.filter(pk__in=result.upload_ids.values()).values_list("project_id", "processing_current_run_id")
        actual_runs = {str(project_id): str(run_id) if run_id is not None else None for project_id, run_id in runs}
        if actual_runs != selection.get("processing_runs", {}):
            raise ValueError("冻结导出处理批次与固定选择不一致。")
        for metric in result.metrics:
            details[(metric.row_code, period)] = _metric_payload(metric)
    return details


def _freeze_preconditions_for_targets(cycle, uploads_by_project):
    blockers = []
    if not cycle.plan_id:
        return blockers
    for project_id, upload in uploads_by_project.items():
        target_set = latest_targets(cycle.plan, upload.project)
        if not target_set or not target_set.constraints.exists():
            continue
        evaluation = (
            TargetEvaluation.objects.filter(upload=upload, target_set=target_set)
            .order_by("-created_at", "-id")
            .first()
        )
        if not evaluation or evaluation.status != "PASS":
            blockers.append(f"{upload.project.code} 目标检查未通过")
    return blockers


update_adjustment_lines = update_adjustment_lines_for_upload



def freeze_preconditions(cycle):
    cycle = BudgetCycle.objects.select_related("plan").get(pk=cycle.pk)
    blockers = []
    if BudgetCycle.objects.filter(pk=cycle.pk, status=BudgetCycle.Status.FROZEN).exists():
        return ["周期已冻结"]
    project_ids = _freeze_project_ids(cycle)
    uploads_by_project = _freeze_project_uploads(cycle, project_ids)
    projects = Project.objects.filter(pk__in=project_ids).order_by("code")
    for project in projects:
        if project.pk not in uploads_by_project:
            blockers.append(f"{project.code} 无正式版本")
    current_uploads = list(uploads_by_project.values())
    if any(has_current_p0(upload) for upload in current_uploads):
        blockers.append("存在 P0 阻断问题")
    if any(has_unacknowledged_current_p1(upload) for upload in current_uploads):
        blockers.append("存在未确认 P1 问题")
    for upload in current_uploads:
        blockers.extend(legacy_adjustment_blockers(upload))
    blockers.extend(_freeze_preconditions_for_targets(cycle, uploads_by_project))
    return blockers




def freeze_cycle(cycle, actor=None):
    cycle = BudgetCycle.objects.select_related("plan", "template").get(pk=cycle.pk)
    blockers = freeze_preconditions(cycle)
    if blockers:
        raise ValueError("；".join(blockers))

    selection = _freeze_selection(cycle)
    snapshot = FreezeSnapshot.objects.create(cycle=cycle, status=FreezeSnapshot.Status.STAGING)
    tmp = settings.BUDGET_STORAGE_ROOT / f"{snapshot.id}.tmp"
    tmp_zip = settings.BUDGET_STORAGE_ROOT / f"{snapshot.id}.tmp.zip"
    final = settings.BUDGET_STORAGE_ROOT / str(snapshot.id)
    zip_path = settings.BUDGET_STORAGE_ROOT / f"{snapshot.id}.zip"
    try:
        tmp.mkdir(parents=True, exist_ok=False)
        _copy_project_originals(cycle, tmp / "项目原件", selection=selection)
        _write_company_report_xlsx(cycle, tmp / "四表汇总.xlsx", selection=selection)
        (tmp / "四表汇总.json").write_text(
            json.dumps(_company_report_payload(cycle, selection=selection), ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        (tmp / "report-selection.json").write_text(
            json.dumps(selection, ensure_ascii=False, indent=2, sort_keys=True),
            encoding="utf-8",
        )
        _write_adjustment_ledger_xlsx(cycle, tmp / "调整台账.xlsx")
        _write_validation_report_xlsx(cycle, tmp / "校验报告.xlsx", selection=selection)
        manifest = _file_manifest(tmp)
        (tmp / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
        if final.exists() or zip_path.exists() or tmp_zip.exists():
            raise ValueError("快照文件已存在")
        with zipfile.ZipFile(tmp_zip, "w", compression=zipfile.ZIP_DEFLATED) as zf:
            for file_path in sorted(tmp.rglob("*")):
                if file_path.is_file():
                    zf.write(file_path, file_path.relative_to(tmp))
        tmp.replace(final)
        tmp_zip.replace(zip_path)
        with transaction.atomic():
            cycle = BudgetCycle.objects.select_for_update().select_related("plan", "template").get(pk=cycle.pk)
            if cycle.status == BudgetCycle.Status.FROZEN:
                raise ValueError("周期已冻结")
            if freeze_preconditions(cycle):
                raise ValueError("冻结提交前版本状态发生变化")
            if _freeze_selection(cycle) != selection:
                raise ValueError("冻结提交前版本状态发生变化")
            snapshot.directory = str(final.relative_to(settings.BUDGET_STORAGE_ROOT))
            snapshot.manifest_path = str((final / "manifest.json").relative_to(settings.BUDGET_STORAGE_ROOT))
            snapshot.history_bindings = selection.get("history_bindings", {})
            snapshot.status = FreezeSnapshot.Status.COMPLETE
            snapshot.completed_at = timezone.now()
            snapshot.save(update_fields=["directory", "manifest_path", "history_bindings", "status", "completed_at"])
            for relative, digest in manifest.items():
                path = final / relative
                SnapshotArtifact.objects.create(
                    snapshot=snapshot,
                    kind=path.suffix.lstrip(".") or "file",
                    relative_path=str(path.relative_to(settings.BUDGET_STORAGE_ROOT)),
                    sha256=digest,
                    size=path.stat().st_size,
                )
            manifest_path = final / "manifest.json"
            SnapshotArtifact.objects.create(
                snapshot=snapshot,
                kind="json",
                relative_path=str(manifest_path.relative_to(settings.BUDGET_STORAGE_ROOT)),
                sha256=sha256_file(manifest_path),
                size=manifest_path.stat().st_size,
            )
            SnapshotArtifact.objects.create(
                snapshot=snapshot,
                kind="zip",
                relative_path=str(zip_path.relative_to(settings.BUDGET_STORAGE_ROOT)),
                sha256=sha256_file(zip_path),
                size=zip_path.stat().st_size,
            )
            cycle.status = BudgetCycle.Status.FROZEN
            cycle.frozen_at = timezone.now()
            cycle.save(update_fields=["status", "frozen_at"])
            audit(actor, "CYCLE_FROZEN", "BudgetCycle", cycle.id, {"snapshot": str(snapshot.id)}, cycle=cycle)
    except Exception as exc:
        if tmp.exists():
            shutil.rmtree(tmp, ignore_errors=True)
        if tmp_zip.exists():
            tmp_zip.unlink(missing_ok=True)
        if final.exists():
            shutil.rmtree(final, ignore_errors=True)
        if zip_path.exists():
            zip_path.unlink(missing_ok=True)
        snapshot.status = FreezeSnapshot.Status.FAILED
        snapshot.error = str(exc)
        snapshot.save(update_fields=["status", "error"])
        raise
    return snapshot



@transaction.atomic
def reopen_cycle(cycle, actor=None, project_ids=None):
    cycle = BudgetCycle.objects.select_for_update().get(pk=cycle.pk)
    later = BudgetCycle.objects.filter(
        budget_year=cycle.budget_year, revision_no=cycle.revision_no + 1
    ).exists()
    if later:
        raise ValueError("下一修订版本已存在，请勿重复开放")
    new_cycle = BudgetCycle.objects.create(
        name=cycle.name,
        budget_year=cycle.budget_year,
        revision_no=cycle.revision_no + 1,
        status=BudgetCycle.Status.OPEN,
        template=cycle.template,
        p1_threshold_cents=cycle.p1_threshold_cents,
    )
    qs = ProjectCycle.objects.filter(cycle=cycle).select_related("project", "current_upload")
    open_ids = set(project_ids or [])
    for pc in qs:
        src = pc.current_upload
        carry = src if src and src.status == UploadVersion.Status.APPROVED else None
        current_upload = _copy_reopen_baseline(carry, new_cycle, actor) if carry else None
        ProjectCycle.objects.create(
            project=pc.project,
            cycle=new_cycle,
            current_upload=current_upload,
            is_open=(pc.project_id in open_ids) or current_upload is None,
        )
    audit(actor, "CYCLE_REOPENED", "BudgetCycle", new_cycle.id, {"source_cycle": cycle.id}, cycle=new_cycle)
    return new_cycle


def _copy_reopen_baseline(source, new_cycle, actor=None):
    copied = UploadVersion.objects.create(
        project=source.project,
        cycle=new_cycle,
        template=source.template,
        status=UploadVersion.Status.APPROVED,
        original_name=source.original_name,
        original_path=source.original_path,
        recalculated_path=source.recalculated_path,
        sha256=source.sha256,
        note=f"由冻结周期 {source.cycle_id} 的版本 {source.id} 复制为修订基线",
        submitted_at=source.submitted_at,
        approved_at=source.approved_at,
    )
    NormalizedValue.objects.bulk_create(
        [
            NormalizedValue(
                upload=copied,
                history_import_id=value.history_import_id,
                report_code=value.report_code,
                row_code=value.row_code,
                row_label=value.row_label,
                period=value.period,
                data_year=value.data_year,
                data_kind=value.data_kind,
                month=value.month,
                unit=value.unit,
                value_int=value.value_int,
                ratio_num=value.ratio_num,
                ratio_den=value.ratio_den,
                source_sheet=value.source_sheet,
                source_cell=value.source_cell,
                source_formula=value.source_formula,
            )
            for value in NormalizedValue.objects.filter(upload=source)
        ]
    )
    for source_run in ValidationRun.objects.filter(upload=source):
        copied_run = ValidationRun.objects.create(
            upload=copied,
            rule_version=source_run.rule_version,
            passed=source_run.passed,
        )
        ValidationIssue.objects.bulk_create(
            [
                ValidationIssue(
                    run=copied_run,
                    severity=issue.severity,
                    code=issue.code,
                    message=issue.message,
                    location=issue.location,
                    actual_value=issue.actual_value,
                    expected_value=issue.expected_value,
                    acknowledged=issue.acknowledged,
                    acknowledgement_note=issue.acknowledgement_note,
                )
                for issue in source_run.issues.all()
            ]
        )
    audit(
        actor,
        "UPLOAD_BASELINE_COPIED",
        "UploadVersion",
        copied.id,
        {"source_upload": str(source.id), "source_cycle": source.cycle_id},
        project=source.project,
        cycle=new_cycle,
        upload=copied,
    )
    from budgeting.services.historical_data import sync_history
    sync_history(copied)
    return copied



def _copy_project_originals(cycle, target_dir, selection=None):
    target_dir.mkdir()
    if selection is not None:
        uploads = _selection_uploads(selection)
        ordered = sorted(uploads.values(), key=lambda upload: upload.project.code)
        for upload in ordered:
            source = settings.BUDGET_STORAGE_ROOT / upload.original_path
            _verify_original_hash(upload, source)
            shutil.copy2(source, target_dir / f"{upload.project.code}_{source.name}")
        return
    for pc in ProjectCycle.objects.filter(cycle=cycle).select_related("project", "current_upload"):
        if pc.current_upload:
            source = settings.BUDGET_STORAGE_ROOT / pc.current_upload.original_path
            _verify_original_hash(pc.current_upload, source)
            shutil.copy2(source, target_dir / f"{pc.project.code}_{source.name}")


def _copy_snapshot_originals(cycle, target_dir, selection=None):
    return _copy_project_originals(cycle, target_dir, selection=selection)



def _record_formula_mapping_issues(run, upload, source_path):
    if not upload.template or not upload.template.manifest_path:
        return
    manifest_path = resolve_template_path(upload.template.manifest_path)
    if not manifest_path.is_absolute():
        manifest_path = Path(settings.BASE_DIR) / manifest_path
    template_path = resolve_template_path(upload.template.file_path)
    if not template_path.is_absolute():
        template_path = Path(settings.BASE_DIR) / template_path
    if not manifest_path.exists() or not template_path.exists():
        return
    try:
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        uploaded_rows, _ = formula_manifest(source_path)
        template_rows, _ = formula_manifest(template_path)
    except Exception:
        return
    uploaded_formulas = {
        (row["sheet"], row["cell"]): row for row in formula_comparison_rows(uploaded_rows, {row["sheet"] for row in uploaded_rows})
    }
    template_formulas = {
        (row["sheet"], row["cell"]): row for row in formula_comparison_rows(template_rows, {row["sheet"] for row in template_rows})
    }
    for report_code, default_sheet in REPORTS.items():
        report = (manifest.get("reports") or {}).get(report_code) or {}
        sheet = report.get("sheet") or default_sheet
        for item in report.get("mapping") or report.get("cells") or []:
            cell_ref = item.get("cell") or item.get("source_cell")
            if not cell_ref:
                continue
            actual = uploaded_formulas.get((sheet, cell_ref))
            expected = template_formulas.get((sheet, cell_ref))
            if actual != expected and (actual is not None or expected is not None):
                ValidationIssue.objects.get_or_create(
                    run=run,
                    code="FORMULA_CHANGED",
                    location=f"{sheet}!{cell_ref}",
                    defaults={
                        "severity": ValidationIssue.Severity.P0,
                        "message": "映射单元格公式与平台模板不一致。",
                        "actual_value": json.dumps(
                            actual or {"formula": ""},
                            ensure_ascii=False,
                            sort_keys=True,
                        ),
                        "expected_value": json.dumps(
                            expected or {"formula": ""},
                            ensure_ascii=False,
                            sort_keys=True,
                        ),
                    },
                )



def _write_company_report_xlsx(cycle, path, selection=None):
    wb = Workbook()
    first = True
    for report_code, title in REPORTS.items():
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = report_code[:31]
        ws.append([title])
        ws.append(["row_code", "period", "unit", "value_int", "ratio_num", "ratio_den"])
        details = _selected_value_details(cycle, selection, report_code) if selection is not None else _company_value_details(cycle, report_code)
        for (row_code, period), detail in sorted(details.items()):
            ws.append([
                row_code,
                period,
                detail["unit"],
                detail["value_int"],
                detail["ratio_num"],
                detail["ratio_den"],
            ])
    wb.save(path)




def _company_report_payload(cycle, selection=None):
    payload = {}
    for report_code in REPORTS:
        details = _selected_value_details(cycle, selection, report_code) if selection is not None else _company_value_details(cycle, report_code)
        payload[report_code] = {
            f"{row_code}|{period}": {
                "unit": detail["unit"],
                "value_int": detail["value_int"],
                "ratio_num": detail["ratio_num"],
                "ratio_den": detail["ratio_den"],
                **({"value": detail["value"], "by_project": detail["by_project"], "missing": detail["missing"]} if "value" in detail else {}),
            }
            for (row_code, period), detail in sorted(details.items())
        }
    return payload



def _write_adjustment_ledger_xlsx(cycle, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "调整台账"
    ws.append(["batch", "project", "report_code", "row_code", "period", "baseline_cents", "allocated_delta_cents", "target_cents", "latest_value_cents", "difference_cents", "status"])
    for line in AdjustmentLine.objects.filter(cycle=cycle).select_related("batch", "project").order_by("batch_id", "project__code"):
        ws.append([line.batch_id, line.project.code, line.report_code, line.row_code, line.period, line.baseline_cents, line.allocated_delta_cents, line.target_cents, line.latest_value_cents, line.difference_cents, line.status])
    wb.save(path)



def _write_validation_report_xlsx(cycle, path, selection=None):
    from budgeting.services.validation_reads import current_validation_issues

    wb = Workbook()
    ws = wb.active
    ws.title = "校验报告"
    ws.append(["upload", "severity", "code", "message", "location", "actual_value", "expected_value", "acknowledged"])
    issues = current_validation_issues().filter(run__upload__cycle=cycle).select_related("run", "run__upload").order_by("id")
    if selection is not None:
        upload_ids = set(_selection_upload_ids(selection))
        rows = []
        for issue in issues:
            upload = issue.run.upload
            if str(upload.pk) not in upload_ids:
                continue
            if not _run_matches(issue.run_id, _selection_expected_run_id(selection, upload)):
                continue
            rows.append(issue)
        issues = rows
    for issue in issues:
        ws.append([str(issue.run.upload_id), issue.severity, issue.code, issue.message, issue.location, issue.actual_value, issue.expected_value, issue.acknowledged])
    wb.save(path)



def _file_manifest(root):
    manifest = {}
    for file_path in sorted(root.rglob("*")):
        if file_path.is_file() and file_path.name != "manifest.json":
            manifest[str(file_path.relative_to(root))] = sha256_file(file_path)
    return manifest
