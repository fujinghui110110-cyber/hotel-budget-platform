"""Read original/recalculated workbook history before independent-history overlay."""
import json
from decimal import Decimal
from pathlib import Path

from django.conf import settings
from budgeting.services.template_paths import resolve_template_path
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple

from budgeting.excel.history import _declarations, _record_value
from budgeting.excel.extract import _history_columns
from budgeting.models import ValidationIssue
from budgeting.services.plan_history import current_binding, VALUE_FIELDS


def validate_history_identity(upload, workbook_path, run):
    from django.core import signing
    from budgeting.excel.ooxml import read_sys_meta
    from budgeting.services.plan_history import history_metadata

    if not upload.cycle.plan_id:
        return 0
    binding = current_binding(upload.cycle.plan, upload.project)
    if not binding:
        return 0
    try:
        meta = read_sys_meta(workbook_path)
        signed = signing.loads(meta.get("project_signature", ""), salt="budget-template-v1")
        expected = history_metadata(binding)
        valid = all(str(meta.get(key, "")) == str(value) and signed.get(key) == value for key, value in expected.items())
    except (signing.BadSignature, ValueError, KeyError):
        valid = False
    if valid:
        return 0
    ValidationIssue.objects.get_or_create(
        run=run, severity="P0", code="HISTORY_BASELINE_STALE", location="SYS_META",
        defaults={"message": "底稿未绑定当前已确认历史版本，请下载最新模板后重新填报。"},
    )
    return 1


def _period_key(value):
    period = str(value["period"])
    if period.startswith(("A", "F")):
        period = period[5:]
        period = period.removeprefix("M") or "YEAR"
    return (value["report_code"], value["row_code"], value["data_year"], value["data_kind"], period)


def template_cell_payload(binding, manifest, workbook=None):
    """Return declared cells; never reuse import-sheet coordinates in another template."""
    sheet, _, declarations = _declarations(manifest)
    locations = {}
    for row, period, cell in declarations:
        _, year, kind, month = period
        if kind in {"ACTUAL", "FORECAST"}:
            key = (row["report_code"], row["row_code"], year, kind, f"{month:02d}" if month else "YEAR")
            locations.setdefault(key, []).append((sheet, cell["cell"], row))
    if workbook is not None:
        for report_code, report in manifest.get("reports", {}).items():
            report_sheet = report.get("sheet")
            if report_sheet not in workbook.sheetnames:
                continue
            header = report.get("region", {}).get("header_row")
            history_cols = _history_columns(workbook[report_sheet], header)
            for item in report.get("mapping", report.get("cells", [])):
                cell = item.get("cell") or item.get("source_cell")
                if not cell or not item.get("row_code"):
                    continue
                row_number = coordinate_to_tuple(cell)[0]
                for col, nature, year in history_cols:
                    if nature not in {"A", "F"}:
                        continue
                    key = (report_code, item["row_code"], year, {"A": "ACTUAL", "F": "FORECAST"}[nature], "YEAR")
                    locations.setdefault(key, []).append((report_sheet, f"{col}{row_number}", item))
    payload = []
    for value in binding.baseline.values.values(*VALUE_FIELDS):
        cells = locations.get(_period_key(value), [])
        if not cells:
            payload.append({"value": value, "sheet": None, "cell": None, "mapping_missing": True})
        seen = set()
        for sheet_name, cell, row in cells:
            if (sheet_name, cell) in seen:
                continue
            seen.add((sheet_name, cell))
            if value["unit"] == "MONEY":
                excel_value = Decimal(value["value_int"]) / 100
            elif value["unit"] == "RATIO":
                den = value["ratio_den"]
                excel_value = Decimal(value["ratio_num"]) / den if den else None
            else:
                excel_value = value["value_int"]
            payload.append({"value": value, "sheet": sheet_name, "cell": cell,
                            "excel_value": excel_value, "row": row, "mapping_missing": False})
    return payload


def validate_workbook_history(upload, workbook_path, run, manifest=None):
    if not upload.cycle.plan_id:
        return 0
    binding = current_binding(upload.cycle.plan, upload.project)
    if not binding:
        return 0
    if manifest is None:
        manifest = json.loads((resolve_template_path(upload.template.manifest_path)).read_text())
    workbook = load_workbook(workbook_path, read_only=True, data_only=True, keep_links=False)
    issues = []
    try:
        for item in template_cell_payload(binding, manifest, workbook):
            expected = item["value"]
            if item["mapping_missing"]:
                issues.append(("HISTORY_MAPPING_MISSING", str(_period_key(expected)), "已确认历史缺少模板映射，不能验证锁定值。"))
                continue
            sheet, cell = item["sheet"], item["cell"]
            raw = workbook[sheet][cell].value if sheet in workbook.sheetnames else None
            converted = _record_value({**item["row"], "unit": expected["unit"]}, raw)
            same = converted is not None
            if same and expected["unit"] == "RATIO":
                # Workbook stores a ratio, not source numerator/denominator pairs.
                den = expected["ratio_den"]
                expected_ratio = Decimal(expected["ratio_num"]) / den if den else None
                same = raw is not None and expected_ratio is not None and abs(Decimal(str(raw)) - expected_ratio) <= Decimal("0.000000000001")
            elif same:
                same = converted["value_int"] == expected["value_int"]
            if not same:
                issues.append(("HISTORY_LOCK_VIOLATION", f"{sheet}!{cell}", "上传底稿历史值与管理员确认基准不一致或为空，请使用当前年度模板。"))
    finally:
        workbook.close()
    for code, location, message in issues:
        ValidationIssue.objects.get_or_create(run=run, severity="P0", code=code, location=location,
                                             defaults={"message": message})
    if not issues:
        upload.history_binding = binding
        upload.history_stale = False
        upload.save(update_fields=["history_binding", "history_stale"])
    return len(issues)
