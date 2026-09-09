from __future__ import annotations

import json
import re
import shutil
import tempfile
import zipfile
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Sequence

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment, Font, PatternFill

from budgeting.models import BudgetCycle, NormalizedValue, Project, REPORTS, UploadVersion
from budgeting.services.budget_versions import REPORTABLE_UPLOAD_STATUSES, selected_project_upload
from budgeting.services.report_labels import report_row_labels
from budgeting.services.workflow import RATIO_SCALE


USABLE_UPLOAD_STATUSES = REPORTABLE_UPLOAD_STATUSES


class VersionReportError(ValueError):
    pass


@dataclass(frozen=True)
class VersionReportArchive:
    path: Path
    manifest: dict


def usable_uploads(cycle: BudgetCycle, project_ids: Sequence[int] | None = None):

    projects = Project.objects.filter(is_active=True).order_by("code")
    if project_ids is not None:
        try:
            ids = [int(value) for value in project_ids]
        except (TypeError, ValueError) as exc:
            raise VersionReportError("项目筛选参数无效。") from exc
        projects = projects.filter(pk__in=ids)
    uploads = UploadVersion.objects.filter(
        cycle=cycle,
        project__in=projects,
        status__in=USABLE_UPLOAD_STATUSES,
    ).exclude(note__startswith="由冻结周期 ").select_related("project", "template")
    latest_by_project = {}
    for upload in uploads.order_by("project__code", "-created_at", "-id"):
        latest_by_project.setdefault(upload.project_id, upload)
    return projects, latest_by_project


def selected_upload(project: Project, cycle: BudgetCycle) -> UploadVersion | None:
    return selected_project_upload(project, cycle)


def latest_upload_status(project: Project, cycle: BudgetCycle) -> UploadVersion | None:
    return (
        UploadVersion.objects.filter(project=project, cycle=cycle)
        .exclude(note__startswith="由冻结周期 ")
        .order_by("-created_at", "-id")
        .first()
    )


def report_codes(upload: UploadVersion) -> list[str]:
    found = set(
        NormalizedValue.objects.filter(upload=upload)
        .values_list("report_code", flat=True)
        .distinct()
    )
    return [*([code for code in REPORTS if code in found]), *sorted(found - set(REPORTS))]


def report_title(report_code: str) -> str:
    return REPORTS.get(report_code, report_code)


def _period_label(value: NormalizedValue, budget_year: int) -> str:
    raw = (value.period or "").upper()
    if raw == "YEAR":
        return f"{budget_year}年预算年度"
    if value.data_year and value.data_kind:
        kind = {"ACTUAL": "实际", "FORECAST": "预测", "BUDGET": "预算"}.get(
            value.data_kind, value.data_kind
        )
        return f"{value.data_year}年{kind}{f'{value.month}月' if value.month else ''}"
    match = re.fullmatch(r"([AFB])(20\d{2})(?:M(0[1-9]|1[0-2]))?", raw)
    if match:
        kind = {"A": "实际", "F": "预测", "B": "预算"}[match.group(1)]
        return f"{match.group(2)}年{kind}{f'{int(match.group(3))}月' if match.group(3) else ''}"
    if re.fullmatch(r"(0[1-9]|1[0-2])", raw):
        return f"{budget_year}年预算{int(raw)}月"
    return raw or "未标注期间"


def _period_order(value: NormalizedValue):
    if value.data_year and value.data_kind:
        return (
            value.data_year,
            {"ACTUAL": 0, "FORECAST": 1, "BUDGET": 2}.get(value.data_kind, 9),
            value.month or 13,
            value.period,
        )
    raw = (value.period or "").upper()
    match = re.fullmatch(r"([AFB])(20\d{2})(?:M(0[1-9]|1[0-2]))?", raw)
    if match:
        return (
            int(match.group(2)),
            {"A": 0, "F": 1, "B": 2}[match.group(1)],
            int(match.group(3) or 13),
            raw,
        )
    if re.fullmatch(r"(0[1-9]|1[0-2])", raw):
        return (9999, 2, int(raw), raw)
    if raw == "YEAR":
        return (9999, 2, 13, raw)
    return (9999, 9, 99, raw)


def _text(cell, value):
    cell.value = "" if value is None else str(value)
    cell.data_type = "s"


def _safe_sheet_name(base: str, existing: set[str]) -> str:
    value = re.sub(r"[\\/*?:\[\]]", "_", base).strip() or "报表"
    value = value[:31]
    candidate, serial = value, 2
    while candidate in existing:
        suffix = f"_{serial}"
        candidate = f"{value[: 31 - len(suffix)]}{suffix}"
        serial += 1
    return candidate


def _value_for_export(value: NormalizedValue):
    if value.unit == NormalizedValue.Unit.MONEY:
        return Decimal(value.value_int) / Decimal("100"), '#,##0.00;[Red]-#,##0.00;-'
    if value.unit == NormalizedValue.Unit.RATIO:
        if value.ratio_num is not None and value.ratio_den is not None:
            ratio = Decimal(value.ratio_num) / Decimal(value.ratio_den) if value.ratio_den else Decimal(0)
            return ratio, "0.00%;[Red]-0.00%;-"
        return Decimal(value.value_int) / Decimal(RATIO_SCALE), "0.00%;[Red]-0.00%;-"
    return int(value.value_int), '#,##0;[Red]-#,##0;-'


def _unit_label(unit: str) -> str:
    return {
        NormalizedValue.Unit.MONEY: "元",
        NormalizedValue.Unit.COUNT: "数量",
        NormalizedValue.Unit.RATIO: "比例",
    }.get(unit, unit)


def _style_sheet(sheet):
    sheet.freeze_panes = "D7"
    sheet.sheet_view.showGridLines = False
    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 30
    sheet.column_dimensions["C"].width = 12
    for index in range(4, sheet.max_column + 1):
        sheet.column_dimensions[get_column_letter(index)].width = 16
    for row in (1, 6):
        for cell in sheet[row]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="34220A")
    for cell in sheet[6]:
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _write_report_sheet(workbook: Workbook, upload: UploadVersion, report_code: str, ordinal: int):
    values = list(
        NormalizedValue.objects.filter(upload=upload, report_code=report_code).order_by(
            "row_code", "period"
        )
    )
    sheet = workbook.create_sheet(
        _safe_sheet_name(f"{ordinal:02d}_{report_title(report_code)}", set(workbook.sheetnames))
    )
    _text(sheet.cell(1, 1), report_title(report_code))
    _text(sheet.cell(2, 1), "来源版本")
    _text(sheet.cell(2, 2), str(upload.id))
    _text(sheet.cell(3, 1), "上传文件")
    _text(sheet.cell(3, 2), upload.original_name or "未记录文件名")
    _text(sheet.cell(4, 1), "数据范围")
    _text(sheet.cell(4, 2), f"{upload.cycle.budget_year}预算及模板内历史实际、预测数据")

    by_period = {value.period: value for value in values}
    period_values = sorted(by_period.values(), key=_period_order)
    periods = [value.period for value in period_values]
    labels = report_row_labels(
        upload.cycle,
        report_code,
        {value.row_code for value in values},
        current_labels=[(value.row_code, value.row_label) for value in values],
    )
    headers = ["科目编码", "科目名称", "单位", *[_period_label(value, upload.cycle.budget_year) for value in period_values]]
    for column, header in enumerate(headers, start=1):
        _text(sheet.cell(6, column), header)

    row_codes = sorted({value.row_code for value in values})
    value_map = {(value.row_code, value.period): value for value in values}
    for row_number, row_code in enumerate(row_codes, start=7):
        values_for_row = [value_map.get((row_code, period)) for period in periods]
        sample = next((value for value in values_for_row if value is not None), None)
        _text(sheet.cell(row_number, 1), row_code)
        _text(sheet.cell(row_number, 2), labels.get(row_code) or (sample.row_label if sample else row_code))
        _text(sheet.cell(row_number, 3), _unit_label(sample.unit) if sample else "")
        for column, value in enumerate(values_for_row, start=4):
            if value is None:
                continue
            exported, number_format = _value_for_export(value)
            cell = sheet.cell(row_number, column, exported)
            cell.number_format = number_format
    _style_sheet(sheet)


def build_project_version_report(upload: UploadVersion, destination: Path | None = None) -> Path:
    if upload.status not in USABLE_UPLOAD_STATUSES:
        raise VersionReportError("该上传版本尚未通过校验，不能导出管理报表。")
    report_list = report_codes(upload)
    temp_dir = None
    if destination is None:
        temp_dir = Path(tempfile.mkdtemp(prefix="budget-version-report-"))
        destination = temp_dir / f"{_safe_filename(upload.project.code)}-{upload.cycle.budget_year}-预算报表.xlsx"
    destination.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    cover = workbook.active
    cover.title = "导出说明"
    cover.sheet_view.showGridLines = False
    details = (
        ("导出类型", "项目预算报表（平台抽取数据）"),
        ("项目", f"{upload.project.code} {upload.project.name}"),
        ("预算周期", f"{upload.cycle.name} R{upload.cycle.revision_no}"),
        ("预算年度", upload.cycle.budget_year),
        ("来源版本 ID", upload.id),
        ("来源文件", upload.original_name or "未记录文件名"),
        ("文件 SHA-256", upload.sha256),
        ("校验状态", upload.get_status_display()),
        ("上传时间", _timestamp(upload.created_at)),
        ("报表数量", len(report_list)),
        ("说明", "本文件为平台抽取报表，不替代原始或重算预算底稿。缺失数据保持为空。"),
    )
    for row, (label, value) in enumerate(details, start=1):
        _text(cover.cell(row, 1), label)
        _text(cover.cell(row, 2), value)
    cover.column_dimensions["A"].width = 20
    cover.column_dimensions["B"].width = 92
    for cell in cover[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="34220A")
    for ordinal, code in enumerate(report_list, start=1):
        _write_report_sheet(workbook, upload, code, ordinal)
    if not report_list:
        _text(cover.cell(13, 1), "提示")
        _text(cover.cell(13, 2), "该版本尚未抽取到报表数据。")
    workbook.save(destination)
    return destination


def build_cycle_version_report_zip(
    cycle: BudgetCycle, *, project_ids: Sequence[int] | None = None
) -> VersionReportArchive:
    projects, selected = usable_uploads(cycle, project_ids)
    all_projects = list(projects)
    if not all_projects:
        raise VersionReportError("当前筛选没有可导出的项目。")
    temp_dir = Path(tempfile.mkdtemp(prefix="budget-version-report-zip-"))
    archive_path = temp_dir / f"{_safe_filename(cycle.name)}-{cycle.budget_year}-项目预算报表.zip"
    manifest = {
        "schema_version": 1,
        "export_type": "budget_version_reports",
        "cycle": {
            "id": cycle.pk,
            "name": cycle.name,
            "budget_year": cycle.budget_year,
            "revision_no": cycle.revision_no,
            "status": cycle.status,
        },
        "generated_at": timezone.now().isoformat(),
        "included_projects": [],
        "omitted_projects": [],
    }
    try:
        with zipfile.ZipFile(archive_path, "w", compression=zipfile.ZIP_DEFLATED, allowZip64=True) as archive:
            for project in all_projects:
                upload = selected.get(project.pk)
                latest = latest_upload_status(project, cycle)
                if upload is None:
                    reason = "未上传" if latest is None else f"最新上传状态为{latest.get_status_display()}，未形成可用版本"
                    manifest["omitted_projects"].append(
                        {
                            "project_id": project.pk,
                            "project_code": project.code,
                            "project_name": project.name,
                            "status": latest.status if latest else "NOT_UPLOADED",
                            "status_label": latest.get_status_display() if latest else "未上传",
                            "reason": reason,
                        }
                    )
                    continue
                output = temp_dir / f"{_safe_filename(project.code)}-{cycle.budget_year}-预算报表.xlsx"
                build_project_version_report(upload, output)
                member = f"{_safe_filename(project.code)}/{output.name}"
                archive.write(output, arcname=member)
                manifest["included_projects"].append(
                    {
                        "project_id": project.pk,
                        "project_code": project.code,
                        "project_name": project.name,
                        "upload_id": str(upload.id),
                        "upload_status": upload.status,
                        "upload_status_label": upload.get_status_display(),
                        "latest_upload_id": str(latest.id) if latest else None,
                        "latest_upload_status": latest.status if latest else None,
                        "latest_upload_status_label": latest.get_status_display() if latest else None,
                        "fallback_from_latest": bool(latest and latest.id != upload.id),
                        "source_file": upload.original_name,
                        "sha256": upload.sha256,
                        "report_count": len(report_codes(upload)),
                        "archive_path": member,
                    }
                )
            archive.writestr(
                "manifest.json",
                json.dumps(manifest, ensure_ascii=False, indent=2, sort_keys=True).encode("utf-8"),
            )
    except Exception:
        shutil.rmtree(temp_dir, ignore_errors=True)
        raise
    if not manifest["included_projects"]:
        shutil.rmtree(temp_dir, ignore_errors=True)
        raise VersionReportError(f"{cycle.name} 没有已校验、已提交或已批准的项目版本可导出。")
    return VersionReportArchive(path=archive_path, manifest=manifest)


def cleanup_version_report(path: Path | str) -> None:
    target = Path(path)
    shutil.rmtree(target.parent, ignore_errors=True)


def _safe_filename(value: str) -> str:
    cleaned = re.sub(r"[^0-9A-Za-z_.\-\u4e00-\u9fff]+", "_", value or "预算报表")
    return cleaned.strip("._") or "预算报表"


def _timestamp(value: datetime | None) -> str:
    return timezone.localtime(value).strftime("%Y-%m-%d %H:%M:%S") if value else "—"
