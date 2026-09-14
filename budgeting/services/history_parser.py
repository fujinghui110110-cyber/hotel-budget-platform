from __future__ import annotations

import re
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any, Iterable, Mapping

from openpyxl import load_workbook

from budgeting.services.workbook_reference import read_workbook


MONTHS = set(range(1, 13))
KIND_ALIASES = {
    "ACTUAL": {"actual", "实际", "实绩"},
    "FORECAST": {"forecast", "预测", "预估", "预计"},
    "BUDGET": {"budget", "预算"},
}
REVERSE_KIND_ALIASES = {
    alias: kind for kind, aliases in KIND_ALIASES.items() for alias in aliases
}
VALUE_HEADER_RE = re.compile(r"(金额|数值|本期|发生额|合计|全年|年度|累计|actual|forecast|budget)", re.I)
LABEL_HEADER_RE = re.compile(r"(科目|项目|指标|名称|内容|account|item|subject)", re.I)
YEAR_HEADER_RE = re.compile(r"^(年份|年度|年|year)$", re.I)
KIND_HEADER_RE = re.compile(r"^(口径|类型|类别|kind|type)$", re.I)
MONTH_RE = re.compile(r"(?<!\d)(1[0-2]|0?[1-9])\s*月|M(0?[1-9]|1[0-2])\b", re.I)
YEAR_RE = re.compile(r"(20\d{2})")
UNSAFE_VALUE_HEADER_RE = re.compile(
    r"(同比|环比|增长|增幅|占比|率差|差异|完成率|达成率|预算率|forecast\s*variance|variance|%)",
    re.I,
)
IRRELEVANT_SHEET_RE = re.compile(r"(附表|明细|说明|封面|目录|台账|调整|拆分|拆中智)")
PNL_SHEET_RE = re.compile(r"(损益|利润|经营|p\s*&?\s*l|profit|income\s*statement)", re.I)


@dataclass(frozen=True)
class ColumnMeta:
    column: int
    header_row: int | None
    month: int | None
    year: int | None
    kind: str | None
    annual: bool
    value_like: bool
    source_cell: str | None
    header_text: str


def propose_history(
    path: str | Path,
    canonical_rows: Iterable[Mapping[str, Any]],
    data_year: int,
    data_kind: str,
    report_code: str = "PL_TOTAL_WINE",
    money_unit: str = "yuan",
) -> dict[str, Any]:
    workbook = _read_history_workbook(path)
    canonical = [
        {
            "row_code": str(row.get("row_code") or ""),
            "row_label": str(row.get("row_label") or row.get("label") or ""),
            "unit": str(row.get("unit") or "MONEY").upper(),
            "aggregation": str(row.get("aggregation") or "").upper(),
        }
        for row in canonical_rows
        if row.get("row_code")
    ]
    target_kind = _normalize_kind(data_kind)
    issues: list[dict[str, Any]] = []
    rows: list[dict[str, Any]] = []
    selected_sheets = _select_sheets(workbook.get("sheets") or [], issues, report_code)
    for sheet in selected_sheets:
        rows.extend(
            _propose_sheet(sheet, canonical, int(data_year), target_kind, report_code, money_unit, issues)
        )
    _flag_duplicate_sources(rows, issues)
    if not rows:
        issues.append({"code": "NO_HISTORY_ROWS", "message": "未找到可解析的历史损益科目行。"})
    return {"rows": rows, "issues": issues}


def _propose_sheet(
    sheet: Mapping[str, Any],
    canonical: list[dict[str, str]],
    data_year: int,
    data_kind: str,
    report_code: str,
    money_unit: str,
    issues: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    cells = list(sheet.get("cells") or [])
    by_row: dict[int, list[dict[str, Any]]] = {}
    by_pos: dict[tuple[int, int], dict[str, Any]] = {}
    for cell in cells:
        row = int(cell.get("row") or 0)
        col = int(cell.get("column") or 0)
        by_row.setdefault(row, []).append(cell)
        by_pos[(row, col)] = cell
    if not by_row:
        return []
    label_col = _detect_label_column(by_row)
    if label_col is None:
        issues.append({"code": "SUBJECT_COLUMN_MISSING", "sheet": sheet.get("name"), "message": "未识别科目列。"})
        return []
    subject_header_row = _detect_subject_header_row(by_row, label_col)
    year_col = _detect_metadata_column(by_row, YEAR_HEADER_RE)
    kind_col = _detect_metadata_column(by_row, KIND_HEADER_RE)
    value_columns = _detect_value_columns(
        by_row,
        label_col,
        year_col,
        kind_col,
        data_year,
        data_kind,
        issues,
        sheet.get("name"),
        subject_header_row,
    )
    if not value_columns:
        issues.append({"code": "VALUE_COLUMN_MISSING", "sheet": sheet.get("name"), "message": "未识别与目标年度/口径匹配的金额列。"})
        return []

    proposals: list[dict[str, Any]] = []
    for row_num in sorted(by_row):
        label_cell = by_pos.get((row_num, label_col))
        label = _cell_text(label_cell)
        if not label or LABEL_HEADER_RE.search(label):
            continue
        row_year = _parse_year(_cell_text(by_pos.get((row_num, year_col))) if year_col else "")
        if row_year is not None and row_year != data_year:
            issues.append(
                {
                    "code": "YEAR_MISMATCH_ROW",
                    "sheet": sheet.get("name"),
                    "row": row_num,
                    "source_label": label,
                    "message": f"行年份 {row_year} 与目标年度 {data_year} 不一致，已跳过。",
                }
            )
            continue
        row_kind = _normalize_kind(_cell_text(by_pos.get((row_num, kind_col))) if kind_col else "")
        if row_kind and row_kind != data_kind:
            continue
        match = _match_label(label, canonical)
        canonical_row = _canonical_for_match(match, canonical)
        unit_hint = _best_unit_hint(label, canonical, match)
        aggregation = (canonical_row or {}).get("aggregation") or _infer_aggregation(label, unit_hint)
        values = _row_values(
            sheet_name=str(sheet.get("name") or ""),
            row_num=row_num,
            by_pos=by_pos,
            columns=value_columns,
            label=label,
            unit_hint=unit_hint,
            aggregation=aggregation,
            data_year=data_year,
            money_unit=money_unit,
            issues=issues,
        )
        if not values:
            code = "UNRECOGNIZED_VALUE" if any(by_pos.get((row_num, meta.column)) for meta in value_columns) else "MISSING_VALUE"
            issues.append(
                {
                    "code": code,
                    "sheet": sheet.get("name"),
                    "row": row_num,
                    "source_label": label,
                    "message": "科目行没有可采集的数值，已保留为空值行供审计。",
                }
            )
        row = {
            "report_code": report_code,
            "source_key": f"{sheet.get('name')}!R{row_num}",
            "source_label": label,
            "suggested_code": match["suggested_code"],
            "confidence": match["confidence"],
            "candidates": match["candidates"],
            "aggregation": aggregation,
            "values": values,
            "include": bool(match["suggested_code"] and values),
        }
        if not match["suggested_code"]:
            issues.append(
                {
                    "code": "UNKNOWN_SUBJECT",
                    "sheet": sheet.get("name"),
                    "row": row_num,
                    "source_label": label,
                    "candidates": match["candidates"],
                    "message": "科目相似度不足，未默认匹配。",
                }
            )
        proposals.append(row)
    return proposals


def _select_sheets(
    sheets: list[Mapping[str, Any]], issues: list[dict[str, Any]], report_code: str
) -> list[Mapping[str, Any]]:
    visible = [sheet for sheet in sheets if sheet.get("visible", True)]
    if len(visible) <= 1:
        if visible and _sheet_contradicts_report(str(visible[0].get("name") or ""), report_code):
            issues.append(
                {
                    "code": "PNL_SHEET_REPORT_MISMATCH",
                    "sheet": visible[0].get("name"),
                    "message": "工作表名称与目标报表口径明显不一致，已拒绝导入。",
                }
            )
            return []
        return visible
    exact = [sheet for sheet in visible if _sheet_matches_report(str(sheet.get("name") or ""), report_code)]
    if exact:
        return exact
    selected = []
    for sheet in visible:
        name = str(sheet.get("name") or "")
        if (
            PNL_SHEET_RE.search(name)
            and not IRRELEVANT_SHEET_RE.search(name)
            and not _sheet_contradicts_report(name, report_code)
        ):
            selected.append(sheet)
    if selected:
        return selected
    scored = []
    for sheet in visible:
        text = " ".join(_cell_text(cell) for cell in (sheet.get("cells") or [])[:200])
        score = len(PNL_SHEET_RE.findall(str(sheet.get("name") or "") + " " + text))
        name = str(sheet.get("name") or "")
        if not IRRELEVANT_SHEET_RE.search(name) and not _sheet_contradicts_report(name, report_code):
            scored.append((score, sheet))
    best_score = max((score for score, _ in scored), default=0)
    chosen = [sheet for score, sheet in scored if score and score == best_score]
    if chosen:
        return chosen
    issues.append({"code": "PNL_SHEET_MISSING", "message": "多 sheet 工作簿中未识别损益相关表。"})
    return []


def _sheet_matches_report(name: str, report_code: str) -> bool:
    if report_code == "PL_TOTAL_WINE":
        return "酒店损益总表" in name and "含名酒" in name and "不含名酒" not in name
    if report_code == "PL_TOTAL_NOWINE":
        return "酒店损益总表" in name and "不含名酒" in name
    if report_code == "PL_ZZ_WINE":
        return "损益表" in name and "拆中智" in name and "含名酒" in name and "不含名酒" not in name
    if report_code == "PL_ZZ_NOWINE":
        return "损益表" in name and "拆中智" in name and "不含名酒" in name
    return False


def _sheet_contradicts_report(name: str, report_code: str) -> bool:
    if report_code in {"PL_TOTAL_WINE", "PL_ZZ_WINE"}:
        return "不含名酒" in name
    if report_code in {"PL_TOTAL_NOWINE", "PL_ZZ_NOWINE"}:
        return "含名酒" in name and "不含名酒" not in name
    return False


def _read_history_workbook(path: str | Path) -> dict[str, Any]:
    workbook = read_workbook(path, include_cells=True)
    if workbook.get("sheets"):
        return workbook
    return _read_openpyxl_workbook(path)


def _read_openpyxl_workbook(path: str | Path) -> dict[str, Any]:
    source = Path(path).expanduser()
    formulas = load_workbook(source, data_only=False, read_only=True, keep_links=False)
    values = load_workbook(source, data_only=True, read_only=True, keep_links=False)
    try:
        sheets = []
        for formula_sheet in formulas.worksheets:
            value_sheet = values[formula_sheet.title]
            cells = []
            for row in formula_sheet.iter_rows():
                for formula_cell in row:
                    raw_value = formula_cell.value
                    if raw_value is None:
                        continue
                    value_cell = value_sheet[formula_cell.coordinate]
                    is_formula = isinstance(raw_value, str) and raw_value.startswith("=")
                    cached_value = value_cell.value
                    cell_value = cached_value if is_formula else raw_value
                    cells.append(
                        {
                            "coordinate": formula_cell.coordinate.upper(),
                            "row": formula_cell.row,
                            "column": formula_cell.column,
                            "display_value": None if cell_value is None else str(cell_value),
                            "cached_value": cell_value,
                            "cache_status": "missing" if is_formula and cached_value is None else "cached" if is_formula else "source",
                            "formula": raw_value if is_formula else None,
                            "is_formula": is_formula,
                            "is_error": isinstance(cell_value, str) and cell_value.startswith("#"),
                            "error_status": cell_value if isinstance(cell_value, str) and cell_value.startswith("#") else None,
                        }
                    )
            sheets.append(
                {
                    "id": formula_sheet.title,
                    "name": formula_sheet.title,
                    "state": "visible",
                    "visible": True,
                    "row_count": formula_sheet.max_row,
                    "col_count": formula_sheet.max_column,
                    "cells": cells,
                }
            )
        return {"source_path": str(source), "sheets": sheets}
    finally:
        formulas.close()
        values.close()


def _detect_label_column(by_row: Mapping[int, list[dict[str, Any]]]) -> int | None:
    scores: dict[int, float] = {}
    for row_num, cells in by_row.items():
        for cell in cells:
            col = int(cell["column"])
            text = _cell_text(cell)
            if not text:
                continue
            score = 0.0
            if LABEL_HEADER_RE.search(text):
                score += 8
            if _looks_like_subject(text):
                score += 1.5
            if _cell_number(cell) is None and len(text) <= 80:
                score += 0.2
            if row_num <= 5:
                score += 0.1
            scores[col] = scores.get(col, 0.0) + score
    return max(scores, key=scores.get) if scores else None


def _detect_metadata_column(by_row: Mapping[int, list[dict[str, Any]]], pattern: re.Pattern[str]) -> int | None:
    for cells in by_row.values():
        for cell in cells:
            text = _cell_text(cell)
            if text and pattern.search(_normalize_spaces(text)):
                return int(cell["column"])
    return None


def _detect_subject_header_row(by_row: Mapping[int, list[dict[str, Any]]], label_col: int) -> int | None:
    candidates = []
    for row_num, cells in by_row.items():
        label_text = ""
        row_has_value_header = False
        for cell in cells:
            text = _cell_text(cell)
            if int(cell["column"]) == label_col:
                label_text = text
            elif text and _cell_number(cell) is None:
                row_has_value_header = row_has_value_header or bool(
                    MONTH_RE.search(text) or VALUE_HEADER_RE.search(text) or _parse_year(text) or _normalize_kind(text)
                )
        if label_text and LABEL_HEADER_RE.search(label_text):
            candidates.append((0 if row_has_value_header else 1, row_num))
    if candidates:
        return min(candidates)[1]
    return None


def _detect_value_columns(
    by_row: Mapping[int, list[dict[str, Any]]],
    label_col: int,
    year_col: int | None,
    kind_col: int | None,
    data_year: int,
    data_kind: str,
    issues: list[dict[str, Any]],
    sheet_name: Any,
    subject_header_row: int | None,
) -> list[ColumnMeta]:
    metas: list[ColumnMeta] = []
    max_row = max(by_row)
    header_rows = _candidate_header_rows(by_row, subject_header_row, max_row)
    columns = sorted({int(cell["column"]) for cells in by_row.values() for cell in cells})
    for col in columns:
        if col in {label_col, year_col, kind_col}:
            continue
        header = _header_for_column(by_row, header_rows, col)
        if not header:
            continue
        if UNSAFE_VALUE_HEADER_RE.search(header):
            continue
        meta = _column_meta(col, header)
        if not meta.value_like:
            continue
        if meta.year is not None and meta.year != data_year:
            issues.append(
                {
                    "code": "YEAR_MISMATCH_COLUMN",
                    "sheet": sheet_name,
                    "column": _column_label(col),
                    "header": header,
                    "message": f"列年份 {meta.year} 与目标年度 {data_year} 不一致，已跳过。",
                }
            )
            continue
        if meta.kind is not None and meta.kind != data_kind:
            continue
        metas.append(meta)
    return metas


def _header_for_column(
    by_row: Mapping[int, list[dict[str, Any]]], header_rows: list[int], col: int
) -> str:
    pieces = []
    last_cell = None
    for row in header_rows:
        cell = next((item for item in by_row.get(row, []) if int(item["column"]) == col), None)
        text = _cell_text(cell)
        if text and _cell_number(cell) is None and _looks_like_header_text(text):
            pieces.append(text)
            last_cell = cell
    return " ".join(pieces[-4:]) if last_cell else ""


def _candidate_header_rows(
    by_row: Mapping[int, list[dict[str, Any]]], subject_header_row: int | None, max_row: int
) -> list[int]:
    if subject_header_row is not None:
        lower = max(1, subject_header_row - 6)
        upper = min(max_row, subject_header_row + 1)
        return [row for row in sorted(by_row) if lower <= row <= upper]
    rows = []
    for row_num, cells in by_row.items():
        text_cells = [_cell_text(cell) for cell in cells if _cell_text(cell) and _cell_number(cell) is None]
        if any(_looks_like_header_text(text) for text in text_cells):
            rows.append(row_num)
    return sorted(rows[:8])


def _looks_like_header_text(text: str) -> bool:
    return bool(
        LABEL_HEADER_RE.search(text)
        or MONTH_RE.search(text)
        or VALUE_HEADER_RE.search(text)
        or _parse_year(text)
        or _normalize_kind(text)
    )


def _column_meta(col: int, header: str) -> ColumnMeta:
    month_match = MONTH_RE.search(header)
    month = int(month_match.group(1) or month_match.group(2)) if month_match else None
    year = _parse_year(header)
    kind = _normalize_kind(header) or None
    annual = bool(re.search(r"(全年|年度|年合计|合计|累计|FY|annual|total)", header, re.I)) and month is None
    value_like = bool(month or annual or VALUE_HEADER_RE.search(header) or year or kind)
    return ColumnMeta(
        column=col,
        header_row=None,
        month=month,
        year=year,
        kind=kind,
        annual=annual,
        value_like=value_like,
        source_cell=None,
        header_text=header,
    )


def _row_values(
    *,
    sheet_name: str,
    row_num: int,
    by_pos: Mapping[tuple[int, int], Mapping[str, Any]],
    columns: list[ColumnMeta],
    label: str,
    unit_hint: str,
    aggregation: str,
    data_year: int,
    money_unit: str,
    issues: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    values: list[dict[str, Any]] = []
    monthly_months: set[int] = set()
    for meta in columns:
        cell = by_pos.get((row_num, meta.column))
        if cell is None:
            continue
        if cell.get("is_formula") and cell.get("cache_status") == "missing":
            issues.append(
                {
                    "code": "FORMULA_CACHE_MISSING",
                    "sheet": sheet_name,
                    "cell": cell.get("coordinate"),
                    "source_label": label,
                    "message": "公式单元格没有缓存值，无法安全采集。",
                }
            )
            continue
        number = _cell_number(cell)
        if number is None:
            continue
        converted = _convert_number(
            number, unit_hint, money_unit, label, issues, sheet_name, cell.get("coordinate")
        )
        if converted is None:
            continue
        month = meta.month
        values.append(
            {
                "month": month,
                "value_int": converted["value_int"],
                "ratio_num": converted["ratio_num"],
                "ratio_den": converted["ratio_den"],
                "unit": unit_hint,
                "source_sheet": sheet_name,
                "source_cell": cell.get("coordinate"),
                "source_formula": cell.get("formula") or "",
            }
        )
        if month in MONTHS:
            monthly_months.add(month)
    if monthly_months == MONTHS and unit_hint != "RATIO" and not any(item["month"] is None for item in values):
        if aggregation != "SUM":
            issues.append(
                {
                    "code": "ANNUAL_AGGREGATION_UNCERTAIN",
                    "sheet": sheet_name,
                    "row": row_num,
                    "source_label": label,
                    "message": "该科目不是明确 SUM 聚合，未由 12 个月自动合成年值。",
                }
            )
            return sorted(values, key=lambda item: (99 if item["month"] is None else int(item["month"])))
        values.append(
            {
                "month": None,
                "value_int": sum(int(item["value_int"] or 0) for item in values if item["month"] in MONTHS),
                "ratio_num": None,
                "ratio_den": None,
                "unit": unit_hint,
                "source_sheet": sheet_name,
                "source_cell": "HISTORY",
                "source_formula": f"SUM({data_year}M01:{data_year}M12)",
            }
        )
    return sorted(values, key=lambda item: (99 if item["month"] is None else int(item["month"])))


def _best_unit_hint(label: str, canonical: list[dict[str, str]], match: Mapping[str, Any] | None = None) -> str:
    match = match or _match_label(label, canonical)
    if match["suggested_code"]:
        for row in canonical:
            if row["row_code"] == match["suggested_code"]:
                return row["unit"]
    if re.search(r"(率|OCC|百分比|%)", label, re.I) and not _is_unit_price_label(label):
        return "RATIO"
    if re.search(r"(房晚|间夜|人数|数量|人次|天数|间数)", label) and not _is_unit_price_label(label):
        return "COUNT"
    return "MONEY"


def _canonical_for_match(match: Mapping[str, Any], canonical: list[dict[str, str]]) -> dict[str, str] | None:
    code = match.get("suggested_code")
    if not code:
        return None
    return next((row for row in canonical if row["row_code"] == code), None)


def _infer_aggregation(label: str, unit: str) -> str:
    if unit == "RATIO" or _is_average_or_stock_label(label):
        return ""
    return ""


def _is_average_or_stock_label(label: str) -> bool:
    return bool(re.search(r"(平均|均价|单价|ADR|RevPAR|房间数|房间总数|可售房间|总房数)", label, re.I))


def _is_unit_price_label(label: str) -> bool:
    return bool(re.search(r"(平均|均价|单价|ADR|RevPAR)", label, re.I))


def _convert_number(
    value: Decimal,
    unit: str,
    money_unit: str,
    label: str = "",
    issues: list[dict[str, Any]] | None = None,
    sheet_name: str = "",
    cell_ref: Any = None,
) -> dict[str, int | None] | None:
    try:
        decimal = Decimal(value)
    except (InvalidOperation, TypeError, ValueError):
        return None
    if not decimal.is_finite():
        return None
    if unit == "RATIO":
        if abs(decimal) > 1:
            decimal = decimal / Decimal("100")
        return {
            "value_int": 0,
            "ratio_num": int((decimal * Decimal("1000000")).quantize(Decimal("1"), rounding=ROUND_HALF_UP)),
            "ratio_den": 1000000,
        }
    if unit == "COUNT":
        return {
            "value_int": int(decimal.quantize(Decimal("1"), rounding=ROUND_HALF_UP)),
            "ratio_num": None,
            "ratio_den": None,
        }
    scale = Decimal("100")
    if str(money_unit).lower() in {"wan", "wanyuan", "wan_yuan", "万元", "10k_yuan"}:
        if _is_unit_price_label(label):
            if issues is not None:
                issues.append(
                    {
                        "code": "MONEY_UNIT_PRICE_AS_YUAN",
                        "sheet": sheet_name,
                        "cell": cell_ref,
                        "source_label": label,
                        "message": "该金额科目属于单价/ADR/RevPAR，按元采集，未按万元放大。",
                    }
                )
        else:
            scale *= Decimal("10000")
    return {
        "value_int": int((decimal * scale).quantize(Decimal("1"), rounding=ROUND_HALF_UP)),
        "ratio_num": None,
        "ratio_den": None,
    }


def _match_label(label: str, canonical: list[dict[str, str]]) -> dict[str, Any]:
    normalized = _normalize_label(label)
    scored = []
    for row in canonical:
        target = _normalize_label(row["row_label"])
        alias_score = _alias_score(normalized, target)
        score = max(SequenceMatcher(None, normalized, target).ratio(), alias_score)
        if normalized == target:
            score = 1.0
        if _semantic_conflict(normalized, target):
            score = min(score, 0.65)
        scored.append({"row_code": row["row_code"], "row_label": row["row_label"], "score": round(score, 4)})
    scored.sort(key=lambda item: item["score"], reverse=True)
    top = scored[0] if scored else None
    runner_up = scored[1]["score"] if len(scored) > 1 else 0
    suggested = ""
    confidence = float(top["score"]) if top else 0.0
    if top and confidence >= 0.9 and confidence - float(runner_up) >= 0.1:
        suggested = str(top["row_code"])
    return {
        "suggested_code": suggested,
        "confidence": round(confidence, 4),
        "candidates": scored[:3],
    }


def _alias_score(source: str, target: str) -> float:
    source_aliases = _label_aliases(source)
    target_aliases = _label_aliases(target)
    if source_aliases & target_aliases:
        return 0.96
    best = 0.0
    for left in source_aliases or {source}:
        for right in target_aliases or {target}:
            best = max(best, SequenceMatcher(None, left, right).ratio())
    return best


def _label_aliases(label: str) -> set[str]:
    aliases = {label}
    replacements = {
        "房费收入": "客房收入",
        "房务收入": "客房收入",
        "已出租房晚": "已售房晚",
        "售出房晚": "已售房晚",
        "出租房晚": "已售房晚",
        "可供出租房晚": "可售房晚",
        "可出租房晚": "可售房晚",
        "入住率": "出租率",
        "occ": "出租率",
        "利润总额": "税前利润",
        "净利润": "税后利润",
    }
    changed = True
    while changed:
        changed = False
        for old, new in replacements.items():
            for item in list(aliases):
                if old in item:
                    candidate = item.replace(old, new)
                    if candidate not in aliases:
                        aliases.add(candidate)
                        changed = True
    return {_strip_generic_words(item) for item in aliases if item}


def _semantic_conflict(source: str, target: str) -> bool:
    conflict_pairs = [
        ("收入", "成本费用支出税金"),
        ("税前", "税后净利润"),
        ("含税", "不含税"),
        ("已售已出租售出出租", "可售可出租可供"),
    ]
    for left_group, right_group in conflict_pairs:
        source_left = any(token in source for token in _tokens(left_group))
        source_right = any(token in source for token in _tokens(right_group))
        target_left = any(token in target for token in _tokens(left_group))
        target_right = any(token in target for token in _tokens(right_group))
        if (source_left and target_right) or (source_right and target_left):
            return True
    return False


def _tokens(text: str) -> list[str]:
    if text == "成本费用支出税金":
        return ["成本", "费用", "支出", "税金"]
    if text == "税后净利润":
        return ["税后", "净利润"]
    if text == "已售已出租售出出租":
        return ["已售", "已出租", "售出", "出租房晚"]
    if text == "可售可出租可供":
        return ["可售", "可出租", "可供"]
    return [text]


def _flag_duplicate_sources(rows: list[dict[str, Any]], issues: list[dict[str, Any]]) -> None:
    labels: dict[str, list[str]] = {}
    for row in rows:
        labels.setdefault(_normalize_label(row["source_label"]), []).append(row["source_key"])
    for label, keys in labels.items():
        if label and len(keys) > 1:
            issues.append(
                {
                    "code": "DUPLICATE_SUBJECT",
                    "source_keys": keys,
                    "message": "发现重复科目，已分别保留，未自动合并。",
                }
            )


def _normalize_kind(value: str) -> str:
    text = _normalize_label(value)
    for alias, kind in REVERSE_KIND_ALIASES.items():
        if _normalize_label(alias) in text:
            return kind
    upper = str(value or "").strip().upper()
    return upper if upper in KIND_ALIASES else ""


def _parse_year(value: str) -> int | None:
    match = YEAR_RE.search(str(value or ""))
    return int(match.group(1)) if match else None


def _cell_text(cell: Mapping[str, Any] | None) -> str:
    if not cell:
        return ""
    value = cell.get("display_value")
    if value is None:
        value = cell.get("cached_value")
    return _normalize_spaces(str(value)) if value is not None else ""


def _cell_number(cell: Mapping[str, Any] | None) -> Decimal | None:
    if not cell or cell.get("is_error"):
        return None
    value = cell.get("cached_value")
    if isinstance(value, bool) or value is None:
        return None
    text = str(value)
    if isinstance(value, str):
        text = value.strip().replace(",", "").replace("，", "").replace("¥", "").replace("￥", "")
        percent = text.endswith("%")
        text = text.rstrip("%")
    else:
        percent = False
    try:
        number = Decimal(text)
    except (InvalidOperation, TypeError, ValueError):
        return None
    if not number.is_finite():
        return None
    return number / Decimal("100") if percent else number


def _looks_like_subject(text: str) -> bool:
    if len(text) > 80:
        return False
    if _parse_year(text) or MONTH_RE.search(text) or UNSAFE_VALUE_HEADER_RE.search(text):
        return False
    return bool(re.search(r"[\u4e00-\u9fffA-Za-z]", text)) and not VALUE_HEADER_RE.fullmatch(text)


def _normalize_label(value: str) -> str:
    text = str(value or "").lower()
    text = re.sub(r"\s+", "", text)
    text = re.sub(r"[()（）【】\[\]{}·、,，.。:：;；/_\\-]", "", text)
    text = text.replace("％", "%")
    return _strip_generic_words(text)


def _strip_generic_words(value: str) -> str:
    text = value
    for token in ("酒店", "项目", "科目", "本期", "本年", "累计"):
        text = text.replace(token, "")
    return text


def _normalize_spaces(value: str) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip())


def _column_label(number: int) -> str:
    result = ""
    while number:
        number, remainder = divmod(number - 1, 26)
        result = chr(65 + remainder) + result
    return result or "A"
