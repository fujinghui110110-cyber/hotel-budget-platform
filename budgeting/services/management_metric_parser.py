"""Read-only parser for the management metric history workbook.

The management workbook is deliberately treated as a source of analytical
history rather than as a budget template.  Sheet names identify metrics and
the first contiguous ``YYYYMM`` header group identifies the period columns.
No formula is evaluated here: a formula is accepted only when Excel has
stored a cached numeric result in the workbook.
"""

from __future__ import annotations

import hashlib
import re
import unicodedata
import zipfile
from collections.abc import Iterable
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from budgeting.services.metrics import METRICS


_YEAR_MONTH_RE = re.compile(r"^(20\d{2})(0[1-9]|1[0-2])$")
_YEAR_RE = re.compile(r"(20\d{2})\s*年?")
_ERROR_RE = re.compile(r"#(?:REF!|DIV/0!|VALUE!|NAME\?|N/A|NUM!|NULL!)", re.I)
_FORMULA_RE = re.compile(r"^\s*=")
_OOXML_FORMULA_RE = re.compile(rb"<f(?:\s[^>]*)?>(.*?)</f>", re.I | re.S)
_XML_LIMIT = 64 * 1024 * 1024
_ZIP_LIMIT = 50 * 1024 * 1024
_UNZIPPED_LIMIT = 500 * 1024 * 1024

_METRIC_ALIASES: dict[str, tuple[str, ...]] = {
    "revenue_total": ("总收入", "营业收入", "total revenue", "revenue total", "收入"),
    "revenue_fb": ("餐饮部收入合计", "餐饮收入", "food and beverage", "food beverage"),
    "revenue_room": ("客房收入", "房费收入", "room revenue"),
    "occ": ("出租率", "入住率", "occ", "occupancy"),
    "adr": ("平均房价", "adr", "average daily rate"),
    "revpar": ("每可卖房收入", "revpar", "revenue per available room"),
    "profit_npi": ("npi", "净经营收益", "经营净收益"),
    "revenue_banquet": ("宴会收入", "宴会", "banquet revenue", "banquet"),
    "revenue_seasonal": ("季节性产品", "季节性收入", "seasonal product", "seasonal revenue"),
    "profit_gop": ("经营毛利润", "gop", "gross operating profit"),
    "cost_labor": ("人工成本", "人工", "labor cost", "labour cost", "labor"),
    "cost_energy": ("能耗成本", "能耗", "能源成本", "energy cost", "energy"),
}

_METRIC_UNITS = {code: str(METRICS.get(code, {}).get("unit") or "MONEY").upper() for code in _METRIC_ALIASES}
_NON_ADDITIVE = {"occ", "adr", "revpar"}
_FIXED_YUAN = {"adr", "revpar"}
_SUM_METRICS = set(_METRIC_ALIASES) - _NON_ADDITIVE
_PLACEHOLDER_VALUES = {"-", "－", "–", "—"}

_KIND_ALIASES: dict[str, tuple[str, ...]] = {
    "ACTUAL": ("actual", "实际", "实绩", "历史", "去年同期"),
    "FORECAST": ("forecast", "预测", "预计", "预估"),
    "BUDGET": ("budget", "预算"),
}


class _CachedCell:
    __slots__ = ("value",)

    def __init__(self, value: Any):
        self.value = value


class _CachedSheet:
    """Materialize one read-only sheet so repeated cell reads do not rescan XML."""

    __slots__ = ("title", "_rows", "max_row", "max_column")

    def __init__(self, sheet: Any):
        self.title = sheet.title
        self._rows = tuple(tuple(row) for row in sheet.iter_rows(values_only=True))
        self.max_row = len(self._rows)
        self.max_column = max((len(row) for row in self._rows), default=0)

    def cell(self, row: int, column: int) -> _CachedCell:
        if row < 1 or column < 1 or row > self.max_row:
            return _CachedCell(None)
        values = self._rows[row - 1]
        return _CachedCell(values[column - 1] if column <= len(values) else None)


def parse_workbook(
    path: str | Path,
    *,
    money_unit: str = "WAN",
    data_kind: str = "ACTUAL",
) -> dict[str, Any]:
    """Parse management metric history into project-period rows.

    ``money_unit`` describes the source unit of total monetary metrics.  The
    supported values are ``WAN``/``万元`` and ``YUAN``/``元``.  ADR and RevPAR
    are always read as yuan.  Values are strings so callers can pass them to
    a ``DecimalField`` without a float round trip.
    """

    source = Path(path).expanduser()
    result: dict[str, Any] = {
        "valid": False,
        "errors": [],
        "warnings": [],
        "rows": [],
        "projects": [],
        "years": [],
        "sha256": _sha256(source),
        "sheets": [],
    }

    normalized_unit = _normalize_money_unit(money_unit)
    if normalized_unit is None:
        result["errors"].append(
            {"code": "MONEY_UNIT_INVALID", "message": "金额单位只支持 WAN/万元或 YUAN/元。"}
        )
    target_kind = _normalize_kind(data_kind)
    if target_kind is None:
        result["errors"].append(
            {"code": "DATA_KIND_INVALID", "message": "历史指标口径只支持 ACTUAL、FORECAST 或 BUDGET。"}
        )

    if result["errors"] or not source.is_file():
        if not source.is_file():
            result["errors"].append({"code": "SOURCE_MISSING", "message": f"找不到工作簿：{source}"})
        return result

    preflight_errors = _preflight_xlsx(source, warnings=result["warnings"])
    if preflight_errors:
        result["errors"].extend(preflight_errors)
        return result

    formulas = values = None
    try:
        formulas = load_workbook(source, read_only=True, data_only=False, keep_links=False)
        values = load_workbook(source, read_only=True, data_only=True, keep_links=False)
        for raw_formula_sheet in formulas.worksheets:
            formula_sheet = _CachedSheet(raw_formula_sheet)
            value_sheet = _CachedSheet(values[raw_formula_sheet.title])
            _parse_sheet(
                formula_sheet,
                value_sheet,
                target_kind,
                normalized_unit,
                result,
            )
    except (OSError, ValueError, zipfile.BadZipFile) as exc:
        result["errors"].append(
            {"code": "WORKBOOK_READ_FAILED", "message": f"工作簿读取失败：{exc}"}
        )
    finally:
        if formulas is not None:
            formulas.close()
        if values is not None:
            values.close()

    if not result["sheets"]:
        result["errors"].append({"code": "NO_RECOGNIZED_SHEETS", "message": "未识别到管理指标工作表。"})
    result["valid"] = not result["errors"]
    return result


def _parse_sheet(
    formula_sheet: Any,
    value_sheet: Any,
    target_kind: str,
    money_unit: str,
    result: dict[str, Any],
) -> None:
    sheet_name = str(formula_sheet.title)
    metric = _match_metric(sheet_name)
    if metric is None:
        result["warnings"].append(
            {"code": "UNKNOWN_METRIC_SHEET", "sheet": sheet_name, "message": "工作表名称未匹配到管理指标，已跳过。"}
        )
        return

    sheet_kind, kind_cell = _sheet_kind(formula_sheet)
    if sheet_kind is not None and target_kind is not None and sheet_kind != target_kind:
        result["warnings"].append(
            {
                "code": "DATA_KIND_MISMATCH",
                "sheet": sheet_name,
                "source_cell": kind_cell,
                "message": f"工作表口径为 {sheet_kind}，与所选 {target_kind} 不同，已跳过。",
            }
        )
        return
    if sheet_kind is None:
        result["warnings"].append(
            {"code": "DATA_KIND_MISSING", "sheet": sheet_name, "message": "工作表未明确实际/预测口径，按所选口径读取。"}
        )

    header_row, all_headers = _find_month_header_row(formula_sheet)
    if header_row is None or not all_headers:
        result["errors"].append(
            {"code": "MONTH_HEADER_MISSING", "sheet": sheet_name, "message": "未找到 YYYYMM 月份表头。"}
        )
        return

    sheet_year = _sheet_year(sheet_name)
    header_years = {year for _, year, _ in all_headers}
    selected_year = sheet_year if sheet_year in header_years else min(header_years)
    if sheet_year is not None and sheet_year not in header_years:
        result["errors"].append(
            {
                "code": "SHEET_YEAR_MISMATCH",
                "sheet": sheet_name,
                "sheet_year": sheet_year,
                "header_years": sorted(header_years),
                "message": "工作表年度与月份表头年度不一致。",
            }
        )
        return
    if len(header_years) > 1 and sheet_year is None:
        result["errors"].append(
            {
                "code": "HEADER_YEAR_AMBIGUOUS",
                "sheet": sheet_name,
                "header_years": sorted(header_years),
                "message": "同一工作表存在多个年度且名称未明确年度。",
            }
        )
        return

    month_columns = _select_month_group(all_headers, selected_year)
    if not month_columns:
        result["errors"].append(
            {"code": "MONTH_HEADER_INCOMPLETE", "sheet": sheet_name, "message": "未找到目标年度的月份列。"}
        )
        return
    if len(month_columns) < 12:
        result["warnings"].append(
            {
                "code": "MONTH_HEADER_INCOMPLETE",
                "sheet": sheet_name,
                "message": f"目标年度只识别到 {len(month_columns)} 个月份列。",
            }
        )

    annual_column = _find_annual_column(formula_sheet, header_row, month_columns)
    project_column = _find_project_column(formula_sheet, header_row)
    title_label = _source_title(formula_sheet)
    title_metric = _match_metric(title_label or "")
    if title_metric != metric:
        result["warnings"].append(
            {
                "code": "TITLE_MISMATCH",
                "sheet": sheet_name,
                "source_label": title_label,
                "message": "表内标题与工作表指标不一致，已按工作表名称取数。",
            }
        )

    metadata = {
        "name": sheet_name,
        "metric": metric,
        "year": selected_year,
        "data_kind": sheet_kind or target_kind,
        "source_label": title_label or _metric_label(metric),
        "header_row": header_row,
        "project_column": get_column_letter(project_column),
        "month_columns": {month: get_column_letter(column) for month, column in month_columns.items()},
        "annual_column": get_column_letter(annual_column) if annual_column else None,
    }
    result["sheets"].append(metadata)
    if selected_year not in result["years"]:
        result["years"].append(selected_year)

    # The first total row ends the project block.  Rows after it are groups,
    # notes, or other analytical totals and must not become projects.
    project_rows = _project_rows(formula_sheet, header_row, project_column)
    if not project_rows:
        result["warnings"].append(
            {"code": "PROJECT_ROWS_MISSING", "sheet": sheet_name, "message": "工作表未找到项目数据行。"}
        )
        return

    for row_number, project_name in project_rows:
        if project_name not in result["projects"]:
            result["projects"].append(project_name)
        for month, column in month_columns.items():
            source_cell = f"{get_column_letter(column)}{row_number}"
            raw_value = _read_cell(
                formula_sheet,
                value_sheet,
                row_number,
                column,
                source_cell,
                sheet_name,
                result,
            )
            result["rows"].append(
                {
                    "project_name": project_name,
                    "metric": metric,
                    "year": selected_year,
                    "month": month,
                    "value": _convert_value(raw_value, metric, money_unit),
                    "source_sheet": sheet_name,
                    "source_cell": source_cell,
                    "unit": _METRIC_UNITS[metric],
                    "source_label": title_label or _metric_label(metric),
                }
            )

        annual_value = None
        if annual_column is not None:
            source_cell = f"{get_column_letter(annual_column)}{row_number}"
            raw_annual = _read_cell(
                formula_sheet,
                value_sheet,
                row_number,
                annual_column,
                source_cell,
                sheet_name,
                result,
            )
            annual_value = _convert_value(raw_annual, metric, money_unit)
            result["rows"].append(
                {
                    "project_name": project_name,
                    "metric": metric,
                    "year": selected_year,
                    "month": 0,
                    "value": annual_value,
                    "source_sheet": sheet_name,
                    "source_cell": source_cell,
                    "unit": _METRIC_UNITS[metric],
                    "source_label": title_label or _metric_label(metric),
                }
            )
        _check_annual(
            metric,
            project_name,
            selected_year,
            row_number,
            month_columns,
            annual_value,
            formula_sheet,
            value_sheet,
            money_unit,
            sheet_name,
            result,
        )


def _read_cell(
    formula_sheet: Any,
    value_sheet: Any,
    row: int,
    column: int,
    source_cell: str,
    sheet_name: str,
    result: dict[str, Any],
) -> Decimal | None:
    formula_cell = formula_sheet.cell(row=row, column=column)
    value_cell = value_sheet.cell(row=row, column=column)
    formula = formula_cell.value
    cached = value_cell.value

    if isinstance(formula, str) and _FORMULA_RE.match(formula):
        if _ERROR_RE.search(formula):
            _issue(result, "FORMULA_ERROR", sheet_name, source_cell, "公式包含 Excel 错误引用。")
            return None
        if isinstance(cached, str) and _ERROR_RE.search(cached):
            _issue(result, "FORMULA_ERROR", sheet_name, source_cell, "公式缓存为 Excel 错误值。")
            return None
        if _is_placeholder(cached):
            _placeholder_warning(result, sheet_name, source_cell)
            return None
        if cached is None:
            _issue(result, "FORMULA_CACHE_MISSING", sheet_name, source_cell, "公式没有已保存的 Excel 缓存值。")
            return None
        return _numeric_value(cached, result, sheet_name, source_cell)
    if isinstance(formula, str) and formula.startswith("#"):
        _issue(result, "CELL_ERROR", sheet_name, source_cell, "单元格为 Excel 错误值。")
        return None
    if formula is None or formula == "":
        return None
    if _is_placeholder(formula):
        _placeholder_warning(result, sheet_name, source_cell)
        return None
    return _numeric_value(formula, result, sheet_name, source_cell)


def _is_placeholder(value: Any) -> bool:
    return isinstance(value, str) and value.strip() in _PLACEHOLDER_VALUES


def _placeholder_warning(result: dict[str, Any], sheet_name: str, source_cell: str) -> None:
    if any(
        item.get("code") == "PLACEHOLDER_VALUE"
        and item.get("sheet") == sheet_name
        and item.get("source_cell") == source_cell
        for item in result["warnings"]
    ):
        return
    result["warnings"].append(
        {
            "code": "PLACEHOLDER_VALUE",
            "sheet": sheet_name,
            "source_cell": source_cell,
            "message": "指标单元格为 '-' 占位符，按缺失值读取，不补为 0。",
        }
    )


def _numeric_value(
    value: Any,
    result: dict[str, Any],
    sheet_name: str,
    source_cell: str,
) -> Decimal | None:
    if isinstance(value, bool):
        _issue(result, "NON_NUMERIC_VALUE", sheet_name, source_cell, "布尔值不能作为指标数值。")
        return None
    try:
        numeric = Decimal(str(value).strip())
    except (InvalidOperation, ValueError, TypeError, AttributeError):
        _issue(result, "NON_NUMERIC_VALUE", sheet_name, source_cell, "指标单元格不是数值。")
        return None
    if not numeric.is_finite():
        _issue(result, "NON_NUMERIC_VALUE", sheet_name, source_cell, "指标单元格不是有限数值。")
        return None
    return numeric


def _convert_value(value: Decimal | None, metric: str, money_unit: str) -> str | None:
    if value is None:
        return None
    if _METRIC_UNITS[metric] == "RATIO":
        return _decimal_text(value)
    factor = Decimal("1")
    if metric not in _FIXED_YUAN and money_unit == "WAN":
        factor = Decimal("10000")
    return _decimal_text((value * factor).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def _check_annual(
    metric: str,
    project_name: str,
    year: int,
    row_number: int,
    month_columns: dict[int, int],
    annual_value: str | None,
    formula_sheet: Any,
    value_sheet: Any,
    money_unit: str,
    sheet_name: str,
    result: dict[str, Any],
) -> None:
    if metric not in _SUM_METRICS or annual_value is None or len(month_columns) != 12:
        return
    monthly: list[str | None] = []
    for month in range(1, 13):
        column = month_columns.get(month)
        if column is None:
            return
        source_cell = f"{get_column_letter(column)}{row_number}"
        value = _read_cell(formula_sheet, value_sheet, row_number, column, source_cell, sheet_name, result)
        monthly.append(_convert_value(value, metric, money_unit))
    if any(value is None for value in monthly):
        return
    annual_decimal = Decimal(annual_value)
    monthly_sum = sum((Decimal(value) for value in monthly), Decimal("0"))
    difference = abs(annual_decimal - monthly_sum)
    if difference == 0:
        return
    code = "ANNUAL_ROUNDING_MISMATCH" if difference <= Decimal("0.01") else "ANNUAL_SUM_MISMATCH"
    result["warnings"].append(
        {
            "code": code,
            "sheet": sheet_name,
            "project_name": project_name,
            "year": year,
            "annual": annual_value,
            "monthly_sum": _decimal_text(monthly_sum),
            "difference": _decimal_text(difference),
            "message": "年度原值与月度合计存在差异，保留年度原值，不做补差。",
        }
    )


def _find_month_header_row(sheet: Any) -> tuple[int | None, list[tuple[int, int, int]]]:
    best_row: int | None = None
    best_headers: list[tuple[int, int, int]] = []
    for row in range(1, min(int(sheet.max_row or 0), 30) + 1):
        headers: list[tuple[int, int, int]] = []
        for column in range(1, int(sheet.max_column or 0) + 1):
            parsed = _parse_year_month(sheet.cell(row=row, column=column).value)
            if parsed is not None:
                year, month = parsed
                headers.append((column, year, month))
        if len(headers) > len(best_headers):
            best_row, best_headers = row, headers
    return best_row, best_headers


def _select_month_group(headers: list[tuple[int, int, int]], year: int) -> dict[int, int]:
    candidates = sorted((column, month) for column, header_year, month in headers if header_year == year)
    if not candidates:
        return {}
    groups: list[list[tuple[int, int]]] = []
    current: list[tuple[int, int]] = []
    for item in candidates:
        if current and item[0] != current[-1][0] + 1:
            groups.append(current)
            current = []
        current.append(item)
    if current:
        groups.append(current)
    complete = [group for group in groups if [month for _, month in group] == list(range(1, 13))]
    chosen = complete[0] if complete else groups[0]
    selected: dict[int, int] = {}
    for column, month in chosen:
        selected.setdefault(month, column)
    return selected


def _find_annual_column(sheet: Any, header_row: int, month_columns: dict[int, int]) -> int | None:
    if not month_columns:
        return None
    last_month = max(month_columns.values())
    for column in range(last_month + 1, min(int(sheet.max_column or 0), last_month + 8) + 1):
        value = sheet.cell(row=header_row, column=column).value
        text = _normalize_text(value)
        if text in {"合计", "累计", "全年", "年度", "年合计", "全年合计", "年度合计", "总计", "fy", "annual"}:
            return column
    return None


def _find_project_column(sheet: Any, header_row: int) -> int:
    for column in range(1, int(sheet.max_column or 0) + 1):
        text = _normalize_text(sheet.cell(row=header_row, column=column).value)
        if text in {"项目", "项目名称", "酒店", "酒店名称", "project", "projectname", "hotel", "hotelname"}:
            return column
    return 1


def _project_rows(sheet: Any, header_row: int, project_column: int) -> list[tuple[int, str]]:
    rows: list[tuple[int, str]] = []
    for row in range(header_row + 1, int(sheet.max_row or 0) + 1):
        value = sheet.cell(row=row, column=project_column).value
        if value is None or not str(value).strip():
            continue
        name = str(value).strip()
        normalized = _normalize_text(name)
        if normalized in {"总计", "合计", "总计不含雄安", "总计不含", "小计"} or normalized.startswith("总计"):
            break
        if normalized in {"项目", "项目名称", "project", "hotel"}:
            continue
        rows.append((row, name))
    return rows


def _sheet_kind(sheet: Any) -> tuple[str | None, str | None]:
    for row in range(1, min(int(sheet.max_row or 0), 4) + 1):
        for column in range(1, min(int(sheet.max_column or 0), 4) + 1):
            value = sheet.cell(row=row, column=column).value
            kind = _normalize_kind(value)
            if kind is not None:
                return kind, f"{get_column_letter(column)}{row}"
    return None, None


def _source_title(sheet: Any) -> str | None:
    for row in range(1, min(int(sheet.max_row or 0), 2) + 1):
        for column in range(1, min(int(sheet.max_column or 0), 4) + 1):
            value = sheet.cell(row=row, column=column).value
            text = str(value).strip() if value is not None else ""
            if not text or _normalize_kind(text) is not None or _parse_year_month(value) is not None:
                continue
            if _normalize_text(text) in {"项目", "项目名称", "project", "hotel"}:
                continue
            return text
    return None


def _match_metric(value: Any) -> str | None:
    text = _metric_text(value)
    if not text:
        return None
    aliases: list[tuple[int, str, str]] = []
    for metric, values in _METRIC_ALIASES.items():
        for alias in values:
            normalized = _metric_text(alias)
            aliases.append((len(normalized), metric, normalized))
    for _, metric, alias in sorted(aliases, reverse=True):
        if text == alias:
            return metric
    for _, metric, alias in sorted(aliases, reverse=True):
        if alias and alias in text:
            return metric
    return None


def _metric_label(metric: str) -> str:
    return str(METRICS.get(metric, {}).get("label") or metric)


def _metric_text(value: Any) -> str:
    text = _normalize_text(value)
    text = _YEAR_RE.sub("", text)
    return text


def _normalize_text(value: Any) -> str:
    if value is None:
        return ""
    text = unicodedata.normalize("NFKC", str(value)).strip().lower()
    return re.sub(r"[\s\u3000()（）\[\]【】_\-—–:：/\\]+", "", text)


def _normalize_kind(value: Any) -> str | None:
    text = _normalize_text(value)
    if not text:
        return None
    for kind, aliases in _KIND_ALIASES.items():
        for alias in aliases:
            normalized = _normalize_text(alias)
            if text == normalized or normalized in text:
                return kind
    return None


def _normalize_money_unit(value: Any) -> str | None:
    text = _normalize_text(value)
    if text in {"wan", "万元", "w"}:
        return "WAN"
    if text in {"yuan", "元", "人民币", "rmb"}:
        return "YUAN"
    return None


def _parse_year_month(value: Any) -> tuple[int, int] | None:
    if value is None or isinstance(value, bool):
        return None
    text = str(value).strip()
    match = _YEAR_MONTH_RE.fullmatch(text)
    if not match:
        return None
    return int(match.group(1)), int(match.group(2))


def _sheet_year(sheet_name: str) -> int | None:
    match = _YEAR_RE.search(unicodedata.normalize("NFKC", sheet_name))
    return int(match.group(1)) if match else None


def _decimal_text(value: Decimal) -> str:
    normalized = value.normalize()
    if normalized == 0:
        return "0"
    return format(normalized, "f")


def _sha256(path: Path) -> str:
    if not path.is_file():
        return ""
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _preflight_xlsx(
    path: Path,
    *,
    warnings: list[dict[str, str]] | None = None,
) -> list[dict[str, str]]:
    errors: list[dict[str, str]] = []
    if path.suffix.lower() != ".xlsx":
        return [{"code": "EXTENSION", "message": "只允许上传 .xlsx 文件。"}]
    if path.stat().st_size > _ZIP_LIMIT:
        return [{"code": "ZIP_SIZE", "message": "工作簿压缩大小超过 50 MiB。"}]
    try:
        package = zipfile.ZipFile(path, "r")
    except (OSError, zipfile.BadZipFile):
        return [{"code": "BAD_ZIP", "message": "文件不是有效的 XLSX/OOXML 压缩包。"}]

    seen: set[str] = set()
    expanded = 0
    has_external_link_parts = False
    has_external_formula = False
    try:
        names = {item.filename for item in package.infolist()}
        if not {"[Content_Types].xml", "xl/workbook.xml"}.issubset(names):
            errors.append({"code": "OOXML_STRUCTURE", "message": "缺少 XLSX 必要结构。"})
        for info in package.infolist():
            normalized = info.filename.replace("\\", "/")
            while normalized.startswith("./"):
                normalized = normalized[2:]
            if normalized.startswith("/") or normalized == ".." or normalized.startswith("../") or "/../" in normalized:
                _append_unique(errors, "ZIP_TRAVERSAL", f"压缩包路径不安全：{info.filename}")
            if normalized.lower() in seen:
                _append_unique(errors, "ZIP_DUPLICATE_ENTRY", f"压缩包包含重复条目：{info.filename}")
            seen.add(normalized.lower())
            expanded += int(info.file_size or 0)
            lower = normalized.lower()
            if lower.startswith("xl/externallinks/"):
                has_external_link_parts = True
            if lower.startswith(("xl/embeddings/", "xl/connections", "xl/activex/", "xl/ctrlprops/")) or lower.endswith(("vbaproject.bin", ".bin")):
                _append_unique(errors, "BLOCKED_PART", f"工作簿包含不允许导入的 OOXML 部件：{info.filename}")
            if lower.endswith((".xml", ".rels")) and info.file_size <= _XML_LIMIT:
                raw = package.read(info)
                if b"<!doctype" in raw.lower() or b"<!entity" in raw.lower():
                    _append_unique(errors, "XML_ENTITY", f"XML 部件包含不允许的实体声明：{info.filename}")
                if lower.startswith("xl/worksheets/") and lower.endswith(".xml"):
                    has_external_formula = has_external_formula or _contains_external_formula(raw)
        if expanded > _UNZIPPED_LIMIT:
            _append_unique(errors, "UNZIPPED_SIZE", "工作簿解压大小超过 500 MiB。")
        if has_external_formula:
            _append_unique(errors, "EXTERNAL_LINK", "工作表公式包含外部工作簿引用，无法安全导入。")
        elif has_external_link_parts and warnings is not None:
            _append_unique(
                warnings,
                "EXTERNAL_LINK_METADATA",
                "工作簿仍含外部链接元数据，但工作表公式未发现外部引用；已按缓存数据读取，不访问外部链接。",
            )
    except (KeyError, OSError, RuntimeError, zipfile.BadZipFile) as exc:
        _append_unique(errors, "ZIP_READ_ERROR", f"读取 XLSX 压缩包失败：{exc}")
    finally:
        package.close()
    return errors


def _contains_external_formula(raw_xml: bytes) -> bool:
    """Return whether worksheet XML contains an external-workbook formula.

    External references are represented by square brackets in OOXML formula
    text, for example ``'[history.xlsx]Sheet1'!A1``.  Restricting the check
    to ``<f>`` elements avoids treating bracketed labels or cell values as
    external references.
    """

    return any(b"[" in formula for formula in _OOXML_FORMULA_RE.findall(raw_xml))


def _append_unique(errors: list[dict[str, str]], code: str, message: str) -> None:
    if any(item.get("code") == code for item in errors):
        return
    errors.append({"code": code, "message": message})


def _issue(result: dict[str, Any], code: str, sheet: str, source_cell: str, message: str) -> None:
    if any(item.get("code") == code and item.get("sheet") == sheet and item.get("source_cell") == source_cell for item in result["errors"]):
        return
    result["errors"].append(
        {"code": code, "sheet": sheet, "source_cell": source_cell, "message": message}
    )
