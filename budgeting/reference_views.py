from django.http import Http404
from django.shortcuts import render
from openpyxl.utils.cell import get_column_letter

from budgeting.views import role_required
from budgeting.services.workbook_reference import list_workbooks, load_sheet


@role_required("ADMIN")
def reference_catalog(request):
    workbooks = list_workbooks()
    return render(request, "budgeting/reference_catalog.html", {
        "workbooks": workbooks,
        "total_sheets": sum(book["sheet_count"] for book in workbooks),
    })


@role_required("ADMIN")
def reference_sheet(request, workbook_id, sheet_id):
    workbook = next((book for book in list_workbooks() if book["id"] == workbook_id), None)
    if workbook is None:
        raise Http404
    try:
        start_row = max(1, int(request.GET.get("row", 1)))
        start_col = max(1, int(request.GET.get("col", 1)))
        sheet = load_sheet(workbook_id, sheet_id, start_row=start_row, max_rows=60,
                           start_col=start_col, max_cols=36)
        if not request.GET.get("row") and not request.GET.get("col"):
            candidates = {}
            for cell in sheet["cells"]:
                if str(cell.get("cached_value")) in {f"{month:02d}" for month in range(1, 13)}:
                    candidates.setdefault(cell["row"], set()).add(str(cell.get("cached_value")))
            header = next((row for row, months in sorted(candidates.items()) if len(months) >= 10), None)
            if header:
                label_columns = [cell["column"] for cell in sheet["cells"] if cell["row"] == header
                                 and cell.get("cached_value") in ("项目", "科目", "项目名称", "费用项目")]
                start_row, start_col = header, min(label_columns, default=1)
                sheet = load_sheet(workbook_id, sheet_id, start_row=start_row, max_rows=60,
                                   start_col=start_col, max_cols=36)
    except (KeyError, ValueError, FileNotFoundError):
        raise Http404
    summary = sheet["sheet"]
    end_row = min(start_row + 59, summary["row_count"])
    end_col = min(start_col + 35, summary["col_count"])
    by_position = {(cell["row"], cell["column"]): cell for cell in sheet["cells"]}
    for cell in by_position.values():
        if cell.get("cache_status") == "missing":
            cell["display_value"] = "缺少缓存"
    columns = list(range(start_col, end_col + 1))
    rows = [{"number": row, "cells": [by_position.get((row, col), {
        "coordinate": f"{get_column_letter(col)}{row}", "display_value": "", "blank": True,
    }) for col in columns]} for row in range(start_row, end_row + 1)]
    return render(request, "budgeting/reference_sheet.html", {
        "workbook": workbook, "sheet": sheet,
        "rows": rows, "columns": [get_column_letter(col) for col in columns],
        "start_row": start_row, "end_row": end_row, "start_col": start_col,
        "end_col": end_col, "previous_row": max(1, start_row - 60),
        "previous_col": max(1, start_col - 36), "next_row": end_row + 1,
        "next_col": end_col + 1,
    })
