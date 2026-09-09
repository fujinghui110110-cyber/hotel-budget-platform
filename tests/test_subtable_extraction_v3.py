from pathlib import Path
from tempfile import TemporaryDirectory

from django.db import transaction
from django.test import TestCase
from openpyxl import Workbook

from budgeting.excel.extract import extract_sub_table_values
from budgeting.models import BudgetCycle, NormalizedValue, Project, UploadVersion


class _Rollback(Exception):
    pass


class SubtableExtractionV3Tests(TestCase):
    def setUp(self):
        project = Project.objects.create(code="SUBV3", name="子表测试")
        cycle = BudgetCycle.objects.create(name="2027预算", budget_year=2027)
        self.upload = UploadVersion.objects.create(
            project=project,
            cycle=cycle,
            original_path="test.xlsx",
            sha256="0" * 64,
        )

    def _workbook_path(self, rows):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "经营明细"
        sheet["A1"] = "项目"
        for month in range(1, 13):
            sheet.cell(row=1, column=month + 1, value=f"{month}月")
        for row_number, (label, values) in enumerate(rows, start=2):
            sheet.cell(row=row_number, column=1, value=label)
            for month, value in enumerate(values, start=2):
                sheet.cell(row=row_number, column=month, value=value)
        directory = TemporaryDirectory()
        path = Path(directory.name) / "subtables.xlsx"
        workbook.save(path)
        workbook.close()
        return directory, path

    def test_normal_decimal_text_is_converted_to_cents(self):
        directory, path = self._workbook_path([("客房收入", ["1,234.56"] + [0] * 11)])
        self.addCleanup(directory.cleanup)

        result = extract_sub_table_values(self.upload, path)

        value = NormalizedValue.objects.get(
            upload=self.upload,
            report_code="经营明细",
            row_code="R0002",
            period="01",
        )
        self.assertEqual(value.value_int, 123456)
        self.assertEqual(value.source_cell, "B2")
        self.assertEqual(result, 13)

    def test_source_text_and_errors_are_skipped_without_zero_filling(self):
        directory, path = self._workbook_path([
            ("奖金", [f"{month}月" for month in range(1, 13)]),
            ("人工成本", ["#DIV/0!"] + [0] * 11),
            ("食品成本率", ["12.5%"] + ["0%"] * 11),
        ])
        self.addCleanup(directory.cleanup)

        result = extract_sub_table_values(self.upload, path)

        self.assertEqual(result, 24)
        self.assertFalse(NormalizedValue.objects.filter(upload=self.upload, row_code="R0002").exists())
        self.assertEqual(
            NormalizedValue.objects.filter(upload=self.upload, row_code="R0003").count(),
            11,
        )
        self.assertFalse(
            NormalizedValue.objects.filter(upload=self.upload, row_code="R0003", period="YEAR").exists()
        )
        ratio = NormalizedValue.objects.get(upload=self.upload, row_code="R0004", period="01")
        self.assertEqual((ratio.ratio_num, ratio.ratio_den), (1250, 10000))

    def test_real_recalculated_workbook_extracts_detail_rows(self):
        path = Path(
            "storage/uploads/DEMO01/e2128407-a666-4db0-a627-b48fbd293d25/"
            "original.recalculated.xlsx"
        )
        if not path.exists():
            self.skipTest("本机真实 V3 重算工作簿不存在")

        try:
            with transaction.atomic():
                result = extract_sub_table_values(self.upload, path)
                report_codes = set(
                    NormalizedValue.objects.filter(upload=self.upload).values_list("report_code", flat=True)
                )
                self.assertEqual(result, NormalizedValue.objects.filter(upload=self.upload).count())
                self.assertGreater(result, 0)
                self.assertGreaterEqual(len(report_codes), 50)
                self.assertIn("工资福利费", report_codes)
                self.assertIn("B16月饼亭", report_codes)
                raise _Rollback
        except _Rollback:
            pass
