from datetime import date

from django.test import SimpleTestCase
from openpyxl import Workbook

from budgeting.excel.template_v3 import (
    _replace_year,
    add_supplementary_sheet,
    add_history_sheet,
    close_report_mapping_inputs,
    history_rows,
    input_manifest_hash,
    normalize_blank_template_formulas,
    parameterize_year,
    protect_workbook,
    repair_confirmed_findings,
    recover_dense_row_numeric_inputs,
    rolling_history_columns,
)


class TemplateV3Tests(SimpleTestCase):
    def test_history_window_is_36_months_and_year_relative(self):
        columns_2027 = rolling_history_columns(2027)
        columns_2028 = rolling_history_columns(2028)
        self.assertEqual(len(columns_2027), 36)
        self.assertEqual((columns_2027[0]["period"], columns_2027[-1]["period"]), ("A2024M01", "F2026M12"))
        self.assertEqual((columns_2028[0]["period"], columns_2028[-1]["period"]), ("A2025M01", "F2027M12"))
        self.assertFalse(any(item["kind"] == "BUDGET" for item in columns_2027))

    def test_history_units_and_labels_match_business_meaning(self):
        rows = {(item["report_code"], item["row_code"]): item for item in history_rows()}
        self.assertEqual(rows[("PL_TOTAL_WINE", "R0028")]["unit"], "COUNT")
        self.assertEqual(rows[("PL_TOTAL_WINE", "R0024")]["unit"], "COUNT")
        self.assertEqual(rows[("PL_TOTAL_WINE", "R0045")]["unit"], "MONEY")
        self.assertEqual(rows[("PL_TOTAL_WINE", "R9003")]["unit"], "MONEY")
        self.assertNotIn(("PL_TOTAL_WINE", "R0029"), rows)
        workbook = Workbook()
        workbook.active.title = "原表"
        _, layout_rows, _ = add_history_sheet(workbook, 2027)
        target = next(item for item in layout_rows if item["report_code"] == "PL_TOTAL_WINE" and item["row_code"] == "R0028")
        label = workbook["历史月度输入"][f"A{target['row_number']}"].value
        self.assertIn("酒店损益总表（含名酒）", label)
        self.assertIn("房晚/数量", label)

    def test_year_parameterization_preserves_non_year_numbers(self):
        self.assertEqual(_replace_year("2023实际 / 2025预测 / 2026预算", 2028), "2025实际 / 2027预测 / 2028预算")
        self.assertEqual(_replace_year(2026, 2028), 2028)
        self.assertEqual(_replace_year(199, 2028), 199)
        self.assertEqual(_replace_year(date(2024, 2, 29), 2027), date(2024, 2, 29))

    def test_year_parameterization_uses_delivery_signature_protocol(self):
        workbook = Workbook()
        workbook.active.title = "SYS_META"
        meta = workbook["SYS_META"]
        for row, (key, value) in enumerate(
            (
                ("template_version", "V1"),
                ("budget_year", 2026),
                ("project_code", ""),
                ("rule_version", "R1"),
                ("formula_manifest_hash", ""),
                ("signature_token", ""),
            ),
            start=1,
        ):
            meta.cell(row=row, column=1, value=key)
            meta.cell(row=row, column=2, value=value)

        parameterize_year(workbook, 2027)

        self.assertEqual(meta["A6"].value, "project_signature")
        self.assertEqual(meta["B1"].value, "V3-2027")
        self.assertEqual(meta["B2"].value, 2027)
        self.assertEqual(meta["B4"].value, "R3")

    def test_input_hash_is_deterministic_and_only_inputs_unlock(self):
        first = {"填报": ["B2", "B1"]}
        second = {"填报": ["B2", "B1"]}
        self.assertEqual(input_manifest_hash(first), input_manifest_hash(second))
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "填报"
        sheet["A1"] = "=1+1"
        sheet["B1"] = 3
        protect_workbook(workbook, {"填报": ["B1"]})
        self.assertTrue(sheet["A1"].protection.locked)
        self.assertFalse(sheet["B1"].protection.locked)
        self.assertTrue(sheet.protection.sheet)

    def test_dense_month_row_recovers_year_like_numeric_values(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "客房"
        inputs = {"客房": [f"{column}7" for column in "KLMNOPRSUV"]}
        sheet["Q7"] = 1900
        sheet["T7"] = 2000

        recovered = recover_dense_row_numeric_inputs(workbook, inputs)

        self.assertEqual(recovered, {"客房": ["Q7", "T7"]})

    def test_report_mapping_nonformulas_are_closed_as_inputs(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "管理表"
        sheet["B2"] = "=1+1"
        reports = {
            "PL": {
                "sheet": "管理表",
                "mapping": [
                    {"cell": "A2", "classification": "formula", "formula": "", "shared_attributes": {"t": "shared"}, "allowed_functions": []},
                    {"cell": "B2", "classification": "formula", "formula": "1+1", "shared_attributes": {}, "allowed_functions": []},
                ],
            }
        }

        inputs = close_report_mapping_inputs(workbook, reports)

        self.assertEqual(inputs, {"管理表": ["A2"]})
        self.assertEqual(reports["PL"]["mapping"][0]["classification"], "input")
        self.assertIsNone(reports["PL"]["mapping"][0]["formula"])
        self.assertEqual(reports["PL"]["mapping"][1]["classification"], "formula")

    def test_confirmed_findings_become_formula_fix_or_explicit_input(self):
        workbook = Workbook()
        workbook.active.title = "B12中餐厅"
        for sheet_name in (
            "损益表（含名酒）（拆中智）",
            "损益表（不含名酒）（拆中智）",
            "酒店损益总表（不含名酒）",
            "A1客房收入(新)",
            "A21前台",
            "B1餐厅汇总",
            "B3宴会厅",
            "B14送餐",
            "B餐饮部汇总",
            "C3康体中心",
            "C5车队",
            "E11行政办",
            "E12人力资源",
            "E2信息及通讯",
            "E31市场销售部",
            "E32公关部",
            "E4工程部",
            "E5能耗",
            "F固定支出",
            "G1员工餐厅分摊",
            "工资福利费",
        ):
            workbook.create_sheet(sheet_name)
        workbook["B12中餐厅"]["K98"] = "=IFERROR(#REF!,0)"
        required, repaired = repair_confirmed_findings(workbook)
        self.assertIsNone(workbook["B12中餐厅"]["K98"].value)
        self.assertIn("K98", required["B12中餐厅"])
        for sheet_name in ("损益表（含名酒）（拆中智）", "损益表（不含名酒）（拆中智）"):
            self.assertEqual(workbook[sheet_name]["S124"].value, "=ROUND(S113+S121-SUM(S115:S120,S122)+S123,2)")
        self.assertEqual(workbook["A1客房收入(新)"]["AE38"].value, "=IF(AE69=0,0,AE100/AE69)")
        self.assertEqual(workbook["A1客房收入(新)"]["X43"].value, "=IF(X74=0,0,X105/X74)")
        self.assertEqual(workbook["酒店损益总表（不含名酒）"]["L66"].value, "=L44+L56+L61+L65")
        self.assertEqual(workbook["酒店损益总表（不含名酒）"]["W66"].value, "=W44+W56+W61+W65")
        self.assertIsNone(workbook["E5能耗"]["K25"].value)
        self.assertIn("K25", required["E5能耗"])
        self.assertEqual(
            workbook["B餐饮部汇总"]["K68"].value,
            "=ROUND((B1餐厅汇总!K68+B3宴会厅!K68),2)+'经营补充指标'!D5",
        )
        self.assertEqual(
            workbook["F固定支出"]["V34"].value,
            "=(+工资福利费!M129)+'经营补充指标'!O6",
        )
        self.assertEqual(
            workbook["A21前台"]["K66"].value,
            "=ROUND('A1客房收入(新)'!K99*A21前台!$W$66,2)+'经营补充指标'!D7",
        )
        self.assertEqual(
            workbook["B14送餐"]["Q31"].value,
            "=ROUND(Q25*$W$31,2)+'经营补充指标'!J8",
        )
        self.assertEqual(len(repaired), 81)

    def test_supplementary_wine_inputs_are_separate_from_pnl(self):
        workbook = Workbook()
        workbook.active.title = "原表"
        inputs = add_supplementary_sheet(workbook)
        sheet = workbook["经营补充指标"]
        self.assertEqual(len(inputs), 60)
        self.assertEqual(inputs[:12], [f"{column}3" for column in "DEFGHIJKLMNO"])
        self.assertEqual(inputs[-12:], [f"{column}8" for column in "DEFGHIJKLMNO"])
        self.assertEqual(sheet["C3"].value, "=SUM(D3:O3)")
        self.assertEqual(sheet["C5"].value, "=SUM(D5:O5)")
        self.assertIn("不得重复加总", sheet["A2"].value)

    def test_empty_denominator_uses_explicit_guard_without_iferror(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet["A1"] = "=B1/C1"
        sheet["A2"] = "=IF(C2=0,0,B2/C2)"
        sheet["A3"] = "="
        guarded, blanked = normalize_blank_template_formulas(workbook)
        self.assertEqual(sheet["A1"].value, "=IF(C1=0,0,B1/C1)")
        self.assertEqual(sheet["A2"].value, "=IF(C2=0,0,B2/C2)")
        self.assertIsNone(sheet["A3"].value)
        self.assertNotIn("IFERROR", sheet["A1"].value)
        self.assertEqual((len(guarded), len(blanked)), (1, 1))
