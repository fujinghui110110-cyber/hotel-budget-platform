from datetime import date

from django.test import SimpleTestCase
from openpyxl import Workbook

from budgeting.excel.template_v3 import _replace_year, parameterize_year


class YearWhitelistTests(SimpleTestCase):
    def workbook(self):
        book = Workbook()
        book.active.title = "SYS_META"
        for row, (key, value) in enumerate([
            ("template_version", "V1"), ("budget_year", 2026),
            ("project_code", "项目2026"), ("rule_version", "R1"),
        ], start=1):
            book.active.cell(row, 1, key)
            book.active.cell(row, 2, value)
        return book

    def test_only_registered_headers_and_metadata_change(self):
        book = self.workbook()
        report = book.create_sheet("酒店损益总表（含名酒）")
        report["Y21"] = "2025年预测"
        report["AA21"] = "2026vs2025"
        report["J41"] = 2026
        report["L41"] = "=2026+2025"
        notes = book.create_sheet("自由说明")
        for row, value in enumerate([2026, "2026元", "编号2026", date(2026, 1, 1), "2026年预算"], 1):
            notes.cell(row, 1, value)
        before = [notes.cell(row, 1).value for row in range(1, 6)]

        changed = parameterize_year(book, 2027)

        self.assertEqual(report["Y21"].value, "2026年预测")
        self.assertEqual(report["AA21"].value, "2027vs2026")
        self.assertEqual(report["J41"].value, 2026)
        self.assertEqual(report["L41"].value, "=2026+2025")
        self.assertEqual([notes.cell(row, 1).value for row in range(1, 6)], before)
        self.assertEqual(book["SYS_META"]["B2"].value, 2027)
        self.assertEqual(book["SYS_META"]["B3"].value, "项目2026")
        self.assertEqual(set(changed), {f"{report.title}!Y21", f"{report.title}!AA21"})

    def test_year_labels_do_not_replace_amounts_even_in_title(self):
        book = self.workbook()
        sheet = book.create_sheet("预算编制说明")
        sheet["A1"] = "2026年预算；金额2026元；编号2026"
        parameterize_year(book, 2027)
        self.assertEqual(sheet["A1"].value, "2027年预算；金额2026元；编号2026")

    def test_year_shift_is_simultaneous_in_both_directions(self):
        self.assertEqual(_replace_year("2026预算vs2025预测vs2024实际", 2025), "2025预算vs2024预测vs2023实际")
        self.assertEqual(_replace_year("2026vs2024说明", 2028), "2028vs2026说明")

    def test_registered_numeric_timeline_is_shifted(self):
        book = self.workbook()
        sheet = book.create_sheet("五年铺排")
        sheet["J21"] = 2023
        sheet["P21"] = 2026
        sheet["J22"] = 2023
        parameterize_year(book, 2028)
        self.assertEqual(sheet["J21"].value, 2025)
        self.assertEqual(sheet["P21"].value, 2028)
        self.assertEqual(sheet["J22"].value, 2023)
