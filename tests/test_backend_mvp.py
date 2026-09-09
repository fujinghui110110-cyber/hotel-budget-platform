import hashlib
import json
import threading
import tempfile
import zipfile
from io import StringIO
from pathlib import Path
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.contrib.messages import get_messages
from django.core.management import call_command
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import close_old_connections
from django.test import Client, TestCase, TransactionTestCase, override_settings
from django.urls import reverse
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from budgeting.excel.extract import extract_report_values
from budgeting.forms import UploadForm
from budgeting.management.commands.budget_worker import Command
from budgeting.models import (
    AdjustmentBatch,
    AdjustmentLine,
    BudgetCycle,
    FreezeSnapshot,
    NormalizedValue,
    ProcessingJob,
    Project,
    ProjectCycle,
    SnapshotArtifact,
    TemplateVersion,
    UploadVersion,
    ValidationIssue,
    ValidationRun,
)
from budgeting.services.files import sha256_file
from budgeting.services.template_delivery import signed_template_copy
from budgeting.services.workflow import (
    InfrastructureProcessingError,
    approve_upload,
    create_adjustment_batch,
    freeze_cycle,
    issue_adjustment,
    _claim_inline_upload_job,
    process_upload_now,
    reopen_cycle,
    submit_upload,
)


class BackendMvpTests(TestCase):
    def setUp(self):
        self.User = get_user_model()
        self.project = Project.objects.create(code="P001", name="示例项目")
        self.other = Project.objects.create(code="P002", name="其他项目")
        self.cycle = BudgetCycle.objects.create(
            name="2026预算", budget_year=2026, status=BudgetCycle.OPEN
        )
        self.template = TemplateVersion.objects.create(
            version="V1",
            budget_year=2026,
            file_path="template.xlsx",
            manifest_path="manifest.json",
            formula_manifest_hash="0" * 64,
        )
        self.admin = self.User.objects.create_user(
            "admin", password="x", role="ADMIN", is_staff=True, is_superuser=True
        )
        self.user = self.User.objects.create_user(
            "p001", password="x", role="PROJECT", project=self.project
        )

    def test_canonical_route_names_resolve(self):
        upload_id = "00000000-0000-0000-0000-000000000001"
        snapshot_id = "00000000-0000-0000-0000-000000000002"
        expected = {
            "project_dashboard": [],
            "project_template_download": [],
            "project_upload_new": [],
            "project_upload_detail": [upload_id],
            "project_upload_submit": [upload_id],
            "project_adjustments": [],
            "project_history": [],
            "upload_status_api": [upload_id],
            "management_cycles": [],
            "management_projects": [],
            "management_report": ["PL_TOTAL_WINE"],
            "management_report_drilldown": ["PL_TOTAL_WINE"],
            "management_upload_approve": [upload_id],
            "management_adjustments": [],
            "management_adjustment_issue": [1],
            "management_freeze": [],
            "management_snapshot_download": [snapshot_id],
            "management_audit": [],
            "healthz": [],
        }
        for name, args in expected.items():
            self.assertTrue(reverse(name, args=args))

    @override_settings(DEBUG=True)
    def test_init_demo_is_idempotent_with_existing_project_account(self):
        stdout = StringIO()
        call_command("init_demo", stdout=stdout)
        call_command("init_demo", stdout=stdout)
        self.assertEqual(self.User.objects.filter(project__code="P001").count(), 1)
        self.assertTrue(self.User.objects.get(username="admin").is_superuser)
        self.assertNotIn("admin123", stdout.getvalue())
        self.assertNotIn("project123", stdout.getvalue())

    def test_worker_retries_infrastructure_error_once(self):
        upload = UploadVersion.objects.create(
            project=self.project,
            cycle=self.cycle,
            template=self.template,
            status=UploadVersion.RECEIVED,
            original_path="x.xlsx",
            sha256="0" * 64,
        )
        ProcessingJob.objects.create(upload=upload, idempotency_key="upload:test")
        with patch(
            "budgeting.management.commands.budget_worker.process_upload",
            side_effect=InfrastructureProcessingError("lo down"),
        ):
            job = Command()._claim_job(lease_seconds=30)
            Command()._run_job(job)
        job.refresh_from_db()
        upload.refresh_from_db()
        self.assertEqual(job.status, ProcessingJob.Status.QUEUED)
        self.assertEqual(upload.status, UploadVersion.Status.PROCESSING)

    def test_inline_upload_does_not_process_job_already_claimed_by_worker(self):
        upload = UploadVersion.objects.create(
            project=self.project,
            cycle=self.cycle,
            template=self.template,
            status=UploadVersion.RECEIVED,
            original_path="x.xlsx",
            sha256="0" * 64,
        )
        ProcessingJob.objects.create(upload=upload, idempotency_key="upload:inline-worker")
        claimed = Command()._claim_job(lease_seconds=30)
        self.assertEqual(claimed.status, ProcessingJob.Status.RUNNING)

        with patch("budgeting.services.workflow.process_upload") as process:
            observed = process_upload_now(upload)

        self.assertEqual(observed.status, ProcessingJob.Status.RUNNING)
        process.assert_not_called()

    def test_inline_upload_claim_is_idempotent_after_processing(self):
        upload = UploadVersion.objects.create(
            project=self.project,
            cycle=self.cycle,
            template=self.template,
            status=UploadVersion.RECEIVED,
            original_path="x.xlsx",
            sha256="0" * 64,
        )
        def mark_validated(item):
            UploadVersion.objects.filter(pk=item.pk).update(status=UploadVersion.VALIDATED)

        with patch("budgeting.services.workflow.process_upload", side_effect=mark_validated) as process:
            first = process_upload_now(upload)
            second = process_upload_now(upload)

        self.assertEqual(first.status, ProcessingJob.Status.DONE)
        self.assertEqual(second.status, ProcessingJob.Status.DONE)
        process.assert_called_once()
        job = ProcessingJob.objects.get(upload=upload)
        self.assertEqual(job.attempts, 1)

    def test_worker_ignores_unclaimed_job(self):
        upload = UploadVersion.objects.create(
            project=self.project,
            cycle=self.cycle,
            template=self.template,
            status=UploadVersion.RECEIVED,
            original_path="x.xlsx",
            sha256="0" * 64,
        )
        job = ProcessingJob.objects.create(upload=upload, idempotency_key="upload:unclaimed")

        with patch("budgeting.management.commands.budget_worker.process_upload") as process:
            Command()._run_job(job)

        process.assert_not_called()
        job.refresh_from_db()
        self.assertEqual(job.status, ProcessingJob.Status.QUEUED)

    def test_upload_view_reports_running_job_without_completion_message(self):
        upload = UploadVersion.objects.create(
            project=self.project,
            cycle=self.cycle,
            template=self.template,
            status=UploadVersion.RECEIVED,
            original_path="x.xlsx",
            sha256="0" * 64,
        )
        job = ProcessingJob.objects.create(
            upload=upload,
            idempotency_key="upload:view-running",
            status=ProcessingJob.Status.RUNNING,
        )
        client = Client()
        client.force_login(self.user)

        with (
            patch("budgeting.views.save_upload", return_value=upload),
            patch("budgeting.views.process_upload_now", return_value=job),
        ):
            response = client.post(
                reverse("project_upload_new"),
                {"file": SimpleUploadedFile("budget.xlsx", b"placeholder")},
            )

        self.assertRedirects(
            response,
            reverse("project_upload_detail", args=[upload.pk]),
            fetch_redirect_response=False,
        )
        message_text = " ".join(str(message) for message in get_messages(response.wsgi_request))
        self.assertIn("正在运行", message_text)
        self.assertNotIn("完成校验", message_text)

    def test_signed_template_copy_injects_project_cycle_metadata_without_changing_source(
        self,
    ):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            wb = Workbook()
            wb.active.title = "Sheet1"
            meta = wb.create_sheet("SYS_META")
            meta.sheet_state = "hidden"
            meta["A1"] = "template_version"
            meta["B1"] = "old"
            source = root / "template.xlsx"
            wb.save(source)
            (root / "manifest.json").write_text(json.dumps({"template_version": "V1"}))
            with override_settings(BASE_DIR=root, BUDGET_STORAGE_ROOT=root / "storage"):
                out = signed_template_copy(self.template, self.project, self.cycle)
                copied = load_workbook(out, data_only=True)
                original = load_workbook(source, data_only=True)
        self.assertEqual(original["SYS_META"]["B1"].value, "old")
        self.assertEqual(copied["SYS_META"]["B1"].value, "V1")
        self.assertEqual(copied["SYS_META"]["B3"].value, "P001")
        self.assertTrue(copied["SYS_META"]["B6"].value)

    def test_adjustment_confirmed_after_approved_reupload(self):
        upload = self._approved_upload(self.project, 10000)
        ProjectCycle.objects.create(
            project=self.project, cycle=self.cycle, current_upload=upload
        )
        batch = create_adjustment_batch(
            self.cycle, "PL_TOTAL_WINE", "ROOM", "01", 100, "test", self.admin
        )
        issue_adjustment(batch, self.admin)
        new_upload = UploadVersion.objects.create(
            project=self.project,
            cycle=self.cycle,
            template=self.template,
            status=UploadVersion.VALIDATED,
            original_path="new.xlsx",
            sha256="2" * 64,
        )
        NormalizedValue.objects.create(
            upload=new_upload,
            report_code="PL_TOTAL_WINE",
            row_code="ROOM",
            period="01",
            unit=NormalizedValue.MONEY,
            value_int=10100,
            source_sheet="s",
            source_cell="A1",
        )
        submit_upload(new_upload, self.user)
        approve_upload(new_upload, self.admin)
        line = AdjustmentLine.objects.get(batch=batch, project=self.project)
        batch.refresh_from_db()
        self.assertEqual(line.difference_cents, 0)
        self.assertEqual(line.status, AdjustmentLine.CONFIRMED)
        self.assertEqual(batch.status, AdjustmentBatch.COMPLETED)

    def test_freeze_creates_complete_snapshot_zip(self):
        with tempfile.TemporaryDirectory() as tmp:
            storage = Path(tmp)
            upload = self._approved_upload(
                self.project, 10000, original_path="uploads/P001/original.xlsx"
            )
            ProjectCycle.objects.create(
                project=self.project, cycle=self.cycle, current_upload=upload
            )
            Project.objects.filter(pk=self.other.pk).update(is_active=False)
            source = storage / upload.original_path
            source.parent.mkdir(parents=True)
            workbook = Workbook()
            workbook.save(source)
            upload.sha256 = sha256_file(source)
            upload.save(update_fields=["sha256"])
            with override_settings(BUDGET_STORAGE_ROOT=storage):
                snapshot = freeze_cycle(self.cycle, self.admin)
                zip_path = storage / f"{snapshot.id}.zip"
                manifest_path = storage / snapshot.manifest_path
                self.assertTrue(zip_path.exists())
                self.assertTrue(manifest_path.exists())
        snapshot.refresh_from_db()
        self.cycle.refresh_from_db()
        self.assertEqual(snapshot.status, FreezeSnapshot.COMPLETE)
        self.assertEqual(self.cycle.status, BudgetCycle.FROZEN)

    def test_upload_form_accepts_valid_extension_when_browser_mime_is_generic(self):
        form = UploadForm(
            files={
                "file": SimpleUploadedFile(
                    "budget.xlsx", b"PK\x03\x04", content_type="text/plain"
                )
            }
        )
        self.assertTrue(form.is_valid(), form.errors)

    def test_extraction_rebuilds_annual_money_and_occ_from_months(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            workbook_path = root / "annual.xlsx"
            manifest_path = root / "manifest.json"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "酒店损益总表（含名酒）"
            mapping = []
            for month, column_number in enumerate(range(12, 24), start=1):
                column = get_column_letter(column_number)
                sheet[f"{column}41"] = None if month == 12 else 0.01
                sheet[f"{column}24"] = 1 if month == 1 else 100
                sheet[f"{column}28"] = 1 if month == 1 else 0
                sheet[f"{column}29"] = 1 if month == 1 else 0
                sheet[f"{column}30"] = 178
                sheet[f"{column}50"] = 10
                sheet[f"{column}51"] = 2
                sheet[f"{column}52"] = 5
                sheet[f"{column}60"] = 999
                mapping.extend(
                    [
                        {
                            "row_code": "REV",
                            "row_label": "收入",
                            "period": f"{month:02d}",
                            "cell": f"{column}41",
                            "unit": "MONEY",
                        },
                        {
                            "row_code": "OCC",
                            "row_label": "OCC",
                            "period": f"{month:02d}",
                            "cell": f"{column}29",
                            "unit": "RATIO",
                            "aggregation": "RATIO",
                        },
                        {
                            "row_code": "ROOMS",
                            "row_label": "房间数",
                            "period": f"{month:02d}",
                            "cell": f"{column}30",
                            "unit": "COUNT",
                            "aggregation": "AVERAGE",
                        },
                        {
                            "row_code": "ADR",
                            "row_label": "平均房价 ADR",
                            "period": f"{month:02d}",
                            "cell": f"{column}52",
                            "unit": "MONEY",
                            "aggregation": "DERIVED",
                            "numerator_cell": f"{column}50",
                            "denominator_cell": f"{column}51",
                        },
                        {
                            "row_code": "TECHNICAL",
                            "row_label": "技术行",
                            "period": f"{month:02d}",
                            "cell": f"{column}60",
                            "unit": "MONEY",
                            "aggregation": "EXCLUDE",
                        },
                    ]
                )
            sheet["J41"] = 0.12
            sheet["J24"] = 12
            sheet["J28"] = 1
            sheet["J29"] = 1 / 12
            sheet["J30"] = 178
            sheet["J50"] = 120
            sheet["J51"] = 24
            sheet["J52"] = 5
            sheet["J60"] = 11988
            mapping.extend(
                [
                    {
                        "row_code": "REV",
                        "row_label": "收入",
                        "period": "FY",
                        "cell": "J41",
                        "unit": "MONEY",
                    },
                    {
                        "row_code": "OCC",
                        "row_label": "OCC",
                        "period": "FY",
                        "cell": "J29",
                        "unit": "RATIO",
                        "aggregation": "RATIO",
                    },
                    {
                        "row_code": "ROOMS",
                        "row_label": "房间数",
                        "period": "FY",
                        "cell": "J30",
                        "unit": "COUNT",
                        "aggregation": "AVERAGE",
                    },
                    {
                        "row_code": "ADR",
                        "row_label": "平均房价 ADR",
                        "period": "FY",
                        "cell": "J52",
                        "unit": "MONEY",
                        "aggregation": "DERIVED",
                        "numerator_cell": "J50",
                        "denominator_cell": "J51",
                    },
                    {
                        "row_code": "TECHNICAL",
                        "row_label": "技术行",
                        "period": "FY",
                        "cell": "J60",
                        "unit": "MONEY",
                        "aggregation": "EXCLUDE",
                    },
                ]
            )
            workbook.save(workbook_path)
            manifest_path.write_text(
                json.dumps(
                    {
                        "reports": {
                            "PL_TOTAL_WINE": {"sheet": sheet.title, "mapping": mapping}
                        }
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            template = TemplateVersion.objects.create(
                version="ANNUAL",
                budget_year=2026,
                file_path=workbook_path.name,
                manifest_path=manifest_path.name,
                formula_manifest_hash="a" * 64,
            )
            upload = UploadVersion.objects.create(
                project=self.project,
                cycle=self.cycle,
                template=template,
                status=UploadVersion.PROCESSING,
                original_path=workbook_path.name,
                sha256="a" * 64,
            )
            run = ValidationRun.objects.create(upload=upload, rule_version="R1")

            with override_settings(BASE_DIR=root, BUDGET_STORAGE_ROOT=root):
                extract_report_values(upload, workbook_path, validation_run=run)

            annual_money = NormalizedValue.objects.get(
                upload=upload,
                report_code="PL_TOTAL_WINE",
                row_code="REV",
                period="YEAR",
            )
            annual_occ = NormalizedValue.objects.get(
                upload=upload,
                report_code="PL_TOTAL_WINE",
                row_code="OCC",
                period="YEAR",
            )
            annual_rooms = NormalizedValue.objects.get(
                upload=upload,
                report_code="PL_TOTAL_WINE",
                row_code="ROOMS",
                period="YEAR",
            )
            annual_adr = NormalizedValue.objects.get(
                upload=upload,
                report_code="PL_TOTAL_WINE",
                row_code="ADR",
                period="YEAR",
            )
            self.assertEqual(annual_money.value_int, 11)
            self.assertEqual((annual_occ.ratio_num, annual_occ.ratio_den), (1, 1101))
            self.assertEqual(annual_rooms.value_int, 178)
            self.assertEqual(
                (annual_adr.value_int, annual_adr.ratio_num, annual_adr.ratio_den),
                (500, 12000, 24),
            )
            self.assertFalse(
                NormalizedValue.objects.filter(upload=upload, row_code="TECHNICAL").exists()
            )
            self.assertEqual(
                ValidationIssue.objects.filter(
                    run=run, code="ANNUAL_RECONCILIATION"
                ).count(),
                2,
            )

    def test_project_adjustment_view_prefetches_only_the_logged_in_project_line(self):
        for project, cents in ((self.project, 10000), (self.other, 20000)):
            upload = self._approved_upload(project, cents)
            ProjectCycle.objects.create(
                project=project, cycle=self.cycle, current_upload=upload
            )
        create_adjustment_batch(
            self.cycle, "PL_TOTAL_WINE", "ROOM", "01", 300, "test", self.admin
        )

        client = Client()
        client.force_login(self.user)
        response = client.get(reverse("project_adjustments"))

        self.assertEqual(response.status_code, 200)
        batch = list(response.context["adjustment_batches"])[0]
        self.assertEqual(
            list(batch.lines.all().values_list("project_id", flat=True)),
            [self.project.id],
        )

    def test_freeze_manifest_matches_zip_and_keeps_zip_hash_out_of_manifest(self):
        with tempfile.TemporaryDirectory() as tmp:
            storage = Path(tmp)
            upload = self._approved_upload(
                self.project, 10000, original_path="uploads/P001/original.xlsx"
            )
            ProjectCycle.objects.create(
                project=self.project, cycle=self.cycle, current_upload=upload
            )
            Project.objects.filter(pk=self.other.pk).update(is_active=False)
            source = storage / upload.original_path
            source.parent.mkdir(parents=True)
            Workbook().save(source)
            upload.sha256 = sha256_file(source)
            upload.save(update_fields=["sha256"])

            with override_settings(BUDGET_STORAGE_ROOT=storage):
                snapshot = freeze_cycle(self.cycle, self.admin)
                manifest_path = storage / snapshot.manifest_path
                manifest_bytes = manifest_path.read_bytes()
                manifest = json.loads(manifest_bytes)
                zip_path = storage / f"{snapshot.id}.zip"
                with zipfile.ZipFile(zip_path) as archive:
                    self.assertEqual(archive.read("manifest.json"), manifest_bytes)
                self.assertNotIn(f"{snapshot.id}.zip", manifest)
                zip_artifact = SnapshotArtifact.objects.get(
                    snapshot=snapshot, kind="zip"
                )
                self.assertEqual(
                    zip_artifact.sha256,
                    hashlib.sha256(zip_path.read_bytes()).hexdigest(),
                )

    def test_reopen_cycle_copies_baselines_and_opens_only_selected_projects(self):
        with tempfile.TemporaryDirectory() as tmp:
            storage = Path(tmp)
            originals = {}
            for project, cents in ((self.project, 10000), (self.other, 20000)):
                path = f"uploads/{project.code}/original.xlsx"
                upload = self._approved_upload(project, cents, original_path=path)
                originals[project.id] = upload
                ProjectCycle.objects.create(
                    project=project, cycle=self.cycle, current_upload=upload
                )
                target = storage / path
                target.parent.mkdir(parents=True, exist_ok=True)
                Workbook().save(target)
                upload.sha256 = sha256_file(target)
                upload.save(update_fields=["sha256"])

            with override_settings(BUDGET_STORAGE_ROOT=storage):
                snapshot = freeze_cycle(self.cycle, self.admin)
                zip_path = storage / f"{snapshot.id}.zip"
                before = hashlib.sha256(zip_path.read_bytes()).hexdigest()
                reopened = reopen_cycle(
                    self.cycle, self.admin, project_ids=[self.project.id]
                )

                self.assertEqual(reopened.revision_no, 2)
                project_cycles = {
                    pc.project_id: pc
                    for pc in ProjectCycle.objects.filter(
                        cycle=reopened
                    ).select_related("current_upload")
                }
                self.assertTrue(project_cycles[self.project.id].is_open)
                self.assertFalse(project_cycles[self.other.id].is_open)
                for project_id, project_cycle in project_cycles.items():
                    copied = project_cycle.current_upload
                    self.assertNotEqual(copied.id, originals[project_id].id)
                    self.assertEqual(copied.cycle, reopened)
                    self.assertEqual(copied.status, UploadVersion.APPROVED)
                    self.assertEqual(
                        list(
                            copied.normalizedvalue_set.values_list(
                                "value_int", flat=True
                            )
                        ),
                        list(
                            originals[project_id].normalizedvalue_set.values_list(
                                "value_int", flat=True
                            )
                        ),
                    )
                self.assertEqual(
                    hashlib.sha256(zip_path.read_bytes()).hexdigest(), before
                )

                with self.assertRaisesRegex(ValueError, "项目周期已关闭"):
                    from budgeting.services.workflow import save_upload

                    save_upload(
                        self.other,
                        reopened,
                        SimpleUploadedFile("closed.xlsx", b"PK\x03\x04"),
                    )

    def _approved_upload(self, project, cents, original_path="original.xlsx"):
        upload = UploadVersion.objects.create(
            project=project,
            cycle=self.cycle,
            template=self.template,
            status=UploadVersion.APPROVED,
            original_path=original_path,
            sha256="1" * 64,
        )
        NormalizedValue.objects.create(
            upload=upload,
            report_code="PL_TOTAL_WINE",
            row_code="ROOM",
            period="01",
            unit=NormalizedValue.MONEY,
            value_int=cents,
            source_sheet="s",
            source_cell="A1",
        )
        return upload


class ProcessingJobConcurrencyTests(TransactionTestCase):
    reset_sequences = True

    def setUp(self):
        self.project = Project.objects.create(code="P001", name="示例项目")
        self.cycle = BudgetCycle.objects.create(
            name="2026预算", budget_year=2026, status=BudgetCycle.OPEN
        )
        self.template = TemplateVersion.objects.create(
            version="V1",
            budget_year=2026,
            file_path="template.xlsx",
            manifest_path="manifest.json",
            formula_manifest_hash="0" * 64,
        )

    def test_two_sqlite_consumers_claim_one_job(self):
        upload = UploadVersion.objects.create(
            project=self.project,
            cycle=self.cycle,
            template=self.template,
            status=UploadVersion.RECEIVED,
            original_path="x.xlsx",
            sha256="0" * 64,
        )
        ProcessingJob.objects.create(upload=upload, idempotency_key="upload:threaded")
        barrier = threading.Barrier(2)
        claimed = []
        errors = []

        def claim_inline_job():
            close_old_connections()
            try:
                barrier.wait(timeout=10)
                _, did_claim = _claim_inline_upload_job(
                    UploadVersion.objects.get(pk=upload.pk)
                )
                claimed.append(did_claim)
            except BaseException as exc:
                errors.append(exc)
            finally:
                close_old_connections()

        def claim_worker_job():
            close_old_connections()
            try:
                barrier.wait(timeout=10)
                claimed.append(Command()._claim_job(lease_seconds=30) is not None)
            except BaseException as exc:
                errors.append(exc)
            finally:
                close_old_connections()

        threads = [
            threading.Thread(target=claim_inline_job),
            threading.Thread(target=claim_worker_job),
        ]
        for thread in threads:
            thread.start()
        for thread in threads:
            thread.join(timeout=15)

        self.assertFalse(errors, errors)
        self.assertEqual(sorted(claimed), [False, True])
        job = ProcessingJob.objects.get(upload=upload)
        self.assertEqual(job.status, ProcessingJob.Status.RUNNING)
        self.assertEqual(job.attempts, 1)
