import io
import unittest

from openpyxl import Workbook, load_workbook

from budgeting.excel.wine_sources import COST, REVENUE, discover_wine_sources


def _month_header(sheet, row, first_column, suffix=""):
    for month in range(1, 13):
        value = f"{month:02d}{suffix}" if suffix else f"{month:02d}"
        sheet.cell(row=row, column=first_column + month - 1, value=value)


def _source_row(sheet, row, label, first_column, value=1):
    sheet.cell(row=row, column=8, value=label)
    for month in range(12):
        sheet.cell(row=row, column=first_column + month, value=value + month)


class WineSourceDiscoveryTests(unittest.TestCase):
    def _read_only(self, workbook):
        buffer = io.BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        return load_workbook(buffer, read_only=True, data_only=True)

    def test_dedicated_detail_selects_leaf_income_and_cost(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "OOD-其他 (名酒)"
        _month_header(sheet, 21, 11)
        _source_row(sheet, 24, "OOD收入", 11, 9000)  # section total: ignored
        _source_row(sheet, 62, "OOD部门成本", 11, 8000)  # section total: ignored
        _source_row(sheet, 61, "OOD收入-其他", 11, 100)
        _source_row(sheet, 71, "OOD成本-其他", 11, 50)

        read_only = self._read_only(workbook)
        try:
            sources, diagnostics = discover_wine_sources(read_only)
        finally:
            read_only.close()

        self.assertEqual(diagnostics, [])
        self.assertEqual(sources, {
            REVENUE: {"sheet": "OOD-其他 (名酒)", "row": 61, "first_column": 11, "label": "OOD收入-其他"},
            COST: {"sheet": "OOD-其他 (名酒)", "row": 71, "first_column": 11, "label": "OOD成本-其他"},
        })

    def test_inserted_row_does_not_break_label_or_header_discovery(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "OOD-其他 (名酒)"
        _month_header(sheet, 5, 11)
        _source_row(sheet, 10, "OOD收入-其他", 11, 100)
        _source_row(sheet, 15, "OOD成本-其他", 11, 50)
        sheet.insert_rows(9, 3)

        read_only = self._read_only(workbook)
        try:
            sources, diagnostics = discover_wine_sources(read_only)
        finally:
            read_only.close()

        self.assertEqual(diagnostics, [])
        self.assertEqual(sources[REVENUE]["row"], 13)
        self.assertEqual(sources[COST]["row"], 18)
        self.assertEqual(sources[REVENUE]["first_column"], 11)

    def test_blank_identified_months_keep_revenue_and_cost_sources(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "OOD-其他"
        _month_header(sheet, 21, 11)
        # The labels and twelve-month header identify both sources.  All
        # monthly input cells are intentionally blank; source discovery must
        # not turn that into WINE_SOURCE_MISSING.
        sheet["H55"] = "OOD收入-商品零售（名酒）"
        sheet["H69"] = "OOD成本-商品零售（名酒）"

        sources, diagnostics = discover_wine_sources(workbook)

        self.assertEqual(diagnostics, [])
        self.assertEqual(sources[REVENUE]["row"], 55)
        self.assertEqual(sources[COST]["row"], 69)
        self.assertEqual(sources[REVENUE]["first_column"], 11)
        self.assertEqual(sources[COST]["first_column"], 11)

    def test_explicit_generic_revenue_is_allowed_but_generic_cost_is_not_inferred(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "OOD-其他"
        _month_header(sheet, 21, 11)
        sheet["H55"] = "OOD收入-商品零售(名酒)"  # explicit source, but blank months
        _source_row(sheet, 69, "OOD成本-商品零售", 11, 50)

        sources, diagnostics = discover_wine_sources(workbook)

        self.assertEqual(sources[REVENUE]["row"], 55)
        self.assertNotIn(COST, sources)
        self.assertTrue(any(item["code"] == "WINE_SOURCE_MISSING" and item["location"] == COST
                            for item in diagnostics))

    def test_same_tier_duplicate_candidates_block_metric_without_summing(self):
        workbook = Workbook()
        workbook.remove(workbook.active)
        for sheet_name, amount in (("OOD-其他", 100), ("OOD-其他-副本", 200)):
            sheet = workbook.create_sheet(sheet_name)
            _month_header(sheet, 21, 11)
            _source_row(sheet, 55, "OOD收入-商品零售（名酒）", 11, amount)

        sources, diagnostics = discover_wine_sources(workbook)

        self.assertNotIn(REVENUE, sources)
        ambiguous = [item for item in diagnostics if item["code"] == "WINE_SOURCE_AMBIGUOUS"]
        self.assertEqual(len(ambiguous), 1)
        self.assertIn("OOD-其他!H55", ambiguous[0]["location"])
        self.assertIn("OOD-其他-副本!H55", ambiguous[0]["location"])

    def test_supplementary_revenue_is_fallback_when_original_detail_is_absent(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "经营补充指标"
        _month_header(sheet, 2, 4, "月")
        sheet["A3"] = "名酒收入（元）"
        for month in range(12):
            sheet.cell(row=3, column=4 + month, value=10 + month)

        sources, diagnostics = discover_wine_sources(workbook)

        self.assertEqual(sources[REVENUE], {
            "sheet": "经营补充指标", "row": 3, "first_column": 4, "label": "名酒收入（元）",
        })
        self.assertNotIn(COST, sources)
        self.assertFalse(any(item["code"] == "WINE_SOURCE_AMBIGUOUS" for item in diagnostics))

    def test_missing_sources_have_metric_level_diagnostics_only(self):
        workbook = Workbook()
        workbook.active.title = "无关工作表"

        sources, diagnostics = discover_wine_sources(workbook)

        self.assertEqual(sources, {})
        self.assertEqual({item["location"] for item in diagnostics}, {REVENUE, COST})
        self.assertTrue(all(":" not in item["location"] for item in diagnostics))


if __name__ == "__main__":
    unittest.main()
