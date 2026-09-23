import json
import tempfile
from pathlib import Path
from decimal import Decimal
from django.test import TestCase
from budgeting.models import BudgetCycle, NormalizedValue, Project, TemplateVersion, UploadVersion
from budgeting.services.report_context import ReportContext
from budgeting.services.report_query import query_report


class ReportQueryTests(TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        path = Path(self.tmp.name) / 'manifest.json'
        mapping = [dict(row_code='ROOMS',row_label='房晚',unit='COUNT',aggregation='SUM',cell='A1'),
                   dict(row_code='ADR',row_label='ADR',unit='MONEY',aggregation='DERIVED',cell='A2',numerator_cell='A3',denominator_cell='A1'),
                   dict(row_code='REVENUE',row_label='收入',unit='MONEY',aggregation='SUM',cell='A3')]
        path.write_text(json.dumps({'template_version':'test','reports':{'PL_TOTAL_WINE':{'mapping':mapping}}}))
        self.template = TemplateVersion.objects.create(version='test',budget_year=2027,manifest_path=str(path))
        self.cycle = BudgetCycle.objects.create(name='2027 R1',budget_year=2027,template=self.template)
        self.projects = [Project.objects.create(code=f'p{i}',name=f'酒店{i}') for i in range(2)]
        for p, rooms, revenue in zip(self.projects,[10,40],[80000,400000]):
            upload = UploadVersion.objects.create(project=p,cycle=self.cycle,template=self.template,status='APPROVED')
            for code,unit,value,num,den in [('ROOMS','COUNT',rooms,None,None),('REVENUE','MONEY',revenue,None,None),('ADR','MONEY',0,revenue,rooms)]:
                NormalizedValue.objects.create(upload=upload,report_code='PL_TOTAL_WINE',row_code=code,period='YEAR',data_year=2027,data_kind='BUDGET',unit=unit,value_int=value,ratio_num=num,ratio_den=den)
        self.context = ReportContext(2027,self.cycle.pk,'PL_TOTAL_WINE','WORKING',tuple(p.pk for p in self.projects))

    def test_weighted_price_and_query_count(self):
        with self.assertNumQueries(4):
            result = query_report(self.context)
        values = {m.row_code:m.value for m in result.metrics}
        self.assertEqual(values['ADR'],Decimal('96.00'))
        self.assertEqual(values['ROOMS'],Decimal(50))
        self.assertEqual(values['REVENUE'],Decimal('4800.00'))

    def test_no_history_or_other_cycle_leak(self):
        upload = UploadVersion.objects.filter(project=self.projects[0]).get()
        NormalizedValue.objects.create(upload=upload,report_code='PL_TOTAL_WINE',row_code='ROOMS',period='A2025',data_year=2025,data_kind='ACTUAL',unit='COUNT',value_int=99999)
        result = query_report(self.context)
        self.assertEqual(next(m.value for m in result.metrics if m.row_code=='ROOMS'),50)

    def test_rejected_latest_does_not_silently_fallback(self):
        UploadVersion.objects.create(project=self.projects[0],cycle=self.cycle,template=self.template,status='REJECTED')
        result = query_report(self.context)
        self.assertEqual(result.missing_projects[self.projects[0].pk],'REJECTED')
        partial = {m.row_code:m for m in result.metrics}
        self.assertEqual(partial['ROOMS'].value,40)
        self.assertEqual(partial['ADR'].value,100)
        self.assertEqual(partial['REVENUE'].value,4000)
        self.assertIn(self.projects[0].pk,partial['ADR'].missing)
        approved = ReportContext(2027,self.cycle.pk,'PL_TOTAL_WINE','APPROVED',self.context.project_ids)
        self.assertEqual(query_report(approved).coverage['available'],2)

    def test_missing_components_not_zero(self):
        NormalizedValue.objects.filter(row_code='ADR').update(ratio_num=None)
        self.assertIsNone(next(m.value for m in query_report(self.context).metrics if m.row_code=='ADR'))

    def test_rooms_not_currency_scaled(self):
        NormalizedValue.objects.filter(row_code='ROOMS').update(value_int=5000)
        self.assertEqual(next(m.value for m in query_report(self.context).metrics if m.row_code=='ROOMS'),10000)

    def test_year_mismatch_rejected(self):
        with self.assertRaisesMessage(ValueError,'年度与轮次不一致'):
            query_report(ReportContext(2028,self.cycle.pk,'PL_TOTAL_WINE','WORKING',self.context.project_ids))

    def test_table_is_batched_and_export_matches(self):
        from budgeting.services.report_query import query_report_table
        from budgeting.services.report_presenter import render_report, render_drilldown
        from django.test import RequestFactory
        from django.contrib.auth import get_user_model
        from openpyxl import load_workbook
        from io import BytesIO
        with self.assertNumQueries(4):
            table = query_report_table(self.context, include_history=True)
        self.assertEqual(len(table),52)
        request = RequestFactory().get('/', {'download':'xlsx'})
        request.user = get_user_model()(role='ADMIN')
        response = render_report(request,self.cycle,'PL_TOTAL_WINE')
        self.assertEqual(response.status_code,200)
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(workbook.active.cell(5,14).value,96)
        request = RequestFactory().get('/', {'row_code':'ADR','period':'YEAR','cycle':self.cycle.pk})
        request.user = get_user_model()(role='ADMIN')
        response = render_drilldown(request,self.cycle,'PL_TOTAL_WINE')
        self.assertContains(response,'80.00')
        self.assertContains(response,'100.00')
        self.assertContains(response,f'cycle={self.cycle.pk}')

    def test_frozen_set_does_not_follow_later_uploads(self):
        import hashlib
        from django.test import override_settings
        from budgeting.models import FreezeSnapshot, SnapshotArtifact
        from budgeting.services.report_query import freeze_selection_payload
        context = ReportContext(2027,self.cycle.pk,'PL_TOTAL_WINE','APPROVED',self.context.project_ids)
        payload = freeze_selection_payload(context)
        snapshot = FreezeSnapshot.objects.create(cycle=self.cycle,status='COMPLETE',directory='snapshot')
        folder = Path(self.tmp.name)/'snapshot'; folder.mkdir()
        raw = json.dumps(payload).encode(); (folder/'report-selection.json').write_bytes(raw)
        SnapshotArtifact.objects.create(snapshot=snapshot,kind='json',relative_path='snapshot/report-selection.json',sha256=hashlib.sha256(raw).hexdigest())
        UploadVersion.objects.filter(project=self.projects[0]).update(status='SUPERSEDED')
        UploadVersion.objects.create(project=self.projects[0],cycle=self.cycle,template=self.template,status='APPROVED')
        frozen = ReportContext(2027,self.cycle.pk,'PL_TOTAL_WINE','FROZEN',self.context.project_ids,snapshot_id=str(snapshot.pk))
        with override_settings(BUDGET_STORAGE_ROOT=self.tmp.name):
            result = query_report(frozen)
            self.assertEqual(next(m.value for m in result.metrics if m.row_code=='ADR'),96)
            SnapshotArtifact.objects.filter(snapshot=snapshot).update(relative_path='snapshot\\report-selection.json')
            result = query_report(frozen)
            self.assertEqual(next(m.value for m in result.metrics if m.row_code=='ADR'),96)
            (folder/'report-selection.json').write_text('{}')
            with self.assertRaisesMessage(ValueError,'冻结集合文件校验失败'):
                query_report(frozen)

    def test_missing_controlled_manifest_shows_gap_and_blocks_export(self):
        from budgeting.services.report_presenter import render_report
        from django.test import RequestFactory
        from django.contrib.auth import get_user_model
        Path(self.template.manifest_path).unlink()
        request=RequestFactory().get('/')
        request.user=get_user_model()(role='ADMIN')
        response=render_report(request,self.cycle,'PL_TOTAL_WINE')
        self.assertContains(response,'模板或冻结来源文件缺失')
        request=RequestFactory().get('/',{'download':'xlsx'})
        request.user=get_user_model()(role='ADMIN')
        self.assertEqual(render_report(request,self.cycle,'PL_TOTAL_WINE').status_code,400)

    def test_manifest_resolves_original_install_root(self):
        from unittest.mock import patch
        from budgeting.services.metric_definitions import load_definitions
        self.template.manifest_path='manifest.json'
        with patch.dict('os.environ',{'BUDGET_TEMPLATE_ROOT':self.tmp.name}):
            definitions,_=load_definitions(self.template,'PL_TOTAL_WINE')
        self.assertIn('ADR',definitions)

    def test_partial_coverage_is_visible_and_never_fills_from_old_upload(self):
        from budgeting.services.report_presenter import render_report
        from django.test import RequestFactory
        from django.contrib.auth import get_user_model
        UploadVersion.objects.create(project=self.projects[0],cycle=self.cycle,template=self.template,status='REJECTED')
        request=RequestFactory().get('/')
        request.user=get_user_model()(role='ADMIN')
        response=render_report(request,self.cycle,'PL_TOTAL_WINE')
        self.assertContains(response,'已上报合计（非完整年度总数）')
        self.assertContains(response,'项目覆盖 1 / 2')
        self.assertContains(response,'4,000.00')
        self.assertNotContains(response,'4,800.00')
        self.assertContains(response,'（不完整）')

    def test_exact_eight_of_ten_coverage_preserves_missing_projects(self):
        from budgeting.models import BudgetPlan, PlanProject
        from budgeting.services.report_presenter import render_report
        from django.test import RequestFactory
        from django.contrib.auth import get_user_model
        from openpyxl import load_workbook
        from io import BytesIO
        plan = BudgetPlan.objects.create(budget_year=2027)
        self.cycle.plan = plan
        self.cycle.save(update_fields=['plan'])
        projects = self.projects + [Project.objects.create(code=f'p{i}', name=f'酒店{i}') for i in range(2, 10)]
        PlanProject.objects.bulk_create([PlanProject(plan=plan, project=p) for p in projects])
        old_cycle = BudgetCycle.objects.create(name='旧轮次', budget_year=2027, revision_no=9, template=self.template, plan=plan)
        for i, project in enumerate(projects[2:], start=2):
            upload = UploadVersion.objects.create(project=project, cycle=self.cycle if i < 8 else old_cycle,
                template=self.template, status='APPROVED')
            for code, unit, amount, num, den in [('ROOMS', 'COUNT', 1, None, None),
                ('REVENUE', 'MONEY', 10000, None, None), ('ADR', 'MONEY', 10000, 10000, 1)]:
                NormalizedValue.objects.create(upload=upload, report_code='PL_TOTAL_WINE', row_code=code,
                    period='YEAR', data_year=2027, data_kind='BUDGET', unit=unit, value_int=amount,
                    ratio_num=num, ratio_den=den)
        context = ReportContext(2027, self.cycle.pk, 'PL_TOTAL_WINE', 'WORKING', tuple(p.pk for p in projects))
        result = query_report(context)
        self.assertEqual(result.coverage, {'available': 8, 'expected': 10})
        self.assertEqual(set(result.missing_projects), {p.pk for p in projects[8:]})
        metrics = {m.row_code: m for m in result.metrics}
        self.assertEqual(metrics['ROOMS'].value, 56)
        self.assertEqual(metrics['REVENUE'].value, Decimal('5400.00'))
        self.assertEqual(metrics['ADR'].value, Decimal('96.43'))
        for download in (False, True):
            request = RequestFactory().get('/', {'projects': ','.join(str(p.pk) for p in projects),
                'source_mode': 'WORKING', **({'download': 'xlsx'} if download else {})})
            request.user = get_user_model()(role='ADMIN')
            response = render_report(request, self.cycle, 'PL_TOTAL_WINE')
            self.assertEqual(response.status_code, 200)
            if not download:
                self.assertContains(response, '项目覆盖 8 / 10')
                self.assertContains(response, '非完整年度总数')
                for missing in projects[8:]:
                    self.assertContains(response, missing.name)
            else:
                workbook = load_workbook(BytesIO(response.content), data_only=True)
                self.assertEqual(list(workbook.active.values)[1], ('覆盖项目', 8, '应报项目', 10) + (None,) * 10)
                provenance = list(workbook['数据来源'].values)
                self.assertEqual(sum(row[2] == 'NOT_UPLOADED' for row in provenance if len(row) > 2), 2)
