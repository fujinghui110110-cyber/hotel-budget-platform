import json
import tempfile
from pathlib import Path

from django.test import TestCase, override_settings
from openpyxl import Workbook, load_workbook

from budgeting.models import BudgetCycle, Project, TemplateVersion, UploadVersion, User, ValidationRun
from budgeting.services.plan_history import ensure_plan, confirm_history, current_binding
from budgeting.services.template_delivery import signed_template_copy
from budgeting.services.history_workbook import validate_history_identity, validate_workbook_history


class HistoryTemplateDeliveryTests(TestCase):
    def test_download_backfills_locked_history_and_signed_binding_then_rejects_stale_revision(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source = root / 'template.xlsx'
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = '损益'
            sheet.append(['2024年实际', '2025年实际', '2026年预测', '预算'])
            sheet.append([0, 0, 0, 9])
            meta = workbook.create_sheet('SYS_META')
            meta.sheet_state = 'hidden'
            for key in ('template_version', 'budget_year', 'project_code', 'rule_version', 'formula_manifest_hash', 'project_signature'):
                meta.append([key, ''])
            workbook.save(source)
            workbook.close()
            manifest = {'management_v2': True, 'input_cells': {'损益': ['D2']}, 'reports': {'PL_TOTAL_WINE': {'sheet': '损益', 'region': {'header_row': 1},
                'mapping': [{'row_code': 'R001', 'cell': 'D2', 'unit': 'MONEY'}]}}}
            manifest_path = root / 'manifest.json'
            manifest_path.write_text(json.dumps(manifest))
            admin = User.objects.create_user(username='delivery-admin', role='ADMIN')
            project = Project.objects.create(code='DELIVERY', name='历史回填验收')
            template = TemplateVersion.objects.create(version='delivery', budget_year=2027,
                file_path=str(source), manifest_path=str(manifest_path), formula_manifest_hash='a' * 64)
            cycle = BudgetCycle.objects.create(name='2027', budget_year=2027, template=template)
            plan = ensure_plan(cycle)
            values = [dict(report_code='PL_TOTAL_WINE', row_code='R001', data_year=year,
                data_kind=kind, period='YEAR', unit='MONEY', value_int=value,
                ratio_num=None, ratio_den=None) for year, kind, value in
                [(2024, 'ACTUAL', 10001), (2025, 'ACTUAL', 10002), (2026, 'FORECAST', 10003)]]
            confirm_history(plan=plan, project=project, values=values, actor=admin,
                reason='管理员确认历史', source_identity={'file_hash': 'b' * 64}, expected_revision=0)
            with override_settings(BUDGET_STORAGE_ROOT=root):
                downloaded = signed_template_copy(template, project, cycle)
                result = load_workbook(downloaded, data_only=True)
                self.assertEqual([result['损益'][f'{col}2'].value for col in 'ABC'], [100.01, 100.02, 100.03])
                self.assertIsNone(result['损益']['D2'].value)
                result.close()
                upload = UploadVersion.objects.create(project=project, cycle=cycle, template=template,
                    original_path=str(downloaded), sha256='c' * 64)
                from budgeting.excel.structure_v2 import validate_v2_structure
                self.assertEqual(validate_v2_structure(upload, downloaded), [])
                run = ValidationRun.objects.create(upload=upload)
                self.assertEqual(validate_history_identity(upload, downloaded, run), 0)
                self.assertEqual(validate_workbook_history(upload, downloaded, run, manifest), 0)
                upload.refresh_from_db()
                self.assertEqual(upload.history_binding_id, current_binding(plan, project).pk)
                values[0]['value_int'] = 10004
                confirm_history(plan=plan, project=project, values=values, actor=admin,
                    reason='更正有依据的历史', source_identity={'file_hash': 'd' * 64}, expected_revision=1)
                self.assertEqual(validate_history_identity(upload, downloaded, run), 1)
                self.assertTrue(run.issues.filter(code='HISTORY_BASELINE_STALE').exists())
