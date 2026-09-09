import io
import json
import tempfile
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from django.contrib.auth import get_user_model
from django.test import TestCase, override_settings
from django.urls import reverse

from budgeting.models import BudgetCycle, NormalizedValue, Project, UploadVersion


class VersionReportTests(TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.settings = override_settings(BUDGET_STORAGE_ROOT=Path(self.temp.name))
        self.settings.enable()
        self.cycle = BudgetCycle.objects.create(name="2027年度预算", budget_year=2027, revision_no=2, status="OPEN")
        self.other_cycle = BudgetCycle.objects.create(name="2028年度预算", budget_year=2028, status="OPEN")
        self.project_a = Project.objects.create(code="A01", name="酒店A")
        self.project_b = Project.objects.create(code="B01", name="酒店B")
        self.project_c = Project.objects.create(code="C01", name="酒店C")
        self.admin = get_user_model().objects.create_user("version-admin", password="secret", role="ADMIN")
        self.project_user = get_user_model().objects.create_user("version-project", password="secret", role="PROJECT", project=self.project_b)

    def tearDown(self):
        self.settings.disable()
        self.temp.cleanup()

    def _upload(self, project, cycle=None, status="VALIDATED", name="budget.xlsx"):
        upload = UploadVersion.objects.create(project=project, cycle=cycle or self.cycle, status=status, original_name=name, original_path="not-needed.xlsx", sha256="a" * 64)
        return upload

    def _value(self, upload, **overrides):
        data = {
            "upload": upload, "report_code": "PL_TOTAL_WINE", "row_code": "R0010", "row_label": "=恶意名称", "period": "01", "unit": "MONEY", "value_int": 12345, "source_sheet": "总表", "source_cell": "B2",
        }
        data.update(overrides)
        return NormalizedValue.objects.create(**data)

    def _login(self):
        self.assertTrue(self.client.login(username="version-admin", password="secret"))

    def test_project_export_is_cycle_scoped_multisheet_and_preserves_blank(self):
        old = self._upload(self.project_a, status="APPROVED", name="old.xlsx")
        current = self._upload(self.project_a, status="SUBMITTED", name="current.xlsx")
        self._value(old, value_int=99999)
        self._value(current, value_int=12345)
        self._value(current, row_code="R0029", row_label="出租率", period="YEAR", unit="RATIO", value_int=0, ratio_num=1000, ratio_den=3100)
        self._value(current, report_code="A21前台", row_code="R0020", row_label="前台人工", period="A2025M01", value_int=500)
        self._value(current, report_code="A21前台", row_code="R0020", row_label="前台人工", period="YEAR", value_int=6000)
        self._value(current, report_code="A21前台", row_code="R0021", row_label="前台人数", period="YEAR", unit="COUNT", value_int=8)
        self._login()
        response = self.client.get(reverse("management_version_report_download", args=[self.project_a.pk]), {"cycle": self.cycle.pk})
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(io.BytesIO(b"".join(response.streaming_content)), data_only=False)
        self.assertEqual(workbook["导出说明"]["B5"].value, str(current.id))
        self.assertIn("01_酒店损益总表（含名酒）", workbook.sheetnames)
        self.assertEqual(workbook["01_酒店损益总表（含名酒）"]["D7"].value, 123.45)
        self.assertAlmostEqual(workbook["01_酒店损益总表（含名酒）"]["E8"].value, 1000 / 3100)
        self.assertEqual(workbook["01_酒店损益总表（含名酒）"]["B7"].data_type, "s")
        self.assertEqual(workbook["01_酒店损益总表（含名酒）"]["B7"].value, "=恶意名称")
        self.assertIn("02_A21前台", workbook.sheetnames)
        self.assertIsNone(workbook["02_A21前台"]["D8"].value)
        response.close()

    def test_batch_export_manifest_names_omitted_projects_and_never_crosses_cycles(self):
        usable = self._upload(self.project_a, status="VALIDATED", name="usable.xlsx")
        self._value(usable)
        failed = self._upload(self.project_b, status="REJECTED", name="failed.xlsx")
        self._value(failed, value_int=9900)
        other = self._upload(self.project_c, cycle=self.other_cycle, status="APPROVED", name="other-cycle.xlsx")
        self._value(other, value_int=8800)
        self._login()
        response = self.client.get(reverse("management_version_report_batch_download"), {"cycle": self.cycle.pk})
        self.assertEqual(response.status_code, 200)
        archive = zipfile.ZipFile(io.BytesIO(b"".join(response.streaming_content)))
        manifest = json.loads(archive.read("manifest.json"))
        self.assertEqual([item["project_code"] for item in manifest["included_projects"]], ["A01"])
        omitted = {item["project_code"]: item for item in manifest["omitted_projects"]}
        self.assertEqual(omitted["B01"]["status"], "REJECTED")
        self.assertEqual(omitted["C01"]["status"], "NOT_UPLOADED")
        self.assertTrue(any(member.endswith(".xlsx") for member in archive.namelist()))
        response.close()

    def test_page_reports_latest_usable_and_denies_project_account(self):
        earlier = self._upload(self.project_a, status="VALIDATED", name="earlier.xlsx")
        later_failed = self._upload(self.project_a, status="REJECTED", name="bad.xlsx")
        self._value(earlier)
        self._value(later_failed)
        self._login()
        page = self.client.get(reverse("management_version_report"), {"cycle": self.cycle.pk, "project_id": self.project_a.pk})
        self.assertEqual(page.status_code, 200)
        self.assertEqual(page.context["selected_version"], earlier)
        self.assertContains(page, "最新上传为已拒绝；当前导出最近可用版本")
        self.client.logout()
        self.client.force_login(self.project_user)
        self.assertEqual(self.client.get(reverse("management_version_report"), {"cycle": self.cycle.pk}).status_code, 403)

    def test_report_period_scope_uses_matching_order_for_headers_and_values(self):
        upload = self._upload(self.project_a, status="VALIDATED")
        self._value(upload, report_code="A21前台", period="A2024", data_year=2024, data_kind="ACTUAL", value_int=100)
        self._value(upload, report_code="A21前台", period="A2025M01", data_year=2025, data_kind="ACTUAL", month=1, value_int=200)
        self._value(upload, report_code="A21前台", period="F2026M01", data_year=2026, data_kind="FORECAST", month=1, value_int=300)
        self._value(upload, report_code="A21前台", period="YEAR", data_year=2027, data_kind="BUDGET", value_int=400)
        self._login()
        base = {"cycle": self.cycle.pk, "project_id": self.project_a.pk, "report_code": "A21前台"}
        annual = self.client.get(reverse("management_version_report"), base)
        self.assertEqual(annual.context["period_scope"], "annual")
        self.assertEqual(annual.context["periods"], ["A2024", "YEAR"])
        self.assertEqual(annual.context["period_labels"], {"A2024": "2024年实际", "YEAR": "2027年预算年度"})
        self.assertEqual([value["display"] for value in annual.context["rows"][0]["values"]], ["1.00", "4.00"])
        complete = self.client.get(reverse("management_version_report"), {**base, "period_scope": "complete"})
        self.assertEqual(complete.context["periods"], ["A2024", "A2025M01", "F2026M01", "YEAR"])
        self.assertEqual([value["display"] for value in complete.context["rows"][0]["values"]], ["1.00", "2.00", "3.00", "4.00"])

    def test_ratio_display_prefers_numerator_and_denominator_over_placeholder_value(self):
        upload = self._upload(self.project_a, status="VALIDATED")
        self._value(upload, report_code="A21前台", row_code="R0029", row_label="出租率", period="YEAR", unit="RATIO", value_int=0, ratio_num=1000, ratio_den=3100)
        self._login()
        page = self.client.get(reverse("management_version_report"), {"cycle": self.cycle.pk, "project_id": self.project_a.pk, "report_code": "A21前台"})
        self.assertEqual(page.context["rows"][0]["values"][0]["display"], "32.26%")
