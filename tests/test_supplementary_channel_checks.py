from pathlib import Path
from tempfile import TemporaryDirectory

from django.test import TestCase
from openpyxl import Workbook

from budgeting.excel.channel_checks import validate_channel_values
from budgeting.excel.supplementary import extract_supplementary_values
from budgeting.models import (
    BudgetCycle,
    NormalizedValue,
    Project,
    UploadVersion,
    ValidationIssue,
    ValidationRun,
)


class SupplementaryExtractionTests(TestCase):
    def setUp(self):
        project = Project.objects.create(code="SUPP", name="补充指标测试")
        cycle = BudgetCycle.objects.create(name="2027预算", budget_year=2027)
        self.upload = UploadVersion.objects.create(
            project=project,
            cycle=cycle,
            original_path="test.xlsx",
            sha256="0" * 64,
        )

    def _workbook_path(self):
        workbook = Workbook()
        workbook.active.title = "B1餐厅汇总"
        for sheet_name in ("B2宴会收入", "B3宴会厅", "B16月饼亭", "经营补充指标"):
            workbook.create_sheet(sheet_name)

        restaurant = workbook["B1餐厅汇总"]
        for month, column in enumerate(range(11, 23), start=1):
            restaurant.cell(row=24, column=column, value=month + 0.25)

        banquet = workbook["B2宴会收入"]
        banquet["K24"] = 123.45
        workbook["B3宴会厅"]["K24"] = 999.99

        wine = workbook["经营补充指标"]
        wine['A3'] = '名酒收入（元）'
        for month, column in enumerate(range(4, 16), start=1):
            wine.cell(row=2, column=column, value=f'{month:02d}月')
            wine.cell(row=3, column=column, value=8 + month / 100)

        workbook["B16月饼亭"]["K24"] = 45.67
        directory = TemporaryDirectory()
        path = Path(directory.name) / "supplementary.xlsx"
        workbook.save(path)
        workbook.close()
        return directory, path

    def test_sources_are_preserved_and_blank_months_are_zero_filled(self):
        directory, path = self._workbook_path()
        self.addCleanup(directory.cleanup)

        extracted = extract_supplementary_values(self.upload, path)

        restaurant = NormalizedValue.objects.get(
            upload=self.upload,
            report_code="PL_TOTAL_WINE",
            row_code="R9001",
            period="01",
        )
        self.assertEqual(extracted, 4 * (12 + 1) * 4)
        self.assertEqual(restaurant.value_int, 125)
        self.assertEqual((restaurant.source_sheet, restaurant.source_cell), ("B1餐厅汇总", "K24"))

        annual = NormalizedValue.objects.get(
            upload=self.upload,
            report_code="PL_TOTAL_WINE",
            row_code="R9001",
            period="YEAR",
        )
        self.assertEqual(annual.value_int, sum(int((month + 0.25) * 100) for month in range(1, 13)))
        self.assertIn("SUM(12 months)", annual.source_formula)

        banquet_values = NormalizedValue.objects.filter(
            upload=self.upload,
            row_code="R9002",
            report_code="PL_TOTAL_WINE",
        )
        self.assertEqual(banquet_values.count(), 13)
        self.assertEqual(banquet_values.get(period="01").value_int, 12345)
        self.assertEqual(banquet_values.get(period="01").source_sheet, "B2宴会收入")
        self.assertEqual(banquet_values.get(period="02").value_int, 0)
        self.assertEqual(banquet_values.get(period="02").source_cell, "L24")
        self.assertEqual(banquet_values.get(period="YEAR").value_int, 12345)
        self.assertFalse(
            NormalizedValue.objects.filter(upload=self.upload, source_sheet="B3宴会厅").exists()
        )

        seasonal = NormalizedValue.objects.get(
            upload=self.upload,
            report_code="PL_TOTAL_WINE",
            row_code="R9005",
            period="01",
        )
        self.assertEqual(seasonal.value_int, 4567)
        seasonal_values = NormalizedValue.objects.filter(
            upload=self.upload,
            report_code="PL_TOTAL_WINE",
            row_code="R9005",
        )
        self.assertEqual(seasonal_values.count(), 13)
        self.assertEqual(seasonal_values.get(period="02").value_int, 0)
        self.assertEqual(seasonal_values.get(period="YEAR").value_int, 4567)


class WineDetailIntegrationTests(TestCase):
    def test_original_detail_preserves_cents_and_reads_blank_month_as_zero(self):
        project = Project.objects.create(code='WINE', name='名酒明细隔离测试')
        cycle = BudgetCycle.objects.create(name='预算', budget_year=2027)
        upload = UploadVersion.objects.create(project=project, cycle=cycle,
                                             original_path='test.xlsx', sha256='1' * 64)
        run = ValidationRun.objects.create(upload=upload, rule_version='test')
        book = Workbook()
        sheet = book.active
        sheet.title = 'OOD-其他 (名酒)'
        sheet['H61'] = 'OOD收入-其他'
        sheet['H71'] = 'OOD成本-其他'
        for month, column in enumerate(range(11, 23), 1):
            sheet.cell(21, column, f'{month:02d}')
            sheet.cell(61, column, 100.01)
            if month != 2:
                sheet.cell(71, column, 80.02)
        # A mixed other-business sheet must never replace missing wine values.
        mixed = book.create_sheet('OOD-其他')
        mixed['H61'] = 'OOD收入-其他'
        mixed['K61'] = 99999
        with TemporaryDirectory() as directory:
            path = Path(directory) / 'detail.xlsx'
            book.save(path)
            book.close()
            extract_supplementary_values(upload, path, validation_run=run)
        income = NormalizedValue.objects.get(upload=upload, report_code='PL_TOTAL_WINE',
                                             row_code='R9003', period='YEAR')
        self.assertEqual(income.value_int, 120012)
        cost = NormalizedValue.objects.get(upload=upload, report_code='PL_TOTAL_WINE',
                                           row_code='R9006', period='01')
        self.assertEqual(cost.value_int, 8002)
        self.assertEqual((cost.source_sheet, cost.source_cell), ('OOD-其他 (名酒)', 'K71'))
        blank_month = NormalizedValue.objects.get(upload=upload, report_code='PL_TOTAL_WINE',
                                                  row_code='R9006', period='02')
        self.assertEqual(blank_month.value_int, 0)
        self.assertEqual((blank_month.source_sheet, blank_month.source_cell), ('OOD-其他 (名酒)', 'L71'))
        cost_year = NormalizedValue.objects.get(upload=upload, report_code='PL_TOTAL_WINE',
                                                row_code='R9006', period='YEAR')
        self.assertEqual(cost_year.value_int, 11 * 8002)
        self.assertFalse(run.issues.filter(code='WINE_DETAIL_INCOMPLETE').exists())

        # Retrying after a source disappears must not retain stale wine amounts.
        with TemporaryDirectory() as directory:
            path = Path(directory) / 'missing-detail.xlsx'
            book = Workbook()
            book.save(path)
            book.close()
            extract_supplementary_values(upload, path, validation_run=run)
            extract_supplementary_values(upload, path, validation_run=run)
        self.assertFalse(NormalizedValue.objects.filter(
            upload=upload, row_code__in=['R9003', 'R9006']).exists())
        self.assertEqual(run.issues.filter(code='WINE_SOURCE_MISSING').count(), 2)
        self.assertFalse(run.issues.filter(code='WINE_DETAIL_INCOMPLETE').exists())

    def test_formula_without_cache_is_not_converted_to_zero(self):
        project = Project.objects.create(code='WINE-FORMULA', name='名酒公式缓存验证')
        cycle = BudgetCycle.objects.create(name='预算', budget_year=2027)
        upload = UploadVersion.objects.create(project=project, cycle=cycle,
                                             original_path='test.xlsx', sha256='2' * 64)
        run = ValidationRun.objects.create(upload=upload, rule_version='test')
        book = Workbook()
        sheet = book.active
        sheet.title = 'OOD-其他 (名酒)'
        sheet['H61'] = 'OOD收入-其他'
        sheet['H71'] = 'OOD成本-其他'
        for month, column in enumerate(range(11, 23), 1):
            sheet.cell(21, column, f'{month:02d}')
            sheet.cell(61, column, 100.01)
            sheet.cell(71, column, '=1+1' if month == 2 else '#REF!' if month == 3 else 80.02)
        with TemporaryDirectory() as directory:
            path = Path(directory) / 'formula-detail.xlsx'
            book.save(path)
            book.close()
            extract_supplementary_values(upload, path, validation_run=run)
        self.assertFalse(NormalizedValue.objects.filter(upload=upload, row_code='R9006', period='02').exists())
        self.assertFalse(NormalizedValue.objects.filter(upload=upload, row_code='R9006', period='03').exists())
        self.assertTrue(run.issues.filter(code='WINE_DETAIL_INCOMPLETE').exists())


class ChannelCheckTests(TestCase):
    def setUp(self):
        project = Project.objects.create(code="CHANNEL", name="渠道校验测试")
        cycle = BudgetCycle.objects.create(name="2027预算", budget_year=2027)
        upload = UploadVersion.objects.create(
            project=project,
            cycle=cycle,
            original_path="test.xlsx",
            sha256="0" * 64,
        )
        self.run = ValidationRun.objects.create(upload=upload, rule_version="test")

    def _workbook_path(self, revenue):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "A1客房收入(新)"
        sheet["K37"] = 123.45
        sheet["K69"] = 2
        sheet["K101"] = revenue
        sheet["V130"] = 0
        directory = TemporaryDirectory()
        path = Path(directory.name) / "channel.xlsx"
        workbook.save(path)
        workbook.close()
        return directory, path

    def test_adr_times_room_nights_exact_to_cent_passes(self):
        directory, path = self._workbook_path(246.90)
        self.addCleanup(directory.cleanup)

        validate_channel_values(path, self.run)

        self.assertFalse(ValidationIssue.objects.filter(run=self.run).exists())

    def test_adr_times_room_nights_difference_is_supplementary_notice(self):
        directory, path = self._workbook_path(246.91)
        self.addCleanup(directory.cleanup)

        validate_channel_values(path, self.run)

        issue = ValidationIssue.objects.get(run=self.run, code="CHANNEL_REVENUE_RECONCILIATION")
        self.assertEqual(issue.severity, "P2")
        self.assertEqual(issue.location, "A1客房收入(新)!K101")
        self.assertEqual((issue.actual_value, issue.expected_value), ("246.91", "246.90"))
