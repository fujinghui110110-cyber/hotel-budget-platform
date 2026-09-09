import unittest
from xml.etree import ElementTree as ET

from openpyxl import Workbook

from budgeting.excel.extract import _detect_grid, _month_number, sheet_slug
from budgeting.services.template_delivery import MAIN_NS, _blank_input_cells


class SheetSlugTests(unittest.TestCase):
    def test_keeps_chinese_and_replaces_specials(self):
        self.assertEqual(sheet_slug("客房市场说明 (250928)"), "客房市场说明_250928")

    def test_plain_chinese_passes_through(self):
        self.assertEqual(sheet_slug("工资福利费"), "工资福利费")

    def test_blank_falls_back_to_sheet(self):
        self.assertEqual(sheet_slug("   "), "sheet")

    def test_long_names_are_truncated(self):
        self.assertLessEqual(len(sheet_slug("很" * 60)), 40)


class MonthNumberTests(unittest.TestCase):
    def test_month_labels(self):
        self.assertEqual(_month_number("1月"), 1)
        self.assertEqual(_month_number("12月"), 12)
        self.assertEqual(_month_number("01"), 1)
        self.assertEqual(_month_number(5), 5)

    def test_non_months(self):
        self.assertIsNone(_month_number("合计"))
        self.assertIsNone(_month_number(15))
        self.assertIsNone(_month_number(None))


class DetectGridTests(unittest.TestCase):
    def test_grid_layout(self):
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "行标签"
        for m in range(1, 13):
            ws.cell(row=1, column=1 + m, value=f"{m}月")
        ws.cell(row=1, column=14, value="合计")
        ws["A2"] = "客房收入"
        ws["A3"] = "小计"

        label_col, month_cols, header_idx = _detect_grid(ws)

        self.assertEqual(header_idx, 1)
        self.assertEqual(label_col, 1)
        self.assertEqual(month_cols, list(range(2, 14)))


class BlankInputCellsTests(unittest.TestCase):
    def _worksheet(self, body):
        return f'<worksheet xmlns="{MAIN_NS}"><sheetData>{body}</sheetData></worksheet>'

    def test_blanks_only_numeric_inputs(self):
        data = self._worksheet(
            '<row r="1">'
            '<c r="A1" t="s"><v>0</v></c>'
            '<c r="B1"><f>SUM(C1:D1)</f><v>5</v></c>'
            '<c r="C1"><v>42</v></c>'
            '<c r="D1" t="inlineStr"><is><t>hi</t></is></c>'
            "</row>"
        )
        root = ET.fromstring(_blank_input_cells(data))

        cells = {c.attrib.get("r"): c for c in root.findall(".//m:c", {"m": MAIN_NS})}
        self.assertIsNotNone(cells["A1"].find("m:v", {"m": MAIN_NS}))
        self.assertIsNotNone(cells["B1"].find("m:f", {"m": MAIN_NS}))
        self.assertIsNotNone(cells["B1"].find("m:v", {"m": MAIN_NS}))
        self.assertIsNone(cells["C1"].find("m:v", {"m": MAIN_NS}))
        self.assertIsNotNone(cells["D1"].find("m:is", {"m": MAIN_NS}))


if __name__ == "__main__":
    unittest.main()
