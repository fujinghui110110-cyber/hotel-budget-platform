import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from budgeting.services.history_parser import propose_history


CANONICAL_ROWS = [
    {"row_code": "R0010", "row_label": "客房收入", "unit": "MONEY", "aggregation": "SUM"},
    {"row_code": "R0020", "row_label": "餐饮收入", "unit": "MONEY", "aggregation": "SUM"},
    {"row_code": "R0030", "row_label": "已售房晚", "unit": "COUNT", "aggregation": "SUM"},
    {"row_code": "R0040", "row_label": "可售房晚", "unit": "COUNT", "aggregation": "SUM"},
    {"row_code": "R0050", "row_label": "出租率", "unit": "RATIO"},
    {"row_code": "R0060", "row_label": "税前利润", "unit": "MONEY", "aggregation": "SUM"},
    {"row_code": "R0070", "row_label": "平均房价 ADR", "unit": "MONEY", "aggregation": "AVERAGE"},
    {"row_code": "R0080", "row_label": "RevPAR", "unit": "MONEY", "aggregation": "AVERAGE"},
    {"row_code": "R0090", "row_label": "房间数", "unit": "COUNT"},
]


def _save(workbook):
    directory = tempfile.TemporaryDirectory()
    path = Path(directory.name) / "history.xlsx"
    workbook.save(path)
    return directory, path


class HistoryParserTests(unittest.TestCase):
    def test_reordered_rows_and_columns_match_by_subject_and_header_semantics(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "2024实际损益"
        sheet["D3"] = "2月"
        sheet["B3"] = "科目"
        sheet["E3"] = "同比"
        sheet["C3"] = "1月"
        sheet["B5"] = "房费收入"
        sheet["D5"] = 20
        sheet["C5"] = 10
        sheet["E5"] = 999
        sheet["B4"] = "餐饮收入"
        sheet["D4"] = 8
        sheet["C4"] = 7
        directory, path = _save(workbook)
        self.addCleanup(directory.cleanup)

        result = propose_history(path, CANONICAL_ROWS, 2024, "ACTUAL")

        room = next(row for row in result["rows"] if row["source_label"] == "房费收入")
        self.assertEqual(room["suggested_code"], "R0010")
        self.assertTrue(room["include"])
        self.assertEqual([(item["month"], item["value_int"]) for item in room["values"]], [(1, 1000), (2, 2000)])
        self.assertNotIn("E5", {item["source_cell"] for item in room["values"]})

    def test_exact_and_alias_matching_prefer_clear_canonical_rows(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "损益表"
        sheet["A1"] = "科目"
        sheet["B1"] = "2024年实际全年"
        sheet["A2"] = "已出租房晚"
        sheet["B2"] = 123.4
        sheet["A3"] = "入住率"
        sheet["B3"] = 0.8
        directory, path = _save(workbook)
        self.addCleanup(directory.cleanup)

        result = propose_history(path, CANONICAL_ROWS, 2024, "ACTUAL")

        sold = next(row for row in result["rows"] if row["source_label"] == "已出租房晚")
        occ = next(row for row in result["rows"] if row["source_label"] == "入住率")
        self.assertEqual(sold["suggested_code"], "R0030")
        self.assertEqual(sold["values"][0]["value_int"], 123)
        self.assertEqual(occ["suggested_code"], "R0050")
        self.assertEqual(occ["values"][0]["ratio_num"], 800000)
        self.assertEqual(occ["values"][0]["ratio_den"], 1000000)
        self.assertEqual(occ["values"][0]["value_int"], 0)

    def test_semantic_conflict_blocks_dangerous_similarity(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "损益表"
        sheet["A1"] = "科目"
        sheet["B1"] = "2024年实际全年"
        sheet["A2"] = "客房成本"
        sheet["B2"] = 100
        sheet["A3"] = "税后利润"
        sheet["B3"] = 50
        directory, path = _save(workbook)
        self.addCleanup(directory.cleanup)

        result = propose_history(path, CANONICAL_ROWS, 2024, "ACTUAL")

        by_label = {row["source_label"]: row for row in result["rows"]}
        self.assertEqual(by_label["客房成本"]["suggested_code"], "")
        self.assertFalse(by_label["客房成本"]["include"])
        self.assertEqual(by_label["税后利润"]["suggested_code"], "")
        self.assertLess(by_label["税后利润"]["confidence"], 0.9)

    def test_monthly_values_create_annual_total_except_ratio(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "损益表"
        sheet["A1"] = "科目"
        for month in range(1, 13):
            sheet.cell(row=1, column=month + 1, value=f"{month}月")
            sheet.cell(row=2, column=month + 1, value=month)
            sheet.cell(row=3, column=month + 1, value=0.5)
        sheet["A2"] = "客房收入"
        sheet["A3"] = "出租率"
        directory, path = _save(workbook)
        self.addCleanup(directory.cleanup)

        result = propose_history(path, CANONICAL_ROWS, 2024, "ACTUAL", money_unit="万元")

        room = next(row for row in result["rows"] if row["source_label"] == "客房收入")
        occ = next(row for row in result["rows"] if row["source_label"] == "出租率")
        annual = next(item for item in room["values"] if item["month"] is None)
        self.assertEqual(annual["value_int"], 78 * 10000 * 100)
        self.assertEqual(len([item for item in occ["values"] if item["month"] is None]), 0)

    def test_year_kind_filter_and_single_pnl_sheet_selection(self):
        workbook = Workbook()
        appendix = workbook.active
        appendix.title = "附表"
        appendix["A1"] = "科目"
        appendix["B1"] = "2024年实际全年"
        appendix["A2"] = "客房收入"
        appendix["B2"] = 999
        pnl = workbook.create_sheet("酒店损益")
        pnl["A1"] = "年份"
        pnl["B1"] = "口径"
        pnl["C1"] = "科目"
        pnl["D1"] = "金额"
        pnl["A2"] = 2023
        pnl["B2"] = "实际"
        pnl["C2"] = "客房收入"
        pnl["D2"] = 1
        pnl["A3"] = 2024
        pnl["B3"] = "预算"
        pnl["C3"] = "客房收入"
        pnl["D3"] = 2
        pnl["A4"] = 2024
        pnl["B4"] = "实际"
        pnl["C4"] = "客房收入"
        pnl["D4"] = 3
        directory, path = _save(workbook)
        self.addCleanup(directory.cleanup)

        result = propose_history(path, CANONICAL_ROWS, 2024, "ACTUAL")

        self.assertEqual(len(result["rows"]), 1)
        self.assertEqual(result["rows"][0]["values"][0]["value_int"], 300)
        self.assertEqual(result["rows"][0]["values"][0]["source_sheet"], "酒店损益")
        self.assertTrue(any(issue["code"] == "YEAR_MISMATCH_ROW" for issue in result["issues"]))

    def test_formula_without_cache_is_reported_and_not_zero_filled(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "损益表"
        sheet["A1"] = "科目"
        sheet["B1"] = "2024年实际全年"
        sheet["A2"] = "客房收入"
        sheet["B2"] = "=1+1"
        directory, path = _save(workbook)
        self.addCleanup(directory.cleanup)

        result = propose_history(path, CANONICAL_ROWS, 2024, "ACTUAL")

        self.assertEqual(len(result["rows"]), 1)
        self.assertEqual(result["rows"][0]["source_label"], "客房收入")
        self.assertEqual(result["rows"][0]["values"], [])
        self.assertFalse(result["rows"][0]["include"])
        self.assertTrue(any(issue["code"] == "FORMULA_CACHE_MISSING" for issue in result["issues"]))
        self.assertTrue(any(issue["code"] == "UNRECOGNIZED_VALUE" for issue in result["issues"]))

    def test_late_header_with_many_subject_rows_does_not_drop_header_context(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "损益表"
        sheet["B25"] = "2024年实际"
        sheet["A26"] = "科目"
        sheet["B26"] = "1月"
        for offset in range(30):
            row = 27 + offset
            sheet.cell(row=row, column=1, value=f"未知科目{offset}")
            sheet.cell(row=row, column=2, value=offset)
        sheet["A57"] = "客房收入"
        sheet["B57"] = 88
        directory, path = _save(workbook)
        self.addCleanup(directory.cleanup)

        result = propose_history(path, CANONICAL_ROWS, 2024, "ACTUAL")

        room = next(row for row in result["rows"] if row["source_label"] == "客房收入")
        self.assertEqual(room["values"][0]["source_cell"], "B57")
        self.assertEqual(room["values"][0]["month"], 1)
        self.assertEqual(room["values"][0]["value_int"], 8800)

    def test_blank_subject_value_is_kept_for_audit(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "损益表"
        sheet["A1"] = "科目"
        sheet["B1"] = "2024年实际全年"
        sheet["A2"] = "客房收入"
        directory, path = _save(workbook)
        self.addCleanup(directory.cleanup)

        result = propose_history(path, CANONICAL_ROWS, 2024, "ACTUAL")

        self.assertEqual(len(result["rows"]), 1)
        self.assertEqual(result["rows"][0]["source_label"], "客房收入")
        self.assertEqual(result["rows"][0]["values"], [])
        self.assertFalse(result["rows"][0]["include"])
        self.assertTrue(any(issue["code"] == "MISSING_VALUE" for issue in result["issues"]))

    def test_single_sheet_opposite_wine_scope_is_rejected(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "酒店损益总表（不含名酒）"
        sheet["A1"] = "科目"
        sheet["B1"] = "2024年实际全年"
        sheet["A2"] = "客房收入"
        sheet["B2"] = 1
        directory, path = _save(workbook)
        self.addCleanup(directory.cleanup)

        result = propose_history(path, CANONICAL_ROWS, 2024, "ACTUAL", report_code="PL_TOTAL_WINE")

        self.assertEqual(result["rows"], [])
        self.assertTrue(any(issue["code"] == "PNL_SHEET_REPORT_MISMATCH" for issue in result["issues"]))

    def test_price_metrics_and_room_stock_are_not_summed_or_scaled_as_wan(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "损益表"
        sheet["A1"] = "科目"
        for month in range(1, 13):
            sheet.cell(row=1, column=month + 1, value=f"{month}月")
            sheet.cell(row=2, column=month + 1, value=500 + month)
            sheet.cell(row=3, column=month + 1, value=300 + month)
            sheet.cell(row=4, column=month + 1, value=100)
        sheet["A2"] = "平均房价 ADR"
        sheet["A3"] = "RevPAR"
        sheet["A4"] = "房间数"
        directory, path = _save(workbook)
        self.addCleanup(directory.cleanup)

        result = propose_history(path, CANONICAL_ROWS, 2024, "ACTUAL", money_unit="万元")

        by_label = {row["source_label"]: row for row in result["rows"]}
        adr = by_label["平均房价 ADR"]
        revpar = by_label["RevPAR"]
        rooms = by_label["房间数"]
        self.assertEqual(adr["values"][0]["value_int"], 50100)
        self.assertEqual(revpar["suggested_code"], "R0080")
        self.assertEqual(revpar["values"][0]["unit"], "MONEY")
        self.assertFalse(any(item["month"] is None for item in adr["values"]))
        self.assertFalse(any(item["month"] is None for item in revpar["values"]))
        self.assertFalse(any(item["month"] is None for item in rooms["values"]))
        self.assertTrue(any(issue["code"] == "MONEY_UNIT_PRICE_AS_YUAN" for issue in result["issues"]))
        self.assertGreaterEqual(
            sum(1 for issue in result["issues"] if issue["code"] == "ANNUAL_AGGREGATION_UNCERTAIN"),
            3,
        )


if __name__ == "__main__":
    unittest.main()
