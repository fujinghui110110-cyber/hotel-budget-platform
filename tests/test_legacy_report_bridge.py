import hashlib
import json
from pathlib import Path
from io import BytesIO
from django.test import TestCase, RequestFactory, override_settings
from django.contrib.auth import get_user_model
from openpyxl import load_workbook
from budgeting.models import TemplateVersion, UploadVersion, NormalizedValue
from budgeting.services.report_context import ReportContext
from budgeting.services.report_query import query_report
from budgeting.services.report_presenter import render_report, render_drilldown
from tests.test_report_context import ReportQueryTests


class LegacyReportBridgeTests(TestCase):
    def setUp(self):
        ReportQueryTests.setUp(self)
        self.override=override_settings(BUDGET_STORAGE_ROOT=self.tmp.name)
        self.override.enable(); self.addCleanup(self.override.disable)
        self.cycle.source_budget_year=2026
        self.cycle.save(update_fields=['source_budget_year'])
        self.upload=UploadVersion.objects.get(project=self.projects[0])
        source=Path(self.tmp.name)/'source.xlsx'; source.write_bytes(b'unchanged original workbook')
        self.upload.sha256=hashlib.sha256(source.read_bytes()).hexdigest()
        self.upload.original_path='source.xlsx'
        self.manifest={'legacy_rehearsal':True,'source_sha256':self.upload.sha256,'source_budget_year':2026,
                       'budget_year':2027,'year_offset':1,'reports':{'PL_TOTAL_WINE':{'sheet':'旧损益表','mapping':[
                           {'row_code':code,'row_label':label,'unit':unit,'period':'FY','cell':cell,'aggregation':agg}
                           for code,label,unit,cell,agg in [('ROOMS','房晚','COUNT','A1','SUM'),('ADR','ADR','MONEY','A2','DERIVED'),('REVENUE','收入','MONEY','A3','SUM')]]}}}
        self.manifest_path=Path(self.tmp.name)/'legacy.json';self.write_manifest()
        self.legacy=TemplateVersion.objects.create(version='LEGACY-private',budget_year=2027,
            manifest_path=str(self.manifest_path),file_path=str(source),formula_manifest_hash=self.upload.sha256,rule_version='LEGACY-CACHE-1')
        self.upload.template=self.legacy;self.upload.save()
        for code,label,cell in [('ROOMS','房晚','A1'),('ADR','ADR','A2'),('REVENUE','收入','A3')]:
            NormalizedValue.objects.filter(upload=self.upload,row_code=code).update(source_sheet='旧损益表',source_cell=cell,row_label=label)
        NormalizedValue.objects.filter(upload=self.upload,row_code='ADR').update(ratio_num=None,ratio_den=None,value_int=8000)
        self.context=ReportContext(2027,self.cycle.pk,'PL_TOTAL_WINE','WORKING',(self.projects[0].pk,))

    def write_manifest(self):
        self.manifest_path.write_text(json.dumps(self.manifest,ensure_ascii=False))

    def test_bound_rehearsal_reaches_main_drilldown_and_export(self):
        result=query_report(self.context)
        self.assertFalse(result.missing_projects)
        self.assertEqual(next(m.value for m in result.metrics if m.row_code=='ADR'),80)
        user=get_user_model()(role='ADMIN')
        request=RequestFactory().get('/',{'projects':str(self.projects[0].pk)})
        request.user=user
        self.assertContains(render_report(request,self.cycle,'PL_TOTAL_WINE'),'800.00')
        request=RequestFactory().get('/',{'projects':str(self.projects[0].pk),'row_code':'ADR','period':'YEAR'})
        request.user=user
        self.assertContains(render_drilldown(request,self.cycle,'PL_TOTAL_WINE'),'80.00')
        request=RequestFactory().get('/',{'projects':str(self.projects[0].pk),'download':'xlsx'})
        request.user=user
        response=render_report(request,self.cycle,'PL_TOTAL_WINE')
        book=load_workbook(BytesIO(response.content))
        self.assertEqual(book.active.cell(5,14).value,80)

    def test_changed_source_or_wrong_rehearsal_year_is_rejected(self):
        self.manifest['source_budget_year']=2025;self.write_manifest()
        self.assertEqual(query_report(self.context).missing_projects[self.projects[0].pk],'TEMPLATE_MISMATCH')
        self.manifest['source_budget_year']=2026;self.write_manifest()
        (Path(self.tmp.name)/'source.xlsx').write_bytes(b'changed')
        self.assertEqual(query_report(self.context).missing_projects[self.projects[0].pk],'TEMPLATE_MISMATCH')

    def test_same_row_code_wrong_subject_is_not_imported(self):
        self.manifest['reports']['PL_TOTAL_WINE']['mapping'][2]['row_label']='其他支出'
        self.write_manifest()
        result=query_report(self.context)
        rows={m.row_code:m for m in result.metrics}
        self.assertIsNone(rows['REVENUE'].value)
        self.assertIsNone(rows['ADR'].value)
        self.assertEqual(rows['ROOMS'].value,10)

    def test_modern_mismatch_and_nonrehearsal_remain_blocked(self):
        self.legacy.rule_version='MODERN';self.legacy.save(update_fields=['rule_version'])
        self.assertIn(self.projects[0].pk,query_report(self.context).missing_projects)
        self.legacy.rule_version='LEGACY-CACHE-1';self.legacy.save(update_fields=['rule_version'])
        self.cycle.source_budget_year=None;self.cycle.save(update_fields=['source_budget_year'])
        self.assertIn(self.projects[0].pk,query_report(self.context).missing_projects)

    def test_source_verification_runs_once_for_all_table_periods(self):
        from unittest.mock import patch
        from budgeting.services.metric_definitions import verified_legacy_mapping
        from budgeting.services.report_query import query_report_table
        with patch('budgeting.services.report_query.verified_legacy_mapping',wraps=verified_legacy_mapping) as verify:
            reports=query_report_table(self.context,include_history=True)
        self.assertEqual(len(reports),52)
        self.assertEqual(verify.call_count,1)
