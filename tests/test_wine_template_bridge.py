from django.test import SimpleTestCase
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from budgeting.excel.template_v3 import add_supplementary_sheet, repair_wine_total_bridge


class WineTemplateBridgeTests(SimpleTestCase):
    def workbook(self):
        book = Workbook()
        book.remove(book.active)
        for title in ('酒店损益总表（含名酒）', '酒店损益总表（不含名酒）', 'OOD-其他', 'C其他运营汇总', '酒店损益总表-明细表（调整用第一轮不用）'):
            book.create_sheet(title)
        book['OOD-其他']['H55'] = 'OOD收入-商品零售(名酒)'
        book['OOD-其他']['H69'] = 'OOD成本-商品零售（名酒)'
        for month in range(12):
            col, target = get_column_letter(11 + month), get_column_letter(12 + month)
            for title in ('OOD-其他', 'C其他运营汇总'):
                book[title][f'{col}24'] = f'=ROUND(({col}25+{col}26+{col}35+{col}41+{col}45+{col}49+{col}55+{col}61+{col}56),2)'
                book[title][f'{col}62'] = f'=ROUND(SUM({col}63:{col}71),2)'
            for row in (55, 69):
                book['C其他运营汇总'][f'{col}{row}'] = f"=ROUND(('C7商品零售'!{col}{row}+'OOD-其他'!{col}{row}+'C8水上乐园'!{col}{row}),2)"
            for parent, total, leaf in ((57, 24, 55), (58, 62, 69)):
                book['酒店损益总表-明细表（调整用第一轮不用）'][f'{target}{parent}'] = f"=-'OOD-其他'!{col}{leaf}"
                for title in ('酒店损益总表（含名酒）', '酒店损益总表（不含名酒）'):
                    book[title][f'{target}{parent}'] = f"=(C其他运营汇总!{col}{total})+('酒店损益总表-明细表（调整用第一轮不用）'!{target}{parent})"
        return book

    def test_income_and_actual_cost_sources_are_subtracted_once(self):
        book = self.workbook()
        # Preserve the existing source cost, whether manually entered or formula driven.
        book['OOD-其他']['K69'] = '=K55*0.98'
        result = repair_wine_total_bridge(book)
        self.assertEqual(result['status'], 'PASS')
        self.assertEqual(len(result['repairs']), 48)
        self.assertEqual(book['酒店损益总表（不含名酒）']['L57'].value,
                         "=ROUND('酒店损益总表（含名酒）'!L57-'OOD-其他'!K55,2)")
        self.assertEqual(book['酒店损益总表（不含名酒）']['W58'].value,
                         "=ROUND('酒店损益总表（含名酒）'!W58-'OOD-其他'!V69,2)")
        self.assertEqual(book['OOD-其他']['K69'].value, '=K55*0.98')
        self.assertIn('K55', result['input_cells']['OOD-其他'])
        self.assertNotIn('K69', result['input_cells']['OOD-其他'])
        self.assertEqual(book['酒店损益总表（含名酒）']['L57'].value, '=ROUND(C其他运营汇总!K24,2)')
        self.assertEqual(book['酒店损益总表（含名酒）']['W58'].value, '=ROUND(C其他运营汇总!V62,2)')
        self.assertEqual(book['酒店损益总表-明细表（调整用第一轮不用）']['L57'].value, "=-'OOD-其他'!K55")
        self.assertEqual(result['expense_allocation'], 'UNKNOWN')
        self.assertEqual(repair_wine_total_bridge(book)['repairs'], [])
        book.close()

    def test_generic_retail_cost_does_not_authorize_a_wine_bridge(self):
        book = self.workbook()
        book['OOD-其他']['H69'] = 'OOD成本-商品零售'
        before = book['酒店损益总表（不含名酒）']['L57'].value
        result = repair_wine_total_bridge(book)
        self.assertEqual(result['status'], 'UNKNOWN')
        self.assertEqual(result['repairs'], [])
        self.assertEqual(book['酒店损益总表（不含名酒）']['L57'].value, before)
        book.close()

    def test_different_existing_scope_prevents_partial_rewrite(self):
        book = self.workbook()
        book['酒店损益总表（不含名酒）']['W58'] = '=W57*0.5'
        before = book['酒店损益总表（不含名酒）']['L57'].value
        result = repair_wine_total_bridge(book)
        self.assertEqual(result['status'], 'UNKNOWN')
        self.assertEqual(book['酒店损益总表（不含名酒）']['L57'].value, before)
        self.assertEqual(book['酒店损益总表（不含名酒）']['W58'].value, '=W57*0.5')
        book.close()

    def test_duplicate_source_in_aggregate_is_not_silently_subtracted_once(self):
        book = self.workbook()
        book['C其他运营汇总']['K55'] = "=ROUND(('C7商品零售'!K55+'OOD-其他'!K55+'OOD-其他'!K55+'C8水上乐园'!K55),2)"
        result = repair_wine_total_bridge(book)
        self.assertEqual(result['status'], 'UNKNOWN')
        self.assertEqual(result['repairs'], [])
        book.close()

    def test_supplementary_wine_is_a_read_only_mirror_not_duplicate_input(self):
        book = self.workbook()
        inputs = add_supplementary_sheet(book, wine_source=('OOD-其他', 55, 11))
        self.assertNotIn('D3', inputs)
        self.assertNotIn('O3', inputs)
        self.assertEqual(book['经营补充指标']['D3'].value,
                         "=IF(COUNT('OOD-其他'!K55)=1,'OOD-其他'!K55,\"\")")
        self.assertEqual(book['经营补充指标']['C3'].value,
                         '=IF(COUNT(D3:O3)=12,SUM(D3:O3),"")')
        book.close()

    def test_unknown_adjustment_is_not_overwritten_or_assumed_zero(self):
        book = self.workbook()
        book['酒店损益总表-明细表（调整用第一轮不用）']['L57'] = '=100'
        before = book['酒店损益总表（含名酒）']['L57'].value
        result = repair_wine_total_bridge(book)
        self.assertEqual(result['status'], 'UNKNOWN')
        self.assertEqual(result['repairs'], [])
        self.assertEqual(book['酒店损益总表（含名酒）']['L57'].value, before)
        book.close()


class WineZZBridgeTests(SimpleTestCase):
    def fixture(self):
        from budgeting.excel.template_v3 import validate_wine_zz_bridge
        book = WineTemplateBridgeTests().workbook()
        for name in ('C7商品零售', 'C8水上乐园'):
            book.create_sheet(name)
        included = '损益表（含名酒）（拆中智）'
        excluded = '损益表（不含名酒）（拆中智）'
        for name in (included, excluded):
            book.create_sheet(name)
        refs = []
        for month in range(12):
            col = get_column_letter(6 + month)
            src = get_column_letter(12 + month)
            for row, parent in ((48, 57), (49, 58)):
                ref = f'{col}{row}'
                refs.append(ref)
                book[included][ref] = f"=+ROUND('酒店损益总表（不含名酒）'!{src}{parent}-'酒店损益总表-明细表（调整用第一轮不用）'!{src}{parent},2)"
                book[excluded][ref] = f"=+ROUND('酒店损益总表（不含名酒）'!{src}{parent},2)"
            refs.append(f'{col}50')
            for name in (included, excluded):
                book[name][f'{col}50'] = '=123'
        reports = {code: {'mapping': [{'cell': ref} for ref in refs]}
                   for code in ('PL_ZZ_WINE', 'PL_ZZ_NOWINE')}
        bridge = repair_wine_total_bridge(book)
        return book, included, excluded, reports, bridge, validate_wine_zz_bridge

    def test_existing_wine_addback_and_shared_costs_are_validated_without_rewrite(self):
        book, yes, no, reports, bridge, validate = self.fixture()
        before = book[yes]['F48'].value
        result = validate(book, bridge, reports)
        self.assertEqual(result['status'], 'PASS')
        self.assertEqual(result['cross_scope_references'], 24)
        self.assertEqual(book[yes]['F48'].value, before)
        book.close()

    def test_missing_mandatory_bridge_is_not_treated_as_equal_blank(self):
        book, yes, no, reports, bridge, validate = self.fixture()
        book[yes]['F48'] = None
        book[no]['F48'] = None
        self.assertEqual(validate(book, bridge, reports)['status'], 'UNKNOWN')
        book.close()

    def test_unconfirmed_shared_cost_difference_is_blocked(self):
        book, yes, no, reports, bridge, validate = self.fixture()
        book[yes]['F50'] = '=234'
        self.assertEqual(validate(book, bridge, reports)['status'], 'UNKNOWN')
        book.close()

    def test_unmapped_unreferenced_scratch_does_not_block_business_rows(self):
        book, yes, no, reports, bridge, validate = self.fixture()
        book[yes]['Q100'] = '=2'
        book[yes]['S100'] = '=Q100'
        result = validate(book, bridge, reports)
        self.assertEqual(result['status'], 'PASS', result['issues'])
        self.assertEqual(result['unmapped_formula_differences'], ['Q100', 'S100'])
        book.close()

    def test_transitive_range_dependency_on_different_scratch_is_blocked(self):
        book, yes, no, reports, bridge, validate = self.fixture()
        book[yes]['Q100'] = '=2'
        for name in (yes, no):
            book[name]['F50'] = '=Z200'
            book[name]['Z200'] = '=SUM(Q99:Q101)'
        result = validate(book, bridge, reports)
        self.assertEqual(result['status'], 'UNKNOWN')
        self.assertTrue(any('辅助差异' in issue for issue in result['issues']))
        book.close()

    def test_reachable_helper_outside_display_columns_is_compared(self):
        book, yes, no, reports, bridge, validate = self.fixture()
        for name in (yes, no):
            book[name]['F50'] = '=R100'
        book[yes]['R100'] = '=2'
        book[no]['R100'] = '=3'
        self.assertEqual(validate(book, bridge, reports)['status'], 'UNKNOWN')
        book.close()

    def test_dynamic_reference_cannot_be_certified_from_equal_formula_text(self):
        book, yes, no, reports, bridge, validate = self.fixture()
        for name in (yes, no):
            book[name]['F50'] = '=INDIRECT("Q100")'
        book[yes]['Q100'] = '=2'
        book[no]['Q100'] = '=3'
        self.assertEqual(validate(book, bridge, reports)['status'], 'UNKNOWN')
        book.close()

    def test_defined_name_is_not_misread_as_a_column_reference(self):
        from openpyxl.workbook.defined_name import DefinedName
        book, yes, no, reports, bridge, validate = self.fixture()
        book.defined_names.add(DefinedName('FOO', attr_text=f"'{yes}'!Q100"))
        for name in (yes, no):
            book[name]['F50'] = '=FOO'
        self.assertEqual(validate(book, bridge, reports)['status'], 'UNKNOWN')
        book.close()

    def test_fixed_sheet_label_indirect_is_followed_without_rewriting(self):
        book, yes, no, reports, bridge, validate = self.fixture()
        detail = book.create_sheet('固定费用')
        detail['A1'] = 'OOD-其他'
        detail['B1'] = '=INDIRECT($A1&"!K55")'
        for sheet in (yes, no):
            book[sheet]['F50'] = "='固定费用'!B1"
        result = validate(book, bridge, reports)
        self.assertEqual(result['status'], 'PASS')
        self.assertEqual(result['resolved_static_indirects'], 1)
        self.assertEqual(detail['B1'].value, '=INDIRECT($A1&"!K55")')
        book.close()

    def test_editable_indirect_sheet_label_remains_unknown(self):
        from openpyxl.styles import Protection
        book, yes, no, reports, bridge, validate = self.fixture()
        detail = book.create_sheet('固定费用')
        detail['A1'] = 'OOD-其他'
        detail['A1'].protection = Protection(locked=False)
        detail['B1'] = '=INDIRECT($A1&"!K55")'
        for sheet in (yes, no):
            book[sheet]['F50'] = "='固定费用'!B1"
        self.assertEqual(validate(book, bridge, reports)['status'], 'UNKNOWN')
        book.close()

    def test_static_indirect_does_not_hide_different_split_view_dependency(self):
        book, yes, no, reports, bridge, validate = self.fixture()
        book[yes]['A1'] = yes
        book[no]['A1'] = no
        for sheet in (yes, no):
            book[sheet]['F50'] = '=INDIRECT($A1&"!R100")'
        book[yes]['R100'] = 1
        book[no]['R100'] = 2
        self.assertEqual(validate(book, bridge, reports)['status'], 'UNKNOWN')
        book.close()

    def test_sum_of_fixed_indirects_follows_both_dependencies(self):
        book, yes, no, reports, bridge, validate = self.fixture()
        detail = book.create_sheet('固定费用')
        detail['A1'] = 'OOD-其他'
        detail['B1'] = '=INDIRECT($A1&"!K55")+INDIRECT($A1&"!K69")'
        for sheet in (yes, no):
            book[sheet]['F50'] = "='固定费用'!B1"
        result = validate(book, bridge, reports)
        self.assertEqual(result['status'], 'PASS')
        self.assertEqual(result['resolved_static_indirects'], 2)
        book.close()
