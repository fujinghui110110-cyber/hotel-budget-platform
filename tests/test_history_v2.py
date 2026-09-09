from django.test import TestCase
from openpyxl import Workbook

from budgeting.excel.history import load_history, validate_history
from budgeting.models import BudgetCycle, NormalizedValue, Project, ProjectCycle, TemplateVersion, UploadVersion, ValidationRun


class HistoryV2Tests(TestCase):
    def setUp(self):
        self.project = Project.objects.create(code="H001", name="历史测试项目")
        self.cycle = BudgetCycle.objects.create(name="2026", budget_year=2026, status=BudgetCycle.Status.OPEN)
        self.template = TemplateVersion.objects.create(
            version="V2-HISTORY",
            budget_year=2026,
            file_path="template.xlsx",
            manifest_path="manifest.json",
            formula_manifest_hash="0" * 64,
        )
        self.upload = UploadVersion.objects.create(
            project=self.project,
            cycle=self.cycle,
            template=self.template,
            status=UploadVersion.Status.VALIDATED,
            original_path="upload.xlsx",
            sha256="1" * 64,
        )
        ProjectCycle.objects.create(project=self.project, cycle=self.cycle, current_upload=self.upload)

    def _layout(self, rows):
        columns = [
            {"column": chr(ord("B") + month - 1), "period": f"A2024M{month:02d}"}
            for month in range(1, 13)
        ]
        return {
            "management_v2": True,
            "history": {
                "sheet": "历史月度输入",
                "columns": columns,
                "rows": rows,
            },
        }

    def _workbook(self):
        workbook = Workbook()
        workbook.active.title = "历史月度输入"
        return workbook

    def test_empty_month_is_missing_and_undeclared_cells_are_ignored(self):
        workbook = self._workbook()
        sheet = workbook["历史月度输入"]
        for month in range(1, 13):
            if month != 2:
                sheet.cell(row=2, column=month + 1, value=month * 10)
        sheet["B3"] = 99999
        manifest = self._layout(
            [{"report_code": "PL_TOTAL_WINE", "row_code": "R0032", "row_label": "酒店总收入", "unit": "MONEY", "row_number": 2}]
        )
        count = load_history(self.upload, workbook, workbook, manifest)
        self.assertEqual(count, 11)
        self.assertEqual(NormalizedValue.objects.filter(upload=self.upload, row_code="R0032", month__isnull=False).count(), 11)
        self.assertFalse(NormalizedValue.objects.filter(upload=self.upload, row_code="R0032", period="A2024").exists())
        self.assertFalse(NormalizedValue.objects.filter(upload=self.upload, value_int=9999900).exists())

    def test_ratio_annual_uses_summed_numerator_and_denominator(self):
        workbook = self._workbook()
        sheet = workbook["历史月度输入"]
        for month in range(1, 13):
            column = month + 1
            sheet.cell(row=2, column=column, value=60)
            sheet.cell(row=3, column=column, value=100)
            sheet.cell(row=4, column=column, value=0.6)
        manifest = self._layout(
            [
                {"report_code": "PL_TOTAL_WINE", "row_code": "R0028", "row_label": "已售房晚", "unit": "COUNT", "row_number": 2},
                {"report_code": "PL_TOTAL_WINE", "row_code": "R0024", "row_label": "可售房晚", "unit": "COUNT", "row_number": 3},
                {"report_code": "PL_TOTAL_WINE", "row_code": "R0029", "row_label": "出租率 OCC", "unit": "RATIO", "metric_code": "occ", "row_number": 4},
            ]
        )
        load_history(self.upload, workbook, workbook, manifest)
        annual = NormalizedValue.objects.get(upload=self.upload, row_code="R0029", period="A2024")
        self.assertEqual(annual.ratio_num, 720)
        self.assertEqual(annual.ratio_den, 1200)
        self.assertEqual(annual.value_int, 0)

    def test_blank_derived_row_is_not_zero_filled_without_sources(self):
        workbook = self._workbook()
        manifest = self._layout(
            [
                {
                    "report_code": "PL_TOTAL_WINE",
                    "row_code": "R0028",
                    "row_label": "分子",
                    "unit": "MONEY",
                    "row_number": 2,
                },
                {
                    "report_code": "PL_TOTAL_WINE",
                    "row_code": "R0024",
                    "row_label": "分母",
                    "unit": "MONEY",
                    "row_number": 3,
                },
                {
                    "report_code": "PL_TOTAL_WINE",
                    "row_code": "R0029",
                    "row_label": "比率",
                    "unit": "RATIO",
                    "aggregation": "RATIO",
                    "numerator_row": "R0028",
                    "denominator_row": "R0024",
                    "row_number": 4,
                },
            ]
        )

        count = load_history(self.upload, workbook, workbook, manifest)

        self.assertEqual(count, 0)
        self.assertFalse(NormalizedValue.objects.filter(upload=self.upload).exists())

    def test_blank_derived_row_uses_both_source_rows(self):
        workbook = self._workbook()
        sheet = workbook["历史月度输入"]
        for month in range(1, 13):
            column = month + 1
            sheet.cell(row=2, column=column, value=60)
            sheet.cell(row=3, column=column, value=100)
        manifest = self._layout(
            [
                {
                    "report_code": "PL_TOTAL_WINE",
                    "row_code": "R0028",
                    "row_label": "分子",
                    "unit": "COUNT",
                    "row_number": 2,
                },
                {
                    "report_code": "PL_TOTAL_WINE",
                    "row_code": "R0024",
                    "row_label": "分母",
                    "unit": "COUNT",
                    "row_number": 3,
                },
                {
                    "report_code": "PL_TOTAL_WINE",
                    "row_code": "R0029",
                    "row_label": "比率",
                    "unit": "RATIO",
                    "aggregation": "RATIO",
                    "numerator_row": "R0028",
                    "denominator_row": "R0024",
                    "row_number": 4,
                },
            ]
        )

        count = load_history(self.upload, workbook, workbook, manifest)
        monthly = NormalizedValue.objects.get(upload=self.upload, row_code="R0029", period="A2024M01")
        annual = NormalizedValue.objects.get(upload=self.upload, row_code="R0029", period="A2024")

        self.assertEqual(count, 39)
        self.assertEqual((monthly.ratio_num, monthly.ratio_den, monthly.value_int), (60, 100, 0))
        self.assertEqual((annual.ratio_num, annual.ratio_den, annual.value_int), (720, 1200, 0))

    def test_existing_budget_rows_receive_dimensions(self):
        value = NormalizedValue.objects.create(
            upload=self.upload,
            report_code="PL_TOTAL_WINE",
            row_code="R0032",
            row_label="酒店总收入",
            period="01",
            unit="MONEY",
            value_int=100,
            source_sheet="损益",
            source_cell="L32",
        )
        load_history(self.upload, self._workbook(), self._workbook(), {"management_v2": True, "history": {"sheet": "历史月度输入"}})
        value.refresh_from_db()
        self.assertEqual((value.data_year, value.data_kind, value.month), (2026, "BUDGET", 1))

    def test_budget_history_columns_are_not_persisted(self):
        workbook = self._workbook()
        sheet = workbook["历史月度输入"]
        sheet["B2"] = 123
        sheet["C2"] = 456
        manifest = {
            "management_v2": True,
            "history": {
                "sheet": "历史月度输入",
                "columns": [
                    {"column": "B", "period": "B2026M01"},
                    {"column": "C", "period": "A2024M01"},
                ],
                "rows": [{
                    "report_code": "PL_TOTAL_WINE",
                    "row_code": "R0032",
                    "row_label": "酒店总收入",
                    "unit": "MONEY",
                    "row_number": 2,
                }],
            },
        }
        count = load_history(self.upload, workbook, workbook, manifest)
        self.assertEqual(count, 1)
        self.assertFalse(NormalizedValue.objects.filter(upload=self.upload, period="B2026M01").exists())
        self.assertTrue(NormalizedValue.objects.filter(upload=self.upload, period="A2024M01").exists())

    def test_validation_reports_formula_in_input_cell(self):
        values = self._workbook()
        formulas = self._workbook()
        values["历史月度输入"]["B2"] = 10
        formulas["历史月度输入"]["B2"] = "=1+1"
        manifest = self._layout(
            [{"report_code": "PL_TOTAL_WINE", "row_code": "R0032", "row_label": "酒店总收入", "unit": "MONEY", "row_number": 2}]
        )
        run = ValidationRun.objects.create(upload=self.upload, rule_version="V2")
        issues = validate_history(run, values, formulas, manifest)
        self.assertTrue(any(issue.code == "HISTORY_INPUT_FORMULA" for issue in issues))
