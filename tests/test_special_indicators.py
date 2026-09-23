import tempfile
from pathlib import Path
from decimal import Decimal

import openpyxl
from django.test import TestCase, override_settings

from budgeting.models import IndicatorProject, Project, SpecialIndicatorBatch, SpecialIndicatorValue, User
from budgeting.services.special_indicators import INDICATORS, comparison_data, confirm_import, preview_import


class SpecialIndicatorTests(TestCase):
    def test_stale_confirmation_cannot_replace_newer_batch(self):
        preview = self.preview()
        first = confirm_import(preview, source_path=self.path, expected_active_batch_ids=[])
        second = confirm_import(preview, source_path=self.path, expected_active_batch_ids=[first.pk])
        with self.assertRaisesMessage(ValueError, '当前批次已发生变化'):
            confirm_import(preview, source_path=self.path, expected_active_batch_ids=[first.pk])
        self.assertEqual(list(SpecialIndicatorBatch.objects.filter(active=True).values_list('pk', flat=True)), [second.pk])

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.settings_override = override_settings(MEDIA_ROOT=self.temp.name)
        self.settings_override.enable()
        self.addCleanup(self.settings_override.disable)
        self.path = Path(self.temp.name) / '指标.xlsx'
        workbook = openpyxl.Workbook()
        workbook.remove(workbook.active)
        for title in INDICATORS.values():
            sheet = workbook.create_sheet(title)
            sheet.append([])
            sheet.append([None, '序号', '项目'] + [f'{m}月' for m in range(1, 13)] + ['合计'])
            sheet.append([None, 1, '深圳凯骊'] + [None] * 13)
        workbook.save(self.path)
        self.project = Project.objects.create(code='SZ', name='深圳凯骊酒店')

    def preview(self, **kwargs):
        return preview_import(self.path, year=2024, data_type='ACTUAL', **kwargs)

    def edit(self, cell, value):
        workbook = openpyxl.load_workbook(self.path)
        workbook.worksheets[0][cell] = value
        workbook.save(self.path)

    def test_empty_template_preserves_null_and_link(self):
        preview = self.preview()
        self.assertTrue(preview['valid'])
        batch = confirm_import(preview, source_path=self.path)
        self.assertEqual(batch.values.count(), 52)
        self.assertFalse(batch.values.exclude(value=None).exists())
        self.assertEqual(IndicatorProject.objects.get().project, self.project)
        data = comparison_data(year=2024, data_type='ACTUAL')
        self.assertIsNone(data['rows'][0]['total'])
        self.assertFalse(data['rows'][0]['complete'])

    def test_units_and_replacement_keep_original_history(self):
        self.edit('D3', 2.5)
        first = confirm_import(self.preview(unit='WAN'), source_path=self.path)
        self.assertEqual(first.values.get(indicator='BANQUET', month=1).value, Decimal('25000'))
        self.edit('D3', 0)
        second = confirm_import(self.preview(), source_path=self.path)
        first.refresh_from_db()
        self.assertFalse(first.active)
        self.assertTrue(Path(first.original.path).exists())
        self.assertEqual(SpecialIndicatorBatch.objects.count(), 2)
        self.assertEqual(second.values.get(indicator='BANQUET', month=1).value, 0)

    def test_validation_errors_are_atomic(self):
        self.edit('D3', '#DIV/0!')
        preview = self.preview()
        self.assertFalse(preview['valid'])
        with self.assertRaises(ValueError):
            confirm_import(preview, source_path=self.path)
        self.assertEqual(SpecialIndicatorValue.objects.count(), 0)

    def test_formula_without_cache_is_not_zero(self):
        self.edit('D3', '=1+2')
        self.assertIn('公式没有缓存值', ' '.join(self.preview()['errors']))

    def test_annual_mismatch_duplicate_and_missing_sheet(self):
        workbook = openpyxl.load_workbook(self.path)
        sheet = workbook.worksheets[0]
        for col in range(4, 16):
            sheet.cell(3, col, 1)
        sheet['P3'] = 11
        sheet['C4'] = '深圳凯骊'
        del workbook['租赁收入']
        workbook.save(self.path)
        errors = ' '.join(self.preview()['errors'])
        self.assertIn('年度合计', errors)
        self.assertIn('重复', errors)
        self.assertIn('缺少工作表', errors)

    def test_scope_and_confirm_authorization(self):
        self.assertFalse(self.preview(allowed_project_ids=[])['valid'])
        user = User.objects.create_user(username='p', project=self.project)
        with self.assertRaises(PermissionError):
            confirm_import(self.preview(), source_path=self.path, user=user)

    def test_changed_source_and_budget_version_rejected(self):
        preview = self.preview()
        self.edit('D3', 12)
        with self.assertRaises(ValueError):
            confirm_import(preview, source_path=self.path)
        self.assertFalse(preview_import(self.path, year=2027, data_type='BUDGET')['valid'])
