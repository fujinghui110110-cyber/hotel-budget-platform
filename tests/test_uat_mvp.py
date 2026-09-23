import json
import shutil
import tempfile
import zipfile
from pathlib import Path
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import IntegrityError, transaction
from django.test import Client, TestCase, override_settings
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from budgeting.excel.extract import extract_report_values
from budgeting.excel.processor import extract_normalized_values
from budgeting.models import (
    AdjustmentLine,
    AuditEvent,
    BudgetCycle,
    FreezeSnapshot,
    NormalizedValue,
    ProcessingJob,
    Project,
    ProjectCycle,
    SnapshotArtifact,
    TemplateVersion,
    UploadVersion,
    ValidationIssue,
    ValidationRun,
)
from budgeting.services.files import check_xlsx_package, sha256_file
from budgeting.services.money import aggregate_occ, cents_to_yuan
from budgeting.services.workflow import (
    approve_upload,
    company_values,
    create_adjustment_batch,
    freeze_cycle,
    freeze_preconditions,
    issue_adjustment,
    process_upload,
    reopen_cycle,
    save_upload,
    submit_upload,
)


class UatMvpTests(TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.storage = self.root / "storage"
        self.storage.mkdir()
        self.settings = override_settings(
            BASE_DIR=self.root,
            BUDGET_STORAGE_ROOT=self.storage,
            SOFFICE_BIN=str(self.root / "soffice"),
        )
        self.settings.enable()
        self.addCleanup(self.settings.disable)

        self.User = get_user_model()
        self.project = Project.objects.create(code="P001", name="项目一")
        self.other_project = Project.objects.create(code="P002", name="项目二")
        self.cycle = BudgetCycle.objects.create(name="2026预算", budget_year=2026, status=BudgetCycle.OPEN)
        self.template_file = self._workbook("template.xlsx")
        self.manifest_file = self._manifest("manifest.json")
        self.template = TemplateVersion.objects.create(
            version="V1",
            budget_year=2026,
            file_path=self.template_file.name,
            manifest_path=self.manifest_file.name,
            formula_manifest_hash="0" * 64,
        )
        self.admin = self.User.objects.create_user("admin", password="admin-pass", role="ADMIN", is_staff=True, is_superuser=True)
        self.user = self.User.objects.create_user("p001", password="project-pass", role="PROJECT", project=self.project)
        self.other_user = self.User.objects.create_user("p002", password="project-pass", role="PROJECT", project=self.other_project)

    def test_uat_01_redirects_anonymous_home_to_login_without_project_data(self):
        response = Client().get("/")
        self.assertEqual(response.status_code, 302)
        self.assertIn("/login/", response["Location"])
        self.assertNotContains(response, "项目一", status_code=302)

    def test_uat_02_logs_admin_into_management_scope(self):
        client = Client()
        self.assertTrue(client.login(username="admin", password="admin-pass"))
        response = client.get("/")
        self.assertRedirects(response, reverse("planning_overview"), fetch_redirect_response=False)

    def test_uat_03_rejects_wrong_password_without_session(self):
        client = Client()
        response = client.post("/login/", {"username": "admin", "password": "wrong"})
        self.assertEqual(response.status_code, 200)
        self.assertNotIn("_auth_user_id", client.session)

    def test_uat_04_allows_first_superuser_admin_login_without_password_audit(self):
        self.User.objects.create_superuser("first-admin", password="new-secret")
        client = Client()
        self.assertTrue(client.login(username="first-admin", password="new-secret"))
        self.assertFalse(AuditEvent.objects.filter(payload__icontains="new-secret").exists())

    def test_uat_05_enforces_unique_project_cycle_and_template_codes(self):
        self.assertEqual(Project.objects.get(code="P001").name, "项目一")
        self.assertEqual(BudgetCycle.objects.get(budget_year=2026, revision_no=1), self.cycle)
        self.assertEqual(TemplateVersion.objects.get(version="V1"), self.template)
        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                Project.objects.create(code="P001", name="重复项目")

    def test_uat_06_requires_project_accounts_to_bind_exactly_one_project(self):
        unbound = self.User(username="unbound", role="PROJECT")
        with self.assertRaises(ValidationError):
            unbound.full_clean()
        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                self.User.objects.create_user("also-p001", role="PROJECT", project=self.project)

    def test_uat_07_denies_other_project_url_api_and_snapshot_file_access(self):
        upload = self._upload(self.other_project, UploadVersion.APPROVED, cents=12345)
        snapshot = FreezeSnapshot.objects.create(cycle=self.cycle, status=FreezeSnapshot.COMPLETE)
        client = Client()
        client.force_login(self.user)
        detail = client.get(reverse("project_upload_detail", args=[upload.id]))
        api = client.get(reverse("upload_status_api", args=[upload.id]))
        download = client.get(reverse("management_snapshot_download", args=[snapshot.id]))
        self.assertEqual((detail.status_code, api.status_code, download.status_code), (403, 403, 404))

    def test_management_project_page_lists_current_uploads_and_admin_opens_details(self):
        upload = self._upload(self.project, UploadVersion.APPROVED, cents=12345)
        ProjectCycle.objects.create(project=self.project, cycle=self.cycle, current_upload=upload)
        client = self._admin_client()

        response = client.get(reverse("management_projects"))
        self.assertContains(response, self.project.name)
        self.assertContains(response, upload.original_name)
        self.assertEqual(
            client.get(reverse("project_upload_detail", args=[upload.id])).status_code,
            200,
        )

    def test_uat_08_accepts_valid_xlsx_upload_and_tracks_hash_job_and_status(self):
        client = Client()
        client.force_login(self.user)
        data = self.template_file.read_bytes()
        response = client.post(reverse("project_upload_new"), {
            "file": SimpleUploadedFile(
                "budget.xlsx",
                data,
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        })
        upload = UploadVersion.objects.get(project=self.project)
        self.assertRedirects(response, reverse("project_upload_detail", args=[upload.id]), fetch_redirect_response=False)
        self.assertEqual(upload.sha256, sha256_file(self.storage / upload.original_path))
        self.assertEqual(ProcessingJob.objects.get(upload=upload).status, ProcessingJob.Status.DONE)

    def test_uat_09_rejects_non_xlsx_before_queueing(self):
        client = Client()
        client.force_login(self.user)
        response = client.post(reverse("project_upload_new"), {"file": SimpleUploadedFile("budget.txt", b"not excel")})
        self.assertEqual(response.status_code, 200)
        self.assertFalse(UploadVersion.objects.exists())
        self.assertFalse(ProcessingJob.objects.exists())

    def test_uat_10_rejects_corrupt_xlsx_without_forged_results(self):
        upload = self._stored_upload("broken.xlsx", b"not a zip file")
        process_upload(upload)
        upload.refresh_from_db()
        self.assertEqual(upload.status, UploadVersion.Status.REJECTED)
        self.assertTrue(ValidationIssue.objects.filter(run__upload=upload, severity=ValidationIssue.Severity.P0).exists())
        self.assertFalse(NormalizedValue.objects.filter(upload=upload).exists())

    def test_uat_11_keeps_duplicate_same_hash_upload_idempotent_and_traceable(self):
        data = self.template_file.read_bytes()
        first = save_upload(self.project, self.cycle, SimpleUploadedFile("budget.xlsx", data))
        second = save_upload(self.project, self.cycle, SimpleUploadedFile("budget.xlsx", data))
        self.assertEqual(first.sha256, second.sha256)
        self.assertEqual(UploadVersion.objects.filter(project=self.project, cycle=self.cycle, sha256=first.sha256).count(), 1)
        self.assertEqual(ProcessingJob.objects.filter(upload__sha256=first.sha256).count(), 1)

    def test_uat_12_rejects_wrong_template_year_version_or_signature_without_replacing_current(self):
        current = self._current_upload(self.project, cents=10000)
        bad_template = TemplateVersion.objects.create(
            version="V0",
            budget_year=2025,
            file_path=self.template_file.name,
            manifest_path=self.manifest_file.name,
            formula_manifest_hash="bad",
        )
        upload = self._upload(self.project, UploadVersion.RECEIVED, template=bad_template)
        shutil.copy2(self.template_file, self.storage / upload.original_path)
        upload.sha256 = sha256_file(self.storage / upload.original_path)
        upload.save(update_fields=["sha256"])
        with patch("budgeting.services.workflow.recalc_with_libreoffice", return_value=self.template_file):
            process_upload(upload)
        current.refresh_from_db()
        upload.refresh_from_db()
        self.assertEqual(upload.status, UploadVersion.Status.REJECTED)
        self.assertTrue(ValidationIssue.objects.filter(run__upload=upload, code="TEMPLATE_SIGNATURE").exists())
        self.assertEqual(ProjectCycle.objects.get(project=self.project, cycle=self.cycle).current_upload, current)

    def test_uat_13_blocks_missing_required_report_sheet_with_location(self):
        workbook = self._workbook("missing-sheet.xlsx", include_reports=["酒店损益总表（含名酒）"])
        upload = self._stored_upload("missing-sheet.xlsx", workbook.read_bytes())
        with patch("budgeting.services.workflow.recalc_with_libreoffice", return_value=workbook):
            process_upload(upload)
        self.assertTrue(ValidationIssue.objects.filter(run__upload=upload, code="MISSING_REQUIRED_SHEET", location__icontains="酒店损益总表（不含名酒）").exists())
        self.assertEqual(UploadVersion.objects.get(pk=upload.pk).status, UploadVersion.Status.REJECTED)

    def test_uat_14_blocks_formula_overwrite_with_actual_expected_location_and_rule_version(self):
        workbook = self._workbook("formula-overwrite.xlsx", formula_overwrite=True)
        upload = self._stored_upload("formula-overwrite.xlsx", workbook.read_bytes())
        with patch("budgeting.services.workflow.recalc_with_libreoffice", return_value=workbook):
            process_upload(upload)
        issue = ValidationIssue.objects.get(
            run__upload=upload,
            code="FORMULA_CHANGED",
            location="酒店损益总表（含名酒）!C3",
        )
        self.assertEqual(issue.severity, ValidationIssue.Severity.P0)
        self.assertIn("酒店损益总表（含名酒）!C3", issue.location)
        self.assertTrue(issue.actual_value)
        self.assertTrue(issue.expected_value)
        self.assertEqual(issue.run.rule_version, self.template.rule_version)

    def test_plan_uat_05_blocks_ref_error_value_before_recalculation(self):
        workbook = self._workbook("ref-error.xlsx")
        editable = load_workbook(workbook)
        editable["酒店损益总表（含名酒）"]["C3"] = "#REF!"
        editable.save(workbook)
        upload = self._stored_upload("ref-error.xlsx", workbook.read_bytes())

        process_upload(upload)

        upload.refresh_from_db()
        self.assertEqual(upload.status, UploadVersion.Status.REJECTED)
        self.assertTrue(
            ValidationIssue.objects.filter(
                run__upload=upload,
                code="EXCEL_ERROR",
                location__contains="酒店损益总表（含名酒）!C3",
            ).exists()
        )

    def test_uat_15_preserves_money_count_ratio_units_periods_and_lineage_without_float_drift(self):
        manifest = self._manifest("mixed-units.json", unit_rows=True)
        template = TemplateVersion.objects.create(
            version="V1-mixed",
            budget_year=2026,
            file_path=self.template_file.name,
            manifest_path=manifest.name,
            formula_manifest_hash="1" * 64,
        )
        workbook = self._workbook("mixed-units.xlsx", mixed_units=True)
        upload = self._upload(self.project, UploadVersion.PROCESSING, template=template)
        extract_normalized_values(upload, workbook)
        money = NormalizedValue.objects.get(upload=upload, report_code="PL_TOTAL_WINE", row_code="MONEY_REV", period="01")
        count = NormalizedValue.objects.get(upload=upload, report_code="PL_TOTAL_WINE", row_code="ROOM_NIGHTS", period="01")
        ratio = NormalizedValue.objects.get(upload=upload, report_code="PL_TOTAL_WINE", row_code="OCC", period="01")
        self.assertEqual((money.unit, money.value_int, money.source_sheet, money.source_cell), (NormalizedValue.MONEY, 124, "酒店损益总表（含名酒）", "C3"))
        self.assertEqual((count.unit, count.value_int), (NormalizedValue.COUNT, 5))
        self.assertEqual((ratio.unit, ratio.ratio_num, ratio.ratio_den), (NormalizedValue.RATIO, 5, 10))

    def test_plan_uat_09_blocks_monthly_occ_cache_that_differs_from_parts(self):
        manifest = self._manifest("monthly-occ.json", unit_rows=True)
        template = TemplateVersion.objects.create(
            version="V1-monthly-occ",
            budget_year=2026,
            file_path=self.template_file.name,
            manifest_path=manifest.name,
            formula_manifest_hash="2" * 64,
        )
        workbook = self._workbook("monthly-occ.xlsx", mixed_units=True)
        editable = load_workbook(workbook)
        for sheet_name in (
            "酒店损益总表（含名酒）",
            "酒店损益总表（不含名酒）",
            "损益表（含名酒）（拆中智）",
            "损益表（不含名酒）（拆中智）",
        ):
            editable[sheet_name]["C5"] = 0.51
        editable.save(workbook)
        upload = self._upload(
            self.project,
            UploadVersion.PROCESSING,
            template=template,
        )
        run = ValidationRun.objects.create(upload=upload, rule_version="R1")

        extract_report_values(upload, workbook, validation_run=run)

        issue = ValidationIssue.objects.filter(
            run=run,
            code="MONTHLY_RATIO_RECONCILIATION",
            location="酒店损益总表（含名酒）!C5",
        ).get()
        self.assertEqual(issue.severity, ValidationIssue.Severity.P0)
        self.assertEqual(issue.actual_value, "0.51")
        self.assertEqual(issue.expected_value, "5/10")

    def test_plan_uat_08_blocks_one_cent_parent_total_formula_override(self):
        workbook = self._workbook("parent-total-overwrite.xlsx")
        editable = load_workbook(workbook)
        editable["酒店损益总表（含名酒）"]["C3"] = 1.25
        editable.save(workbook)
        upload = self._stored_upload(
            "parent-total-overwrite.xlsx",
            workbook.read_bytes(),
        )

        process_upload(upload)

        issue = ValidationIssue.objects.get(
            run__upload=upload,
            code="FORMULA_CHANGED",
            location="酒店损益总表（含名酒）!C3",
        )
        self.assertEqual(issue.severity, ValidationIssue.Severity.P0)
        self.assertTrue(issue.actual_value)
        self.assertTrue(issue.expected_value)

    def test_uat_16_worker_claims_queue_and_records_done_or_failed_terminal_status(self):
        job = ProcessingJob.objects.create(upload=self._upload(self.project, UploadVersion.RECEIVED), idempotency_key="upload:uat16")
        with patch("budgeting.management.commands.budget_worker.process_upload") as fake_process:
            from budgeting.management.commands.budget_worker import Command

            claimed = Command()._claim_job(lease_seconds=30)
            fake_process.side_effect = lambda upload: UploadVersion.objects.filter(pk=upload.pk).update(status=UploadVersion.VALIDATED)
            Command()._run_job(claimed)
        job.refresh_from_db()
        self.assertEqual(job.status, ProcessingJob.Status.DONE)

    def test_uat_17_requeues_expired_running_worker_job_for_restart_recovery(self):
        from django.utils import timezone
        from datetime import timedelta
        from budgeting.management.commands.budget_worker import Command

        job = ProcessingJob.objects.create(
            upload=self._upload(self.project, UploadVersion.PROCESSING),
            idempotency_key="upload:uat17",
            status=ProcessingJob.Status.RUNNING,
            lease_until=timezone.now() - timedelta(seconds=1),
        )
        claimed = Command()._claim_job(lease_seconds=30)
        self.assertEqual(claimed.id, job.id)
        self.assertEqual(claimed.status, ProcessingJob.Status.RUNNING)

    def test_uat_18_healthz_reports_storage_libreoffice_and_worker_status(self):
        ProcessingJob.objects.create(upload=self._upload(self.project), idempotency_key="upload:uat18")
        response = Client().get(reverse("healthz"))
        payload = response.json()
        self.assertEqual(response.status_code, 200)
        self.assertTrue(payload["storage"])
        self.assertIn("libreoffice", payload)
        self.assertIn("worker_last_seen", payload)

    def test_uat_19_recalculation_uses_copy_and_preserves_original_hash(self):
        upload = self._stored_upload("original.xlsx", self.template_file.read_bytes())
        before = sha256_file(self.storage / upload.original_path)
        recalc = self._workbook("recalculated.xlsx", amount=88)
        with patch("budgeting.services.workflow.recalc_with_libreoffice", return_value=recalc):
            process_upload(upload)
        upload.refresh_from_db()
        self.assertEqual(sha256_file(self.storage / upload.original_path), before)
        self.assertNotEqual(Path(upload.recalculated_path), Path(upload.original_path))

    def test_uat_20_freeze_exports_all_four_management_reports_with_correct_periods(self):
        upload = self._upload(self.project, UploadVersion.APPROVED, original_path="uploads/P001/four-reports.xlsx")
        (self.storage / "uploads/P001").mkdir(parents=True, exist_ok=True)
        shutil.copy2(self.template_file, self.storage / upload.original_path)
        upload.sha256 = sha256_file(self.storage / upload.original_path)
        upload.save(update_fields=["sha256"])
        for code in ("PL_TOTAL_WINE", "PL_TOTAL_NOWINE", "PL_ZZ_WINE", "PL_ZZ_NOWINE"):
            NormalizedValue.objects.create(upload=upload, report_code=code, row_code="R0003", period="01", unit=NormalizedValue.MONEY, value_int=100, source_sheet="s", source_cell="A1")
        ProjectCycle.objects.create(project=self.project, cycle=self.cycle, current_upload=upload)
        Project.objects.filter(pk=self.other_project.pk).update(is_active=False)
        snapshot = freeze_cycle(self.cycle, self.admin)
        report = json.loads((self.storage / snapshot.directory / "四表汇总.json").read_text(encoding="utf-8"))
        self.assertEqual(set(report), {"PL_TOTAL_WINE", "PL_TOTAL_NOWINE", "PL_ZZ_WINE", "PL_ZZ_NOWINE"})
        self.assertEqual(
            report["PL_TOTAL_WINE"]["R0003|01"],
            {
                "unit": NormalizedValue.Unit.MONEY,
                "value_int": 100,
                "ratio_num": None,
                "ratio_den": None,
            },
        )

    def test_company_aggregation_recomputes_derived_money_from_project_ratios(self):
        for project, ratio_num, ratio_den, adr_cents in (
            (self.project, 10_000, 10, 1_000),
            (self.other_project, 60_000, 30, 2_000),
        ):
            upload = self._upload(project, UploadVersion.APPROVED)
            NormalizedValue.objects.create(
                upload=upload,
                report_code="PL_TOTAL_WINE",
                row_code="ADR",
                period="01",
                unit=NormalizedValue.Unit.MONEY,
                value_int=adr_cents,
                ratio_num=ratio_num,
                ratio_den=ratio_den,
                source_sheet="s",
                source_cell="A1",
            )
            ProjectCycle.objects.update_or_create(
                project=project,
                cycle=self.cycle,
                defaults={"current_upload": upload},
            )

        self.assertEqual(
            company_values(self.cycle, "PL_TOTAL_WINE")[("ADR", "01")],
            1_750,
        )

    def test_ratio_aggregation_rounds_displays_and_exports_with_unit(self):
        for project, ratio_num, ratio_den in (
            (self.project, 1, 3),
            (self.other_project, 10, 27),
        ):
            upload = self._upload(project, UploadVersion.APPROVED, original_path=f"uploads/{project.code}/original.xlsx")
            NormalizedValue.objects.create(
                upload=upload,
                report_code="PL_TOTAL_WINE",
                row_code="OCC",
                period="01",
                unit=NormalizedValue.Unit.RATIO,
                value_int=0,
                ratio_num=ratio_num,
                ratio_den=ratio_den,
                source_sheet="s",
                source_cell="A1",
            )
            ProjectCycle.objects.update_or_create(
                project=project,
                cycle=self.cycle,
                defaults={"current_upload": upload},
            )
            source = self.storage / upload.original_path
            source.parent.mkdir(parents=True, exist_ok=True)
            Workbook().save(source)
            upload.sha256 = sha256_file(source)
            upload.save(update_fields=["sha256"])

        report = self._admin_client().get(reverse("management_report", args=["PL_TOTAL_WINE"]))
        self.assertContains(report, "36.67%")
        drilldown = self._admin_client().get(
            reverse("management_report_drilldown", args=["PL_TOTAL_WINE"]),
            {"row_code": "OCC", "period": "01"},
        )
        self.assertContains(drilldown, "36.67%")

        snapshot = freeze_cycle(self.cycle, self.admin)
        root = self.storage / snapshot.directory
        payload = json.loads((root / "四表汇总.json").read_text(encoding="utf-8"))
        self.assertEqual(
            payload["PL_TOTAL_WINE"]["OCC|01"],
            {
                "unit": NormalizedValue.Unit.RATIO,
                "value_int": 3667,
                "ratio_num": 11,
                "ratio_den": 30,
            },
        )
        sheet = load_workbook(root / "四表汇总.xlsx", data_only=True)["PL_TOTAL_WINE"]
        headers = [cell.value for cell in sheet[2]]
        self.assertIn("unit", headers)
        self.assertNotIn("value_cents", headers)
        rows = {
            (row[0].value, row[1].value): [cell.value for cell in row]
            for row in sheet.iter_rows(min_row=3)
        }
        self.assertEqual(rows[("OCC", "01")][headers.index("unit")], NormalizedValue.Unit.RATIO)
        self.assertEqual(rows[("OCC", "01")][headers.index("value_int")], 3667)

    def test_uat_21_aggregates_monthly_and_annual_occ_by_numerator_denominator(self):
        month_num, month_den, month_ratio = aggregate_occ([(1, 3), (2, 3)])
        annual_num, annual_den, annual_ratio = aggregate_occ([(month_num, month_den), (3, 6)])
        self.assertEqual((month_num, month_den, str(month_ratio)), (3, 6, "0.5"))
        self.assertEqual((annual_num, annual_den, str(annual_ratio)), (6, 12, "0.5"))

    def test_uat_22_blocks_p0_submit_and_approval_switches_current_version(self):
        blocked = self._upload(self.project, UploadVersion.VALIDATED)
        run = ValidationRun.objects.create(upload=blocked, rule_version="R1", passed=False)
        ValidationIssue.objects.create(run=run, severity=ValidationIssue.Severity.P0, code="P0", message="blocked")
        with self.assertRaises(ValueError):
            submit_upload(blocked, self.user)
        replacement = self._upload(self.project, UploadVersion.VALIDATED, cents=222)
        submit_upload(replacement, self.user)
        approve_upload(replacement, self.admin)
        self.assertEqual(ProjectCycle.objects.get(project=self.project, cycle=self.cycle).current_upload, replacement)

    def test_uat_23_legacy_adjustment_review_blocks_freeze_and_explicit_targets_still_block_approval(self):
        self._current_upload(self.project, cents=10000)
        Project.objects.filter(pk=self.other_project.pk).update(is_active=False)
        batch = create_adjustment_batch(self.cycle, "PL_TOTAL_WINE", "ROOM", "01", 1, "rounding", self.admin)
        issue_adjustment(batch, self.admin)
        blockers = freeze_preconditions(self.cycle)
        self.assertTrue(any("旧下发任务" in blocker for blocker in blockers))
        new_upload = self._upload(self.project, UploadVersion.VALIDATED, cents=10001, original_path="uploads/P001/adjusted.xlsx")
        (self.storage / "uploads/P001").mkdir(parents=True, exist_ok=True)
        shutil.copy2(self.template_file, self.storage / new_upload.original_path)
        new_upload.sha256 = sha256_file(self.storage / new_upload.original_path)
        new_upload.save(update_fields=["sha256"])
        submit_upload(new_upload, self.user)
        approve_upload(new_upload, self.admin)
        freeze_cycle(self.cycle, self.admin)
        with self.assertRaises(ValueError):
            save_upload(self.project, self.cycle, SimpleUploadedFile("late.xlsx", self.template_file.read_bytes()))
        reopened = reopen_cycle(self.cycle, self.admin, project_ids=[self.project.id])
        self.assertEqual(reopened.revision_no, 2)

    def test_uat_24_snapshot_manifest_zip_and_database_artifact_hashes_align_after_restore_copy(self):
        upload = self._current_upload(self.project, original_path="uploads/P001/source.xlsx", cents=10000)
        Project.objects.filter(pk=self.other_project.pk).update(is_active=False)
        (self.storage / "uploads/P001").mkdir(parents=True, exist_ok=True)
        shutil.copy2(self.template_file, self.storage / "uploads/P001/source.xlsx")
        upload.sha256 = sha256_file(self.storage / upload.original_path)
        upload.save(update_fields=["sha256"])
        snapshot = freeze_cycle(self.cycle, self.admin)
        manifest = json.loads((self.storage / snapshot.manifest_path).read_text(encoding="utf-8"))
        copied_root = self.root / "restore"
        shutil.copytree(self.storage, copied_root)
        artifact = SnapshotArtifact.objects.get(snapshot=snapshot, relative_path=f"{snapshot.id}.zip")
        self.assertNotIn(f"{snapshot.id}.zip", manifest)
        self.assertEqual(sha256_file(copied_root / f"{snapshot.id}.zip"), artifact.sha256)

    def test_uat_25_audit_events_include_actor_object_action_time_and_no_password_or_token(self):
        upload = self._upload(self.project, UploadVersion.VALIDATED, cents=1)
        submit_upload(upload, self.user)
        approve_upload(upload, self.admin)
        event = AuditEvent.objects.get(action="UPLOAD_APPROVED")
        payload = json.dumps(event.payload, ensure_ascii=False)
        self.assertEqual(event.actor, self.admin)
        self.assertEqual(event.upload, upload)
        self.assertTrue(event.created_at)
        self.assertNotIn("admin-pass", payload)
        self.assertNotIn("project-pass", payload)
        self.assertNotIn("token", payload.lower())

    def test_external_links_are_blocked_as_p0_before_normalization(self):
        workbook = self._workbook("external-link.xlsx")
        with zipfile.ZipFile(workbook, "a") as zf:
            zf.writestr("xl/externalLinks/externalLink1.xml", "<externalLink />")
        with self.assertRaises(ValidationError):
            check_xlsx_package(workbook)

    def test_one_cent_batch_difference_blocks_issue(self):
        self._current_upload(self.project, cents=10000)
        batch = create_adjustment_batch(self.cycle, "PL_TOTAL_WINE", "ROOM", "01", 1, "rounding", self.admin)
        AdjustmentLine.objects.filter(batch=batch).update(allocated_delta_cents=0)
        with self.assertRaises(ValueError):
            issue_adjustment(batch, self.admin)

    def test_one_cent_reupload_difference_remains_open_for_project_correction(self):
        current = self._current_upload(self.project, cents=10000)
        batch = create_adjustment_batch(self.cycle, "PL_TOTAL_WINE", "ROOM", "01", 1, "rounding", self.admin)
        issue_adjustment(batch, self.admin)
        from budgeting.services.plan_history import ensure_plan
        from budgeting.services.targets import issue_targets
        plan = ensure_plan(self.cycle)
        issue_targets(
            actor=self.admin, plan=plan, project=self.project, origin_cycle=self.cycle,
            expected_plan_revision=plan.revision_token, reason="客房收入最低要求",
            selected_target_rows=[{
                "report_code": "PL_TOTAL_WINE", "row_code": "ROOM", "period": "01",
                "unit": "CNY_CENT", "metric_kind": "REVENUE", "comparator": "GE",
                "target_int": 10001, "sign_multiplier": 1,
                "evidence": "管理员明确选定客房收入最低要求", "rule_version": "test-v1",
            }],
        )
        reupload = self._upload(self.project, UploadVersion.VALIDATED, cents=10000)
        submit_upload(reupload, self.user)
        with self.assertRaisesMessage(ValueError, "TARGET_UNMET"):
            approve_upload(reupload, self.admin)
        self.assertEqual(ProjectCycle.objects.get(project=self.project, cycle=self.cycle).current_upload, current)
        line = AdjustmentLine.objects.get(batch=batch, project=self.project)
        self.assertEqual(line.difference_cents, -1)
        self.assertEqual(line.status, AdjustmentLine.Status.OPEN)

    def test_failed_new_version_retains_old_approved_current_version(self):
        current = self._current_upload(self.project, cents=10000)
        failed = self._upload(self.project, UploadVersion.REJECTED, cents=99999)
        with self.assertRaises(ValueError):
            approve_upload(failed, self.admin)
        self.assertEqual(ProjectCycle.objects.get(project=self.project, cycle=self.cycle).current_upload, current)

    def test_hundred_yuan_three_way_adjustment_allocates_3334_3333_3333_cents(self):
        projects = [self.project, self.other_project, Project.objects.create(code="P003", name="项目三")]
        for project in projects:
            self._current_upload(project, cents=0)
        batch = create_adjustment_batch(self.cycle, "PL_TOTAL_WINE", "ROOM", "01", 10000, "split", self.admin)
        self.assertEqual(
            list(batch.lines.order_by("project__code").values_list("allocated_delta_cents", flat=True)),
            [3334, 3333, 3333],
        )

    def test_management_report_uses_approved_current_uploads_only(self):
        stale = self._upload(self.project, UploadVersion.APPROVED)
        current = self._upload(self.project, UploadVersion.APPROVED, cents=200)
        ProjectCycle.objects.create(project=self.project, cycle=self.cycle, current_upload=current)
        NormalizedValue.objects.create(upload=stale, report_code="PL_TOTAL_WINE", row_code="ROOM", period="01", unit=NormalizedValue.MONEY, value_int=100, source_sheet="old", source_cell="A1")
        response = self._admin_client().get(reverse("management_report", args=["PL_TOTAL_WINE"]))
        self.assertContains(response, "2.00")
        self.assertNotContains(response, "1.00")
        self.assertEqual(response.context["period_scope"], "budget")
        self.assertNotContains(response, 'aria-label="报表摘要"')
        self.assertContains(response, "row_code=ROOM&amp;period=01")

    def test_management_drilldown_total_equals_report_total_and_includes_lineage(self):
        upload = self._current_upload(self.project, report_code="PL_TOTAL_WINE", row_code="ROOM", period="01", cents=12345, source_sheet="房务", source_cell="C7")
        report = self._admin_client().get(reverse("management_report", args=["PL_TOTAL_WINE"]))
        drill = self._admin_client().get(reverse("management_report_drilldown", args=["PL_TOTAL_WINE"]), {"row_code": "ROOM", "period": "01"})
        unfiltered = self._admin_client().get(reverse("management_report_drilldown", args=["PL_TOTAL_WINE"]))
        self.assertContains(report, "123.45")
        self.assertContains(drill, "123.45")
        self.assertContains(drill, self.project.code)
        self.assertContains(drill, str(upload.id))
        self.assertContains(drill, "房务")
        self.assertContains(drill, "C7")
        self.assertContains(unfiltered, "暂无可穿透的项目来源")
        self.assertNotContains(unfiltered, str(upload.id))

    def test_ten_thousand_yuan_display_does_not_change_stored_cents(self):
        upload = self._current_upload(self.project, cents=123456789)
        displayed = cents_to_yuan(upload.normalizedvalue_set.get().value_int) / 10000
        self.assertEqual(upload.normalizedvalue_set.get().value_int, 123456789)
        self.assertEqual(str(displayed), "123.456789")

    def _admin_client(self):
        client = Client()
        client.force_login(self.admin)
        return client

    def _stored_upload(self, filename, data):
        rel = Path("uploads") / self.project.code / filename
        path = self.storage / rel
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(data)
        return UploadVersion.objects.create(
            project=self.project,
            cycle=self.cycle,
            template=self.template,
            status=UploadVersion.RECEIVED,
            original_name=filename,
            original_path=str(rel),
            sha256=sha256_file(path),
        )

    def _upload(self, project, status=UploadVersion.VALIDATED, *, template=None, report_code="PL_TOTAL_WINE", row_code="ROOM", period="01", cents=None, original_path="original.xlsx", source_sheet="Sheet1", source_cell="A1"):
        upload = UploadVersion.objects.create(
            project=project,
            cycle=self.cycle,
            template=template or self.template,
            status=status,
            original_path=original_path,
            sha256=f"{UploadVersion.objects.count():064x}"[-64:],
        )
        if cents is not None:
            NormalizedValue.objects.create(
                upload=upload,
                report_code=report_code,
                row_code=row_code,
                period=period,
                unit=NormalizedValue.MONEY,
                value_int=cents,
                source_sheet=source_sheet,
                source_cell=source_cell,
            )
        return upload

    def _current_upload(self, project, **kwargs):
        upload = self._upload(project, UploadVersion.APPROVED, **kwargs)
        ProjectCycle.objects.update_or_create(project=project, cycle=self.cycle, defaults={"current_upload": upload})
        return upload

    def _manifest(self, name, unit_rows=False):
        path = self.root / name
        mapping = [
            {"row_code": "MONEY_REV", "row_label": "金额", "period": "01", "cell": "C3", "unit": "MONEY"},
            {"row_code": "ROOM_NIGHTS", "row_label": "间夜", "period": "01", "cell": "C4", "unit": "COUNT"},
            {"row_code": "OCC", "row_label": "OCC", "period": "01", "cell": "C5", "unit": "RATIO"},
        ] if unit_rows else [
            {"row_code": "ROOM", "row_label": "客房收入", "period": "01", "cell": "C3", "unit": "MONEY"},
        ]
        manifest = {
            "template_version": "V1",
            "reports": {code: {"sheet": sheet, "mapping": mapping} for code, sheet in {
                "PL_TOTAL_WINE": "酒店损益总表（含名酒）",
                "PL_TOTAL_NOWINE": "酒店损益总表（不含名酒）",
                "PL_ZZ_WINE": "损益表（含名酒）（拆中智）",
                "PL_ZZ_NOWINE": "损益表（不含名酒）（拆中智）",
            }.items()},
        }
        path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
        return path

    def _workbook(self, name, *, include_reports=None, amount=1.24, formula_overwrite=False, mixed_units=False):
        include_reports = include_reports or ["酒店损益总表（含名酒）", "酒店损益总表（不含名酒）", "损益表（含名酒）（拆中智）", "损益表（不含名酒）（拆中智）"]
        wb = Workbook()
        wb.remove(wb.active)
        for sheet in include_reports:
            ws = wb.create_sheet(sheet)
            ws["A1"] = "科目"
            ws["C1"] = "01"
            ws["A3"] = "客房收入"
            ws["C3"] = amount if formula_overwrite or mixed_units else "=1.24"
            if mixed_units:
                ws["A4"] = "售出间夜"
                ws["C4"] = 5
                ws["A5"] = "OCC"
                ws["C5"] = 0.5
                ws["A24"] = "总可售房晚"
                ws["C24"] = 10
                ws["A28"] = "已售房晚"
                ws["C28"] = 5
        meta = wb.create_sheet("SYS_META")
        meta.sheet_state = "hidden"
        meta["A1"] = "template_version"
        meta["B1"] = "V1"
        meta["A2"] = "budget_year"
        meta["B2"] = 2026
        meta["A3"] = "project_code"
        meta["B3"] = self.project.code if hasattr(self, "project") else "P001"
        meta["A6"] = "project_signature"
        meta["B6"] = "signed"
        path = self.root / name
        wb.save(path)
        return path
