from __future__ import annotations

import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import Workbook

from budgeting.services.management_metric_parser import parse_workbook


class ManagementMetricParserTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)

    def _path(self, name: str = "metrics.xlsx") -> Path:
        return Path(self.temp.name) / name

    @staticmethod
    def _sheet(
        workbook: Workbook,
        title: str,
        metric_title: str,
        year: int,
        *,
        projects: tuple[str, ...] = ("深圳凯骊",),
        annual: bool = True,
        header_row: int = 3,
    ):
        sheet = workbook.create_sheet(title)
        sheet.cell(row=header_row - 2, column=1, value=metric_title)
        sheet.cell(row=header_row - 1, column=1, value="实际")
        sheet.cell(row=header_row, column=1, value="项目")
        for month in range(1, 13):
            sheet.cell(row=header_row, column=month + 1, value=year * 100 + month)
        if annual:
            sheet.cell(row=header_row, column=14, value="合计")
        for offset, name in enumerate(projects, header_row + 1):
            sheet.cell(row=offset, column=1, value=name)
        return sheet, header_row + 1

    def test_multi_year_dynamic_headers_and_zero_missing_distinction(self):
        workbook = Workbook()
        workbook.remove(workbook.active)
        sheet, row = self._sheet(workbook, "收入（2027年）", "总收入", 2027)
        sheet.cell(row=row, column=2, value=0)
        sheet.cell(row=row, column=3, value=None)
        sheet.cell(row=row, column=4, value=2)
        sheet.cell(row=row, column=14, value=2)
        second, second_row = self._sheet(workbook, "收入（2028年）", "总收入", 2028)
        second.cell(row=second_row, column=2, value=3)
        second.cell(row=second_row, column=14, value=3)
        path = self._path()
        workbook.save(path)

        result = parse_workbook(path, money_unit="WAN", data_kind="ACTUAL")

        self.assertTrue(result["valid"], result["errors"])
        self.assertEqual(result["years"], [2027, 2028])
        self.assertEqual(result["projects"], ["深圳凯骊"])
        rows = {(row["year"], row["month"]): row for row in result["rows"]}
        self.assertEqual(rows[(2027, 1)]["value"], "0")
        self.assertIsNone(rows[(2027, 2)]["value"])
        self.assertEqual(rows[(2027, 3)]["value"], "20000")
        self.assertEqual(rows[(2027, 0)]["value"], "20000")
        self.assertEqual(rows[(2028, 1)]["value"], "30000")

    def test_sheet_name_wins_over_copied_title_and_total_row_is_not_project(self):
        workbook = Workbook()
        workbook.remove(workbook.active)
        sheet, row = self._sheet(
            workbook,
            "宴会（2027年）",
            "能耗",  # copied title in the source workbook
            2027,
            projects=("深圳凯骊", "总计", "自营酒店"),
        )
        sheet.cell(row=row, column=2, value=1)
        sheet.cell(row=row, column=14, value=12)
        sheet.cell(row=row + 1, column=2, value=99)
        path = self._path()
        workbook.save(path)

        result = parse_workbook(path, money_unit="YUAN")

        self.assertTrue(result["valid"], result["errors"])
        self.assertEqual({row["metric"] for row in result["rows"]}, {"revenue_banquet"})
        self.assertEqual(result["projects"], ["深圳凯骊"])
        self.assertTrue(any(item["code"] == "TITLE_MISMATCH" for item in result["warnings"]))

    def test_formula_cache_missing_and_formula_error_are_errors_not_zero(self):
        workbook = Workbook()
        workbook.remove(workbook.active)
        sheet, row = self._sheet(workbook, "GOP（2027年）", "GOP", 2027)
        sheet.cell(row=row, column=2, value="=1+1")
        sheet.cell(row=row, column=3, value="=#REF!")
        sheet.cell(row=row, column=14, value="=SUM(B4:M4)")
        path = self._path()
        workbook.save(path)

        result = parse_workbook(path, money_unit="YUAN")

        self.assertFalse(result["valid"])
        codes = {item["code"] for item in result["errors"]}
        self.assertIn("FORMULA_CACHE_MISSING", codes)
        self.assertIn("FORMULA_ERROR", codes)
        values = {(row["month"], row["value"]) for row in result["rows"]}
        self.assertIn((1, None), values)
        self.assertIn((2, None), values)

    def test_dash_placeholder_is_missing_not_zero_or_non_numeric_error(self):
        workbook = Workbook()
        workbook.remove(workbook.active)
        sheet, row = self._sheet(workbook, "出租率（2027年）", "出租率", 2027)
        for column in range(2, 15):
            sheet.cell(row=row, column=column, value="-")
        path = self._path()
        workbook.save(path)

        result = parse_workbook(path, money_unit="YUAN")

        self.assertTrue(result["valid"], result["errors"])
        self.assertNotIn("NON_NUMERIC_VALUE", {item["code"] for item in result["errors"]})
        self.assertTrue(all(item["value"] is None for item in result["rows"]))
        self.assertEqual(
            sum(item["code"] == "PLACEHOLDER_VALUE" for item in result["warnings"]),
            13,
        )

    def test_numeric_text_is_parsed_without_float_round_trip(self):
        workbook = Workbook()
        workbook.remove(workbook.active)
        sheet, row = self._sheet(workbook, "收入（2027年）", "总收入", 2027)
        sheet.cell(row=row, column=2, value="123.45")
        sheet.cell(row=row, column=14, value="123.45")
        path = self._path()
        workbook.save(path)

        result = parse_workbook(path, money_unit="YUAN")

        self.assertTrue(result["valid"], result["errors"])
        values = {(item["month"], item["value"]) for item in result["rows"]}
        self.assertIn((1, "123.45"), values)
        self.assertIn((0, "123.45"), values)

    def test_annual_sum_mismatch_keeps_source_annual_value(self):
        workbook = Workbook()
        workbook.remove(workbook.active)
        sheet, row = self._sheet(workbook, "收入（2027年）", "总收入", 2027)
        for month in range(1, 13):
            sheet.cell(row=row, column=month + 1, value=1)
        sheet.cell(row=row, column=14, value=12.01)
        path = self._path()
        workbook.save(path)

        result = parse_workbook(path, money_unit="YUAN")

        self.assertTrue(result["valid"], result["errors"])
        annual = next(row for row in result["rows"] if row["month"] == 0)
        self.assertEqual(annual["value"], "12.01")
        self.assertTrue(any(item["code"] == "ANNUAL_ROUNDING_MISMATCH" for item in result["warnings"]))

    def test_year_mismatch_is_rejected(self):
        workbook = Workbook()
        workbook.remove(workbook.active)
        sheet, row = self._sheet(workbook, "收入（2028年）", "总收入", 2027)
        sheet.cell(row=row, column=2, value=1)
        path = self._path()
        workbook.save(path)

        result = parse_workbook(path)

        self.assertFalse(result["valid"])
        self.assertIn("SHEET_YEAR_MISMATCH", {item["code"] for item in result["errors"]})
        self.assertEqual(result["rows"], [])

    def test_external_link_part_without_formula_is_not_an_error(self):
        workbook = Workbook()
        workbook.active.title = "收入（2027年）"
        path = self._path()
        workbook.save(path)
        linked = self._path("linked.xlsx")
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(linked, "w") as target:
            for item in source.infolist():
                target.writestr(item, source.read(item.filename))
            target.writestr("xl/externalLinks/externalLink1.xml", b"<externalLink/>")

        result = parse_workbook(linked)

        self.assertNotIn("EXTERNAL_LINK", {item["code"] for item in result["errors"]})
        self.assertIn("EXTERNAL_LINK_METADATA", {item["code"] for item in result["warnings"]})

    def test_residual_external_link_metadata_is_warning_without_external_formula(self):
        workbook = Workbook()
        workbook.remove(workbook.active)
        sheet, row = self._sheet(workbook, "收入（2027年）", "总收入", 2027)
        sheet.cell(row=row, column=2, value=1)
        sheet.cell(row=row, column=14, value=1)
        path = self._path()
        workbook.save(path)
        linked = self._path("linked-metadata-only.xlsx")
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(linked, "w") as target:
            for item in source.infolist():
                target.writestr(item, source.read(item.filename))
            target.writestr("xl/externalLinks/externalLink1.xml", b"<externalLink/>")

        result = parse_workbook(linked)

        self.assertTrue(result["valid"], result["errors"])
        self.assertIn("EXTERNAL_LINK_METADATA", {item["code"] for item in result["warnings"]})
        self.assertNotIn("EXTERNAL_LINK", {item["code"] for item in result["errors"]})

    def test_external_formula_is_rejected_even_with_residual_link_metadata(self):
        workbook = Workbook()
        workbook.remove(workbook.active)
        sheet, row = self._sheet(workbook, "收入（2027年）", "总收入", 2027)
        sheet.cell(row=row, column=2, value="='[history.xlsx]Sheet1'!A1")
        path = self._path()
        workbook.save(path)
        linked = self._path("linked-formula.xlsx")
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(linked, "w") as target:
            for item in source.infolist():
                target.writestr(item, source.read(item.filename))
            target.writestr("xl/externalLinks/externalLink1.xml", b"<externalLink/>")

        result = parse_workbook(linked)

        self.assertFalse(result["valid"])
        self.assertIn("EXTERNAL_LINK", {item["code"] for item in result["errors"]})


if __name__ == "__main__":
    unittest.main()
