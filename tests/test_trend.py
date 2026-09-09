import unittest

from django.contrib.auth import get_user_model
from django.test import Client, TestCase
from openpyxl import Workbook

from budgeting.excel.extract import _history_columns, _history_rows
from budgeting.excel.money import yuan_to_cents
from budgeting.management.commands.seed_demo import _generate_values
from budgeting.models import (
    BudgetCycle,
    NormalizedValue,
    Project,
    ProjectCycle,
    TemplateVersion,
    UploadVersion,
)
from budgeting.services.workflow import company_trend


class HistoryColumnsTests(unittest.TestCase):
    def test_detects_year_named_columns_only(self):
        wb = Workbook()
        ws = wb.active
        headers = {
            1: "项目", 10: "全年合计", 12: "01", 23: "12",
            25: "2025年预测", 26: "%", 27: "2026vs2025",
            32: "2024年实际", 33: "2026vs2024说明",
        }
        for col, val in headers.items():
            ws.cell(row=1, column=col, value=val)
        self.assertEqual(_history_columns(ws, 1), [("Y", "F", 2025), ("AF", "A", 2024)])

    def test_none_header_row_yields_empty(self):
        wb = Workbook()
        self.assertEqual(_history_columns(wb.active, None), [])


class HistoryRowsTests(unittest.TestCase):
    def test_builds_history_rows_for_money_and_ratio(self):
        row_meta = {
            "R0001": (23, "MONEY", "客房收入"),
            "R0029": (24, "RATIO", "出租率"),
        }
        history_cols = [("Y", "F", 2025), ("AF", "A", 2024)]
        values = {"Y23": 100.5, "AF23": 90.0, "Y24": 0.65, "AF24": 0.60}
        rows = _history_rows(None, "PL_TOTAL_WINE", "酒店损益总表", row_meta, history_cols, values)

        by_key = {(r.row_code, r.period): r for r in rows}
        self.assertEqual(len(rows), 4)
        money = by_key[("R0001", "F2025")]
        self.assertEqual(money.value_int, yuan_to_cents(100.5))
        self.assertEqual(money.unit, "MONEY")
        ratio = by_key[("R0029", "A2024")]
        self.assertEqual(ratio.ratio_num, 6000)
        self.assertEqual(ratio.ratio_den, 10000)

    def test_skips_empty_and_budget_nature_cells(self):
        row_meta = {"R0001": (23, "MONEY", "客房收入")}
        history_cols = [("Y", "F", 2025), ("AF", "B", 2026)]
        rows = _history_rows(None, "PL_TOTAL_WINE", "s", row_meta, history_cols, {})
        self.assertEqual(rows, [])


class CompanyTrendTests(TestCase):
    def setUp(self):
        self.project1 = Project.objects.create(code="P001", name="项目一")
        self.project2 = Project.objects.create(code="P002", name="项目二")
        self.cycle = BudgetCycle.objects.create(
            name="2026", budget_year=2026, status=BudgetCycle.Status.OPEN
        )
        self.template = TemplateVersion.objects.create(
            version="V1", budget_year=2026, file_path="x.xlsx",
            manifest_path="m.json", formula_manifest_hash="0" * 64,
        )
        self.uploads = []
        for project in (self.project1, self.project2):
            upload = UploadVersion.objects.create(
                project=project, cycle=self.cycle, template=self.template,
                status=UploadVersion.Status.APPROVED, original_path="o.xlsx", sha256="a" * 64,
            )
            ProjectCycle.objects.create(
                project=project, cycle=self.cycle, current_upload=upload, is_open=True
            )
            self.uploads.append(upload)

    def _ratio(self, upload, num, den):
        NormalizedValue.objects.create(
            upload=upload, report_code="PL_TOTAL_WINE", row_code="R0029", period="A2024",
            unit="RATIO", value_int=0, ratio_num=num, ratio_den=den,
            source_sheet="s", source_cell="Y24",
        )

    def test_company_trend_weights_ratio(self):
        self._ratio(self.uploads[0], 6000, 10000)
        self._ratio(self.uploads[1], 4000, 5000)
        detail = company_trend(self.cycle, "PL_TOTAL_WINE", "R0029")["A2024"]
        self.assertEqual(detail["ratio_num"], 10000)
        self.assertEqual(detail["ratio_den"], 15000)
        self.assertEqual(detail["value_int"], 6667)


class ManagementTrendViewTests(TestCase):
    def setUp(self):
        User = get_user_model()
        self.admin = User.objects.create_user("admin", password="x", role="ADMIN", is_staff=True)
        BudgetCycle.objects.create(name="2026", budget_year=2026, status=BudgetCycle.Status.OPEN)

    def test_trend_page_returns_200(self):
        client = Client()
        client.force_login(self.admin)
        response = client.get("/management/trend/", HTTP_HOST="127.0.0.1")
        self.assertEqual(response.status_code, 200)

    def test_invalid_report_falls_back_to_200(self):
        client = Client()
        client.force_login(self.admin)
        response = client.get("/management/trend/?report_code=BOGUS", HTTP_HOST="127.0.0.1")
        self.assertEqual(response.status_code, 200)


class SeedHistoryTests(TestCase):
    def setUp(self):
        self.project = Project.objects.create(code="P001", name="项目一")
        self.cycle = BudgetCycle.objects.create(
            name="2026", budget_year=2026, status=BudgetCycle.Status.OPEN
        )
        self.template = TemplateVersion.objects.create(
            version="V1", budget_year=2026, file_path="x.xlsx",
            manifest_path="m.json", formula_manifest_hash="0" * 64,
        )
        self.upload = UploadVersion.objects.create(
            project=self.project, cycle=self.cycle, template=self.template,
            status=UploadVersion.Status.APPROVED, original_path="o.xlsx", sha256="a" * 64,
        )

    def test_generate_values_emits_history_periods(self):
        manifest = {
            "reports": {
                "PL_TOTAL_WINE": {
                    "sheet": "酒店损益总表",
                    "mapping": [
                        {"row_code": "R0001", "row_label": "客房收入", "unit": "MONEY", "aggregation": "SUM"}
                    ],
                }
            }
        }
        rows = _generate_values(self.upload, manifest, self.cycle.budget_year)
        periods = {r.period for r in rows if r.row_code == "R0001"}
        self.assertIn("A2024", periods)
        self.assertIn("F2025", periods)
        self.assertIn("YEAR", periods)


if __name__ == "__main__":
    unittest.main()
