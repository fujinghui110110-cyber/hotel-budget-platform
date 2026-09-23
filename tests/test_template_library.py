from io import BytesIO
from pathlib import Path
import tempfile
from django.test import TestCase, override_settings
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.exceptions import ValidationError
from django.urls import reverse
from openpyxl import Workbook
from budgeting.models import BudgetCycle, BudgetTemplateFile, User, TemplateVersion, Project, ProjectCycle, REPORTS
from budgeting.excel.ooxml import formula_manifest, validate_upload_contract
from budgeting.services.template_library import download_distribution_template
from types import SimpleNamespace
import json
from openpyxl import load_workbook
from budgeting.services.template_library import publish_template, get_distribution_template, distribution_template_path


class TemplateLibraryTests(TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory()
        self.addCleanup(self.folder.cleanup)
        self.override = override_settings(BUDGET_STORAGE_ROOT=Path(self.folder.name))
        self.override.enable()
        self.addCleanup(self.override.disable)
        self.admin = User.objects.create_user(username='template-admin', role='ADMIN')
        source = Path(self.folder.name) / 'rules.xlsx'
        source.write_bytes(self.workbook().read())
        canonical = load_workbook(source)
        metadata = canonical.create_sheet('SYS_META')
        metadata.sheet_state = 'hidden'
        for key in ('template_version', 'budget_year', 'project_code', 'rule_version', 'formula_manifest_hash', 'project_signature'):
            metadata.append([key, ''])
        canonical.save(source)
        canonical.close()
        manifest = Path(self.folder.name) / 'rules.json'
        manifest.write_text(json.dumps({'management_v2': True, 'input_cells': {v: ['B1'] for v in REPORTS.values()}, 'reports': {k: {'sheet': v} for k, v in REPORTS.items()}}))
        _, digest = formula_manifest(source)
        template = TemplateVersion.objects.create(version='2031-standard', budget_year=2031,
            file_path=str(source), manifest_path=str(manifest), formula_manifest_hash=digest)
        self.cycle = BudgetCycle.objects.create(name='年度预算', budget_year=2031, template=template, status=BudgetCycle.OPEN)
        self.client.force_login(self.admin)

    def workbook(self):
        workbook = Workbook()
        workbook.remove(workbook.active)
        for title in REPORTS.values():
            sheet = workbook.create_sheet(title)
            sheet['A1'] = '收入'
            sheet['B1'] = 120
            sheet['C1'] = '=B1*2'
        workbook.create_sheet('自编附表')['A1'] = 88
        data = BytesIO()
        workbook.save(data)
        return SimpleUploadedFile('预算表.xlsx', data.getvalue())

    def test_publish_keeps_original_and_previous_revision(self):
        uploaded = self.workbook()
        expected = uploaded.read()
        uploaded.seek(0)
        first = publish_template(self.cycle, uploaded, self.admin)
        second = publish_template(self.cycle, self.workbook(), self.admin, '新版填报表')
        self.assertEqual(distribution_template_path(first).read_bytes(), expected)
        self.assertEqual(get_distribution_template(self.cycle), second)
        self.assertEqual(BudgetTemplateFile.objects.count(), 2)
        other = BudgetCycle.objects.create(name='次年预算', budget_year=2032)
        self.assertIsNone(get_distribution_template(other))

    def test_invalid_workbook_is_not_published(self):
        with self.assertRaises(ValidationError):
            publish_template(self.cycle, SimpleUploadedFile('损坏.xlsx', b'not excel'), self.admin)
        self.assertFalse(BudgetTemplateFile.objects.exists())

    def test_frozen_cycle_cannot_publish(self):
        self.cycle.status = BudgetCycle.FROZEN
        self.cycle.save()
        with self.assertRaises(ValidationError):
            publish_template(self.cycle, self.workbook(), self.admin)
        self.assertFalse(BudgetTemplateFile.objects.exists())

    def test_management_page_uses_selected_year(self):
        response = self.client.get(reverse('planning_templates'), {'cycle': self.cycle.pk})
        self.assertContains(response, '2031 年')
        self.assertContains(response, '发布到当前预算版本')
        self.assertNotContains(response, 'V3-2027')

    def test_published_download_can_return_through_upload_contract(self):
        item = publish_template(self.cycle, self.workbook(), self.admin)
        project = Project.objects.create(code='HOTEL', name='新酒店')
        path = download_distribution_template(item, project, self.cycle)
        upload = SimpleNamespace(template=self.cycle.template, project=project, cycle=self.cycle, cycle_id=self.cycle.pk)
        self.assertFalse([i for i in validate_upload_contract(upload, path) if i[0] == 'P0'])
        wb = load_workbook(path)
        self.assertEqual(wb[next(iter(REPORTS.values()))]['B1'].value, 120)
        self.assertEqual(wb['自编附表']['A1'].value, 88)
        wb.close()

    def test_incompatible_summary_formula_cannot_be_published(self):
        data = self.workbook()
        wb = load_workbook(BytesIO(data.read()))
        wb[next(iter(REPORTS.values()))]['C1'] = '=999'
        output = BytesIO()
        wb.save(output)
        wb.close()
        with self.assertRaisesMessage(ValidationError, '汇总表公式'):
            publish_template(self.cycle, SimpleUploadedFile('修改表.xlsx', output.getvalue()), self.admin)
        self.assertFalse(BudgetTemplateFile.objects.exists())

    def test_project_download_selects_open_cycle_and_keeps_identity(self):
        item = publish_template(self.cycle, self.workbook(), self.admin)
        project = Project.objects.create(code='ONLY', name='当前酒店')
        user = User.objects.create_user(username='project-download', role='PROJECT', project=project)
        ProjectCycle.objects.create(project=project, cycle=self.cycle, is_open=True)
        self.client.force_login(user)
        response = self.client.get(reverse('project_template_download'))
        self.assertEqual(response.status_code, 200)
        output = Path(self.folder.name) / 'download.xlsx'
        output.write_bytes(b''.join(response.streaming_content))
        from budgeting.excel.ooxml import read_sys_meta
        self.assertEqual(read_sys_meta(output)['project_code'], 'ONLY')
        self.assertEqual(str(read_sys_meta(output)['budget_year']), '2031')
        self.cycle.status = BudgetCycle.FROZEN
        self.cycle.save()
        response = self.client.get(reverse('project_template_download'))
        self.assertRedirects(response, reverse('project_dashboard'))

    def test_macro_file_is_rejected_before_publishing(self):
        import zipfile
        output = BytesIO(self.workbook().read())
        with zipfile.ZipFile(output, 'a') as archive:
            archive.writestr('xl/vbaProject.bin', b'not executed')
        with self.assertRaises(ValidationError):
            publish_template(self.cycle, SimpleUploadedFile('宏表.xlsm', output.getvalue()), self.admin)
        self.assertFalse(BudgetTemplateFile.objects.exists())
