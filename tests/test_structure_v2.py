import json
import tempfile
from pathlib import Path
from types import SimpleNamespace

from django.test import SimpleTestCase
from django.core import signing
from openpyxl import Workbook, load_workbook

from budgeting.excel.structure_v2 import validate_v2_structure
from budgeting.services.template_delivery import _inject_sys_meta


class StructureV2Tests(SimpleTestCase):
    def test_contract_accepts_located_issue_tuples_without_crashing(self):
        from budgeting.excel.ooxml import validate_upload_contract
        workbook = load_workbook(self.template)
        workbook.create_sheet('SYS_META')
        workbook.save(self.template)
        self.upload.template.version = 'V2'
        self.upload.template.rule_version = 'R2'
        self.upload.template.formula_manifest_hash = 'fingerprint'
        self.upload.project = SimpleNamespace(code='P001')
        self.upload.cycle = SimpleNamespace(budget_year=2027)
        self.upload.cycle_id = 1
        payload = dict(project_code='P001', budget_year=2027, cycle_id=1, template_version='V2')
        meta = dict(payload, rule_version='R2', formula_manifest_hash='fingerprint',
                    project_signature=signing.dumps(payload, salt='budget-template-v1'))
        _inject_sys_meta(self.template, self.submitted, meta, input_cells={})
        issues = validate_upload_contract(self.upload, self.submitted)
        self.assertTrue(any(len(issue) >= 4 for issue in issues))

    def test_signed_v2_download_preserves_numeric_system_constants(self):
        workbook = load_workbook(self.template)
        workbook.active['A1'] = 2027
        workbook.active['B2'] = '=B1*2'
        workbook.create_sheet('SYS_META')
        workbook.save(self.template)
        _inject_sys_meta(self.template, self.submitted, {}, input_cells={'填报': ['B1']})
        result = load_workbook(self.submitted)
        self.assertEqual(result['填报']['A1'].value, 2027)
        self.assertEqual(result['填报']['B2'].value, '=B1*2')
        self.assertIsNone(result['填报']['B1'].value)
        self.assertEqual(validate_v2_structure(self.upload, self.submitted), [])

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.template = self.root / 'template.xlsx'
        self.submitted = self.root / 'upload.xlsx'
        self.manifest = self.root / 'manifest.json'
        workbook = Workbook()
        workbook.active.title = '填报'
        workbook.active.append(['固定指标', 100])
        workbook.save(self.template)
        self.manifest.write_text(json.dumps({'management_v2': True, 'input_cells': {'填报': ['B1']}}))
        self.upload = SimpleNamespace(template=SimpleNamespace(file_path=str(self.template), manifest_path=str(self.manifest)))

    def test_only_explicit_input_changes_allowed(self):
        workbook = load_workbook(self.template)
        workbook.active['B1'] = 300
        workbook.save(self.submitted)
        self.assertEqual(validate_v2_structure(self.upload, self.submitted), [])

    def test_label_and_new_outside_input_are_rejected(self):
        workbook = load_workbook(self.template)
        workbook.active['A1'] = '伪造指标'
        workbook.active['C100'] = 12
        workbook.save(self.submitted)
        issues = validate_v2_structure(self.upload, self.submitted)
        self.assertEqual(len(issues), 2)
        self.assertTrue(all(issue[1] == 'V2_SYSTEM_CELL_CHANGED' for issue in issues))

    def test_extra_sheet_rejected(self):
        workbook = load_workbook(self.template)
        workbook.create_sheet('私加表')
        workbook.save(self.submitted)
        self.assertEqual(validate_v2_structure(self.upload, self.submitted)[0][1], 'V2_SHEET_STRUCTURE')
