import json
import shutil
import tempfile
from pathlib import Path
from types import SimpleNamespace

from django.core import signing
from django.test import SimpleTestCase
from openpyxl import Workbook, load_workbook

from budgeting.excel.ooxml import formula_manifest, validate_upload_contract, validate_xlsx_zip
from budgeting.models import REPORTS
from budgeting.services.template_delivery import _inject_sys_meta


class SummaryUploadScopeTests(SimpleTestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.root = Path(temp.name)
        self.template = self.root / 'template.xlsx'
        self.uploaded = self.root / 'upload.xlsx'
        manifest = self.root / 'manifest.json'
        workbook = Workbook()
        workbook.remove(workbook.active)
        for name in REPORTS.values():
            sheet = workbook.create_sheet(name)
            sheet.append(['收入', 100, '=B1*2'])
        workbook.create_sheet('项目附表')['A1'] = '原说明'
        workbook.create_sheet('SYS_META')
        workbook.save(self.template)
        _, digest = formula_manifest(self.template)
        manifest.write_text(json.dumps({
            'management_v2': True,
            'reports': {code: {'sheet': name} for code, name in REPORTS.items()},
            'input_cells': {name: ['B1'] for name in REPORTS.values()},
        }))
        payload = dict(project_code='P001', budget_year=2027, cycle_id=1, template_version='V3')
        self.upload = SimpleNamespace(
            template=SimpleNamespace(file_path=str(self.template), manifest_path=str(manifest),
                                     version='V3', rule_version='R3', formula_manifest_hash=digest),
            project=SimpleNamespace(code='P001'), cycle=SimpleNamespace(budget_year=2027), cycle_id=1,
        )
        _inject_sys_meta(self.template, self.uploaded, dict(
            payload, rule_version='R3', formula_manifest_hash=digest,
            project_signature=signing.dumps(payload, salt='budget-template-v1'),
        ), input_cells={})
        shutil.copyfile(self.uploaded, self.template)

    def test_freeform_supplementary_sheets_do_not_block_summary(self):
        workbook = load_workbook(self.uploaded)
        workbook['项目附表']['A1'] = '项目自由说明'
        workbook['项目附表']['B1'] = '=2+3'
        sheet = workbook.create_sheet('自定义营销测算', 0)
        sheet['A1'] = '=SUM(1,2)'
        sheet['B1'] = '#REF!'
        workbook.save(self.uploaded)
        issues = validate_upload_contract(self.upload, self.uploaded)
        self.assertFalse([issue for issue in issues if issue[0] == 'P0'], issues)
        self.assertTrue(any(issue[0] == 'P2' and issue[1] == 'EXCEL_ERROR' for issue in issues))

    def test_summary_formula_changes_still_block(self):
        workbook = load_workbook(self.uploaded)
        workbook[next(iter(REPORTS.values()))]['C1'] = '=999'
        workbook.save(self.uploaded)
        issues = validate_upload_contract(self.upload, self.uploaded)
        self.assertTrue(any(issue[0] == 'P0' and issue[1] == 'FORMULA_FINGERPRINT' for issue in issues))

    def test_summary_cached_error_still_blocks(self):
        workbook = load_workbook(self.uploaded)
        workbook[next(iter(REPORTS.values()))]['B1'] = '#REF!'
        workbook.save(self.uploaded)
        issues = validate_upload_contract(self.upload, self.uploaded)
        self.assertTrue(any(issue[0] == 'P0' and issue[1] == 'EXCEL_ERROR' for issue in issues))

    def test_supplementary_sheets_do_not_relax_package_security(self):
        import zipfile
        with zipfile.ZipFile(self.uploaded, 'a') as package:
            package.writestr('xl/externalLinks/externalLink1.xml', '<externalLink/>')
            package.writestr('xl/vbaProject.bin', b'macro')
        codes = {issue[1] for issue in validate_xlsx_zip(self.uploaded)}
        self.assertIn('EXTERNAL_LINK', codes)
        self.assertIn('MACRO_OR_OLE', codes)
